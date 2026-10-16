import json,sys
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(sys.argv[1]);TOP=ROOT.parent;B=TOP/'experiment';R5=B/'round5';R6=B/'round6';R7=B/'round7'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def main():
    s5=read(R5/'results/summary.json');s6=read(R6/'results/summary.json');s7=read(R7/'results/summary.json');audit=read(ROOT/'results/mapping_audit.json');u=read(ROOT/'results/usage_sample_audit.json');sizes=read(ROOT/'docs/official_sizes_by_format.json')['total'];sel=read(R5/'results/selection.json');verify=read(R5/'results/verification.json');scores=s5['scores'];cmp=s5['comparison_vs_previous_best']['selected']
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21.59);sec.page_height=Cm(27.94);sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for st in doc.styles:
        for x in list(st.element.iter(qn('w:pBdr'))):x.getparent().remove(x)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Calibri';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor.from_string('000000')
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.07;doc.styles['Normal'].paragraph_format.space_after=Pt(5);doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 2'].font.size=Pt(12)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;foot.add_run('Google 2019 算电预测研究  |  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    for rr in foot.runs:rr.font.size=Pt(8)
    uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');doc.settings.element.append(uf)
    def p(t,b=False):q=doc.add_paragraph();q.add_run(t).bold=b;return q
    def h(t):doc.add_heading(t,1)
    def page(t):q=doc.add_heading(t,1);q.paragraph_format.page_break_before=True
    def table(head,rows,widths):
        tb=doc.add_table(rows=1,cols=len(head));tb.autofit=False
        for c,w in zip(tb.columns,widths):c.width=Cm(w)
        for c,t,w in zip(tb.rows[0].cells,head,widths):c.text=str(t);c.width=Cm(w)
        repeat=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for c,t,w in zip(tb.add_row().cells,row,widths):c.text=str(t);c.width=Cm(w)
        for i,row in enumerate(tb.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E1EDF1' if i==0 else ('F4F7F8' if i%2==0 else 'FFFFFF'));pr.append(shade);borders=OxmlElement('w:tcBorders')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
                pr.append(borders);mar=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'65');x.set(qn('w:type'),'dxa');mar.append(x)
                pr.append(mar)
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.05
                    for rr in para.runs:rr.font.size=Pt(9.5);rr.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    doc.add_heading('Google PowerData 2019\n算电融合与 ClusterData 扩展研究汇报',0)
    p('研究日期 2026 年 10 月 1 日　｜　第五轮至第七轮实际结果')
    h('当前结论')
    old=scores['old_equal_mix']['wape'];new=scores['selected']['wape'];previous=scores['anchor']['wape'];rel=(old-new)/old*100
    p(f'冻结主方案的次日曲线 WAPE 为 {new:.4f}%，此前最优为 {previous:.4f}%，相对下降 {cmp["relative_reduction_percent"]:.2f}%；相对最早旧组合 {old:.4f}% 下降 {rel:.2f}%。已有小幅实测改善，尚未获得大幅、显著且跨月份可靠提升的证据。',True)
    p('本次实际完成融合网络训练、基础模型微调、残差校准，并从 Google 官方下载 ClusterData2019 的完整机器状态表，将 CPU 和内存样本加入对照实验。数据、模型、程序、预测和报告继续保存在原桌面文件夹。')
    rows=[]
    for label,key in [('最早 Chronos 与 TiRex 等权组合','old_equal_mix'),('此前冻结固定基础组合','anchor'),('第五轮冻结主方案','selected'),('第五轮融合网络 两种子平均','fusion')]:
        q=scores[key];rows.append([label,f'{q["wape"]:.4f}',f'{q["peak_mape"]:.4f}'])
    z=s6['scores']['ridge_resource_1000000'];rows.append(['第六轮资源状态纯候选',f'{z["wape"]:.4f}',f'{z["peak_mape"]:.4f}']);table(['方法','曲线 WAPE %','峰值 MAPE %'],rows,[10.1,3.5,3.5])
    p('2.3668% 是验证期选出的资源家族在纯候选形态下的测试结果；验证期保守组合规则选择该分支权重为 0。因此主方案仍为 2.3742%，不能在看到测试分数后把纯候选改报为已获验证的最终方案。',True)
    p('你的融合架构已形成可运行研究模型，当前曲线 WAPE 为 2.9029%，高于主方案，精度较低。融合网络峰值误差较小，但仍未超过此前前一天同刻基线的 1.0682%。不以局部指标代替整体有效性。')
    page('官方数据下载与覆盖范围')
    p('ClusterData2019 覆盖 2019 年 5 月的八个 Borg 集群 a 至 h，与 PowerData 的同月时间和机器标识可关联。来源是 Google cluster-data 官方仓库、官方说明 PDF 和官方公开存储桶；未采用第三方镜像。数据及说明许可为 CC BY 4.0，需保留来源署名。')
    table(['已保存内容','文件数或记录数','原文件大小'],[['完整 machine_events JSON 与 Parquet','16 个文件','9,541,963 字节'],['完整 machine_attributes JSON 与 Parquet','16 个文件','217,453,745 字节'],['a 集群单个 instance_usage JSON 分片','4,483,198 条记录','432,275,123 字节'],['官方说明与五张表 schema','本地 docs 目录','含 16 页官方 PDF']],[7.2,4.8,5.1])
    p('32 个机器表文件总计 226,995,708 字节，约 216.48 MiB；使用分片约 412.25 MiB。全部下载对象核对官方 MD5 和字节数，JSON 经过 gzip 解压与逐行解析。Parquet 对象虽有 .parquet.gz 后缀，实际以 PAR1 标识起止，没有外层 gzip；保存原名并核对 MD5 与格式标识，尚未解码其行数。建模采用完整 JSON 机器事件。')
    p(f'当前清单中，instance_usage 的 JSON 导出为 {sizes["instance_usage_json"]["files"]:,} 个文件、{sizes["instance_usage_json"]["bytes"]:,} 字节，约 {sizes["instance_usage_json"]["bytes"]/2**40:.2f} TiB；Parquet 为 {sizes["instance_usage_parquet"]["files"]:,} 个文件、{sizes["instance_usage_parquet"]["bytes"]:,} 字节，约 {sizes["instance_usage_parquet"]["bytes"]/2**40:.2f} TiB。两者是不同格式的导出，不能相加解释为一份数据。旧 README 的约 2.4 TiB 与当前对象清单不一致。')
    p('桌面所在 C 盘核查时剩余约 103 GiB，无法容纳任何一个完整使用表版本。当前没有下载完整 collection_events、instance_events 或 instance_usage，没有把局部文件冒称完整数据集。完整负载可通过 BigQuery 聚合后只下载小型序列表；查询方案已保存，尚缺可用项目与账号，没有提交查询或产生查询费用。',True)
    table(['关联审核','处理结果'],[['可关联供电域','55 个；i 与 j 两个场站没有 ClusterData 资源记录'],['机器到域映射歧义','c 中机器 102、124 对应多个域，无有效时间，已排除'],['缺失机器和容量','保留观测覆盖率、容量已知率和异常标记'],['资源指标含义','调度器可用容量与机器在线状态，不是实际使用率']],[5.2,11.9])
    page('建模协议与融合网络实现')
    table(['阶段','范围与用途'],[['训练开发模型','历史截至 05-14；训练起报索引 7 至 13'],['前四个验证日','05-15 至 05-18，选择家族成员、轮数或微调检查点'],['后四个验证日','05-19 至 05-22，选择组合权重；一个标准误规则偏向低权重'],['最终训练','只重训已冻结配置，标签截至 05-22'],['次日测试','05-23 至 05-31；57 域，49,243 个有效点']],[5.2,11.9])
    p('统一主指标为供电域宏平均 WAPE：每域先计算有效时点绝对误差总和除以真值绝对值总和，再对 57 域等权平均。历史输入每次可读取已结束的前一日；在线误差特征只使用当前起报之前已观测的误差。测试未来标签不进入输入或参数更新。')
    p('第五轮采用七天历史功率 P 与生产功率代理 W，带两类质量掩码。RevIN 只用历史计算均值与尺度；5、25 点移动平均形成多尺度分解，历史 FFT 提供频域表示。P、W 各有编码器，交叉注意力融合，16、32、96 点尺度对应三个专家，门控每样本只运行选中的两个专家。输出 P 与 W 的 0.1、0.5、0.9 分位数，再做开发期残差校准。')
    p('机器状态尚未参与第五轮融合网络，随后在第六轮进入资源残差分支；部分真实任务 CPU/内存在第七轮进入对照分支。因此未宣称完整 P、真实 W、动态 R 的端到端联合架构已经胜出。工作量代理不能替代实测算力。',True)
    table(['实际探索','范围与结果'],[['融合网络','宽度 16 与 32；4、10、20 轮；交叉、频域、稠密专家消融，共 15 个开发检查点'],['重复训练','按前四日选择宽度 16 与 10 轮，固定第二随机种子，不按测试挑种子'],['TimesFM 3 微调','输出层及尾部，3 个配置各 50、150、300 步；另有一个组合目标配置的 3 个检查点'],['正则残差校准','Ridge 正则 1,000、10,000、100,000；入选 100,000 并与旧基础组合各半'],['真实资源状态','匹配的三档 Ridge 对照加两档树模型，完整历史资源状态'],['部分任务使用样本','三档 Ridge 匹配对照；掩码标明样本缺失']],[4.5,12.6])
    p('容量输出限制在 [0,1]，代理生产分位数不高于对应总功率分位数；这属于归一化边界与分量顺序约束，不是实际电量守恒。未观测供电容量 kW、冷却与各任务真实用电，不能据此声称发现完整算电物理机制。')
    page('第五轮结果与不确定性')
    p(f'第五轮主方案相对此前固定组合改善 {cmp["difference_pp"]:.6f} 个百分点。57 个域中 {cmp["domains_won"]} 个改善，9 个测试日中 {cmp["days_won"]} 个改善，10 个集群中 {cmp["cells_won"]} 个改善。平均改善并不意味着每个域都改善。')
    table(['评价','实测结果与含义'],[['日期块 95% 改进区间',f'[{cmp["day_block_95"][0]:.6f}, {cmp["day_block_95"][1]:.6f}] 个百分点，跨零'],['日期块与集群联合 95% 区间',f'[{cmp["day_and_cell_95"][0]:.6f}, {cmp["day_and_cell_95"][1]:.6f}] 个百分点，跨零'],['重采样范围','5,000 次；连续三日块和集群重采样；仅九个测试日'],['主方案日均 WAPE',f'{scores["selected"]["daily_wape"]:.4f}%'],['主方案峰值 MAPE',f'{scores["selected"]["peak_mape"]:.4f}%，仍高于前一天峰值基线']],[6.2,10.9])
    p('匹配消融固定宽度 32、10 轮和同一种子，避免把不同模型大小混为结构收益。交叉注意力在本次匹配实验有帮助；去除频域后的测试分数略低，说明不能为每个设计模块都宣称正贡献。不同随机种子的平均也不能代表充分的重复实验。')
    table(['固定宽度 32 的消融','验证 WAPE %','测试 WAPE %'],[['完整模型','3.0992',f'{scores["ablation_fusion32"]["wape"]:.4f}'],['无交叉注意力','3.1076',f'{scores["ablation_no_cross"]["wape"]:.4f}'],['无频域','3.1265',f'{scores["ablation_no_frequency"]["wape"]:.4f}'],['稠密专家','3.1598',f'{scores["ablation_dense_moe"]["wape"]:.4f}']],[8.1,4.5,4.5])
    unc=s5['uncertainty'];table(['名义 80% 区间','测试覆盖率','平均区间宽度'],[[label,f'{unc[k]["coverage"]*100:.2f}%',f'{unc[k]["mean_width"]:.5f}'] for label,k in [('融合网络原始分位数','fusion_raw'),('融合网络校准后','fusion_calibrated'),('主方案残差区间','selected_residual_interval')]],[8.1,4.5,4.5])
    p('区间宽度是利用率单位。开发期拟合的融合校准提高覆盖率，但时间序列相关、分布漂移与反复选模使名义覆盖率没有独立保证。主方案区间作为研究输出，不能直接解释为供电安全保障。')
    p('TimesFM 3 梯度训练最初出现非有限值，失败日志保留。修正只停止迭代归一化统计量的反向梯度；校验微调路径与官方路径的前向输出完全一致、输出层梯度有限。该修正不被当成精度创新。')
    page('ClusterData 资源与算力样本的实测作用')
    p('第六轮用完整机器事件构造每个十五分钟末的在线机器比例、CPU 和内存调度容量、容量已知率及状态异常比例。按事件时间推进状态，只使用起报之前已发生的事件；不拿下一天资源容量当已知协变量，也不向过去补填未来首次出现的机器。')
    rows=[]
    for a in [10000,100000,1000000]:
        c=s6['scores'][f'ridge_control_{a}']['wape'];z=s6['scores'][f'ridge_resource_{a}']['wape'];rows.append([f'{a:,}',f'{c:.4f}',f'{z:.4f}',f'{c-z:+.6f}'])
    table(['相同 Ridge 正则','不加资源 WAPE %','加资源 WAPE %','资源增益 百分点'],rows,[4.1,4.3,4.3,4.4])
    p('在 1,000,000 正则下，资源模型比匹配对照仅改善约 0.0030 个百分点；更强正则本身解释了大部分相对第五轮的变化。不能把纯候选的全部改善归因于新增数据。验证期 50% 混合虽有更低点估计，但收益落在配对集群标准误范围内，保守规则选择权重 0。')
    p('映射审核发现两台机器存在多个供电域归属且缺少有效时间。更正前结果已归档，更正后排除歧义机器重新拟合，原冻结配置和组合权重保持不变，不利用测试结果改选。')
    p(f'第七轮保留 a 集群分片中的 JOB 记录：原始 {u["collection_type_counts"]["0"]:,} 条，经过时间、映射与字段检查使用 {u["used_job_rows"]:,} 条；其余 {u["collection_type_counts"]["1"]:,} 条为 ALLOC_SET，避免与任务层重复计算，未合并成总负载。分片包含全月时间，但各域有观测的十五分钟槽只有 139 至 2,611 个，总槽数为 2,976。')
    p('样本 CPU、内存是已观测任务窗口的时长加权均值，缺失保留掩码；窗口只在 end_time 后成为可用输入。按结束时间分箱会带来最多约五分钟的时序延迟。文件分片的入选概率未知，不能推断总体 CPU 总量或认为样本具有代表性。')
    table(['正则','仅完整资源状态 %','加部分任务样本 %'],[[f'{a:,}',f'{s7["scores"][f"ridge_control_{a}"]["wape"]:.4f}',f'{s7["scores"][f"ridge_resource_{a}"]["wape"]:.4f}'] for a in [10000,100000,1000000]],[4.1,6.5,6.5])
    p('加入部分 CPU/内存样本在三档匹配对照上均未降低测试 WAPE。开发期也选择不含任务样本的家族，最终新增组合权重为 0。该结果支持获取完整且可审核的负载序列，不能支持“大量使用 CPU 信息后精度已经大幅提高”。',True)
    page('本地部署与复核文件')
    p('所有新增原文件、聚合数据、模型参数、实验日志及结果保存在原桌面研究目录 GooglePowerData2019_预测研究_20260929。此前轮次文件和 E 盘原件保留。主运行入口已更新；资源和样本候选单独保留为研究对照。')
    table(['桌面入口','输出与用途'],[['01 研究用新组合预测次日','第五轮主方案；round5/results/selected_deployment_origin31_1day.csv'],['02 TiRex 预测七天','沿用已有七天较优方案，WAPE 2.8803%'],['06 融合网络探索预测次日','两种子功率与生产代理、校准区间；供研究对照'],['07 资源状态候选预测次日','资源纯候选，未按验证规则替换主方案'],['08 部分算力样本对照预测次日','第七轮冻结家族，样本分支未获选']],[6.4,10.7])
    p('默认输入截止 2019-05-31，示例预测 2019-06-01；不是今天的实时负荷。六月无真值，导出成功只证明本机推理可运行。新增分支只评估次日 96 点，没有把次日分数冒用为七天成绩。')
    table(['重点核查位置','内容'],[['round5/docs 与 results','协议、冻结 SHA256、全部开发分数、逐点预测、分域分日误差、区间覆盖'],['round5/models','两种子最终网络、微调检查点、最终校准参数'],['round6 与 round7','资源及样本对照模型、冻结选择、匹配分数、因果性检查'],['google-clusterdata-2019/raw','a 至 h 官方机器文件和 a 的使用分片'],['google-clusterdata-2019/aggregated','15 分钟资源状态与部分任务样本序列'],['google-clusterdata-2019/docs','对象清单、格式拆分大小、官方 schema 与下载校验'],['google-clusterdata-2019/code','可复现实验和尚未执行的完整使用表提取 SQL']],[6.4,10.7])
    p(f'第五轮 {len(verify["checks"])} 项实现核查全部通过；第六轮和第七轮的未来输入扰动与冻结哈希检查全部通过。这些检查验证时间因果性和实现一致性，不能证明没有过拟合。一个月数据此前已反复查看，模型预训练与公开数据是否重叠也未排除。')
    p('TimesFM 3.0 权重采用 TimesFM Non-Commercial v1.0 许可，本目录仅用于非商业本地研究；不用于生产或对外分发。原数据及其他模型的具体许可文件保存在各原目录。',True)
    page('明确限制与下一阶段依据')
    p('当前已取得的结果是小幅改进与可复现的候选发现，没有充分证据声称大幅提高精度。主方案改进的联合区间跨零；资源候选受开发期规则约束未获选；局部 CPU 样本没有带来可靠收益。保留这些负结果有助于避免继续把复杂模型或更多字段当作必然有效的创新。',True)
    p('下一项实质进展应是完整任务 CPU/内存观测，而非再随机下载几个分片。完整提取 SQL 已按八个官方表、机器到域映射和 JOB 层准备；剔除歧义映射，保留观测数量和字段缺失，不把预留实例与任务重复相加。查询模板需要替换授权项目并上传映射表，尚未经过 BigQuery 服务端语法或费用验证，不能称为已经完成全量提取。')
    p('有可用 Google Cloud 项目与账号后，先做 dry run 获取扫描字节数，再按批准上限设置 maximumBytesBilled 执行。把完整资源使用聚合表下载到同一桌面目录，复核机器覆盖、异常窗口、监控遗漏、跨日时序及样本缺失，再替换代理 W。新增输入只用历史，未来 W 必须由模型预测，不能由数据集未来真值直接提供。')
    p('最终研究结论需要新的月份或不参与选模的额外场站确认。现有同月冻结和因果性检查可以约束本轮流程，但不能恢复独立测试。若新数据仍无法取得，报告应维持探索性表述，不能通过不断挑选这九天的最佳分数来证明模型普遍领先。')
    h('来源与实现依据')
    for title,url in [('Google 官方 ClusterData2019 与 CC BY 4.0','https://github.com/google/cluster-data/blob/master/ClusterData2019.md'),('Google cluster usage traces v3 官方说明','https://drive.google.com/file/d/10r6cnJ5cJ89fPWCgj7j4LtLBqYN9RiI9/view'),('Google 官方字段与 JOB ALLOC_SET 定义','https://github.com/google/cluster-data/blob/master/clusterdata_trace_format_v3.proto'),('Google 官方 PowerData2019','https://github.com/google/cluster-data/blob/master/PowerData2019.md'),('Google BigQuery dry run 与费用控制','https://docs.cloud.google.com/bigquery/docs/best-practices-costs'),('RevIN 架构依据','https://github.com/ts-kim/RevIN'),('稀疏专家架构参考','https://github.com/Time-MoE/Time-MoE'),('TimesFM 官方实现与许可','https://github.com/google-research/timesfm')]:
        q=p(title+'\n'+url);q.paragraph_format.space_after=Pt(3)
        for rr in q.runs:rr.font.size=Pt(9)
    p('RevIN 与稀疏专家为实现参考。本研究是针对数据量定制的小网络，并非原论文或 Time-MoE 预训练权重的逐项复现。',False)
    path=TOP/'Google PowerData 2019 算电融合与 ClusterData 扩展研究汇报.docx';doc.save(path);print(path,flush=True)
if __name__=='__main__':main()
