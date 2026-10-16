"""Seven-page report from completed round10 repair1 and verified D-drive move.

Run with the Codex bundled document Python after the authoring marker.  This
builder never runs an experiment, changes a selection, or fabricates a result.
The parent agent must render and inspect every page before delivering the DOCX.
"""
import argparse
import datetime as dt
import hashlib
import json
import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TITLE = "Google PowerData 2019 原生算电融合与D盘迁移研究汇报"
EXPECTED_TOP = Path(r"${PROJECT_ROOT}")
LABELS = {
    "timesfm_usage": "TimesFM 3 CPU 内存",
    "timesfm_usage_resource": "TimesFM 3 CPU 内存 资源",
    "chronos_usage": "Chronos 2 CPU 内存",
    "chronos_usage_resource": "Chronos 2 CPU 内存 资源",
    "timesfm_control": "TimesFM 3 匹配控制",
    "chronos_control": "Chronos 2 匹配控制",
}
METRIC_KEYS = ("wape", "daily_wape", "peak_mape")
FROZEN_SCRIPT_SHA256 = "fd0e90d490ab947dc155aff95e7e646dc9c95f86f3c2404fa6423dfd549cbfeb"
REPAIR_POLICY = {"timesfm_test": "ieee", "chronos_test": "tf32",
                 "timesfm_deploy": "ieee", "chronos_deploy": "ieee"}


def read(path):
    path = Path(path)
    if not path.is_file():
        raise RuntimeError(f"Required completed evidence is absent: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def all_checks(proof, label):
    require(proof.get("all_passed") is True, f"{label} is not complete and passed")
    checks = proof.get("checks")
    if checks is not None:
        require(isinstance(checks, dict) and len(checks) > 0,
                f"{label} has no implementation checks")
        boolean_checks = [v for v in checks.values() if isinstance(v, bool)]
        require(boolean_checks and all(boolean_checks), f"{label} contains a failed check")


def date(value):
    return dt.datetime.fromisoformat(value.replace("Z", "+00:00"))


def checked_metrics(metrics, label):
    for key in METRIC_KEYS:
        value = metrics.get(key)
        require(isinstance(value, (float, int)) and math.isfinite(value) and value >= 0,
                f"Missing or invalid metric {label}/{key}")
    require(int(metrics.get("valid_points", 0)) > 0, f"No valid evaluation points: {label}")


def load_repair_evidence(r, cluster, s, deploy, sel, protocol_hash, freeze_hash):
    """Require actual repair execution and numerical diagnosis before authoring."""
    provenance_path = r / "docs/repair1_provenance.json"
    provenance = read(provenance_path)
    require(sel["candidate"] == "chronos_usage_resource" and sel["chosen"]["alpha"] == 0,
            "This repair report must retain the actual original alpha-zero choice")
    repair_script = cluster / "code/native_covariate_experiment_repair1.py"
    require(sha256(cluster / "code/native_covariate_experiment.py") == FROZEN_SCRIPT_SHA256,
            "Original scientific script changed")
    required = dict(repair_id="repair1", status="approved_for_execution",
                    original_script_sha256=FROZEN_SCRIPT_SHA256,
                    original_selection_sha256=freeze_hash,
                    original_protocol_sha256=protocol_hash,
                    new_script_sha256=sha256(repair_script),
                    scientific_configuration_unchanged=True,
                    source_and_weight_hashes_unchanged=True,
                    test_target_truth_used=False, test_reselection=False,
                    original_candidate=sel["candidate"], original_alpha=sel["chosen"]["alpha"])
    require(all(provenance.get(key) == value for key, value in required.items()),
            "Repair provenance does not retain the real frozen configuration")
    require(provenance.get("global_fp32_policy") == REPAIR_POLICY and
            provenance.get("local_fp32_policy") == "none", "Repair precision policy differs")
    execution = dict(repair_id="repair1", implementation_sha256=sha256(repair_script),
                     provenance_sha256=sha256(provenance_path),
                     original_selection_sha256=freeze_hash,
                     original_protocol_sha256=protocol_hash,
                     global_fp32_policy=REPAIR_POLICY, local_fp32_policy="none")
    require(s.get("execution_fingerprint") == execution and
            deploy.get("execution_fingerprint") == execution,
            "Summary and deployment do not identify the verified repair execution")
    require(Path(s["repair_provenance"]).resolve() == provenance_path.resolve() and
            Path(deploy["repair_provenance"]).resolve() == provenance_path.resolve(),
            "Repair provenance path differs")
    require(s.get("original_failure_evidence_preserved") is True,
            "Original failure evidence was not retained")
    diagnostics = []
    evidence = provenance.get("diagnostic_evidence", [])
    require(evidence, "Repair requires actual diagnostic evidence")
    for item in evidence:
        path = Path(item["path"])
        require(path.is_file() and sha256(path) == item["sha256"],
                "Diagnostic evidence is absent or its SHA256 differs")
        if path.suffix.lower() == ".json":
            value = read(path)
            if value.get("matching_legacy_conditions") is not None:
                diagnostics.append((path, value))
    require(len(diagnostics) == 1, "Require one complete precision diagnosis with repeat results")
    diagnostic_path, diagnostic = diagnostics[0]
    require(diagnostic.get("origin") == 22 and
            diagnostic.get("future_target_truth_scored") is False and
            diagnostic.get("test_reselection") is False and
            diagnostic.get("immutable_preserved") is True and
            diagnostic.get("all_adjacent_repeats_equal") is True,
            "Precision diagnosis must retain inputs, selection and repeated predictions")
    require(diagnostic["immutable_before"] == diagnostic["immutable_after"],
            "Diagnostic immutable fingerprints changed")
    groups = {}
    for mode in ("ieee", "tf32"):
        rows = [row for row in diagnostic["calls"]
                if row.get("flags", {}).get("global_fp32_precision") == mode]
        require(len(rows) >= 4, f"Missing repeated {mode} numerical diagnosis")
        require(all(row["vs_old_test_cache"]["points"] == 480 for row in rows),
                "Precision diagnosis must compare the same five-domain forecast")
        maximum = max(row["vs_old_test_cache"]["maximum_absolute_difference"] for row in rows)
        if mode == "tf32":
            require(maximum == 0 and all(row["vs_old_test_cache"]["bitwise_values_equal"] is True and
                    row["vs_old_test_cache"]["registered_tolerance_passed"] is True for row in rows),
                    "TF32 historical-cache reproduction was not actually exact")
        else:
            require(maximum > 2e-6 and all(row["vs_old_test_cache"]["registered_tolerance_passed"] is False
                                        for row in rows), "Initial IEEE regression failure not documented")
        groups[mode] = dict(repeats=len(rows), maximum_absolute_difference=maximum)
    dev_regression = read(r / "results/dev_control_regression.json")
    test_regression = read(r / "results/repair1/test_control_regression.json")
    for backbone in ("timesfm", "chronos"):
        dev_control = dev_regression[backbone]
        require(set(dev_control["origins"]) == {str(o) for o in range(14, 22)} and
                dev_control["maximum_absolute_difference"] == 0 and dev_control["passed"] is True,
                "The eight original development controls must actually remain exact")
        test_control = test_regression[backbone]
        require(set(test_control["origins"]) == {str(o) for o in range(22, 31)} and
                test_control["passed"] is True and
                test_control["atol"] == 2e-6 and test_control["rtol"] == 1e-6 and
                all(row["passed"] is True for row in test_control["origins"].values()),
                "All nine repaired test controls must pass the original tolerance")
    for split, keys, origins in (
        ("test", list(LABELS), range(22, 31)),
        ("deploy", [sel["candidate"], "chronos_control"], [31]),
    ):
        for key in keys:
            for origin in origins:
                cache = r / "results/cache/repair1" / f"{split}_{key}_origin{origin}.npz"
                meta = read(cache.with_suffix(".json"))
                require(cache.is_file() and sha256(cache) == meta["prediction_archive_sha256"] and
                        meta["identity"]["execution_fingerprint"] == execution and
                        meta["identity"]["origin"] == origin and meta["identity"]["split"] == split,
                        "An actual repair prediction cache is absent or has the wrong execution identity")
    return dict(provenance=provenance, execution=execution, diagnostic=diagnostic,
                diagnostic_path=diagnostic_path, groups=groups,
                dev_regression=dev_regression, test_regression=test_regression)


def load_evidence(top):
    require(top.is_absolute(), "The research root must be absolute")
    top = top.resolve(strict=True)
    require(top == EXPECTED_TOP, "Refuse to create the D-drive report under a different root")
    b = top / "experiment"
    r = b / "round10"
    cluster = top / "google-clusterdata-2019"
    s = read(r / "results/repair1/summary.json")
    dev = read(r / "results/dev_summary.json")
    sel = read(r / "results/selection.json")
    protocol_path = r / "docs/protocol.json"
    protocol = read(protocol_path)
    registration = read(r / "docs/registration.json")
    source = read(r / "docs/source_audit.json")
    audit = read(cluster / "results/complete_projection_audit.json")
    deploy = read(r / "results/repair1/deployment_verification.json")
    migration_root = cluster / "results/migration_to_d"
    migration = read(migration_root / "migration_verification.json")
    compatibility = read(migration_root / "migration_compatibility.json")
    for proof, label in ((s, "round10 test"), (dev, "round10 development"),
                         (audit, "complete source audit"), (deploy, "deployment"),
                         (migration, "D-drive migration"), (compatibility, "migration compatibility")):
        all_checks(proof, label)
    require(bool(s.get("completed_at_utc")), "Test pipeline has not completed")
    require(s["selection"] == sel and dev["selection"] == sel,
            "Summary does not retain the exact frozen selection")
    require(s.get("test_reselection") is False and s.get("independent_confirmation") is False,
            "This report requires the exploratory reused-test protocol")
    require(sel.get("test_scored_before_freeze") is False,
            "The new round selection must precede its test scoring")
    require(date(registration["registered_at_utc"]) <= date(sel["frozen_at_utc"]) <
            date(s["completed_at_utc"]), "Registration, freeze, and completed-test chronology differs")
    freeze_hash = (r / "results/selection.sha256").read_text(encoding="utf-8").strip()
    require(sha256(r / "results/selection.json") == freeze_hash, "Frozen selection SHA256 differs")
    protocol_hash = sha256(protocol_path)
    require(sel["protocol_sha256"] == protocol_hash == registration["protocol_sha256"],
            "Protocol fingerprint differs")
    repair = load_repair_evidence(r, cluster, s, deploy, sel, protocol_hash, freeze_hash)
    require(source["fingerprints"] == sel["source_fingerprints"], "Source fingerprints differ")
    require(source.get("source_complete") is True and source["source_files"] == 12085,
            "Source usage must contain all 12085 cell-a selected-column projections")
    require(len(source["cpu_domains"]) == 5 and audit["source_files"] == 12085,
            "The actual CPU scope must remain five cell-a domains")
    require(list(protocol["candidates"]) == list(LABELS)[:4] and
            list(protocol["matched_controls"]) == list(LABELS)[4:],
            "Registered candidate families differ")
    require(protocol["development_origins"] == list(range(14, 22)) and
            protocol["test_origins"] == list(range(22, 31)), "Time split differs")
    require(dev["uncertainty"].get("calibration_origins") == protocol["development_origins"] and
            s["uncertainty"].get("calibration_origins") == protocol["development_origins"],
            "Residual calibration must identify the actual eight development origins")
    alpha = sel["chosen"]["alpha"]
    require(alpha in protocol["alpha_grid"], "Frozen alpha is outside the declared grid")
    require(sel["retained_previous_best"] == (alpha == 0), "Alpha-zero retention flag differs")
    for tag in ("selected_scores", "prior_scores"):
        for scope in ("all57", "cella5"):
            checked_metrics(s[tag][scope], f"{tag}/{scope}")
    for key in LABELS:
        checked_metrics(s["candidate_scores"][key]["raw_native_cella5"], key)
        if key in protocol["candidates"]:
            for scope in ("all57", "cella5"):
                checked_metrics(s["candidate_scores"][key]["controlled_delta"][scope], f"{key}/{scope}")
    if alpha == 0:
        require(s["selected_scores"] == s["prior_scores"] and
                s["checks"].get("alpha_zero_exact_prior") is True,
                "Alpha-zero prediction must exactly preserve the prior")
    require(deploy["selection_alpha"] == alpha and deploy["candidate"] == sel["candidate"],
            "Deployment does not use the frozen choice")
    require(deploy.get("future_truth_available") is False and deploy.get("not_live_forecast") is True,
            "June1 deployment must not be presented as a verified live forecast")
    require(deploy["input_end"] == "2019-05-31" and deploy["forecast_start"] == "2019-06-01",
            "Deployment time boundary differs")
    output = Path(deploy["output"])
    require(output.resolve() == (r / "results/repair1/selected_deployment_origin31_1day.csv").resolve() and
            output.is_file() and sha256(output) == deploy["output_sha256"],
            "Deployment output is absent or its content fingerprint differs")
    require(migration.get("all_names_lengths_and_sha256_equal") is True and
            Path(migration["destination"]) == top and migration["files"] > 0 and migration["bytes"] > 0,
            "Whole-folder byte-for-byte migration is not verified")
    hash_records = migration_root / "all_files_sha256.jsonl"
    require(hash_records.is_file() and sha256(hash_records) == migration["evidence_sha256"],
            "Full-file migration hash evidence is absent or changed")
    required_compatibility = ("desktop_junction", "real_root_D", "python_runs_from_D",
                              "python_runs_via_old_alias", "torch_cuda_available",
                              "data_source_hashes_equal", "source_shape_names_equal")
    require(all(compatibility["checks"].get(k) is True for k in required_compatibility),
            "Migration compatibility checks are incomplete")
    r8b = read(b / "round8b/results/summary.json")
    r8c = read(b / "round8c/results/summary.json")
    r9 = read(b / "round9/results/summary.json")
    all_checks(r8c, "preserved round8c"); all_checks(r9, "preserved round9")
    return dict(top=top, b=b, r=r, cluster=cluster, s=s, dev=dev, sel=sel,
                protocol=protocol, source=source, audit=audit, deploy=deploy,
                migration=migration, compatibility=compatibility, r8b=r8b, r8c=r8c, r9=r9,
                freeze_hash=freeze_hash, repair=repair)


def f(value):
    return f"{value:.6f}"


def interval(values):
    require(len(values) == 2 and all(math.isfinite(v) for v in values), "Invalid confidence interval")
    return f"[{values[0]:.6f}, {values[1]:.6f}]"


def build(e):
    s, sel, protocol = e["s"], e["sel"], e["protocol"]
    alpha = sel["chosen"]["alpha"]
    selected, prior = s["selected_scores"], s["prior_scores"]
    compare = s["comparison"]
    global_gain = compare["all57_improvement_pp"]
    five_gain = compare["cella5_improvement_pp"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    sec.left_margin = sec.right_margin = Cm(2)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor.from_string("000000")
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Title"].paragraph_format.space_after = Pt(9)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "第十轮冻结回测和整目录D盘迁移"
    doc.core_properties.author = ""
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Google 2019 负荷预测研究  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)

    def p(text, bold=False, size=None):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(str(text)); run.bold = bold
        if size is not None:
            run.font.size = Pt(size)
        return paragraph

    def heading(text, new_page=False):
        paragraph = doc.add_heading(text, 1)
        paragraph.paragraph_format.page_break_before = new_page
        return paragraph

    def table(headers, rows, widths):
        tb = doc.add_table(rows=1, cols=len(headers)); tb.autofit = False
        for column, width in zip(tb.columns, widths):
            column.width = Cm(width)
        for cell, text, width in zip(tb.rows[0].cells, headers, widths):
            cell.text = str(text); cell.width = Cm(width)
        tb.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for values in rows:
            require(len(values) == len(headers), "Report table schema differs")
            for cell, text, width in zip(tb.add_row().cells, values, widths):
                cell.text = str(text); cell.width = Cm(width)
        for number, row in enumerate(tb.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for column, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                fill = OxmlElement("w:shd"); fill.set(qn("w:fill"), "E1EDF1" if number == 0 else "FFFFFF")
                properties.append(fill)
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "bottom", "left", "right"):
                    node = OxmlElement("w:" + side)
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        node.set(qn("w:" + key), value)
                    borders.append(node)
                properties.append(borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    node = OxmlElement("w:" + side); node.set(qn("w:w"), "70"); node.set(qn("w:type"), "dxa")
                    margins.append(node)
                properties.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.06
                    if column > 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5); run.bold = number == 0
        gap = doc.add_paragraph(); gap.paragraph_format.space_after = Pt(1)
        gap.paragraph_format.line_spacing = 0.5
        return tb

    # Page 1: the decision and both evaluation scopes.
    doc.add_heading("Google PowerData 2019\n原生算电融合与D盘迁移研究汇报", 0)
    p("第十轮实际冻结结果  |  2019 年 5 月探索性回测")
    heading("本轮结论")
    p(f"冻结功率方案全 57 域次日曲线 WAPE 为 {f(selected['all57']['wape'])}%，"
      f"上一方案为 {f(prior['all57']['wape'])}%，误差下降 {global_gain:+.6f} 个百分点。"
      f"开发期选定 {LABELS[sel['candidate']]}，新增权重为 {alpha:g}。", True)
    if alpha == 0:
        p("本轮开发期规则保留原最佳模型，预测逐值与上一方案相同。四组原生 CPU 内存候选与两个匹配控制已经实际评价；"
          "本轮没有取得大幅提高，也没有按测试分数重新选择候选。", True)
    elif global_gain > 0:
        p("本轮冻结模型的测试点估计有所改善。改善幅度、五域结果和日期重采样区间见下表及后页；"
          "该月测试已被反复查看，尚不能据此声称取得可靠的大幅提升。", True)
    else:
        p("本轮开发期冻结的新权重未提高整体测试曲线精度。保留真实结果及配置，不根据已看到的测试成绩改选模型或权重。", True)
    p("新方式将真实历史 CPU、内存和资源状态直接送入冻结基础模型，检验其相对同主干控制的增量。"
      "只有 cell-a 五域具有完整真实任务使用信息；其余 52 域保持上一方案。")
    table(["方案及评价范围", "曲线 WAPE %", "日均 WAPE %", "峰值 MAPE %"],
          [[label, *[f(score[k]) for k in METRIC_KEYS]] for label, score in (
              ("上一方案 全部 57 域", prior["all57"]), ("冻结方案 全部 57 域", selected["all57"]),
              ("上一方案 cell-a 五域", prior["cella5"]), ("冻结方案 cell-a 五域", selected["cella5"]))],
          [7.2, 3.4, 3.4, 3.4])
    p(f"五域曲线误差下降 {five_gain:+.6f} 个百分点。全 57 域日期重采样 95% 改进区间为 "
      f"{interval(compare['all57_paired_date_bootstrap95_pp'])} 个百分点；"
      f"五域区间为 {interval(compare['cella5_paired_date_bootstrap95_pp'])}。正值表示误差下降。"
      "本轮权重为零，两组预测完全相同，因此改进区间退化为零；这不表示绝对预测误差或未来不确定性为零。")
    p("全部研究目录已经实体迁移到 D 盘，桌面原路径保留兼容入口。历史负面结果、原始数据、模型权重与部署文件均保留。")

    # Page 2: model inputs and exact controlled correction.
    heading("建模方式与时间边界", True)
    p("本轮冻结已有 TimesFM 3 和 Chronos 2 基础权重，不做梯度微调。每次起报联合输入五条功率目标、"
      "可用生产功率历史代理，以及五域真实 CPU 和内存历史。资源版额外输入在线比例和每映射机器的平均调度 CPU 容量。"
      "最多 30 个变量，最长 14 天历史，预测未来 96 个十五分钟点。")
    p("TimesFM 使用 past_only_covariates，Chronos 使用 past_covariates。CPU 与资源输入严格截止起报边界，"
      "没有送入未来真实 CPU，也没有用未来实际负载填充缺失值。任务 CPU 缺失只存在于没有完整源的其他域，"
      "这些域不作为物理零值输入。生产功率代理保留质量标记与历史可用率筛选。")
    table(["登记分支", "历史输入", "使用目的"], [
        ["两个匹配控制", "功率与可用生产功率代理", "复算同主干原生预测"],
        ["两组 CPU 内存候选", "控制输入加真实 CPU 与内存", "检验直接算电信息增量"],
        ["两组资源候选", "再加在线比例与平均调度容量", "检验额外资源信息增量"],
    ], [4.2, 8.0, 5.2])
    p("正式预测在五域执行：上一方案 + α ×（含协变量的原生功率预测 − 同主干匹配控制预测），再限制到 [0,1]。"
      "两条原生预测使用相同主干、目标、日期和原功率代理。该差分尽量剔除直接换主干的影响，"
      "但仍是探索性的建模对照，不能据此证明 CPU 对功率的物理因果效应。")
    table(["时段", "预测目标日期", "作用"], [
        ["前四开发日", "2019 年 5 月 15 至 18 日", "从四组候选选择一组"],
        ["后四开发日", "2019 年 5 月 19 至 22 日", "选择新增权重"],
        ["全部八开发日", "2019 年 5 月 15 至 22 日", "冻结后校准残差区间"],
        ["九个测试日", "2019 年 5 月 23 至 31 日", "冻结后评价 不重新选参"],
    ], [4.2, 8.0, 5.2])
    p("权重候选仅为 0、0.25、0.5、1。后四开发日采用日期配对的一标准误规则，优先取满足阈值的最小权重。"
      "五域 WAPE 是开发期主指标，全 57 域是辅助指标。没有将九天测试点作为新的选参验证集。")
    p("未来值扰动只检查输入截断实现；同主干控制与旧缓存的一致性只检查预测复现。这些实现检查不能证明没有过拟合。")

    # Page 3: true raw controls and alpha-one attribution candidates.
    heading("原生控制与算电候选的实际结果", True)
    p("下表先列原生匹配控制。它们只预测 cell-a 五域，并不等同于已有最佳组合；因此不能直接将五域控制成绩称作全 57 域成绩。")
    table(["五域原生控制", "曲线 WAPE %", "日均 WAPE %", "峰值 MAPE %"],
          [[LABELS[key], *[f(s["candidate_scores"][key]["raw_native_cella5"][m]) for m in METRIC_KEYS]]
           for key in protocol["matched_controls"]], [7.2, 3.4, 3.4, 3.4])
    p("四组候选使用固定 α=1 的同主干差分修正，供信息贡献对照。正式冻结组合使用前页的开发期权重，"
      "不能按下表最低测试值替换。所有候选对其他 52 域均沿用上一方案。")
    table(["α为1的差分候选", "57 域曲线\nWAPE %", "五域曲线\nWAPE %", "五域日均\nWAPE %", "五域峰值\nMAPE %"],
          [[LABELS[key], f(s["candidate_scores"][key]["controlled_delta"]["all57"]["wape"]),
            *[f(s["candidate_scores"][key]["controlled_delta"]["cella5"][m]) for m in METRIC_KEYS]]
           for key in protocol["candidates"]], [6.2, 2.8, 2.8, 2.8, 2.8])
    p("含算力输入的原生主干本身也已保存，下面列五域曲线结果，便于区分原生输出与差分校正输出。")
    table(["含协变量的原生模型", "五域曲线 WAPE %"],
          [[LABELS[key], f(s["candidate_scores"][key]["raw_native_cella5"]["wape"])]
           for key in protocol["candidates"]], [11.8, 5.6])
    p("WAPE 沿用此前口径：先计算每个供电域在质量掩码下的绝对误差总和与真实值绝对总和之比，再对各域取算术平均。"
      "日均和峰值另行评价，不能以较低曲线误差推断峰值也改善。六位小数仅用于识别微小差异。")

    # Page 4: actual DEV ledger, chronological freeze and uncertainty.
    heading("开发期选择与冻结复核", True)
    table(["前四开发日候选", "五域曲线 WAPE %", "57 域曲线 WAPE %"],
          [[LABELS[row["candidate"]], f(row["scores"]["cella5"]["wape"]),
            f(row["scores"]["all57"]["wape"])] for row in sel["early_development"]],
          [9.0, 4.2, 4.2])
    p(f"前四日选定 {LABELS[sel['candidate']]}。下表为该候选在后四开发日的权重评价，"
      "配对标准误按日期计算；权重选择不读取本轮测试排名。")
    table(["新增权重", "五域曲线 WAPE %", "配对标准误\n百分点"],
          [[f"{row['alpha']:g}", f(row["primary_wape"]), f(row["paired_day_se"])]
           for row in sel["late_development"]], [4.0, 7.6, 5.8])
    late = {row["alpha"]: row for row in sel["late_development"]}
    gain025 = late[0]["primary_wape"] - late[.25]["primary_wape"]
    p(f"α=0.25 的后四日五域改善只有 {gain025:.6f} 个百分点，"
      f"配对标准误为 {late[.25]['paired_day_se']:.6f} 个百分点。"
      "这点改善远小于日间不确定性，因此一标准误规则选择 α=0，保留原预测。", True)
    p(f"实际冻结时间为 {sel['frozen_at_utc']}，测试结果完成时间为 {s['completed_at_utc']}。"
      "选择记录的 SHA256 与汇总及部署均一致，恢复缓存时保留原选择。", size=10)
    p("冻结记录 SHA256  " + e["freeze_hash"], size=8.7)
    p(f"本轮测试 {len(s['checks'])} 项检查、部署 {len(e['deploy']['checks'])} 项检查全部通过。"
      "包括未来输入扰动、匹配控制复现、冻结哈希、功率有限且有界、其他 52 域保持原预测以及六月起报边界。")
    uncertainty = s["uncertainty"]
    p(f"区间宽度由全部 {len(e['dev']['uncertainty']['calibration_origins'])} 个开发日残差相对历史七日均值校准，"
      f"起报编号为 {e['dev']['uncertainty']['calibration_origins']}，名义覆盖为 {uncertainty['nominal_coverage_percent']:g}%。"
      f"开发期经验覆盖 {uncertainty['development_calibration_coverage_percent']:.2f}%，"
      f"测试经验覆盖 {uncertainty['test_empirical_coverage_percent']:.2f}%。"
      "开发期校准覆盖不是独立覆盖保证，模型选择与序列相关也会影响该区间。")
    p("测试重采样以九个预测日期为单位，各域共同重采样，保持域间相关。区间描述该月的结果波动；"
      "由于只有九日且测试早已复用，不能恢复独立显著性检验或证明未来月份的覆盖率。")

    # Page 5: source-backed numerical reproduction repair, separate from modeling.
    heading("计算精度复现与实施修复", True)
    repair = e["repair"]
    p("开发期两主干控制共 16 个起报点与旧档案逐值一致，测试期 TimesFM 控制也复现通过。"
      "最初的 Chronos 测试控制未通过既定容差，所以实验停止并保留失败缓存，未把中断当作模型精度失败。")
    table(["origin22 数值状态", "重复次数", "与旧控制最大绝对差"], [
        ["IEEE 浮点矩阵计算", repair["groups"]["ieee"]["repeats"],
         f"{repair['groups']['ieee']['maximum_absolute_difference']:.9e}"],
        ["TF32 矩阵计算", repair["groups"]["tf32"]["repeats"],
         f"{repair['groups']['tf32']['maximum_absolute_difference']:.9e}"],
    ], [7.3, 3.0, 7.1])
    p("诊断保持同一输入、权重和目标排列，只比较既有预测档案。两种线程数下重复结果一致；"
      "TF32 状态逐值恢复旧 Chronos 测试控制。这里的差值属于归一化功率预测的数值复现，不能称为 WAPE 精度提升。")
    p("旧测试程序先完成两次 Chronos 微调，再保存零样本控制。Transformers 在训练时设置全局 TF32；"
      "Chronos 只保存并恢复局部矩阵和 cuDNN 设置。局部恢复为 none 时仍继承全局 TF32，"
      "因此开发期和旧测试档案具有不同数值执行状态。权重本身仍为 FP32。")
    p("修复版根据源代码、旧运行日志和重复诊断恢复该状态：TimesFM 测试与部署采用 IEEE，"
      "Chronos 测试采用 TF32，Chronos 部署采用 IEEE；局部标记统一继承全局。"
      "各候选与其匹配控制共享同一阶段的状态，退出后恢复调用前标记。")
    p("该修复没有改变来源、权重、四个候选、历史长度、开发规则、冻结选择或原回归容差，"
      "没有读取测试真值选择旗标。不同阶段数值状态的历史差异仍需透明保留；"
      "以后在新数据上确认时，应事先固定统一的数值策略。", True)
    p("原 DEV 与冻结文件保持原样，TEST 和六月部署的修复缓存独立保存于 results/cache/repair1。"
      "结果位于 results/repair1，并绑定实施脚本和修复来源哈希；报告仅在完整 TEST、部署及全部检查通过后生成。")
    p("实施脚本 SHA256  " + repair["execution"]["implementation_sha256"], size=8.5)
    p("修复来源 SHA256  " + repair["execution"]["provenance_sha256"], size=8.5)

    # Page 6: physical scope, prior negative results, and what can be concluded.
    heading("研究范围与精度平台判断", True)
    audit = e["audit"]
    p("PowerData 覆盖 2019 年 5 月 57 个供电域。真实任务使用仅完成官方 cell-a 的 12,085 个 Parquet 文件必要列读取，"
      "五域各有 2,976 个十五分钟槽。未取得其他七个集群完整任务使用，也未读取全部字段。资源事件虽覆盖更多集群，"
      "不能因此称任务使用为八集群全量。")
    p("JOB聚合排除ALLOC_SET；读取列未包括完整任务去重键。使用记录按完成窗口归属，"
      "功率按区间起点标注，两者可能存在时间对齐差异。")
    p(f"源对象全部字段合计 {audit['original_objects_bytes']:,} 字节"
      f"（{audit['original_objects_bytes']/2**30:.2f} GiB），"
      f"取得并保留的官方压缩列及间隙为 {audit['retained_bytes']:,} 字节"
      f"（{audit['retained_bytes']/2**30:.2f} GiB）。原列字节、逐对象聚合与来源身份保持完整，迁移未改变其内容。")
    p("CPU 单位为 NCU，内存为归一化调度单位，功率为供电容量利用率，均不直接等于物理 kW。"
      "调度可用 CPU 容量不能当作实际用电的硬上限。第九轮条件单调回归只检验条件响应假设，不构成能量守恒。")
    table(["已完成的历史探索", "冻结曲线 WAPE %", "保留决定"], [
        ["第八轮最佳组件微调", f(e["r8b"]["scores"]["selected"]["wape"]), "现有固定上一方案"],
        ["第八轮组合权重探索", f(e["r8c"]["scores"]["selected"]["wape"]), "曲线变差 未采用"],
        ["第九轮完整算力融合", f(e["r9"]["scores"]["selected"]["all57"]["wape"]), "验证规则保留上一方案"],
    ], [7.9, 4.6, 4.9])
    p("第八轮组件微调相对第五轮仅改善约 0.000052 个百分点，原组合权重搜索与第九轮完整算力方案均未支持升级。"
      "这些负面结果、模型及已冻结选择继续保留，本轮没有重写其选择规则。")
    if alpha == 0 or global_gain <= 0:
        p("本轮结果继续支持现有短月样本进入工程平台期的判断。原生信息融合已经实际检验，未获得可推荐的大幅提高；"
          "这不等于穷尽所有模型或证明理论最低误差。扩大同一九天的测试搜索更容易产生选择偏差。", True)
    else:
        p("本轮原生融合给出可供后续验证的点估计改善，仍不能认定突破了泛化平台期。"
          "已有候选与本轮选择只使用同一月份，结论需要新的日期或场站确认。", True)
    p("基础监督信息仍只有一个月，连续窗口高度相关。只修正五域的当前设计也限制整体改进范围。"
      "下一步应优先取得未参与选模的新月份或新场站，冻结现有方法做确认，之后再决定是否扩大模型结构或局部适配。")

    # Page 7: physical move, compatibility, actual export, resources and sources.
    heading("D盘迁移与本机交付", True)
    migration = e["migration"]
    p("整个研究文件夹现在实体保存在以下位置，数据、模型、依赖、程序、历史结果与新报告均在其中。", True)
    p(str(e["top"]), size=10)
    table(["迁移核验项目", "实际证据"], [
        ["全目录复制时文件数", f"{migration['files']:,}"],
        ["全目录复制时逻辑字节", f"{migration['bytes']:,}  即 {migration['bytes']/2**30:.3f} GiB"],
        ["逐文件核验", "相对路径 长度 SHA256 全部一致"],
        ["核验时间", migration["verified_utc"]],
    ], [6.4, 11.0])
    p("完整哈希清单与兼容核验保存于 google-clusterdata-2019/results/migration_to_d。"
      "桌面原文件夹是指向 D 的目录联接，原入口仍可使用；新旧路径 Python 运行、CUDA、数据哈希与形状核对通过。"
      "E 盘早期功率副本保留。上表是搬移时清单，不包含随后新增的实验与报告文件。")
    p("本轮预测导出在 experiment/round10/results/repair1/selected_deployment_origin31_1day.csv，"
      "已核对 57 域 × 96 点共 5,472 行。历史截止 2019 年 5 月 31 日，预测 6 月 1 日 00:00 至 23:45。"
      "六月没有真值，该文件证明本机推理可运行，不代表今天的负荷或已经验证的未来精度。")
    p("导出包含冻结预测、α为1的研究候选、原方案及残差区间。CPU 和内存列是历史七日均值，"
      "不是未来 CPU 预测；其他 52 域保留缺失标记。")
    policy = protocol["resources"]
    p(f"两主干串行推理，CPU 线程为 {policy['cpu_threads']}。加载前要求可用物理内存至少 "
      f"{policy['minimum_available_physical_gib_before_load']:g} GiB、可用提交额度至少 "
      f"{policy['minimum_available_commit_gib_before_load']:g} GiB、可用显存至少 "
      f"{policy['minimum_available_cuda_gib_before_load']:g} GiB。资源不足时停止下一操作并保留缓存；"
      "外层接续启动前要求可用物理内存至少 3 GiB，运行中低于物理 1.5 GiB 或提交 2 GiB 连续两次时仅停止本实验。"
      "启动门槛不等于运行内存上限。本轮未新增数据或权重下载。")
    p("Google 数据和说明采用 CC BY 4.0，需保留署名。TimesFM 3 权重采用 TimesFM Non-Commercial v1.0，"
      "仅限非商业 非生产用途。本目录用于本地研究，不能因此获得商业部署或权重分发许可。")
    for label, url in (("Google 功率源与许可", "https://github.com/google/cluster-data/blob/master/PowerData2019.md"),
                       ("Google 任务使用源与许可", "https://github.com/google/cluster-data/blob/master/ClusterData2019.md"),
                       ("TimesFM 官方实现", "https://github.com/google-research/timesfm"),
                       ("Chronos 官方实现", "https://github.com/amazon-science/chronos-forecasting")):
        q = p(label + "  " + url, size=8.5); q.paragraph_format.space_after = Pt(3)
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("top", help="Absolute D-drive research folder")
    parser.add_argument("--validate-only", action="store_true", help="Read evidence without authoring a DOCX")
    args = parser.parse_args()
    e = load_evidence(Path(args.top))
    if args.validate_only:
        print(json.dumps(dict(all_required_evidence_passed=True,
                              frozen_candidate=e["sel"]["candidate"],
                              frozen_alpha=e["sel"]["chosen"]["alpha"],
                              selection_sha256=e["freeze_hash"],
                              execution_fingerprint=e["repair"]["execution"]), ensure_ascii=False), flush=True)
        return
    doc = build(e)
    output = e["top"] / (TITLE + ".docx")
    doc.save(output)
    print(output, flush=True)


if __name__ == "__main__":
    main()
