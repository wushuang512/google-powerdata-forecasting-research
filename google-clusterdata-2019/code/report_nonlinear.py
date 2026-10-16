"""Create the report only from completed and independently verified results."""
import datetime,json,sys
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
TOP=Path(sys.argv[1]).resolve();R=TOP/'experiment/round11'
OUT=TOP/'Google PowerData 2019 最佳模型非线性校准调优汇报.docx'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def main():
    summary=read(R/'results/summary.json');selection=read(R/'results/selection.json');dev=read(R/'results/dev_summary.json')
    verification=read(R/'results/independent_verification.json');run=read(R/'results/guarded_pipeline_run.json')
    assert verification['all_passed'] and read(R/'results/deployment_verification.json')['all_passed']
    assert selection['chosen']['alpha']==0 and not summary['large_gain_rule_met']
    if '--validate-only' in sys.argv:print('Real results and independent reconstruction verified');return
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
        pp=s.element.find(qn('w:pPr'))
        if pp is not None:
            for child in list(pp):
                if child.tag==qn('w:pBdr'):pp.remove(child)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(8)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    def para(t):return doc.add_paragraph(t)
    def h(t):return doc.add_paragraph(t,'Heading 1')
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for i,w in enumerate(widths):t.columns[i].width=Inches(w)
        for j,value in enumerate(headers):t.rows[0].cells[j].text=value
        for row in rows:
            cells=t.add_row().cells
            for j,value in enumerate(row):cells[j].text=str(value)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders);marg=OxmlElement('w:tcMar')
                for side,value in [('top','60'),('bottom','60'),('left','100'),('right','100')]:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');marg.append(e)
                pr.append(marg)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF5');pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:r.font.size=Pt(10.5);r.bold=i==0
        spacer=doc.add_paragraph();spacer.paragraph_format.line_spacing=Pt(6);spacer.paragraph_format.space_after=Pt(6)
        return t
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    doc.add_paragraph('Google PowerData 2019\n最佳模型非线性校准调优汇报','Title')
    para('第十一轮研究记录   2026年10月3日   保存位置 D盘研究目录')
    h('结论与采用方案')
    para('本轮完成了在 round8b 最佳组合上的四组小范围非线性校准实验，没有找到可采纳的全域曲线精度提升。前四个开发日起报选择了 L1 损失和 7 叶节点配置，但后四个开发日未支持加入该修正，冻结权重为 0。正式预测继续采用原最佳模型。')
    table(['57域测试指标','保留的原最佳','未采用的候选'],[
        ['次日曲线 WAPE',f'{summary["scores"]["prior"]["all57"]["wape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["wape"]:.6f}%'],
        ['日均 WAPE',f'{summary["scores"]["prior"]["all57"]["daily_wape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["daily_wape"]:.6f}%'],
        ['峰值 MAPE',f'{summary["scores"]["prior"]["all57"]["peak_mape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["peak_mape"]:.6f}%']], [2.9,1.8,1.8])
    para('未采用候选指完整使用本轮非线性修正的对照，曲线误差增加约 0.005609 个百分点。它的峰值误差有所下降，但本轮预先指定的主要指标是全域曲线 WAPE，不能事后改用峰值名次采纳模型。')
    para(f'候选曲线改善量的按日期共同重采样探索性95%区间为 [{summary["candidate_day_bootstrap95_pp"][0]:.6f}, {summary["candidate_day_bootstrap95_pp"][1]:.6f}] 个百分点，跨越零。最终冻结方案与原最佳逐点相同，改善量为零；其差值区间为零只是方案相同，不表示预测没有误差。')
    h('本轮改动的依据')
    para('已有线性校准、组合权重与基础模型权重微调的增益很小。本轮改用浅层 LightGBM 拟合历史负荷与预测误差之间的非线性关系，并用同一特征集的已冻结 Ridge 校准器作为匹配对照。Chronos 2、TiRex 2、TimesFM 3 及原最佳权重全部保留，没有重新加载或训练基础模型，也没有新增下载。')
    doc.add_page_break()
    h('开发期筛选与参数范围')
    para('数据为2019年5月57个PDU的15分钟功率利用率。起报7至13用于训练，14至17选择配置，18至21选择修正强度，22至30作为冻结后的九个测试起报。起报编号等于已经完成的日数，起报22对应预测5月23日。最终校准器用起报7至21重训；不使用测试目标训练。')
    rows=[['原最佳',f'{dev["prior_early4"]["wape"]:.6f}',f'{dev["prior_late4"]["wape"]:.6f}']]
    labels={'l1_leaves7':'L1损失 7叶','l1_leaves15':'L1损失 15叶','l2_leaves7':'L2损失 7叶','l2_leaves15':'L2损失 15叶'}
    for r in dev['candidates']:rows.append([labels[r['model']],f'{r["early4"]["wape"]:.6f}',f'{r["late4"]["wape"]:.6f}'])
    table(['配置','前4日 WAPE %','后4日 WAPE %'],rows,[2.9,1.8,1.8])
    para('四组配置固定80棵树、学习率0.03、最大深度4、叶节点最少1024行、L2正则100。使用原校准器的30个历史与时间场站特征，并按历史七日均值归一化残差。新增功率修正最多为历史均值的1%，再乘冻结强度，最终限制到0至1。这是数值稳定约束，不能解释为能量守恒。')
    table(['L1 7叶的修正强度','后4日 WAPE %'],[[f'{r["alpha"]:g}',f'{r["wape"]:.6f}'] for r in selection['weights']],[3.4,3.1])
    para('后四日的所有非零强度均比原最佳差；完整修正在四个日期中三日变差。一标准误规则取日期配对误差与场站配对误差中的较大者，优先较小修正强度，最终取0。冻结记录与SHA256在测试评分前生成，未根据测试结果改选其他配置。')
    para('训练阶段只有7个预测日期。虽然包含多域多时间点，这些样本相关，不能等同于大量独立日期。开发期和测试期都已经被历轮研究反复使用，当前结果是探索证据，严格划分与冻结也不能保证不存在适应性过拟合。')
    doc.add_page_break()
    h('实际核查与部署文件')
    para(f'独立重建核查共 {verification["check_count"]} 项全部通过，包括原历史特征逐点一致、冻结筛选重建、测试指标重算、旧文件指纹、未来真值扰动检查，以及部署预测重建。未来真值扰动检查只证明已检查的实现没有读取起报后的真值，不能证明模型不会过拟合。')
    para('正式与未采用候选均已导出，包含57域5472行，时间从2019年6月1日00:00到23:45，数值有限、在0至1内，区间顺序正确。正式功率逐点等于 round8b。80%区间依据开发期残差经验校准，没有独立覆盖率保证。六月没有真值，此输出是历史复现推理，不是今天的电负荷预测。')
    table(['文件位置','用途'],[
        ['experiment/round11/docs/protocol.json','登记参数与数据边界'],
        ['experiment/round11/results/selection.json','开发期冻结选择'],
        ['experiment/round11/results/summary.json','测试指标与负面结论'],
        ['experiment/round11/models','四个开发模型与最终候选'],
        ['experiment/round11/results','预测 逐日逐域表 与审计'],
        ['13_最佳模型非线性校准_预测次日.cmd','复现本轮冻结预测']], [4.3,2.2])
    para(f'全部新增程序和产出保存在 ${PROJECT_ROOT}。桌面原目录仍为兼容链接。本轮使用已有预测缓存和两个CPU线程，受保护实验耗时约10秒，运行期间可用物理内存最低 {run["minimum_available_physical_bytes"]/2**30:.2f} GiB，没有触发内存保护。自动续跑保持停用，当前实验已结束。')
    h('平台判断与后续研究空间')
    para('本轮负面结果进一步支持当前短期数据与输入上的工程平台判断，但没有证明精度已到理论极限。原生CPU内存资源融合、第九轮MoE、第八轮权重微调和本轮非线性校准均没有支持可靠升级。继续大范围搜索同一月份最低测试误差，难以形成可信的突破结论。')
    para('更有价值的下一步是取得独立月份或新的可对齐功率观测，用其确认已冻结方法，再考虑工作负载变化、供电温度冷却等新增解释变量。目前只有a集群五域具有完整真实任务使用特征；本轮没有扩充该范围。功率为归一化利用率，CPU单位是NCU，不能直接当成kW。Google数据为CC BY 4.0；组合中的TimesFM 3权重限非商业、非生产使用。')
    doc.save(OUT);print(OUT)
if __name__=='__main__':main()
