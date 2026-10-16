"""Final research report from real frozen results; refuse absent completed results."""
import json,sys
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(sys.argv[1]);TOP=ROOT.parent;B=TOP/'experiment'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def main():
    r8=read(B/'round8/results/summary.json');old=read(B/'round5/results/summary.json')['scores'];ft=read(B/'round8b/results/summary.json');s=read(B/'round9/results/summary.json');audit=read(ROOT/'results/complete_projection_audit.json');verify=read(ROOT/'results/finetuning_verification.json');assert audit['all_passed'] and s['all_passed'] and verify['all_passed'];sel=s['selection'];scores=s['scores'];previous=ft['scores']['selected']['wape'];current=scores['selected']['all57']['wape'];cmp=s['comparison_vs_previous'];key=sel['config'];alpha=sel['chosen']['alpha'];delta=previous-current
    gate=read(B/'round8c/results/summary.json');assert gate['all_passed'];plateau=read(B/'accuracy_assessment/assessment.json')
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21.59);sec.page_height=Cm(27.94);sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for st in doc.styles:
        for x in list(st.element.iter(qn('w:pBdr'))):x.getparent().remove(x)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Calibri';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor.from_string('000000')
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.08;doc.styles['Normal'].paragraph_format.space_after=Pt(6);doc.styles['Title'].font.size=Pt(21);doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 2'].font.size=Pt(12)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Google 2019 负荷预测研究  |  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for r in footer.runs:r.font.size=Pt(8)
    def p(t,b=False):q=doc.add_paragraph();q.add_run(t).bold=b;return q
    def h(t):return doc.add_heading(t,1)
    def page(t):q=h(t);q.paragraph_format.page_break_before=True
    def table(head,rows,widths):
        tb=doc.add_table(rows=1,cols=len(head));tb.autofit=False
        for c,w in zip(tb.columns,widths):c.width=Cm(w)
        for c,t,w in zip(tb.rows[0].cells,head,widths):c.text=str(t);c.width=Cm(w)
        tb.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in rows:
            for c,t,w in zip(tb.add_row().cells,row,widths):c.text=str(t);c.width=Cm(w)
        for i,row in enumerate(tb.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E1EDF1' if i==0 else ('F4F7F8' if i%2==0 else 'FFFFFF'));pr.append(shade);borders=OxmlElement('w:tcBorders')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
                pr.append(borders);mar=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'65');x.set(qn('w:type'),'dxa');mar.append(x)
                pr.append(mar)
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.05
                    if j>0 and len(str(c.text))<30:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for rr in para.runs:rr.font.size=Pt(9.5);rr.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    doc.add_heading('Google PowerData 2019\n完整算力特征与微调研究汇报',0);p('第八轮至第九轮实际结果　｜　2019 年 5 月同月探索性回测')
    h('当前结论')
    p(f'最新冻结主方案次日曲线 WAPE 为 {current:.6f}%。加入完整算力特征前为 {previous:.6f}%，本轮变化 {delta:+.6f} 个百分点。组合选择来自开发期，未按本轮测试排名重新选择。',True)
    p(f'现有最佳组件微调先把 WAPE 从 2.374180% 降到 2.374128%，改善仅约 0.000052 个百分点，联合区间跨零。第九轮冻结家族为 {key}，新增权重为 {alpha}。'+('验证规则保留上一方案；完整 CPU 方案未获推荐替换。' if alpha==0 else '该权重是按开发期保守规则确定的本轮输出。'))
    p('本轮新增功率模型没有取得大幅或可靠改善。完整 CPU 历史、预测工作量、条件单调回归及多尺度稀疏 MoE 已全部实际评估；最终按验证期规则保留原最佳。',True)
    p('已完整遍历 a 集群全部 12,085 个官方 Parquet 使用文件的必要列，完成真实 CPU 和内存序列、匹配对照、融合网络训练、回测和本机预测。仅 a 的五个供电域有完整真实使用输入，其余域保留缺失。未取得新月份独立确认，不能把同月点估计或实现检查称为无过拟合证明。')
    table(['冻结方案','曲线 WAPE %','峰值 MAPE %'],[['最早 Chronos 与 TiRex 等权',f'{old["old_equal_mix"]["wape"]:.6f}',f'{old["old_equal_mix"]["peak_mape"]:.6f}'],['此前固定基础组合',f'{old["anchor"]["wape"]:.6f}',f'{old["anchor"]["peak_mape"]:.6f}'],['第五轮残差组合',f'{old["selected"]["wape"]:.6f}',f'{old["selected"]["peak_mape"]:.6f}'],['第八轮基础组件微调',f'{previous:.6f}',f'{ft["scores"]["selected"]["peak_mape"]:.6f}'],['第九轮完整算力冻结组合',f'{current:.6f}',f'{scores["selected"]["all57"]["peak_mape"]:.6f}']],[9.3,3.9,3.9])
    p('有较低测试分数的纯候选时，仍分别列出候选与冻结组合。整体曲线、每日峰值和预测区间分别评价。保留六位小数便于看清微小变化，不代表测量或泛化具有相应精度。')
    page('官方数据与实际完整范围')
    p('功率来自 Google 官方 PowerData2019，覆盖 2019 年 5 月的 57 个供电域。使用记录来自 Google 官方 ClusterData2019 a 集群公开 GCS 存储桶，源对象名称均保留。无需可用 Cloud 项目或账号，没有提交 BigQuery 或付费云任务。数据及说明采用 CC BY 4.0，需保留来源署名。')
    table(['实际处理内容','尺寸或记录数'],[['a 集群原 Parquet 对象',f'12,085 个；{audit["original_objects_bytes"]:,} 字节；{audit["original_objects_bytes"]/2**30:.2f} GiB'],['本地保留的列字节及间隙',f'{audit["retained_bytes"]:,} 字节；{audit["retained_bytes"]/2**30:.2f} GiB'],['本地原始范围片段',f'{audit["retained_range_files"]:,} 个 .bin 文件'],['扫描全部原表记录',f'{audit["all_rows"]:,} 条'],['有效映射后的 JOB 窗口',f'{audit["used_job_rows"]:,} 条'],['十五分钟主聚合 NPZ',f'{audit["aggregated_file_bytes"]:,} 字节；另导出 14,880 行 CSV']],[7.0,10.1])
    p('只读取 start_time、end_time、machine_id、collection_type、average_usage 的 CPU 与内存以及 sample_rate。通过官方 HTTP Range 读取 Parquet 页，保留取得的原字节、footer、SHA256 和对象清单。相邻字段合并读取会保留少量非建模间隙字节。未下载原对象的全部列，也未把此范围称为八个集群全量数据。')
    table(['供电域','观测槽','CPU 平均 NCU','CPU 最大 NCU'],[[x['domain'],str(x['observed_slots']),f'{x["cpu_ncu_mean"]:.3f}',f'{x["cpu_ncu_max"]:.3f}'] for x in audit['domains']],[5.1,3.0,4.5,4.5])
    p('每域 2,976 个槽；CPU 是已映射、已监控 JOB 窗口的时长加权总量，不是物理 CPU 核数。完整官方导出仍可能缺少监控与 alloc 层开销，不能等同全部物理工作量。原文件虽以 .parquet.gz 结尾，实际是 PAR1 Parquet，没有外层 gzip；未制作额外巨型压缩包。')
    page('时间对齐与数据审核')
    p('官方功率时间戳表示五分钟区间的起点，600,000,000 微秒对应 2019-05-01 00:00 PT。功率排序后以三个点构成十五分钟平均。使用窗口只在 end_time 已结束后可用，按结束时间归入对应十五分钟区间；对最多 300 秒的窗口计入平均使用量乘持续时间再除以 900 秒。')
    p('该处理守住输入可用时间，但不是逐秒精确重叠重采样：边界窗口可能存在不足五分钟的位移。CPU 与内存累计使用时间在每个分片和最终合并之间核对守恒。机器映射视为固定物理归属，因为原映射不含有效日期；这是建模假设，不能声称已审核动态迁移的真实归属。')
    table(['审核对象','实际核查'],[['文件覆盖','对象清单 12,085 个，逐对象聚合与 JSON 证据均存在'],['取得的字节','每个片段接收时计算 SHA256；再次核对所有片段长度'],['原对象 MD5','与官方 HTTP 对象头和清单比较；未下载全列，因此未独立重算原对象完整 MD5'],['独立重复哈希',f'全部聚合文件重查；{audit["checks"]["raw_hash_objects_independently_rechecked"]} 个预固定随机对象的完整片段重查'],['行数与累计使用','有效 JOB 窗口数量、CPU 与内存累计时间一致'],['缺失域','其他 52 域的使用不存在；模型输入显式标记缺失']],[5.0,12.1])
    p('仅使用 JOB，排除 ALLOC_SET，避免把预留与任务相加。无映射、无有效时间、超长窗口、非有限值和负使用量不进入聚合。没有读取 collection_id 和 instance_index，因此未独立完成全表任务键去重审核；结果依赖官方导出记录的语义。')
    p('质量掩码按官方分析 notebook 的 NOT bad_measurement_data 和 NOT bad_production_power_data 过滤。说明 PDF 中两处真假描述与其相反，本轮核对并采用官方代码示例及字段命名一致的过滤规则。不得把实现核查通过解释为模型有足够泛化证据。')
    page('选模协议与现有最佳模型微调')
    table(['阶段','日期和作用'],[['开发模型训练','标签截至 05-14；起报索引 7 至 13'],['前四个开发日','05-15 至 05-18，选择模型家族和检查点'],['后四个开发日','05-19 至 05-22，选组合权重，偏向较小新增权重'],['最终拟合','只重训已冻结配置，标签截至 05-22'],['次日回测','05-23 至 05-31；57 域，49,243 个有效点']],[5.1,12.0])
    p('第八轮残差系数微调从最佳 Ridge 校准参数开始，按供电域宏平均 WAPE 优化并向原参数做二次惩罚。三组学习率与惩罚系数，各检查 25、100、300 步。验证期新增权重为 0，纯候选测试 WAPE 为 2.377111%，没有改善上一主方案。')
    p('基础模型组件微调直接优化已有三基础模型加残差组合，而不是只优化一个单模型。冻结 Chronos、TiRex 和校准分支；TimesFM 分别尝试仅输出层与输出层加最后一层，学习率 3×10⁻⁵ 与 1×10⁻⁶，教师锚定系数 0.05 与 0.1，各运行 600 步并检查 100、300、600 步。训练包含曲线误差、分位误差、边界、分位顺序与教师约束。')
    fs=ft['selection'];fc=ft['comparison_vs_previous'];p(f'前四日选择尾层配置的 100 步检查点，后四日选择组合权重 {fs["chosen"]["alpha"]}。最终按此配置重训到 05-22，再测试得到 {previous:.6f}%。九日中 {fc["days_won"]} 日、57 域中 {fc["domains_won"]} 域改善，联合 95% 改进区间 [{fc["day_and_cell_95"][0]:.6f}, {fc["day_and_cell_95"][1]:.6f}] 个百分点，跨零。')
    p(f'微调 {len(verify["checks"])} 项复核全部通过，包括冻结哈希、分数复算、未来值扰动、拟合截止日期和参数范围。迭代归一化统计量停止反向梯度，沿用此前前向一致性与有限梯度修正。该实现修正不作为预测精度创新。')
    page('等待数据期间的组合权重微调')
    p('在完整使用数据读取期间，继续训练一个围绕已有最佳预测的小参数组合。全局模式只学习三个偏置；状态模式学习九个历史或已预测特征对应的共享斜率与三个偏置，共三十个参数。状态包括近期均值变化、波动、代理功率可用率、时刻和基础模型分歧，不学习供电域专属参数。')
    p('Chronos、TiRex 和已接受 TimesFM 组件形成三个专家。Softmax 使权重非负且和为一，增量 logit 限制在 ±0.5；功率裁剪在 [0,1]。门控从零校正出发，与已有最佳预测逐值一致。通过宏平均归一化的平滑绝对误差优化，并将权重和斜率收缩到原组合。')
    p('两种模式与 0.1、1、10 三档收缩共六组，分别检查 100、300、600 步，共十八个候选。沿用前四开发日选择配置，后四开发日按日期与集群较大配对标准误选择新增权重。该规则在本轮测试评价前登记和冻结，仍不能消除多轮使用相同开发期的选择偏差。')
    gg=gate['scores'];gc=gate['comparison_vs_previous'];gsel=gate['selection']
    table(['权重微调方案','曲线 WAPE %','峰值 MAPE %'],[['已有最佳固定上一方案',f'{gg["prior"]["wape"]:.6f}',f'{gg["prior"]["peak_mape"]:.6f}'],['新权重纯候选',f'{gg["candidate"]["wape"]:.6f}',f'{gg["candidate"]["peak_mape"]:.6f}'],[f'开发期冻结组合 权重 {gsel["chosen"]["alpha"]}',f'{gg["selected"]["wape"]:.6f}',f'{gg["selected"]["peak_mape"]:.6f}']],[9.3,3.9,3.9])
    p(f'冻结配置为 {gsel["config"]}，最后拟合标签截至 05-22。曲线 WAPE 比已有最佳高 {-gc["difference_pp"]:.6f} 个百分点，九日中仅 {gc["days_won"]} 日改善；联合 95% 改进区间 [{gc["day_and_cell_95"][0]:.6f}, {gc["day_and_cell_95"][1]:.6f}] 个百分点，跨零。峰值微降不能替代曲线精度改善。',True)
    p('本轮没有提高曲线精度，不升级当前最佳部署；完整保留预声明选择、未采用的方法及预测。部署保留是在探索性测试诊断后作出的操作决定，不应被描述为另一次独立验证。第九轮沿用事先登记的 round8b 作为固定上一方案，未因这一负面结果重写其选择协议。')
    p('八项实现检查通过，包括凸组合、有限值、边界、零校正一致性、拟合截止、冻结哈希与三起报日未来值扰动。检查实现正确，不能证明没有过拟合。入口 11 仅复现权重对照；相关文件在 experiment/round8c。')
    page('现有精度平台与可检验的上升空间')
    p('本页是完整 CPU 输入实验之前对第八轮最佳模型的诊断。结论为：在现有功率与生产功率代理信息下已进入明显工程平台期，不能据此认定理论最低误差，也不能把未开始评价的完整真实工作量实验视为已经失败。')
    pc=plateau['expert_error_correlation']['matrix'];pb=plateau['a_only_correction_bound']
    table(['诊断证据','实际结果'],[['当前功率代理最佳 WAPE',f'{plateau["current"]["wape"]:.6f}%'],['第五轮以后组件微调改善',f'{plateau["recent_component_gain_pp"]:.9f} 个百分点，联合区间跨零'],['Chronos 与 TiRex 误差相关',f'{pc[0][1]:.4f}'],['Chronos 与 TimesFM 误差相关',f'{pc[0][2]:.4f}'],['TiRex 与 TimesFM 误差相关',f'{pc[1][2]:.4f}'],['开发期日偏差一日滞后相关',f'{plateau["normalized_daily_bias_lag1"]["development"]:.4f}，持续性弱']],[8.1,9.0])
    p('功率仅三十一天，开发模型监督起报日七个、最终拟合十五个。四万多个有效测试点具有日期与场站相关，不能当作等量独立样本。各基础模型误差相关很高，新权重微调没有提高曲线精度；扩大相似模型搜索更容易产生开发期选择波动。')
    p(f'若事后知道每天每域的最佳固定偏移，WAPE 为 {plateau["hindsight_perfect_daily_offset"]["score"]["wape"]:.6f}%；若事后逐点挑选最接近真值的专家，为 {plateau["hindsight_per_point_expert"]["score"]["wape"]:.6f}%。两者都读取未来真值，只作诊断，不是可部署预测、可达到精度或不可约误差。它们说明残差中存在水平变化与模型分歧，是否能提前识别仍需验证。')
    p(f'a 的五域当前 WAPE 为 {plateau["current_cell_a"]["wape"]:.6f}%。对于只修正这五域的分支，即使把其误差事后降为零，全部五十七域 WAPE 最多降至 {pb["all57_best_possible_wape_if_a_perfect"]:.6f}%；这是局部修正的乐观数学界限。a 误差减少 20% 时，整体仅减少 {pb["if_a_error_reduced_20_percent_global_gain_pp"]:.6f} 个百分点。改变其他域的联合网络不适用该界限，也不能把其他域变化全部归因于 a 的 CPU 信息。')
    p(f'因此继续检验完整真实 CPU、预测未来工作量与无 CPU 匹配对照，不扩展旧输入的在线滤参搜索。本轮第九轮冻结结果实际为 {current:.6f}%，新增权重 {alpha}，是否提升以实测和已登记的选择为准。任何突破结论还需未参与选模的新月份或场站确认。',True)
    p('完整逐域、逐日、逐集群表与独立诊断图保存在 experiment/accuracy_assessment。前十域贡献当时宏平均误差的 24.40%，前三日贡献 38.52%；这反映异质性，未据此按测试排序改参数。')
    page('真实算力融合与物理边界')
    p('第九轮采用七天历史功率、完整 a 集群 JOB CPU 与内存、质量掩码、历史机器状态和时间场站特征。RevIN 仅由历史计算；5 和 25 点平均分解功率，历史 FFT 表示频域；16、32、96 点尺度分别编码功率和工作量，通过交叉注意力融合。资源容量、在线比例和容量已知率进入资源嵌入，三个专家每样本仅运行两个。')
    p('网络宽度固定为 16，从相同随机种子与对应训练截止的第五轮编码器开始；功率输出层重新置零，从当前已校准基础组合预测残差。联合输出功率和真实 CPU 的 0.1、0.5、0.9 分位数，使用负载辅助损失、正则、残差收缩和专家负载平衡。两固定种子先平均后评价，未按测试挑种子。')
    table(['物理量或约束','本轮含义'],[['总功率 P','供电容量利用率，分位输出限制在 [0,1]'],['工作量 W','完成 JOB 窗口的 CPU NCU，分位非负且有序'],['内存与机器容量','原数据归一化 RAM 及调度器可用容量，均只用历史'],['CPU 与功率的关系','单位不同，不采用 CPU≤P；可用调度容量也不是物理硬上限'],['未来 W','先由历史预测或由联合网络输出；不把未来观测作为输入'],['约束范围','边界与排序约束，未宣称完整电量守恒或实际 kW 模型']],[5.2,11.9])
    p('匹配的 Ridge 分支比较仅资源状态、加入完整 CPU 内存历史、再加入预测工作量，正则均为 10,000、100,000、1,000,000。其输出只修正 a 的五域，其他 52 域逐值保持上一方案。基础模型先从最多十四天历史预测五域 CPU、内存与窗口当量；只把这些预测送入功率模型。')
    p('网络比较完整 CPU、无 CPU、无交叉注意力和无频域四种结构，均运行 4、10、20 轮并平均相同两种子。最终匹配比较采用同一冻结轮数。新增权重按开发日期的配对标准误选择；仅一集群有新 CPU 输入，预先声明以日期为单位，统计支持仍有限。')
    page('完整算力特征的匹配结果')
    ra=s['final_matched_ridge_alpha'];ep=s['final_matched_neural_epoch'];labels={'control':'仅机器状态','history':'加完整算力历史','forecast':'加预测工作量'}
    table([f'Ridge 正则 {ra:,}','全部 57 域 WAPE %','a 五域 WAPE %'],[[labels[g],f'{scores[f"{g}_ridge{ra}"]["all57"]["wape"]:.6f}',f'{scores[f"{g}_ridge{ra}"]["cell_a_5"]["wape"]:.6f}'] for g in ['control','history','forecast']],[7.9,4.6,4.6])
    p('上表保持同一正则、拟合时间、功率基础预测和资源特征，历史与预测分支的差异才可用于讨论额外算力信息。纯候选表现不决定最终替换：开发期冻结的新增权重和家族优先于测试排名。')
    nl={'fusion_cpu':'完整真实 CPU 网络','fusion_control':'无 CPU 同宽网络','fusion_no_cross':'无交叉注意力','fusion_no_frequency':'无频域'}
    table([f'宽度 16 两种子平均 {ep} 轮','全部 57 域 WAPE %','a 五域 WAPE %'],[[nl[n],f'{scores[f"{n}_e{ep}"]["all57"]["wape"]:.6f}',f'{scores[f"{n}_e{ep}"]["cell_a_5"]["wape"]:.6f}'] for n in nl],[7.9,4.6,4.6])
    p(f'最终主方案全 57 域 WAPE 为 {current:.6f}%，a 五域为 {scores["selected"]["cell_a_5"]["wape"]:.6f}%；上一方案 a 五域为 {scores["prior"]["cell_a_5"]["wape"]:.6f}%。以上使用同一质量掩码与九个测试日；没有把五域成绩作为 57 域成绩。',True)
    p('需要区分三种情况：模型在测试期点估计变好，开发期支持以某权重组合，以及新算力信息比匹配对照有帮助。这三个条件不能互相替代。完整数据改善了可审核性，是否提高精度以本表和冻结结果为准。')
    table(['CPU 预测仅 a 五域','曲线 WAPE %','日均 WAPE %','峰值 MAPE %'],[[label,*[f'{s["cpu_scores"][k][m]:.4f}' for m in ['wape','daily_wape','peak_mape']]] for label,k in [('前一天同刻','last_day'),('历史基础模型预测','foundation'),('功率 CPU 联合融合网络','fusion')]],[7.1,3.2,3.4,3.4])
    p(f'联合 CPU 原始名义 80% 分位区间覆盖 {s["cpu_raw_nominal80_coverage"]*100:.2f}%，未单独校准。这是已监控 JOB NCU 的预测评价，不是物理核数或功率误差。')
    cpu_gain=100*(s['cpu_scores']['last_day']['wape']-s['cpu_scores']['foundation']['wape'])/s['cpu_scores']['last_day']['wape']
    p(f'基础模型的 CPU 曲线误差相对下降 {cpu_gain:.2f}%，但日均与峰值误差增大，未全面改善。工作量曲线的改善没有转化为本轮冻结功率方案的改善，不能将 CPU 指标当成功率精度。',True)
    page('单调负载响应的物理对照')
    p('该分支只拟合每个 a 供电域的两个斜率和一个截距，以历史基础模型预测的 CPU、内存相对前一天同刻的变化修正冻结功率预测。两个斜率强制非负；预测值裁剪至 [0,1]，只改 a 的五域。固定其他输入时，功率预测随预测 CPU 或内存增加不下降，属于条件单调性。')
    p('这是一项可检验的简化假设：没有温度、业务组成、真实电量或全部任务开销时，不能把非负斜率当成物理因果定律。它也不构成完整能量守恒。匹配控制把负载变化列置零，保留相同拟合日期、基础预测、域内截距和惩罚；三档正则 10、100、1,000 在模型训练前登记。')
    table(['正则','截距控制 WAPE %','单调分支 WAPE %','a 五域单调 WAPE %'],[[str(a),f'{scores[f"physical_control_alpha{a}"]["all57"]["wape"]:.6f}',f'{scores[f"monotone_alpha{a}"]["all57"]["wape"]:.6f}',f'{scores[f"monotone_alpha{a}"]["cell_a_5"]["wape"]:.6f}'] for a in [10,100,1000]],[2.9,4.6,4.8,4.8])
    p('上表属于按预声明范围保留的测试对照，不据此重新选择正则或组合权重。即使某个纯分支分数低于主方案，最终结论仍取冻结家族和权重。第九轮共比较 27 个开发期候选：9 个多特征 Ridge、6 个单调及截距控制、12 个网络检查点；网络检查点已先平均固定两种子。')
    p('该设计用低参数响应分支检验负载预测是否有实用价值，再与多特征残差和融合网络比较。新输入贡献要同时结合匹配对照和开发期选择，不能把增加字段、约束或模型复杂度本身称为精度创新。')
    page('不确定性与研究结论')
    table(['评价','冻结组合实测'],[['相对上一方案变化',f'{cmp["difference_pp"]:+.6f} 个百分点；相对误差下降 {cmp["relative_reduction_percent"]:.4f}%'],['改善数量',f'{cmp["domains_won"]}/57 域，{cmp["days_won"]}/9 日，{cmp["cells_won"]}/10 集群'],['日期块 95% 改进区间',f'[{cmp["day_block_95"][0]:.6f}, {cmp["day_block_95"][1]:.6f}] 个百分点'],['日期与集群联合 95% 区间',f'[{cmp["day_and_cell_95"][0]:.6f}, {cmp["day_and_cell_95"][1]:.6f}] 个百分点'],['每日平均负荷 WAPE',f'{scores["selected"]["all57"]["daily_wape"]:.6f}%'],['每日峰值 MAPE',f'{scores["selected"]["all57"]["peak_mape"]:.6f}%'],['名义 80% 残差区间',f'测试覆盖 {s["uncertainty"]["test_coverage"]*100:.2f}%；平均利用率宽度 {s["uncertainty"]["mean_width"]:.5f}']],[6.2,10.9])
    p('重采样使用九日中的连续三日块和十个集群，5,000 次。区间是同月的描述性不确定性；此前已反复查看该月测试，不能恢复独立检验或给出稳健显著性结论。只有一个集群有新算力，五域相关限制跨场站推广；基础模型预训练是否覆盖公开同类数据也未排除。')
    p('残差区间只在开发期拟合相对历史功率均值的宽度，测试期不重拟合。时间序列相关、样本选择与分布漂移使名义覆盖率没有独立保证。第九轮新增分支评估次日 96 点；既有 TiRex 七天 WAPE 2.8803% 沿用，没有把次日成绩冒用为七天成绩。')
    p('当前结论是完成一条可复现、较少本地存储的官方算力提取路线，并真实检验了现有最佳微调和完整 CPU 的增益。能否进一步取得大幅且可靠的提高，还需新的月份或未参与选模的场站确认。增加模型复杂度和继续挑选这九天的最低分不能替代独立证据。',True)
    page('本地运行与复核位置')
    p('全部新增数据、取得的字节、聚合、模型、依赖、预测、日志与报告继续保存在原桌面文件夹 GooglePowerData2019_预测研究_20260929。E 盘原数据和之前轮次文件保留。')
    table(['运行入口','结果与用途'],[['01 研究用新组合预测次日','第九轮冻结组合；round9/results/selected_deployment_origin31_1day.csv'],['09 现有最佳微调预测次日','第八轮基础组件微调；round8b/results/selected_deployment_origin31_1day.csv'],['10 完整算力融合候选预测次日','完整 CPU 联合网络纯候选；与冻结主方案分开'],['02 至 08 原入口','七天、峰值、代理融合、资源及局部样本研究对照']],[6.1,11.0])
    p('入口 11 是未获曲线改善的权重微调对照，保存在 experiment/round8c。默认输入截止 2019-05-31，预测 2019-06-01；六月没有真值。成功导出只证明本机推理可运行，不代表今天的负荷或未来已验证。候选入口复用上一方案区间宽度仅为研究示例，未单独校准候选区间。')
    table(['目录与文件','核查内容'],[['google-clusterdata-2019/column_cache/a','取得的原 Parquet 列字节与 footer'],['projected_usage/a','逐原文件聚合、行数、原对象身份及 SHA256'],['aggregated/cella_complete_usage_15min.csv','五域 14,880 行真实使用、历史容量与质量字段'],['results/complete_projection_audit.json','完整范围、尺寸、哈希、单位、累计时间与限制'],['experiment/round8 与 round8b','实际微调配置、检查点、冻结选择及预测'],['experiment/round9/docs 与 results','预声明搜索、协议修订、冻结哈希、开发与测试分数'],['experiment/round9/models','最终匹配参数；宽度和种子与报告一致'],['google-clusterdata-2019/code','提取、审核、训练、预测、交付与报告程序']],[7.2,9.9])
    p(f'第九轮 {len(s["checks"])} 项检查全部通过，包括未来 CPU 资源与功率值扰动、组合哈希、分位非负有序以及 Ridge 对其他 52 域保持原预测。检查约束实现和时间因果性，不保证没有过拟合。')
    p('TimesFM 3.0 权重使用 TimesFM Non-Commercial v1.0，本目录用于非商业本地研究，不用于生产或对外分发。CPU 使用信息已进入新分支；第五轮的 W 仍是生产功率代理，旧结果与真实算力结果分别保留。')
    page('来源与可复现依据')
    p('针对内存不足，本轮采用单进程串行与两个 CPU 线程。省内存加载使用原 float32 权重，全部 445 个张量和原配置核对一致；重算开发期五域 480 点与已有功率预测的最大差为零。初次加载错误与修复证据均保留。安装缓存、重复权重副本及旧排版初稿已清理，实际数据、模型与实验记录保留，详见根目录清理说明与资源占用说明。')
    for title,url in [('Google 官方 PowerData2019 与许可','https://github.com/google/cluster-data/blob/master/PowerData2019.md'),('Google 官方 ClusterData2019 与许可','https://github.com/google/cluster-data/blob/master/ClusterData2019.md'),('Google 官方资源使用字段定义','https://github.com/google/cluster-data/blob/master/clusterdata_trace_format_v3.proto'),('官方功率时间与质量过滤示例','https://github.com/google/cluster-data/blob/master/power_trace_analysis_colab.ipynb'),('a 集群官方公开对象来源','https://storage.googleapis.com/clusterdata_2019_a/'),('Apache Arrow 官方 Parquet 读取实现','https://arrow.apache.org/docs/python/parquet.html'),('TimesFM 官方实现及权重许可','https://github.com/google-research/timesfm'),('RevIN 实现参考','https://github.com/ts-kim/RevIN'),('稀疏时间序列专家结构参考','https://github.com/Time-MoE/Time-MoE')]:
        q=p(title+'\n'+url);q.paragraph_format.space_after=Pt(5)
        for rr in q.runs:rr.font.size=Pt(9.5)
    p('本网络是依据少量训练日期定制的研究模型，借鉴 RevIN 和稀疏专家结构；不是 Time-MoE 论文或预训练权重的完整复现。源数据、模型 revision、许可和每轮选择证据均保存在桌面目录。')
    path=TOP/'Google PowerData 2019 完整算力特征与微调研究汇报.docx';doc.save(path);print(path,flush=True)
if __name__=='__main__':main()
