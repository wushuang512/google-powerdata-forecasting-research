"""Create the report only from completed and independently verified results."""
import datetime,json,sys
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
TOP=Path(sys.argv[1]).resolve();R=TOP/'experiment/round13'
OUT=TOP/'Google PowerData 2019 近期加权与稳健损失调优汇报.docx'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def main():
    summary=read(R/'results/summary.json');selection=read(R/'results/selection.json');dev=read(R/'results/dev_summary.json')
    verification=read(R/'results/independent_verification.json');run=read(R/'results/guarded_pipeline_run.json')
    assert verification['all_passed'] and read(R/'results/deployment_verification.json')['all_passed']
    assert selection['chosen']['alpha']==0 and not summary['large_gain_rule_met']
    if '--validate-only' in sys.argv:print('Real results and independent reconstruction verified');return
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
        pp=s.element.find(qn('w:pPr'))
        if pp is not None:
            for child in list(pp):
                if child.tag==qn('w:pBdr'):pp.remove(child)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(8)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    def para(t):return doc.add_paragraph(t)
    def h(t):return doc.add_paragraph(t,'Heading 1')
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for i,w in enumerate(widths):t.columns[i].width=Inches(w)
        for j,value in enumerate(headers):t.rows[0].cells[j].text=value
        for row in rows:
            cells=t.add_row().cells
            for j,value in enumerate(row):cells[j].text=str(value)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders);marg=OxmlElement('w:tcMar')
                for side,value in [('top','60'),('bottom','60'),('left','100'),('right','100')]:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');marg.append(e)
                pr.append(marg)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF5');pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:r.font.size=Pt(10.5);r.bold=i==0
        spacer=doc.add_paragraph();spacer.paragraph_format.line_spacing=Pt(6);spacer.paragraph_format.space_after=Pt(6)
        return t
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    doc.add_paragraph('Google PowerData 2019\n近期加权与稳健损失调优汇报','Title')
    para('第十三轮研究记录   2026年10月4日   D盘研究目录')
    h('本轮结论与正式预测')
    para('本轮实际训练六个小型残差校准器。开发期前四日选出半衰期为3日、较强正则的稳健损失模型，后四日不支持采用，修正权重在测试评分前冻结为0。正式预测继续采用 round8b，逐点数值与原最佳一致。')
    para('未采用候选的测试曲线分数略低，但改善仅约0.000318个百分点，约0.0134%的相对误差下降。开发证据不稳定，日期重采样区间跨零，不能宣称显著或大幅提升，也不能根据这个测试分数改写采用决定。')
    table(['测试指标','正式原最佳','未采用候选'],[
        ['57域曲线 WAPE',f'{summary["scores"]["prior"]["all57"]["wape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["wape"]:.6f}%'],
        ['57域日均 WAPE',f'{summary["scores"]["prior"]["all57"]["daily_wape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["daily_wape"]:.6f}%'],
        ['57域峰值 MAPE',f'{summary["scores"]["prior"]["all57"]["peak_mape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["peak_mape"]:.6f}%'],
        ['a集群5域曲线 WAPE',f'{summary["scores"]["prior"]["a5"]["wape"]:.6f}%',f'{summary["scores"]["candidate"]["a5"]["wape"]:.6f}%'],
        ['a集群5域日均 WAPE',f'{summary["scores"]["prior"]["a5"]["daily_wape"]:.6f}%',f'{summary["scores"]["candidate"]["a5"]["daily_wape"]:.6f}%'],
        ['a集群5域峰值 MAPE',f'{summary["scores"]["prior"]["a5"]["peak_mape"]:.6f}%',f'{summary["scores"]["candidate"]["a5"]["peak_mape"]:.6f}%']], [2.9,1.8,1.8])
    para('候选指完整加入新校准器差值的对照。它改善了全域峰值误差，但日均误差增加，a集群峰值也变差。因此，单个更低的曲线数字不能表示所有预测目标都改善。主指标是57域各PDU的WAPE平均值，不是把所有PDU电量合并计算的误差。')
    para(f'候选曲线改善量的探索性95%日期配对区间为 [{summary["candidate_day_bootstrap95_pp"][0]:.6f}, {summary["candidate_day_bootstrap95_pp"][1]:.6f}] 个百分点。正式方案改善量为0，因其与原最佳相同；这不表示预测误差为0。数据只有同一个月，不能把该区间解释成独立泛化确认。')
    h('重新训练的方法与时间边界').paragraph_format.page_break_before=True
    para('仍使用原校准器的30个历史特征，学习30个系数和1个截距。两个候选保留平方损失，给近期样本更高权重；另外四个候选采用平滑绝对误差损失，测试3日或7日半衰期及两档固定正则。该损失更接近WAPE的绝对误差形式，但训练权重仍是历史功率代理，不能声称精确优化了最终WAPE。')
    para('起报7至13训练，14至17选模型，18至21按一标准误规则选强度，22至30在冻结后评分。每个起报只使用已经完成的历史功率和代理特征。前四个开发日选出的模型在起报7至21重新拟合，然后冻结最终模型再评分测试；没有按测试误差决定迭代次数或重新训练。')
    labels={'recent_ridge_h3':'近期平方损失 3日','recent_ridge_h7':'近期平方损失 7日',
            'smoothL1_h3_p05':'稳健损失 3日 正则0.5','smoothL1_h3_p2':'稳健损失 3日 正则2',
            'smoothL1_h7_p05':'稳健损失 7日 正则0.5','smoothL1_h7_p2':'稳健损失 7日 正则2'}
    rows=[['原最佳',f'{dev["prior_early4"]["wape"]:.6f}',f'{dev["prior_late4"]["wape"]:.6f}']]
    for r in dev['candidates']:rows.append([labels[r['model']],f'{r["early4"]["wape"]:.6f}',f'{r["late4"]["wape"]:.6f}'])
    table(['重新训练配置','前4日 WAPE %','后4日 WAPE %'],rows,[2.9,1.8,1.8])
    para('近期平方损失模型在后段有极小改善，但前段不及被选出的稳健模型，不能事后改用后段赢家。预登记的顺序筛选保留全部六组结果，反映模型选择的不稳定性。')
    para('所选模型的修正强度与后四日WAPE分别为：'+ '；'.join(f'{r["alpha"]:g} 对应 {r["wape"]:.6f}%' for r in selection['weights'])+'。')
    para('所选候选的所有非零强度都比原最佳差。新校准器替换原组合中50%的旧校准分量，新增差值限制在历史七日均值的正负1%，最后乘采用强度并限制功率在0至1。强度0由开发期决定，测试上略低的候选仅作为保留的探索性对照。')
    h('真实核查与历轮研究证据').paragraph_format.page_break_before=True
    para(f'本轮{verification["check_count"]}项核查全部通过，覆盖文件指纹、时间边界、冻结选择、优化收敛、指标独立重算、区间与部署。原Ridge从训练数据重新拟合，与旧保存预测最大差值为 {dev["control"]["max_absolute_difference"]:.3g}，比较基准一致。')
    para('稳健候选均在12次迭代内收敛，最终拟合用了11次。开发训练有38298个有效点，最终训练82067个；它们在时间和场站上相关，不是同等数量的独立样本。未来真值扰动通过，只证明核查的实现未读取未来真值，不能保证没有过拟合。')
    history=[
        ['第8b轮','组件微调','保留 原最佳'],
        ['第8c轮','权重门控','曲线变差 未采用'],
        ['第9轮','完整CPU对照 条件单调与MoE','开发权重0'],
        ['第10轮','原生CPU内存资源协变量','开发权重0'],
        ['第11轮','小型非线性残差树','开发权重0'],
        ['第12轮','专家中位数与离群截断','开发权重0'],
        ['第13轮','近期加权与稳健损失','开发权重0']]
    table(['轮次','研究方向','正式采用结果'],history,[.8,3.3,2.4])
    para('完整CPU与原生协变量研究均已完成，在当前对齐、覆盖和设置下未支持升级，不证明算电信息永远无效。条件单调约束不等于能量守恒。历轮模型、协议和未采用结果全部保留。')
    h('对精度平台的判断')
    para('现有31天数据上，多种方向未获稳定开发收益，支持工程平台判断，未证明理论误差下限。反复查看开发期与测试期会累积选择偏差，先冻结再评分不能恢复数据独立性。')
    para('进一步突破应优先补充未查看过的真实功率月份，或与同一PDU对齐的工作量、设备功耗和资源状态。新模型须先提出可检验的误差机制，登记有限候选，再独立确认。现有CPU记录不能替代新功率真值，没有这些条件就不能承诺大幅提升。')
    h('部署文件与数据限制').paragraph_format.page_break_before=True
    para('正式部署输出57域5472行，时间为2019年6月1日00:00至23:45。数值有限且在0至1内，经验区间顺序正确，所选功率与原最佳逐点相同。六月没有功率真值，这一输出是历史复现推理，不是今天的电负荷。80%区间由开发期残差校准，未保证独立数据覆盖率。')
    table(['保存位置','用途'],[
        ['experiment/round13/docs/protocol.json','六组配置与时间边界'],
        ['experiment/round13/results/selection.json','测试前冻结的选择'],
        ['experiment/round13/models','六个开发模型与最终候选'],
        ['experiment/round13/results','候选 正式预测 逐日逐域表 审计'],
        ['15_近期加权稳健损失_预测次日.cmd','复现本轮正式冻结预测']],[4.2,2.3])
    para('所有新增程序、模型和产出均保存在 ${PROJECT_ROOT}，桌面同名目录仍为兼容链接。09保留原最佳入口，01至14及此前正式汇报仍可使用。开始阅读、当前进展和完整文件清单随本轮交付更新，未删除负面结果。')
    h('资源与来源范围')
    para(f'本轮未下载数据、未加载基础模型权重，CPU最多2线程。可用内存达到3GiB物理与6GiB提交条件后自动启动，运行中物理储备最低 {run["minimum_available_physical_bytes"]/2**30:.3f} GiB，没有触发停止。运行仍保留1.5GiB物理与2GiB提交保护，只管理本实验子进程。先前授权的定时自动续跑保持停用。')
    para('官方cell a的12085个Parquet必要列已完整读取和审计，剩余下载为0；真实完整CPU特征覆盖a集群五个功率域，不能表述成读取全部八个集群或全部字段。84.78GiB保留压缩列、407.55GiB对应原对象总量与105.481GiB迁移时逻辑文件总量是三种不同口径。后续增加的报告和实验文件不属于当时迁移快照。')
    para('功率是归一化容量利用率，CPU为NCU，不能直接换算成kW。任务窗口只统计JOB，缺少完整任务去重键和严格一一对应的功率观测，时间对齐也有限制。Google数据许可为CC BY 4.0；现有TimesFM 3权重限定非商业、非生产使用。')
    doc.save(OUT);print(OUT)
if __name__=='__main__':main()

