"""Create the report only from completed and independently verified results."""
import datetime,json,sys
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
TOP=Path(sys.argv[1]).resolve();R=TOP/'experiment/round12'
OUT=TOP/'Google PowerData 2019 鲁棒专家融合精度探索汇报.docx'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def main():
    summary=read(R/'results/summary.json');selection=read(R/'results/selection.json');dev=read(R/'results/dev_summary.json')
    verification=read(R/'results/independent_verification.json');run=read(R/'results/guarded_pipeline_run.json')
    assert verification['all_passed'] and read(R/'results/deployment_verification.json')['all_passed']
    assert selection['chosen']['alpha']==0 and not summary['large_gain_rule_met']
    if '--validate-only' in sys.argv:print('Real results and independent reconstruction verified');return
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
        pp=s.element.find(qn('w:pPr'))
        if pp is not None:
            for child in list(pp):
                if child.tag==qn('w:pBdr'):pp.remove(child)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(8)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    def para(t):return doc.add_paragraph(t)
    def h(t):return doc.add_paragraph(t,'Heading 1')
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for i,w in enumerate(widths):t.columns[i].width=Inches(w)
        for j,value in enumerate(headers):t.rows[0].cells[j].text=value
        for row in rows:
            cells=t.add_row().cells
            for j,value in enumerate(row):cells[j].text=str(value)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders);marg=OxmlElement('w:tcMar')
                for side,value in [('top','60'),('bottom','60'),('left','100'),('right','100')]:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');marg.append(e)
                pr.append(marg)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF5');pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:r.font.size=Pt(10.5);r.bold=i==0
        spacer=doc.add_paragraph();spacer.paragraph_format.line_spacing=Pt(6);spacer.paragraph_format.space_after=Pt(6)
        return t
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    training=read(R/'results/training_diagnostic.json')
    doc.add_paragraph('Google PowerData 2019\n鲁棒专家融合精度探索汇报','Title')
    para('第十二轮研究记录   2026年10月3日   保存位置 D盘研究目录')
    h('结论与正式模型')
    para('本轮完成四组鲁棒专家融合探索，未找到可采纳的全域曲线提升。前四个开发日选出围绕专家中位数截断的 winsor01 方案，后四个开发日未支持加入修正，冻结强度为0。正式预测继续使用 round8b 最佳模型，逐点数值保持一致。')
    table(['测试指标','保留的原最佳','未采用的候选'],[
        ['57域曲线 WAPE',f'{summary["scores"]["prior"]["all57"]["wape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["wape"]:.6f}%'],
        ['57域日均 WAPE',f'{summary["scores"]["prior"]["all57"]["daily_wape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["daily_wape"]:.6f}%'],
        ['57域峰值 MAPE',f'{summary["scores"]["prior"]["all57"]["peak_mape"]:.6f}%',f'{summary["scores"]["candidate"]["all57"]["peak_mape"]:.6f}%'],
        ['a集群5域曲线 WAPE',f'{summary["scores"]["prior"]["a5"]["wape"]:.6f}%',f'{summary["scores"]["candidate"]["a5"]["wape"]:.6f}%']], [2.9,1.8,1.8])
    para('未采用候选指完整加入 winsor01 修正的对照，全域曲线误差增加约0.004525个百分点。其a集群曲线和全域峰值指标略有改善，但不能用事后出现的局部指标替代预先指定的全域曲线目标。')
    para(f'候选曲线改善量的按日期共同重采样探索性95%区间为 [{summary["candidate_day_bootstrap95_pp"][0]:.6f}, {summary["candidate_day_bootstrap95_pp"][1]:.6f}] 个百分点，跨越零。正式方案改善量为零，差值区间为零是因为与原模型相同，不表示没有预测误差。')
    h('本轮方向的依据')
    para(f'仅使用训练起报7至13检查原始固定专家组合，均值组合 WAPE 为 {training["prior"]["wape"]:.6f}%，逐点中位数为 {training["unbounded_median"]["wape"]:.6f}%。训练期信号支持进一步检查离群专家的影响，但这个指标属于未微调的旧骨干组合，不是当前最佳模型的训练分数，也不能当作验证收益。')
    para('本轮使用已保存的 Chronos 2、TiRex 2 和 TimesFM 3 专家预测，调整它们的聚合方式，然后把相对旧聚合的差值有限地加入现有最佳预测。原最佳的权重微调和残差校准保留。本轮没有新增可学习参数，没有加载基础模型或下载新权重。')
    h('开发期筛选与约束').paragraph_format.page_break_before=True
    para('数据为2019年5月57个PDU的15分钟功率利用率。起报7至13只作训练期信号诊断，14至17选择聚合方式，18至21选择修正强度，22至30为冻结后九个测试起报。起报22预测5月23日。任何起报都只使用已经完成的历史七日均值及该起报的模型预测，不输入未来真实负荷或CPU。')
    labels={'point_median':'逐点专家中位数','daily_consensus':'专家日均共识','winsor01':'离群截断 1%','winsor03':'离群截断 3%'}
    rows=[['原最佳',f'{dev["prior_early4"]["wape"]:.5f}',f'{dev["prior_late4"]["wape"]:.5f}']]
    for r in dev['candidates']:rows.append([labels[r['model']],f'{r["early4"]["wape"]:.5f}',f'{r["late4"]["wape"]:.5f}'])
    table(['融合方式','前4日 WAPE %','后4日 WAPE %'],rows,[2.9,1.8,1.8])
    para('逐点中位数限制单个离群专家的影响；日均共识只改变预测日的整体水平；离群截断把各专家限制在中位数周围，再按原25%、25%、50%权重求平均。1%和3%为截断半径相对于历史七日均值的比例。所有新增功率差值限制在历史均值的正负1%，最后乘冻结强度并将功率限制在0至1。')
    table(['winsor01 修正强度','后4日 WAPE %'],[[f'{r["alpha"]:g}',f'{r["wape"]:.5f}'] for r in selection['weights']],[3.4,3.1])
    para('后四日的所有非零修正强度都比原最佳差，完整修正在两个日期改善、两个日期变差。采用日期配对与场站配对标准误中较大者的一标准误规则，优先较小强度，结果为0。原选择及其SHA256在测试评分前冻结，没有按照测试结果再选其他方案。')
    para('开发期和测试期均已被历轮研究反复使用。新的冻结规则可以限制本轮测试驱动的改选，不能恢复数据独立性或保证不存在适应性过拟合。所列区间只描述当前九个日期的探索性差值，没有跨月泛化保证。')
    h('独立核查与数值审计').paragraph_format.page_break_before=True
    para(f'独立核查共 {verification["check_count"]} 项全部通过，包括逐点重建四种聚合方式、开发期选择与强度、测试指标、日期重采样区间、旧文件指纹、部署专家均值对照，以及未来真值扰动检查。未来扰动只证明已检查的实现没有读取起报后的真值，不能证明模型不会过拟合。')
    audit=verification['dev_arithmetic_audit']
    para(f'开发期组合重建因数组合并后的精度提升，与原程序存在最大 {audit["max_power_util_difference_vs_canonical_round8b"]:.10f} 的利用率舍入差异。独立审计保留原数据类型重算，配置与强度0的决策不变；没有改写原冻结记录。初次审计失败记录保留。测试与正式预测直接采用原最佳数组，逐点一致。')
    para('部署包含57域5472行，时间从2019年6月1日00:00到23:45，数值有限、在0至1内，区间顺序正确。Chronos部署分量由已存旧等权组合和TiRex分量代数恢复；与原组合的对照最大差值为0。六月没有真值，输出是历史复现推理，不是今天的电负荷。80%区间只按开发期残差经验校准。')
    table(['文件位置','用途'],[
        ['experiment/round12/docs/protocol.json','候选 参数与数据边界'],
        ['experiment/round12/results/selection.json','原始冻结选择'],
        ['experiment/round12/models/robust_fusion.json','冻结融合器配置'],
        ['experiment/round12/results','预测 指标与完整审计'],
        ['14_鲁棒专家融合_预测次日.cmd','复现冻结预测']], [4.3,2.2])
    h('平台判断与研究空间')
    para('本轮继续支持现有数据与输入上的工程平台判断，但没有证明精度已到理论极限。近期非线性校准和鲁棒融合都没有支持可靠升级。继续搜索同一月份最低测试误差容易积累选择偏差；下一步应优先获取独立月份或新的可对齐功率观测，确认冻结方法的泛化能力。')
    para('全部新增文件保存在D盘研究目录，桌面原路径为兼容链接，自动续跑保持停用。本轮未增加下载，也未扩充CPU覆盖：仍只有a集群五域拥有完整任务使用特征。功率是归一化容量利用率，CPU单位是NCU，不能直接当成kW。Google数据为CC BY 4.0；TimesFM 3权重限非商业、非生产使用。')
    doc.save(OUT);print(OUT)
if __name__=='__main__':main()
