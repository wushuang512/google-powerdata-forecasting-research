"""Create the report only from completed and independently verified results."""
import datetime,json,sys
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
TOP=Path(sys.argv[1]).resolve();R=TOP/'experiment/round14'
OUT=TOP/'Google PowerData 2019 直接预测网络与共享状态研究汇报.docx'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def main():
    summary=read(R/'results/summary.json');selection=read(R/'results/selection.json');dev=read(R/'results/dev_summary.json')
    verification=read(R/'results/independent_verification.json');run=read(R/'results/guarded_pipeline_run.json')
    assert verification['all_passed'] and read(R/'results/deployment_verification.json')['all_passed']

    if '--validate-only' in sys.argv:print('Real results and independent reconstruction verified');return
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
        pp=s.element.find(qn('w:pPr'))
        if pp is not None:
            for child in list(pp):
                if child.tag==qn('w:pBdr'):pp.remove(child)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(8)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    def para(t):return doc.add_paragraph(t)
    def h(t):return doc.add_paragraph(t,'Heading 1')
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for i,w in enumerate(widths):t.columns[i].width=Inches(w)
        for j,value in enumerate(headers):t.rows[0].cells[j].text=value
        for row in rows:
            cells=t.add_row().cells
            for j,value in enumerate(row):cells[j].text=str(value)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders);marg=OxmlElement('w:tcMar')
                for side,value in [('top','60'),('bottom','60'),('left','100'),('right','100')]:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');marg.append(e)
                pr.append(marg)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF5');pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:r.font.size=Pt(10.5);r.bold=i==0
        spacer=doc.add_paragraph();spacer.paragraph_format.line_spacing=Pt(6);spacer.paragraph_format.space_after=Pt(6)
        return t
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    alpha=selection['chosen']['alpha']; adopted=alpha>0
    large=summary['large_gain_rule_met'];prior=summary['scores']['prior']['all57'];selected=summary['scores']['selected']['all57'];candidate=summary['scores']['candidate']['all57']
    verdict='达到预登记的探索性大幅改善标准' if large else ('开发期支持加入新模型 但未达到大幅改善标准' if adopted else '开发期未支持替代原最佳')
    doc.add_paragraph('Google PowerData 2019\n直接预测网络与共享状态研究汇报','Title')
    para('第十四轮研究记录   2026年10月4日   D盘研究目录')
    h('本轮结论与采用决定')
    para(f'本轮把研究扩大到直接预测完整日曲线，实际训练6个配置、每个2个随机种子，比较20、60、120轮检查点的双种子平均，共18个候选。前四个开发日选出 {selection["candidate"]}，后四日冻结新模型权重为 {alpha:g}。结果为{verdict}。')
    para(f'正式57域曲线WAPE为 {selected["wape"]:.6f}%，相对原最佳 {prior["wape"]:.6f}% 的误差下降为 {summary["selected_relative_reduction_percent"]:.4f}%。未采用或未完全采用的直接模型对照为 {candidate["wape"]:.6f}%。不能把训练损失下降、某个更低的局部指标或事后看到的测试名次当成正式提升。')
    table(['测试指标','原最佳','本轮正式结果'],[
        ['57域曲线 WAPE',f'{prior["wape"]:.6f}%',f'{selected["wape"]:.6f}%'],
        ['57域日均 WAPE',f'{prior["daily_wape"]:.6f}%',f'{selected["daily_wape"]:.6f}%'],
        ['57域峰值 MAPE',f'{prior["peak_mape"]:.6f}%',f'{selected["peak_mape"]:.6f}%'],
        ['a集群5域曲线 WAPE',f'{summary["scores"]["prior"]["a5"]["wape"]:.6f}%',f'{summary["scores"]["selected"]["a5"]["wape"]:.6f}%'],
        ['a集群5域日均 WAPE',f'{summary["scores"]["prior"]["a5"]["daily_wape"]:.6f}%',f'{summary["scores"]["selected"]["a5"]["daily_wape"]:.6f}%'],
        ['a集群5域峰值 MAPE',f'{summary["scores"]["prior"]["a5"]["peak_mape"]:.6f}%',f'{summary["scores"]["selected"]["a5"]["peak_mape"]:.6f}%']], [2.9,1.8,1.8])
    ci=summary['selected_day_bootstrap95_pp'];cci=summary['candidate_day_bootstrap95_pp']
    para(f'正式曲线改善量的探索性95%日期配对区间为 [{ci[0]:.6f}, {ci[1]:.6f}] 个百分点；直接模型对照为 [{cci[0]:.6f}, {cci[1]:.6f}]。大幅改善预定义为相对WAPE下降至少5%且改善区间下界大于0。若正式权重为0，零差值区间只说明与原模型相同，不代表预测误差为0。')
    para('2019年5月的开发期与测试期已被反复查看。即使达到本轮大幅改善规则，也只能是当前数据上的探索性证据，必须再用未查看过的真实功率数据确认。冻结规则不能恢复数据独立性或保证没有适应性过拟合。')
    h('十八个开发期候选的完整结果').paragraph_format.page_break_before=True
    para('主指标为57个PDU各自WAPE的平均值。前四日起报14至17只选网络与训练轮次，后四日起报18至21只选采用权重。每个分数由两个固定种子等权平均；没有挑选测试最好的种子或检查点。')
    labels={'dense_local64':'稠密 本域 64','dense_shared64':'稠密 共享 64','mixer_local32':'混合 本域 32','mixer_shared32':'混合 共享 32','mixer_local64':'混合 本域 64','mixer_shared64':'混合 共享 64'}
    rows=[]
    for r in dev['candidates']:
        name,ep=r['model'].rsplit('_e',1)
        rows.append([labels[name],ep,f'{r["early4"]["wape"]:.5f}',f'{r["late4"]["wape"]:.5f}'])
    table(['网络与宽度','训练轮数','前4日 WAPE %','后4日 WAPE %'],rows,[2.6,.8,1.55,1.55])
    h('扩大建模改动的方法与限制').paragraph_format.page_break_before=True
    para(f'原最佳前四日WAPE为 {dev["prior_early4"]["wape"]:.6f}%，后四日为 {dev["prior_late4"]["wape"]:.6f}%。上一页列出直接模型完整输出的分数，与部分权重的正式分数不同。')
    para('被选候选的后四日采用强度为：'+'；'.join(f'{r["alpha"]:g} 对应 {r["wape"]:.5f}%' for r in selection['weights'])+'。按日期和场站配对标准误中较大者的一标准误规则，优先较小权重。')
    import re
    from docx.opc.constants import RELATIONSHIP_TYPE
    p=doc.add_paragraph()
    source_urls={'TiDE':'https://arxiv.org/abs/2304.08424','TSMixer':'https://arxiv.org/abs/2303.06053'}
    t='稠密网络使用残差编码器，把历史小时片段和场站信息直接映射到未来96点；时间混合网络分别沿时间和特征维混合历史表示，再解码完整曲线。设计参考TiDE与TSMixer，属于本数据的专用实现，并非官方模型复现。两类模型都允许学习完整功率差值，没有此前历史均值正负1%的微小修正限制，也没有向旧模型或前一天曲线收缩的训练惩罚。'
    for part in re.split('(TiDE|TSMixer)',t):
        if part in source_urls:
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(source_urls[part],RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
            word_run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');pr.append(color)
            underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');pr.append(underline);word_run.append(pr)
            text=OxmlElement('w:t');text.text=part;word_run.append(text);link.append(word_run);p._p.append(link)
        else:p.add_run(part)
    para('共享状态来自其余PDU以及同场站其余PDU的历史标准化功率变化，排除目标PDU自身。三个本域与共享网络对照对使用相同宽度、参数数目和训练样本，仅共享输入是否置零不同。该状态不是物理功率总和，不能用来主张场站能量守恒。生产功率代理及其缺失掩码仍来自已完成历史；本轮没有增加未来真实CPU。')
    para('输入为过去672点，输出未来96点，按4点即一小时移动训练窗口。开发模型只用完成到第14天的数据，最终所选网络按固定轮数重训到第22天。开发阶段145个重叠窗口形成8265个PDU样本，窗口目标全部截止在训练边界内；重叠样本并不等于独立样本数。按训练窗口内有效历史均值与标准差归一化，最终输出限制在0至1。')
    rows=[]
    for name in ['dense_local64','dense_shared64','mixer_local32','mixer_shared32','mixer_local64','mixer_shared64']:
        meta=read(R/f'results/dev_{name}_seed20261004_training_meta.json')
        rows.append([labels[name],f'{meta["parameters"]:,}',meta['samples']])
    table(['架构','可学习参数','每个种子的训练样本'],rows,[2.7,1.5,2.3])
    h('独立核查与候选对照').paragraph_format.page_break_before=True
    para(f'核查共 {verification["check_count"]} 项全部通过，包括时间边界、18组双种子平均、前后开发期选择、模型文件与冻结记录、CPU重算已保存网络、独立指标与区间重算，以及未来扰动和共享状态排除自身检查。网络重算容差为利用率1e-6，仅处理CPU与原训练设备的数值差异，没有据此改写原选择。')
    table(['完整直接模型测试指标','原最佳','直接模型对照'],[
        ['57域曲线 WAPE',f'{prior["wape"]:.6f}%',f'{candidate["wape"]:.6f}%'],
        ['57域日均 WAPE',f'{prior["daily_wape"]:.6f}%',f'{candidate["daily_wape"]:.6f}%'],
        ['57域峰值 MAPE',f'{prior["peak_mape"]:.6f}%',f'{candidate["peak_mape"]:.6f}%'],
        ['a集群5域曲线 WAPE',f'{summary["scores"]["prior"]["a5"]["wape"]:.6f}%',f'{summary["scores"]["candidate"]["a5"]["wape"]:.6f}%']], [2.9,1.8,1.8])
    baseline=read(TOP/'experiment/accuracy_assessment/assessment.json')['baseline']
    para(f'直接候选的全域峰值误差有所改善，但前一天同刻基线的峰值MAPE为 {baseline["peak_mape"]:.6f}%，仍优于该候选。峰值下降伴随整体曲线和日均误差增加，不能称为整体精度大幅提升，也不能宣称峰值已达到现有方法最优。')
    para('训练损失由有效点绝对误差、日均误差、峰值误差和越界惩罚组成；实际训练与模型都已保留。停止与恢复保留优化器及随机数状态，避免重复启动。基础模型权重没有重新下载或加载；训练的是本轮新网络，而不是再次微调原基础模型。')
    h('对继续提高精度的判断')
    para('此前第9轮完整CPU与MoE、第10轮原生协变量、第11轮残差树、第12轮鲁棒专家融合均未支持替代原最佳。第13轮候选的测试小幅改善也因开发后段变差而未采用。本轮进一步扩大网络和替代幅度，结论仍以冻结结果为准；不能用模型更复杂证明方法更有效。')
    para('当前数据只有31天，缺少独立月份。即使本轮未改善，也只支持当前输入上的工程平台判断，未证明理论误差极限。后续高价值条件是未查看过的真实功率数据，或严格对齐的工作量与设备功耗观测；单纯反复挑同九天的最低测试值会累积选择偏差。')
    h('部署交付与来源范围').paragraph_format.page_break_before=True
    para('冻结预测已实际导出57域5472行，时间为2019年6月1日00:00至23:45。数值有限且在0至1内，经验区间顺序正确。六月没有真值，输出是历史复现推理，不是今天的电负荷；80%区间仅按开发残差经验校准，不保证独立数据覆盖率。')
    table(['保存位置','用途'],[
        ['experiment/round14/docs/protocol.json','候选 训练边界与约束'],
        ['experiment/round14/models','双种子模型与恢复状态'],
        ['experiment/round14/results/selection.json','测试前冻结的选择'],
        ['experiment/round14/results','全体开发分数 正式预测与审计'],
        ['16_直接预测共享状态_预测次日.cmd','复现正式冻结预测']],[4.2,2.3])
    para('全部新增程序、模型与产出在 ${PROJECT_ROOT}，桌面同名目录为兼容链接。历轮权重、原始数据、负面结果和正式汇报均保留。开始阅读、当前进展及完整文件清单随本轮更新。')
    h('资源与数据口径')
    runtime=read(R/'results/runtime.json')
    para(f'训练设备为 {runtime["gpu"] or runtime["device"]}，CPU最多2线程。启动仍需3GiB可用物理与6GiB提交额度；运行中低于1.5GiB物理或2GiB提交连续两次，仅停止本研究子进程并保留恢复状态。实际运行未触发资源停止，采样可用物理最低 {run["minimum_available_physical_bytes"]/2**30:.3f}GiB。定时自动续跑保持停用。')
    para('官方cell a的12085个Parquet必要列已完整读取和审计，剩余下载0；完整真实CPU只覆盖a集群五个功率域。不能说已读取全部八个集群或全部字段。本轮共享功率历史覆盖57域，不等于增加了57域CPU数据。')
    para('84.78GiB为保留压缩列，407.55GiB为对应原对象总量，105.481GiB为迁移时逻辑文件总量，口径不同。功率为归一化容量利用率，CPU为NCU，均不能直接当作kW。任务窗口仅统计JOB，缺少完整去重键与严格一一对应的物理观测，时间对齐存在限制。Google数据为CC BY 4.0；原方案使用的TimesFM 3权重仍限非商业、非生产使用。')
    doc.save(OUT);print(OUT)
if __name__=='__main__':main()

