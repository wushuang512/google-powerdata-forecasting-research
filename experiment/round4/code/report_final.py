from pathlib import Path
import json,sys,csv
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
R=Path(sys.argv[1]);B=R.parent;TOP=B.parent
def main():
    r3=json.loads((B/'round3/results/summary.json').read_text());s=r3['scores'];r2=json.loads((B/'round2/results/summary.json').read_text());seven=json.loads((R/'results/seven_day_summary.json').read_text())['scores'];rob=json.loads((R/'results/robustness.json').read_text());verify=json.loads((R/'results/verification.json').read_text());r4=json.loads((R/'results/summary.json').read_text());comp=r3['comparison_vs_old']['secondary_equal_mix']
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21.59);sec.page_height=Cm(27.94);sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for st in doc.styles:
        for x in list(st.element.iter(qn('w:pBdr'))):x.getparent().remove(x)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Calibri';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor.from_string('000000')
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.15;doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 2'].font.size=Pt(12)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;foot.add_run('Google PowerData 2019 · 持续探索汇报  |  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';run.append(text);fld.append(run);foot._p.append(fld)
    for rr in foot.runs:rr.font.size=Pt(8)
    uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');doc.settings.element.append(uf)
    def p(t,b=False):q=doc.add_paragraph();q.add_run(t).bold=b;return q
    def h(t):doc.add_heading(t,1)
    def page(t):doc.add_page_break();h(t)
    def table(head,rows,widths):
        tb=doc.add_table(rows=1,cols=len(head));tb.autofit=False
        for c,w in zip(tb.columns,widths):c.width=Cm(w)
        for c,t,w in zip(tb.rows[0].cells,head,widths):c.text=str(t);c.width=Cm(w)
        repeat=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for c,t,w in zip(tb.add_row().cells,row,widths):c.text=str(t);c.width=Cm(w)
        for i,row in enumerate(tb.rows):
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for j,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E1EDF1' if i==0 else ('F4F7F8' if i%2==0 else 'FFFFFF'));pr.append(shade)
                borders=OxmlElement('w:tcBorders')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
                pr.append(borders);mar=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');mar.append(x)
                pr.append(mar)
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.08;para.alignment=WD_ALIGN_PARAGRAPH.CENTER if j>0 else WD_ALIGN_PARAGRAPH.LEFT
                    for rr in para.runs:rr.font.size=Pt(9.5);rr.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        return tb
    doc.add_heading('Google PowerData 2019\n持续探索与精度改进汇报',0)
    p('研究日期 2026 年 9 月 30 日　｜　第三轮与第四轮结果及累计结论')
    h('结论先行')
    p('已经跑出比旧最优更低的次日曲线误差：WAPE 从 2.4165% 降至 2.3848%，相对下降 1.31%。这是小幅、探索性的实测改进，尚不能称为跨集群稳定的显著提升。',True)
    p('表现最好的方案是测试前冻结的第三轮次要组合：25% Chronos-2 单变量、25% TiRex-2 单变量、50% TimesFM 3.0 集群与历史生产特征，最终利用率限制在 [0,1]。第三轮主方案为 2.3914%。第四轮新增调参与微调未获得足够可靠的额外收益，按保守验证规则保留该组合。')
    rows=[]
    for label,key in [('前一天同刻','last'),('TiRex-2 单变量','tirex_native'),('复现旧等权组合','old_equal_mix'),('TimesFM 3.0 集群与生产','tf3_cellprod'),('第三轮冻结主方案','new_curve'),('第三轮冻结次要组合','secondary_equal_mix')]:
        x=s[key];rows.append([label,f'{x["wape"]:.4f}',f'{x["daily_wape"]:.4f}',f'{x["peak_mape"]:.4f}'])
    table(['方法','曲线 WAPE %','日均 WAPE %','峰值 MAPE %'],rows,[7,3.2,3.2,3.7])
    p('统一测试为 2019-05-23 至 05-31，共 9 天、57 个供电域。七天曲线仍以 TiRex-2 的 2.8803% 较优；新组合为 2.9032%。不同目标应分别选择方案。')
    p('目标状态：未达到“显著优于已有方式”的充分证据标准。同月已反复探索，日期块与集群联合区间跨零；不能保证没有选模过拟合，也不能据此承诺跨月精度。',True)
    page('数据与评价协议')
    p('数据来自 Google 官方 PowerData 2019，而非第三方镜像。样本覆盖 2019 年 5 月的 31 天、10 个集群与 57 个供电域；每域 8,928 条五分钟记录，总计 508,896 条。先按时间排序，再将连续三个五分钟点聚合为十五分钟点，每日 96 点。')
    p('目标是按供电容量归一化的功率利用率，不是 kW 或 MW。原始总功率坏标记有 24 条；生产功率坏标记约 22.90%。按官方 Notebook 使用 bad_* 为真即无效的规则，评分排除无效总功率，历史输入因果前向填充，生产特征保持缺失掩码。官方 PDF 与 Notebook 的标记文字差异已记录。')
    table(['阶段','使用范围','作用'],[['开发训练','截至 05-14','拟合局部参数与残差模型'],['开发验证','预测 05-15—05-22','选择配置、权重与停止点'],['最终训练','截至 05-22','只重训已冻结的入选配置'],['次日回测','预测 05-23—05-31','9 个起报点，49,243 个有效点'],['七天回测','05-22、23、24 日结束起报','3 个重叠窗口，114,902 个有效点']],[3,5,9.1])
    p('主指标为供电域宏平均 WAPE：每个域先将全部有效时点的绝对误差之和除以真值绝对值之和，再对 57 个域等权平均。日均与峰值指标使用相同有效时点，避免缺失导致不公平比较。报告中的改善“百分点”与相对下降百分比明确区分。')
    p('测试期每次预测只读取起报时刻以前的数据。已经观测到的前一日可以成为下一日的历史输入；残差在线校准也只使用已结束预测的误差。模型参数、组合权重不能根据测试期未来标签更新。七天递推不读入中途真实未来值。')
    p('这不是从未见过的独立测试集：第二轮以前已查看过整月表现。第三、四轮内部冻结流程能够约束进一步调参，但不能撤销历史查看带来的适应性偏差。基础模型预训练是否包含本公开数据也尚未排除。',True)
    page('探索内容与保留依据')
    table(['探索方向','实际配置范围','结果处理'],[['第二轮基础与微调','32 个成功候选，15 次开发期微调','曲线主方案 2.4463%，未超过旧最优'],['第三轮多尺度与变换','Chronos 7 个、TiRex 9 个配置','含降采样、上下文、镜像与差分'],['残差学习与在线校准','11 个正则残差模型、18 个校准配置','形成主方案的部分组成'],['TimesFM 3.0','8 个单变量、多变量与协变量配置','集群与生产组合进入次要方案'],['层级一致性','12 个共享集群指数与残差配置','不解释为实际电量守恒'],['第四轮目标对齐微调','4 种训练范围 × 3 个检查点','新增权重按验证规则取 0'],['TimesFM 2.5','3 种上下文 × 2 种翻转设置','入选对照测试 WAPE 2.4771%']],[4.3,7.6,5.2])
    p('第三轮主方案将残差 Ridge、因果校准及基础模型组合。模型家族成员先在前四个验证日选择，组合权重在时间与集群划分上检查，再用完整八日开发期冻结。集群划分仅约束组合权重，不能表述为所有底层模型都未见过该集群。主方案测试误差 2.3914%；固定半旧组合、半 TimesFM 3 的次要方案为 2.3848%。')
    p('第四轮直接优化掩码后的点预测 WAPE 近似目标，以冻结教师输出限制参数漂移，并加入生产特征辅助损失、生产不高于总量的惩罚和容量边界惩罚。尝试只训练输出层、尾部网络和全量参数；学习率为 5×10⁻⁷ 至 10⁻⁵，检查点为 50、150、300 步。全部开发权重和损失记录保留。')
    p('前四个验证日选出的微调候选，在后四日以 10% 混合仅改善 0.000755 个百分点，配对集群标准误为 0.005155 个百分点。按一个标准误规则优先更简单、较低权重的方案，最终新增权重为零，未再为追求更好测试分数强行采用微调。')
    p('最终胜出组合只保留容量边界约束；分量守恒与生产顺序惩罚属于未胜出的探索分支。生产特征帮助预测不等于证明物理机理，剩余分量也不是实测冷却功率。',True)
    page('改进的稳定性与不确定性')
    p('固定新组合向更早日期补充回测，没有依据该结果再改权重。四个时间块整体均改善，但这属于选模后的稳健性审计，不能当作新增独立验证。下表按各时间块重新计算宏平均 WAPE，未把每日 WAPE 简单平均。')
    table(['目标日期','旧组合 WAPE %','新组合 WAPE %','相对下降 %'],[[x['period'].replace('May','05/'),f'{x["old"]:.4f}',f'{x["new"]:.4f}',f'{x["relative_percent"]:.2f}'] for x in rob['time_blocks']],[4.3,4.3,4.3,4.2])
    doc.add_picture(str(R/'figures/test_daily.png'),width=Cm(16.8))
    p('图 1　九个测试日的曲线误差；新组合并非每天都更优。')
    p('测试期新组合在 32/57 个供电域、6/9 个日期、6/10 个集群获益。用 5,000 次配对重采样估计旧减新误差：连续三日块的 95% 区间为 [0.0118, 0.0656] 个百分点；同时重采样日期块和集群后为 [-0.0211, 0.1068] 个百分点。后者跨零，因此不作跨集群显著优势结论。')
    p('区间是小样本、探索性的描述，未校正多轮多配置筛选，也不是跨月误差保证。集群仅 10 个、测试仅 9 天；不能只引用正的日期区间而忽略联合区间。',True)
    page('不同目标的结果与防泄漏检查')
    table(['七天方法','曲线 WAPE %','日均 WAPE %','峰值 MAPE %'],[[label,f'{seven[k]["wape"]:.4f}',f'{seven[k]["daily_wape"]:.4f}',f'{seven[k]["peak_mape"]:.4f}'] for label,k in [('TiRex-2 单变量','tirex_uni'),('新固定组合','best_mix'),('TimesFM 3.0 集群与生产','tf3_cellprod'),('旧等权组合','old_equal_mix')]],[7,3.2,3.2,3.7])
    p('七天组合直接沿用次日权重，没有看七天测试结果再调权重。TiRex-2 官方单次预测上限为 320 点，本地使用 320＋320＋32 点递推，追加预测而非未来真值。三个七天窗口存在大量重叠，不能当作 21 个独立测试日。')
    p('次日峰值方面，新组合 MAPE 为 2.0687%，旧组合为 2.1185%，但前一天同刻只有 1.0682%。第二轮峰值修正为 1.6138%，其曲线 WAPE 为 2.4603%。因此该研究的主要收益是整体曲线，不能把它宣传为峰值预测全面领先。')
    h('实现检查结果')
    p(f'最新 {len(verify["checks"])} 项检查全部通过：在多个起报日扰动未来总功率、生产功率及有效标记，集群输入、残差特征和在线校准结果保持不变；训练抽样不越过 14 日或 22 日截止线；第三、四轮冻结文件哈希一致；固定组合可精确重建；指标复算一致，预测满足容量边界。')
    p('检查证明这些已覆盖路径没有使用被扰动的未来数据，不能证明所有可能形式的过拟合都不存在。第二轮另有 10,000 组随机物理投影检查；完整记录位于 round2/results/verification.json 与 round4/results/verification.json。')
    h('真实可用性边界')
    p('桌面已安装本机可运行的模型与依赖，实际导出次日和七天示例。默认历史截止 2019-05-31，输出从 2019-06-01 开始，数据没有六月真值，故示例只能验证推理流程，不能据此报告预测精度。')
    p('新组合包含 TimesFM 3.0 的非商业、非生产许可权重，只作为本地研究入口。现有运行脚本依赖本机 CUDA、绝对路径和本数据形状，尚不是面向任意月份输入的生产服务。')
    page('桌面交付与后续研究')
    p('统一保存位置：${PROJECT_ROOT}。打开“开始阅读.md”可看到使用说明；01 入口运行次日新组合，02 入口运行七天 TiRex，03 保留旧组合，04 保留峰值修正，05 用于新组合七天对照。E 盘原件仍保留。')
    p('google-powerdata-2019 保存官方原始 gzip、解压 CSV、来源与许可；runtime/python 和 experiment/.venv 保存本机运行环境；experiment/models 保存基础权重。round2、round3、round4 分别保留各轮代码、开发权重、预测、参数、成功与失败日志、图表和报告校验文件。最新结论以本报告为准，前两轮报告保留为研究过程记录。')
    p('重点复核文件：round4/results/forecasts_1day.csv 为逐点新旧预测与真值；test7.csv 为七天结果；time_blocks.csv 为时间块统计；round3/results/per_domain.csv 与 per_day.csv 为细分结果；selection.json 及 SHA256 为冻结记录。全部依赖版本见 requirements-lock.txt，权重校验与目录大小见“交付文件清单.json”。')
    h('来源与许可')
    p('数据：github.com/google/cluster-data/blob/master/PowerData2019.md；下载自 storage.googleapis.com/powerdata_2019。许可 CC BY 4.0。58 个 gzip 共 3,254,733 字节（3.104 MiB），解压 CSV 共 24,763,748 字节；包括 57 域和 1 个机器映射表，官方不提供总 ZIP。')
    p('官方模型：huggingface.co/amazon/chronos-2、NX-AI/TiRex-2、google/timesfm-2.5-200m-pytorch，均保留 Apache 2.0 许可。google/timesfm-3.0-pytorch 使用 TimesFM Non-Commercial License v1.0，仅限许可允许的非商业、非生产用途，且禁止按该许可再分发模型。模型目录保留各自原始 LICENSE 与固定版本记录。')
    h('下一阶段的判断标准')
    p('停止把同月测试区间当作可无限调参的目标。若获得新月份或真实独立集群，先锁定本报告的新旧方案、清洗方法和指标，再一次性评价；分离训练、校准和独立评估，按连续日期与集群报告不确定性，并单独考核峰值与七天任务。')
    p('可把“相对旧方案至少下降 5%，联合日期与集群区间下界大于零，且峰值无实质恶化”作为下一阶段预先约定的实践门槛。本月最优下降 1.31%，尚未达到该门槛。当前应采用的研究结论是：找到了可复现的小幅改进组合，同时明确保留未取得显著突破的结论。',True)
    for e in list(doc.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
    target=TOP/'Google PowerData 2019 持续探索与精度改进汇报.docx';doc.save(target);print(target)
if __name__=='__main__':main()
