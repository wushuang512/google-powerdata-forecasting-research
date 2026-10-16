from pathlib import Path
from zipfile import ZipFile
import sys,json,hashlib,datetime
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
T=Path(sys.argv[1]);R=T/'experiment/round44';W=R/'docs/integrated_template';Q=R/'results';src=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx';out=T/'Google PowerData 2019 三项预测任务与模型比较汇报_加入HDSF.docx'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
e=read(W/'template_evidence.json');a=read(Q/'hdsf_comparison.json');log=read(W/'authoring_log.json');old=Document(src);new=Document(out);checks=[]
def check(n,v):checks.append(dict(name=n,passed=bool(v)));assert v,n
def num(s):return s.replace('★','').strip()
check('original_word_unchanged',sha(src)==e['source_sha256']);zs=ZipFile(src);zn=ZipFile(out);check('same_package_parts',set(zs.namelist())==set(zn.namelist()))
for p in e['parts']:
 if p['scope']=='preserve_only':check('preserved '+p['path'],hashlib.sha256(zn.read(p['path'])).hexdigest()==p['sha256'])
check('same_section_XML',etree.tostring(old.sections[0]._sectPr)==etree.tostring(new.sections[0]._sectPr));check('same_section_and_table_count',len(old.sections)==len(new.sections)==1 and len(old.tables)==len(new.tables)==11)
for i in [3,7,8]:check('unchanged source table '+str(i),etree.tostring(old.tables[i]._tbl)==etree.tostring(new.tables[i]._tbl))
oldhead=[p.text for p in old.paragraphs if p.style.name=='Heading 1'];newhead=[p.text for p in new.paragraphs if p.style.name=='Heading 1'];check('same_chapter_order',len(newhead)==8 and oldhead[1:]==newhead[1:])
for index in [4,5,6]:
 for i,row in enumerate(old.tables[index].rows[1:],1):
  newi=i if i==1 else i+1;check(f'T{index} oldrow{i} unchangednumeric',[num(c.text) for c in row.cells[1:]]==[num(c.text) for c in new.tables[index].rows[newi].cells[1:]])
fields_curve=['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae','rmse','peak_mape','p95_ape','hit2_percent','hit5_percent'];fields_daily=['daily_wape','relative_daily_error_reduction_percent','daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent']
for index,H,fields in [(4,1,fields_curve),(5,7,fields_curve),(6,7,fields_daily)]:
 keys=log['table_model_keys'][str(index)];bank=[r for r in a['unique_rows'] if r['horizon_days']==H];check(f'T{index} each oldmodelplusHDSF',len(new.tables[index].rows)==len(keys)+1 and any('HDSF半日单纯形融合' in row.cells[0].text for row in new.tables[index].rows))
 for i,key in enumerate(keys,1):
  r=next(r for r in bank if r['model']==key or key in r['aliases'])
  for j,k in enumerate(fields,1):
   v=r.get(k);cell=new.tables[index].rows[i].cells[j];check(f'T{index} row{i} {k} actual_value',num(cell.text)==('—' if v is None else f'{v:.6f}'))
   winning=v is not None and k!='development_wape_original_grid' and abs(v-(max if 'hit' in k or k.startswith('relative') else min)(r[k] for r in bank))<1e-10
   check(f'T{index} row{i} {k} updated_best_mark',('★' in cell.text)==winning and (cell._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill'))=='E9F4E5')==winning)
for i in range(1,17):
 check(f'multidimensional_oldrow{i}_values_retained',[num(c.text) for c in old.tables[1].rows[i].cells]==[num(new.tables[1].rows[i].cells[j].text) for j in [0,1,2,4,5]])
for H,pos in [(1,2),(7,3)]:
 r=next(r for r in a['unique_rows'] if r['model']==f'r44_hdsf_H{H}');curve=r['wape'];check(f'summary_HDSF_H{H}',new.tables[0].rows[pos-1].cells[2].text==f'{curve:.6f}')
check('summary_HDSF_daily',new.tables[0].rows[3].cells[2].text==f'{next(r for r in a["unique_rows"] if r["model"]=="r44_hdsf_H7")["daily_wape"]:.6f}')
prior_iv=[[c.text for c in row.cells[1:]] for row in old.tables[10].rows[1:]];new_iv=[[c.text for c in row.cells[1:]] for row in new.tables[10].rows[1:] if 'HDSF' not in row.cells[0].text];check('old_intervals_preserved',prior_iv==new_iv)
for H in [1,7]:
 iv=next(v for v in a['new_intervals'] if v['horizon_days']==H and v['method'].startswith('HDSF'));row=next(row for row in new.tables[10].rows if row.cells[0].text.startswith(f'{H}天 HDSF'));check(f'actual_HDSF_interval_H{H}',[c.text for c in row.cells[1:]]==['80',f'{iv["coverage_percent"]:.3f}',f'{iv["mean_width_normalized"]:.6f}',f'{iv["winkler80"]:.6f}'])
check('canonical_evaluation_source_unmodified',sha(T/'experiment/round29/code/aligned_metrics.py')==a['reference_audit']['canonical_code_sha256']);check('new_word_hash_authoring_matches',sha(out)==log['output_sha256'])
part_old=etree.fromstring(zs.read('word/footer1.xml'));part_new=etree.fromstring(zn.read('word/footer1.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};check('PAGE_fields_preserved',part_old.xpath('.//w:fldSimple/@w:instr',namespaces=ns)==part_new.xpath('.//w:fldSimple/@w:instr',namespaces=ns))
print(json.dumps(dict(all_passed=True,checks=checks,check_count=len(checks),source_sha256=sha(src),output_sha256=sha(out),source_preserved=True,canonical_evaluation_unchanged=True,verified_utc=datetime.datetime.now(datetime.timezone.utc).isoformat()),ensure_ascii=False,indent=2))
