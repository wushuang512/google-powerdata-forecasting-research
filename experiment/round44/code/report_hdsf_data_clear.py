from pathlib import Path
from copy import deepcopy
import argparse,csv,hashlib,json,zipfile,datetime
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ap=argparse.ArgumentParser();ap.add_argument('root');ap.add_argument('--validate-only',action='store_true');args=ap.parse_args()
T=Path(args.root);B=T/'experiment';R=B/'round44';Q=R/'results';D=R/'docs/causes_data_report_20261009'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
a=read(D/'data_audit.json');ev=read(D/'template_evidence.json');c=read(Q/'hdsf_comparison.json');s=read(Q/'summary.json');mechanism=read(Q/'mechanism_analysis_20261008/analysis.json')
assert a['all_passed'] and read(Q/'mechanism_analysis_20261008/verification.json')['all_passed']
assert a['power_columns']==['time','cell','pdu','measured_power_util','production_power_util','bad_measurement_data','bad_production_power_data']
assert sha(D/'reference.docx')==ev['reference_sha256']==sha(Path(ev['reference']))
for p,h in a['source_file_sha256'].items():assert sha(T/p)==h
assert a['invalid_measured_blocks']==a['partial_measured_blocks']==24 and a['all_empty_measured_blocks']==0
assert a['local24_feature_count']==64 and a['target_valid_slots']==169608
assert a['production_causal_median_filled_blocks']+a['production_causal_median_remaining_nonfinite']==38556
if args.validate_only:
    print(json.dumps(dict(all_passed=True,real_results=True,raw_fields=7,mapping_fields=3,local_features=64,expected_pages=10)));raise SystemExit

ref=Document(D/'reference.docx');doc=Document(D/'reference.docx');template=deepcopy(ref.tables[4]._tbl)
body=doc._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)
output=T/'Google PowerData 2019 HDSF结果原因与数据处理清晰版汇报.docx'
width=doc.sections[0].page_width.inches-doc.sections[0].left_margin.inches-doc.sections[0].right_margin.inches
headings=[];tables=[];field_rows=[]
def para(text,bold=False,size=None,after=7):
    p=doc.add_paragraph(style='Normal');p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.1
    r=p.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    r.font.color.rgb=RGBColor(0,0,0);return p
def heading(text,page=False,title=False):
    p=doc.add_paragraph(text,style='Title' if title else 'Heading 1');p.paragraph_format.page_break_before=page;p.paragraph_format.keep_with_next=True
    for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
    headings.append(text);return p
def sub(text):return para(text,True,12,5)
def grid(headers,rows,widths,highlight=None):
    tbl=doc.add_table(rows=1,cols=len(headers));tbl.autofit=False
    # Clone the actual source table property system; content and grid are user-authorized rewrite slots.
    pr=tbl._tbl.tblPr;tbl._tbl.remove(pr);tbl._tbl.insert(0,deepcopy(template.tblPr))
    for i,w in enumerate(widths):tbl.columns[i].width=Inches(w)
    for ri,values in enumerate([headers]+rows):
        row=tbl.rows[0] if ri==0 else tbl.add_row()
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,value in enumerate(values):
            cell=row.cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for k,val in [('top',55),('bottom',55),('left',75),('right',75)]:
                m=OxmlElement('w:'+k);m.set(qn('w:w'),str(val));m.set(qn('w:type'),'dxa');margins.append(m)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for k in ['top','left','bottom','right','insideH','insideV']:
                v=OxmlElement('w:'+k);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'D9D9D9');borders.append(v)
            tcpr.append(borders)
            fill='DCEAF2' if ri==0 else 'F5F7F9' if ri%2==0 else 'FFFFFF'
            best=ri>0 and highlight and highlight(ri-1,j)
            if best:fill='E9F4E5'
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);tcpr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(13)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or not isnumber(value) else WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(str(value));r.bold=ri==0 or bool(best);r.font.name='Calibri';r.font.size=Pt(10.5);r.font.color.rgb=RGBColor(0,0,0)
            r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    tables.append(dict(headers=headers,rows=rows))
    gap=para('',after=6);gap.paragraph_format.line_spacing=Pt(1)
    return tbl
def isnumber(v):
    try:float(str(v).replace('%','').replace('+',''));return True
    except ValueError:return False
def n(v):return f'{v:.6f}'
prior={H:next(r for r in c['unique_rows'] if r['horizon_days']==H and r['status']=='当前综合推荐') for H in [1,7]}
hd={H:s['scores'][str(H)]['hdsf'] for H in [1,7]}

# Page 1: results and explanations before any background.
heading('Google PowerData 2019 HDSF结果与数据处理汇报',title=True)
para('24小时曲线和7天日均领先 7天曲线尚未领先',True,15,6)
para('2026年10月9日  基于已完成的真实回测与10月8日固定消融  评价沿用9.28统一口径',size=9.5,after=7)
rows=[]
for label,H,k in [('24小时完整曲线 WAPE%',1,'wape'),('未来7天完整曲线 WAPE%',7,'wape'),('未来7天每日平均负荷 WAPE%',7,'daily_wape'),('24小时日最大值 MAPE%',1,'peak_mape'),('7天日最大值 MAPE%',7,'peak_mape')]:
    rows.append([label,n(prior[H][k]),n(hd[H][k]),f'{hd[H][k]-prior[H][k]:+.6f}', '改善' if hd[H][k]<prior[H][k] else '退步'])
grid(['主要结果','此前推荐','HDSF','变化 百分点','判断'],rows,[3.3,1.7,1.7,1.7,width-8.4],lambda i,j:j==2 and rows[i][-1]=='改善')
para('当前比较池中，HDSF的24小时曲线为56组不同预测中最低，7天日均为58组中最低；7天曲线排第6，R40的2.694096%仍更低。三项相对误差降幅分别约1.63%、0.18%、1.86%，不称为大幅提高。',True,11.5)
sub('主要原因及已有证据')
para('分时使用专家。24小时前半天主要由Chronos-2和Toto2.0预测，后半天主要由两种TimesFM3模式预测；总改善的约86.52%发生在后半天。同一成员池中，分段优于等权平均及整段共享权重。')
para('7天日均更稳。每日平均水平MSE下降6.35%，日内形状MSE反而上升1.14%；历史水平与日型平滑平均占29.01%权重，删除后日均WAPE变差。这解释了日均优势大于完整曲线优势。')
para('成员误差有限互补。非负组合能抵消部分一高一低的误差，但主要成员残差相关仍约0.93至0.96，所以改善有限；24小时本地LightGBM成员平均仅占0.63%，不是主要来源。')
para('证据边界：同月测试多轮查看、7天窗口重叠，探索置信区间跨零；峰值、cell-a五域等仍有退步。结论为当前回测中的局部优势，未证明独立月份的全面领先。',True)

# Page 2: complete metric matrix.
heading('三项任务及多维评价结果',True)
para('浅绿色表示优于同跨度的原推荐。数值均由同一有效真值和掩码评分，不混入原项目池化WAPE。MAE和RMSE单位均为归一化功率利用率。')
fields=[('曲线 WAPE%','wape'),('曲线较历史基线误差减少%','relative_curve_error_reduction_percent'),('曲线 MAE','mae'),('曲线 RMSE','rmse'),('逐点 APE 第95百分位%','p95_ape'),('逐点 APE≤2% 占比%','hit2_percent'),('逐点 APE≤5% 占比%','hit5_percent'),('每日均值 WAPE%','daily_wape'),('日均较历史基线误差减少%','relative_daily_error_reduction_percent'),('每日均值 MAE','daily_mae'),('每日均值 RMSE','daily_rmse'),('日均 APE≤2% 占比%','daily_hit2_percent'),('日均 APE≤5% 占比%','daily_hit5_percent'),('日最大值 MAPE%','peak_mape'),('峰时平均绝对误差 分钟','peak_time_mae_minutes'),('全局绝对偏差%','absolute_bias_percent'),('cell-a五域曲线 WAPE%','cella5_wape'),('其余52域曲线 WAPE%','other52_wape')]
rows=[[label,n(prior[1][k]),n(hd[1][k]),n(prior[7][k]),n(hd[7][k])] for label,k in fields]
def improved(i,j):
    if j not in [2,4]:return False
    H=1 if j==2 else 7;k=fields[i][1];higher='hit' in k or k.startswith('relative_');return hd[H][k]>prior[H][k] if higher else hd[H][k]<prior[H][k]
grid(['评价指标','原24小时 R32','HDSF 24小时','原7天 R36','HDSF 7天'],rows,[3.0]+[(width-3)/4]*4,improved)
para('R32为Chronos-2、TiRex-2、TimesFM3、Ridge与末层微调TinyTimeMixer的稀疏峰值10%方案。R36为Moirai2同cell14天10%加原TimesFM3、Ridge1000、HGB、TinyTimeMixer90%，作全57域预测日均协调。',size=10)
para('HDSF对原推荐的曲线改善覆盖24小时39/57域与5/9起报，7天29/57域与1/3起报。7天日均误差改善602/1187条完整日记录；重复日期不是独立样本。',size=10)

# Page 3: actual named models and weights.
heading('HDSF实际使用哪些模型',True)
para('将预测按12小时分段。源实现先拟合基础模型A和加入本地成员B的两套权重，再按开发期冻结比例混合。最终每段权重非负且总和为1，跨57域及起报共享；预测时不根据未来真值换专家。')
sub('24小时两个时段的实际有效权重')
cfg=s['selection']['configs']['24h'];aw=cfg['component_a']['2'];bw=cfg['component_b']['2'];mix=cfg['mix']
eff=[[(1-mix)*(aa+[0])[k]+mix*bb[k] for k in range(len(bb))] for aa,bb in zip(aw,bw)]
names24={'chronos2_long':'Amazon Chronos-2 最长28天历史','toto2_joint':'Datadog Toto2.0 313M 同cell和设备类型调用 14天','timesfm30':'Google TimesFM3 57域联合调用 14天','timesfm30_uni':'Google TimesFM3 逐域调用 最长21天','local24':'四种子LightGBM残差与加权滞后回归各50%'}
rows=[]
for k,method in enumerate(cfg['methods']):
    if max(eff[0][k],eff[1][k])>0:rows.append([names24[method],f'{100*eff[0][k]:.2f}%',f'{100*eff[1][k]:.2f}%'])
grid(['模型及输入方式','00时至11时45分','12时至23时45分'],rows,[6.4,(width-6.4)/2,(width-6.4)/2])
para('TiRex-2、TimesFM2.5、Toto2.0逐域及TimesFM3联合长历史也是候选，本轮24小时有效权重为0。A/B混合比例7.5%不是本地模型直接占7.5%；本地模型全天平均有效权重仅0.63%。')
sub('7天五个成员的平均有效权重')
weights=mechanism['findings']['7']['mean_effective_weights'];names7=['Google TimesFM3 57域联合调用 最长21天','Datadog Toto2.0 313M 同cell和设备类型调用 14天','Google TimesFM3 逐域调用 最长21天','IBM Granite PatchTST-FM-r2 385M 最长21天','历史水平与日内形状平滑 最多24天']
grid(['实际成员','14个半日段的平均权重'],[[name,f'{100*w:.2f}%'] for name,w in zip(names7,weights)],[7.4,width-7.4])
para('7天各段权重仍不同，上表平均值不用于替代逐段计算。基础模型使用已完成的历史实测功率；Toto按cell和PDU/MVPP类型分组。没有新增基础模型梯度训练；HDSF未读取ClusterData真实CPU或内存。local24使用历史production_power_util估算代理。')

# Page 4: measured mechanisms.
heading('消融与误差分解怎样支持原因判断',True)
abl=mechanism['ablations']
def case(H,label):return next(r for r in abl if r['horizon_days']==H and r['case']==label)
rows=[]
for label,key in [('原冻结HDSF','HDSF frozen'),('同一成员全部等权','All members uniform'),('整段一组权重 开发重拟合','DEV-refit 1 segments, original mix fixed')]:
    x=case(1,key);y=case(7,key);rows.append([label,n(x['test_wape']),n(y['test_wape']),n(y['test_daily_wape'])])
x=case(7,'DEV-refit 7 segments, original mix fixed');rows.append(['7天每天一组权重 开发重拟合','—',n(x['test_wape']),n(x['test_daily_wape'])])
grid(['固定诊断','24小时曲线 WAPE%','7天曲线 WAPE%','7天日均 WAPE%'],rows,[4.2]+[(width-4.2)/3]*3)
para('分段优于整段共享；但每天分段的7天曲线略优于半日，日均与峰值却较差。12小时粒度没有同时统治所有指标。结构对照只用开发期重拟合权重，固定原A/B比例；所有结果为事后诊断，不据测试名次替换冻结方案。')
sub('7天日均优势来自水平误差而不是形状全面改善')
ff=mechanism['findings']['7'];rows=[]
for label,key in [('日均水平MSE','level_mse'),('去日均后的形状MSE','shape_mse'),('完整有效日总体MSE','mse')]:
    old=ff['prior_mse_decomposition'][key];new=ff['hdsf_mse_decomposition'][key];rows.append([label,f'{old:.9f}',f'{new:.9f}',f'{100*(new/old-1):+.3f}%'])
grid(['平方误差组成','原R36','HDSF','相对变化'],rows,[4.2]+[(width-4.2)/3]*3)
para('仅在1187条96点均有效的记录上分解。每日水平误差下降，形状误差略升；日均误差标准差降低4.41%。全局偏差仍增大，不能把日均MAE/RMSE改善解释为所有偏差都减小。')
para('固定删除历史平滑并归一剩余权重，7天曲线变为2.731997%，日均1.910905%；删除PatchTST-FM后为2.711203%和1.884065%。独立成员较差也可提供互补；删除后未重拟合，其差值不是严格贡献分摊。')
para('成员误差一高一低时组合可抵消。24小时加权绝对误差诊断量2.367305%降至组合2.335113%；7天2.795796%降至2.701623%。诊断量依赖真值，不能当可部署预测。主要24小时成员残差相关0.934和0.956，仍常共同犯错。')
para('峰值取舍：24小时峰值低估比例由87.72%增至90.64%。平均化可能压低错位尖峰，HDSF曲线损失也没有独立峰值目标；旧R32专门修正峰值。相关机制与观察一致，尚非独立因果证明。')

# Page 5: complete PowerData raw dictionary.
heading('PowerData原始字段分别表示什么',True)
para('主数据为2019年5月31天、10个cell、57个供电域，其中55个传统PDU、2个MVPP。57份功率CSV加1份机器映射CSV，各自GZIP压缩；每域8928条5分钟记录，共508896条。原始raw字节保持不变，csv目录仅解压。')
power=[('time','UINT64 微秒','5分钟平均功率窗口的起点；相对trace时间，不是Unix时间戳。600000000对应2019年5月1日00时PT。','排序 聚合 定位日期'),('cell','字符串','匿名集群标识，如a。不能据此推断真实地理位置。','分组与元数据'),('pdu','字符串','匿名供电域标识，如pdu6或mvpp1；与cell共同确定一条序列。','序列识别与设备类型'),('measured_power_util','小数 利用率','该5分钟窗口总功率除以本域容量，包含机房冷却等非IT设备。0.75表示约75%容量利用。','预测目标与历史输入'),('production_power_util','小数 利用率','归于生产工作负载的估算功率；官方用CPU利用率及机器闲忙功率等估计，含非IT功率分摊。不是实测CPU或NCU。','local24历史代理特征'),('bad_measurement_data','布尔值','总功率质量标记。本研究按官方Notebook处理：true为坏点，false可用，并同时检查数值有限。','评分掩码与聚合'),('bad_production_power_data','布尔值','生产功率质量标记，同样true为坏点；不自动等同于总功率无效。','生产代理独立处理')]
grid(['原始字段','类型与单位','含义','实际用途'],power,[2.15,1.25,5.3,width-8.7])
sub('机器到供电域映射表的三个原始字段')
mapping=[('machine_id','匿名机器ID，按cell和machine_id连接，避免跨cell同名ID混淆。'),('pdu','该机器所属供电域。'),('cell','该机器所在cell；映射仅覆盖传统架构a至h，不覆盖i和j的MVPP。')]
grid(['字段','含义与用法'],mapping,[2.15,width-2.15])
para('映射有96616行、96614个不同的cell与machine_id组合；2个组合对应多个PDU。映射没有生效时间，不能据此恢复机器迁移轨迹。原官方PDF对bad标记的文字存在反向描述，本研究以官方Notebook的NOT bad筛选和代码为准。',size=10)

# Page 6: actual rather than hypothetical imputation.
heading('功率原始数据实际做了哪些处理',True)
steps=[
('1 读入与排序','解压后按cell与pdu识别序列，按time升序；检查时间唯一、5分钟等间隔及57域对齐。原始行序并非时间序，raw文件不改写。'),
('2 质量筛选','总功率和生产功率分别按自身bad标记及数值有限性筛选。没有按测试结果删异常，也没有把缺失统一改成0。'),
('3 五分钟聚合到十五分钟','连续3条记录合为一块，取有效点平均，标签保留块起点；每天288点变96点。'),
('4 输入与评分分开','部分坏点用剩余好点均值作历史输入，target_valid仅在3点均有效时为true；坏块不评分。'),
('5 整块缺失的处理分支','若3点全坏，代码取块前最后12个有效原始点中位数；本次0块触发。'),
('6 生产代理补值','为空块另建production_causal_median，用块前最后12个好点中位数；无历史保留NaN，特征层再按前一天同刻向前补。原production数组不覆盖。'),
('7 单位与因果边界','目标保持官方利用率，未统一再做Min-Max或换成kW；成员内部处理由原生实现决定。外部补值和特征只读已完成历史。')]
for title,text in steps:para(title+'。'+text)
rows=[['总功率坏点','24 / 508896 原始点'],['含坏点的15分钟块','24块 全部为部分坏点'],['整块总功率中位数补值','0块 实际没有触发'],['生产功率坏点','116517原始点 约22.90%'],['生产功率聚合为空','38556个15分钟块'],['生产中位数代理实际补值','20452块 另有18104块无历史仍为NaN'],['最终功率张量与有效标签','57×31×96 共169632点 有效169608点']]
grid(['实际处理统计','本地核对结果'],rows,[4.2,width-4.2])
para('示例：一块有两个好点0.70、0.72和一个坏点，历史输入为0.71，评分target_valid仍为false。与此前方案的有效真值和评分掩码完全相同；24个无效上下文块的处理不同，最大差约0.005167。',size=10)

# Page 7: derived variables and split.
heading('处理后的字段怎样进入模型',True)
para('基础模型读历史实测功率，local24构造64项历史特征。下表为派生量，完整名称与解释见本报告证据目录的字段字典CSV。')
features=[['同刻滞后与滚动统计','lag1/2/3/7/14 med3/5/7 mean3','过去不同天同一15分钟时刻的功率与近期均值中位数。'],['差分与相对比值','lag1_minus_med3等6项','历史变化与相对水平。源mean3_minus_mean7实际为mean3减med7，保留原命名并注明。'],['扩展统计与指数平均','run_mean run_med run_std ewm0.3/0.6等8项','只对当前预测日前已完成历史统计，估计长期水平及近期变化。'],['日级与时刻趋势','lvl1/3/7 lvla trend curv等12项','跨域有效点的历史日水平、趋势、曲率与该时刻的整体变化。'],['域自身历史水平','ser_mean_run及比值差分6项','各域各时刻的历史均值与偏离。'],['其他域与同cell信息','loo_lag1 loo_med3 cell_hist等6项','排除自身后的其他域摘要和同cell历史摘要。'],['生产功率估算代理','plag1 plag2 p_lvl p_ratio p_diff p_trend','只读历史production代理，形成1/2天滞后、占比和差值。'],['时间相位与编号','dow_sin/cos is_weekend slot_sin/cos1至3 series_id cell_id','固定7日相位、日内周期与匿名编号；源7日偏移不是准确公历星期。']]
grid(['特征组','字段示例','实际含义'],features,[2.35,4.05,width-6.4])
para('Toto实际分组用cell与is_mvpp。内部域编号、映射机器数和PDU计数属于派生元数据，不是实测工作负载；完整定义见字段字典。')
sub('训练开发测试的时间边界')
rows=[['local24训练','开发拟合只到第14天 测试最终拟合只到第22天'],['24小时开发与测试','开发在14/16/18/20天末起报；测试22至30天末起报，预测23至31天'],['7天开发与测试','开发14/15天末起报；测试22/23/24天末起报，预测23至29、24至30、25至31天'],['融合权重与预测区间','权重开发选择后冻结；经验80%区间仅用开发残差10/90分位校准']]
grid(['阶段','已执行的规则'],rows,[2.35,width-2.35])
para('24小时输出96点，7天输出672点，日均每96点取平均。源local24首日残差使用训练前缀均值的问题保留披露，未称为严格逐日起报因果处理。',size=10)

# Page 8: fields of the selected ClusterData export and complete machine-event leaf schema.
heading('ClusterData本地字段及用途',True)
para('它与PowerData是配套但不同的数据。instance_usage仅完成cell-a的12085个Parquet必要列，真实CPU内存覆盖a的5个PDU；机器事件来自a至h，可形成传统55域资源状态。以下字段用于此前CPU资源对照，未额外输入本轮HDSF。')
usage_fields=[('start_time','微秒','实例使用统计窗口开始时间。'),('end_time','微秒','窗口结束时间，只把已完成窗口纳入历史。'),('machine_id','匿名ID','结合cell，通过机器映射归属到PDU。'),('collection_type','整数','0为JOB，1为ALLOC SET；本次只保留JOB，避免把两类叠加。'),('average_usage.cpus','NCU使用率','实例窗口内平均CPU使用，NCU为统一缩放后的计算单位，不是CPU功率。'),('average_usage.memory','归一化内存','按全trace最大机器内存容量缩放的内存量，不能当GB。'),('sample_rate','次每秒','窗口内平均采样频率，名义目标约1Hz；用于采样覆盖诊断。')]
sub('实例使用记录实际投影的七个字段')
grid(['原始字段','单位','含义与处理'],usage_fields,[3.0,1.6,width-4.6])
machine_fields=[('time','微秒','机器状态事件的时间；0可表示trace开始前已存在的机器。'),('machine_id','匿名ID','连接映射，按各cell分别维护状态。'),('type','事件枚举','本地导出1为ADD、2为REMOVE、3为UPDATE，用于维护可用状态。'),('capacity.cpus','NCU','机器CPU容量，不是实际CPU使用率。'),('capacity.memory','归一化内存','机器内存容量，与使用量区分。'),('missing_data_reason','原因枚举','记录事件信息缺失原因；本地统计非0的比例。'),('platform_id','匿名字符串','机器微架构和芯片组类别，本次资源聚合未使用。'),('switch_id','匿名字符串','机器连接的交换机标识，本次资源聚合未使用。')]
sub('机器事件JSON的全部叶字段')
grid(['原始字段','单位','含义与处理'],machine_fields,[3.0,1.6,width-4.6])

# Page 9: actual CPU/resource transformation definitions.
heading('CPU与资源记录实际怎样聚合',True)
para('按cell和machine_id连接PDU，排除归属歧义；只保留JOB、可映射机器、有效非负CPU/内存、0至300秒窗口和月内时间。每个窗口按结束时间分配到15分钟区间，CPU和内存乘窗口秒数再除900后求和。未按区间真实重叠长度拆窗，可能有不足5分钟对齐偏差。')
usage_defs=[('completed_job_cpu_NCU','该区间JOB的平均CPU使用加权和，即逐窗CPU×时长/900求和。PDU总量可大于1，不能当kW。'),('completed_job_memory_normalized','同样按时长加权的JOB归一化内存使用总量，不能直接换算GB。'),('job_window_equivalent','窗口时长/900之和，表示JOB窗口等效量；不是CPU方差或独立任务数。'),('job_window_count','纳入该区间的JOB窗口记录数；同一任务可出现多窗，不是去重任务数量。'),('sample_rate_duration_sum','采样频率乘时长后求和，用于估计采样覆盖。'),('duration_seconds_sum','纳入窗口的时长秒数总和；与上一项相除可得到平均采样频率。')]
grid(['CPU聚合后的字段','实际定义'],usage_defs,[4.1,width-4.1])
para('机器事件按time排序，各cell单独维护状态，只处理区间终点之前的事件并向前保持状态；REMOVE令不可用，ADD/UPDATE更新容量。缺容量时保留既有已知容量。没有拿未来事件回填过去。以下比例均以可确定归属的映射机器数为分母。')
resource_defs=[('mapped_machine_observed_fraction','已出现可识别状态事件的映射机器比例。'),('active_fraction','当前处于可用状态的映射机器比例。'),('mean_cpu_capacity_NCU','活跃且容量已知机器的CPU容量之和，除全部映射机器数。'),('mean_memory_capacity_normalized','活跃且容量已知机器的内存容量之和，除全部映射机器数。'),('known_cpu_capacity_fraction','活跃且CPU容量已知的机器比例。'),('known_memory_capacity_fraction','活跃且内存容量已知的机器比例。'),('state_missing_reason_fraction','当前状态带非0缺失原因的机器比例。'),('mapping_available','该域是否有可确定归属的机器映射，取0或1。')]
grid(['资源聚合后的字段','实际定义'],resource_defs,[4.1,width-4.1])
para('真实CPU只覆盖a五域，其他52域的CPU数组0为缺失占位，不是实测CPU为0；资源数据不覆盖i/j两个MVPP。JOB-only遗漏alloc层额外开销与未监控工作负载；固定映射无生效日期。上述限制会影响算电协变量研究，不能据此把功率与CPU建立完整物理守恒。',size=10)

# Page 10: definitions, reproducibility and sources.
heading('统一评价与研究结论的边界',True)
sub('与9月28日评价保持一致的计算方式')
para('曲线WAPE：每个PDU在有效点上，绝对误差总和除以真实值绝对总和并乘100%，再对57域等权平均。24小时用9个起报，7天用3个起报；项目原公布的全部点池化WAPE只作独立复现核对，不替换本报告分数。')
para('MAE、RMSE、APE95及2%/5%命中使用相同有效点。每日均值要求96点均有效，完整日记录为24小时508条、7天1187条；重复目标日不是独立样本。')
para('峰值MAPE每天在有效点中取真值和预测最大值比较，峰时取首个最大值。日均利用率不是日电量kWh；历史基线为前一天同刻与前3天同刻均值。')
para('经验80%区间仅由开发残差10/90分位校准。HDSF覆盖率为24小时75.568%、7天76.009%，未达到80%；宽度与Winkler80仍按统一定义计算。')
sub('结果应怎样理解')
para('开发期只有4个24小时起报、2个7天起报，7天窗口重叠。日期配对bootstrap改善区间分别为−0.016550至+0.093224、−0.056384至+0.072188个百分点，均跨0。源设计明确接触过历史测试，本地同月测试也多轮查看；缺少真正独立的确认月份。')
para('非负共享权重提供参数限制，不能保证没有过拟合。原流程同时约束峰值、分组和偏差，HDSF主要优化池化绝对误差；统一评分相同，开发目标不同。两种历史无效值处理也有差异，所以跨旧新方案总增益不能全部归因为融合。')
para('本次仅核查并重排报告，冻结选择、部署、原始数据和评价实现保持原样。独立确认后才可判断分时专家与7天水平稳定化的泛化优势。')
sub('来源与复查位置')
para('Google官方字段说明：google-powerdata-2019/docs/power_trace_documentation.pdf；质量标记筛选：同目录power_trace_analysis_colab.ipynb。官方页面 https://github.com/google/cluster-data/blob/master/PowerData2019.md',size=9.5)
para('ClusterData官方字段文档：google-clusterdata-2019/docs/cluster_usage_traces_v3.pdf。HDSF固定源commit为648777507ebc57e7498ce61f4ae2236ab179e6a8，来源 https://github.com/daifeizhai/PowerData_Prediction',size=9.5)
para('完整字典与数据审计：experiment/round44/docs/causes_data_report_20261009；消融：experiment/round44/results/mechanism_analysis_20261008。未读取完整usage主键或做完整去重；必要列84.78GiB与原对象407.55GiB范围不同。',size=9.5)
para('数据及文档CC BY4.0，需署名并注明修改；TimesFM3权重限非商业非生产研究。D盘为实体目录，桌面为兼容入口，原三任务比较Word保留。',size=9.5)

for element in doc.part.element.xpath('//w:pPr/w:pBdr'):
    element.getparent().remove(element)
# Preserve the original PAGE field and only change footer text.
for t in doc.sections[0].footer._element.xpath('.//w:t'):
    if '2026' in (t.text or ''):t.text='2026 年 10 月 9 日  HDSF结果原因与数据处理 '
draft=D/'authoring_tmp.docx';doc.save(draft)
with zipfile.ZipFile(D/'reference.docx') as src,zipfile.ZipFile(draft) as revised,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as dst:
    allowed={'word/document.xml','word/footer1.xml'}
    for item in src.infolist():dst.writestr(deepcopy(item),revised.read(item.filename) if item.filename in allowed else src.read(item.filename))
with zipfile.ZipFile(D/'reference.docx') as src,zipfile.ZipFile(output) as dst:
    assert src.namelist()==dst.namelist()
    assert all(src.read(k)==dst.read(k) for k in src.namelist() if k not in allowed)
assert sha(Path(ev['reference']))==ev['reference_sha256']
dictionary=[]
for name,unit,meaning,use in power:dictionary.append(dict(table='PowerData power CSV',field=name,unit=unit,meaning=meaning,used=use))
for name,meaning in mapping:dictionary.append(dict(table='PowerData machine mapping',field=name,unit='identifier',meaning=meaning,used='metadata or historical ClusterData join'))
for name,unit,meaning in usage_fields:dictionary.append(dict(table='ClusterData selected instance_usage',field=name,unit=unit,meaning=meaning,used='previous CPU contrast; not HDSF'))
for name,unit,meaning in machine_fields:dictionary.append(dict(table='ClusterData machine_events',field=name,unit=unit,meaning=meaning,used='previous resource contrast; platform/switch not used'))
for name,meaning in usage_defs+resource_defs:dictionary.append(dict(table='derived ClusterData',field=name,unit='see meaning',meaning=meaning,used='previous CPU/resource contrast; not HDSF'))
local_meaning={
'lag1':'前1天同刻功率','lag2':'前2天同刻功率','lag3':'前3天同刻功率','lag7':'前7天同刻功率','lag14':'前14天同刻功率','med3':'前3天同刻中位数','mean3':'前3天同刻均值','med5':'前5天同刻中位数','med7':'前7天同刻中位数',
'lag1_minus_med3':'lag1减med3','lag1_over_med3':'lag1除以med3，分母至少1e-6','med3_minus_med7':'med3减med7','mean3_minus_mean7':'实际代码为mean3减med7，名称与实际中位数存在偏差','lag1_minus_lag2':'lag1减lag2','lag7_minus_lag14':'lag7减lag14',
'run_mean':'预测日前各天同刻的扩展均值','run_med':'预测日前各天同刻的扩展中位数','run_std':'预测日前各天同刻的扩展标准差','lag1_minus_run_mean':'lag1减run_mean','lag1_over_run_mean':'lag1除以run_mean','run_std_over_run_mean':'run_std除以run_mean','ewm0.3':'按alpha0.3递推的同刻历史指数均值','ewm0.6':'按alpha0.6递推的同刻历史指数均值',
'lvl1':'前1天跨57域有效点功率均值','lvl3':'前3天跨57域有效点功率均值','lvl7':'前7天跨57域有效点功率均值','lvla':'全部已完成历史跨域有效点均值','lvl_trend1':'lvl1减lvl3','lvl_trend2':'lvl3减lvl7','lvl_trend3':'lvl7减lvla','lvl_curv':'lvl1减lvl3再减lvl7减lvla','lvl_z':'lvl1减lvla再除以abs(lvla)','lvl_ratio':'lvl1除以lvla','slot_trend':'各域当前时刻平均曲线的前1天减前7天摘要再除前7天摘要','slot_dev':'各域当前时刻平均曲线的前1天摘要减全部历史摘要',
'ser_mean_run':'该域该时刻全部过去天的均值','ser_mean_lag1':'ser_mean_run滞后1天','ser_mean_diff':'ser_mean_lag1减ser_mean_run','lag1_minus_sermean':'lag1减ser_mean_run','lag1_over_sermean':'lag1除以ser_mean_run','sermean_over_lvla':'ser_mean_run除以跨域历史lvla',
'loo_lag1':'除自身外其他域lag1的均值','loo_med3':'除自身外其他域med3的均值','series_over_loo_lag1':'lag1除以loo_lag1','loo_trend':'其他域lag1均值减其他域lag2均值','cell_hist':'同cell各域ser_mean_run的均值','sermean_over_cell':'ser_mean_run除以cell_hist',
'plag1':'历史生产功率代理滞后1天','plag2':'历史生产功率代理滞后2天','p_lvl':'预测日前生产功率代理全历史均值','p_ratio':'plag1除以lag1','p_diff':'plag1减lag1','p_trend':'plag1减plag2','dow_sin':'固定偏移3的7日相位正弦，非准确公历星期','dow_cos':'同7日相位余弦','is_weekend':'源固定相位下5或6的指标，不能直接称真实周末','series_id':'内部域序号','cell_id':'内部cell类别序号'}
for k in (1,2,3):
    local_meaning[f'slot_sin{k}']=f'日内96时刻第{k}谐波正弦';local_meaning[f'slot_cos{k}']=f'日内96时刻第{k}谐波余弦'
assert set(local_meaning)==set(a['derived_feature_names'])
for name in a['derived_feature_names']:dictionary.append(dict(table='derived local24 feature',field=name,unit='power ratio or encoded scalar',meaning=local_meaning[name],used='historical only; LightGBM accepts remaining NaN'))
for name,meaning in [('series_key','cell和pdu组成的序列键'),('series_id','自然文件顺序的内部序号'),('cell','原匿名cell'),('pdu','原匿名供电域'),('machine_count','映射表按域统计的不同机器ID数，元数据统计'),('is_mvpp','pdu名称是否以mvpp开头，设备类型'),('cell_pdu_count','同cell传统非MVPP供电域数')]:dictionary.append(dict(table='derived metadata',field=name,unit='identifier or count',meaning=meaning,used='metadata; cell/type grouping'))
with (D/'数据字段完整说明.csv').open('w',encoding='utf-8-sig',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=['table','field','unit','meaning','used']);writer.writeheader();writer.writerows(dictionary)
save(D/'authoring_log.json',dict(output=str(output),output_sha256=sha(output),reference_sha256=ev['reference_sha256'],headings=headings,tables=tables,
    intended_pages=10,field_dictionary_rows=len(dictionary),preserve_only_parts_unchanged=True,allowed_parts=sorted(allowed),source_data_audit_sha256=sha(D/'data_audit.json'),comparison_sha256=sha(Q/'hdsf_comparison.json'),mechanism_sha256=sha(Q/'mechanism_analysis_20261008/analysis.json')))
print(json.dumps(dict(created=True,output=str(output),sha256=sha(output),tables=len(tables),sections=len(doc.sections),planned_pages=10,field_dictionary_rows=len(dictionary)),ensure_ascii=True))
