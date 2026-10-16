from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from copy import deepcopy
import sys,json,hashlib
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree
T=Path(sys.argv[1]);B=T/'experiment';R=B/'round44';W=R/'docs/integrated_template';Q=R/'results';source=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx';ref=W/'reference.docx';out=T/'Google PowerData 2019 三项预测任务与模型比较汇报_加入HDSF.docx'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
evidence=read(W/'template_evidence.json');assert sha(source)==sha(ref)==evidence['source_sha256'];assert (W/'artifact.md').exists();a=read(Q/'hdsf_comparison.json');assert len(a['rows'])==142 and len(a['unique_rows'])==114 and read(Q/'final_delivery_verification.json')['all_passed'];assert sha(B/'round29/code/aligned_metrics.py')==a['reference_audit']['canonical_code_sha256']
if '--validate-only' in sys.argv:print('Originaltemplate unchanged; HDSFactual metrics and142/114 comparison pool verified');sys.exit(0)
doc=Document(ref);pp=doc.paragraphs;tables=doc.tables;changes=[]
def paragraph_text(p,text,bold=None):
 rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 for child in list(p._p):
  if child.tag!=qn('w:pPr'):p._p.remove(child)
 run=p.add_run(text)
 if rpr is not None:run._r.insert(0,rpr)
 if bold is not None:run.bold=bold
def para(i,text):paragraph_text(pp[i],text);changes.append('paragraph '+str(i))
def set_cell(cell,text,*,bold=False,fill='FFFFFF',alignment=None):
 for extra in cell.paragraphs[1:]:cell._tc.remove(extra._p)
 p=cell.paragraphs[0];paragraph_text(p,str(text),bold)
 if alignment is not None:p.alignment=alignment
 pr=cell._tc.get_or_add_tcPr();shd=pr.find(qn('w:shd'))
 if shd is None:shd=OxmlElement('w:shd');pr.append(shd)
 shd.set(qn('w:fill'),fill)
def widths(table):return [int(c.get(qn('w:w')))/1440 for c in table._tbl.tblGrid]
def rebuild(idx,headers,rows,ww=None,marks=None):
 table=tables[idx];ww=ww or widths(table);assert len(headers)==len(ww);hdr=deepcopy(table.rows[0]._tr);body=deepcopy(table.rows[1]._tr)
 for tr in list(table._tbl.tr_lst):table._tbl.remove(tr)
 grid=table._tbl.tblGrid
 for g in list(grid):grid.remove(g)
 for w in ww:g=OxmlElement('w:gridCol');g.set(qn('w:w'),str(round(w*1440)));grid.append(g)
 for rid,values in enumerate([headers,*rows]):
  tr=deepcopy(hdr if rid==0 else body);cells=tr.findall(qn('w:tc'))
  while len(cells)<len(headers):new=deepcopy(cells[-1]);tr.append(new);cells.append(new)
  while len(cells)>len(headers):tr.remove(cells.pop())
  table._tbl.append(tr);row=table.rows[-1]
  for j,(cell,value) in enumerate(zip(row.cells,values)):
   winning=rid>0 and marks and (rid-1,j) in marks;fill='DCEAF2' if rid==0 else'E9F4E5' if winning else'F5F7F9' if rid%2==0 else'FFFFFF';set_cell(cell,value,bold=bool(rid==0 or winning),fill=fill,alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER);cell._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(round(ww[j]*1440)))
 changes.append('table '+str(idx))
def rows(H):return [r for r in a['unique_rows'] if r['horizon_days']==H]
def get(H,key):return next(r for r in rows(H) if key==r['model'] or key in r['aliases'])
def current(H):return next(r for r in rows(H) if r['status']=='当前综合推荐')
def hdsf(H):return get(H,f'r44_hdsf_H{H}')
def n(v):return f'{v:.6f}'
def star(H,k,v):return abs(v-(max if 'hit' in k or k.startswith('relative') else min)(r[k] for r in rows(H)))<1e-10
def rank(H,k,v):return 1+sum(r[k]<v-1e-10 for r in rows(H))
def label(r):
 if r['model'].startswith('r44_hdsf'):return 'HDSF半日单纯形融合\nD14/D22固定训练'
 if r['model']=='r32_selected_H1':return '24小时当前推荐\nChronos2 TiRex2 TimesFM3 Ridge TTM'
 if r['model']=='r36_selected_H7':return '7天当前推荐\nMoirai2同场站10% 全域日均协调'
 return r['report_name']
para(0,'Google PowerData 2019 三项预测任务与模型比较 加入HDSF')
para(1,'HDSF半日单纯形融合已加入原模型对照。按9月28日统一评价，24小时曲线WAPE2.335113%、7天曲线2.701623%、7天每日均值1.868830%。24小时曲线与7天日均为全部实测最低；7天曲线排第6，尚未实现同一方案全面或大幅领跑。')
summary=[]
for title,H,k in [('未来24小时完整曲线',1,'wape'),('未来7天完整曲线',7,'wape'),('未来7天每日平均负荷',7,'daily_wape')]:
 cur=current(H);hh=hdsf(H);win=min(rows(H),key=lambda r:r[k]);summary.append([title,n(cur[k]),n(hh[k]),n(win[k]),label(win),f'{100*(1-win[k]/cur[k]):.3f}%'])
rebuild(0,['任务','当前推荐\nWAPE %','HDSF\nWAPE %','全部实测最低\nWAPE %','最低值的明确模型','最低值相对当前\n误差减少'],summary,[1.55,1.10,1.10,1.25,4.05,1.30])
para(2,'本报告保留原有模型和研究章节，增加HDSF及其成员实测记录后的完整CSV共142条，合并相同预测为114组。正文对照表展示原报告模型与新增HDSF，★及排名均按114组计算。HDSF是测试前已开发冻结的复现候选；正式部署仍保留第32轮与第36轮，未按测试重选参数或替换入口。')
para(3,'模型名称说明：Moirai2为Salesforce Moirai 2.0 R small；TTM为IBM Granite TinyTimeMixer R2；HGB为HistGradientBoosting。HDSF使用Chronos-2长历史、TiRex-2、TimesFM2.5、Toto2.0 313M、TimesFM3与局部成员按12小时分段拟合非负且和为1的权重；7天另含IBM PatchTST-FM-r2。开发A/B比例分别为0.075与1.0，D14/D22固定训练。')
para(4,pp[4].text.replace('当前方案是后续四轮的固定参照。','当前方案作为后续研究的固定参照。'))
para(6,'HDSF与当前推荐及已有研究候选的多维比较')
fields=[('wape','曲线WAPE %↓'),('daily_wape','日均WAPE %↓'),('mae','曲线MAE↓'),('rmse','曲线RMSE↓'),('peak_mape','日最大值MAPE %↓'),('p95_ape','APE95分位 %↓'),('hit2_percent','APE≤2%占比 %↑'),('hit5_percent','APE≤5%占比 %↑'),('daily_mae','日均MAE↓'),('daily_rmse','日均RMSE↓'),('daily_hit2_percent','日均APE≤2%占比 %↑'),('daily_hit5_percent','日均APE≤5%占比 %↑'),('cella5_wape','a5曲线WAPE %↓'),('other52_wape','其余52域曲线WAPE %↓'),('absolute_bias_percent','绝对总偏差 %↓'),('peak_time_mae_minutes','峰值时刻误差 分钟↓')]
cases=[current(1),get(1,'r37_diagnostic_H1'),hdsf(1),current(7),get(7,'r40_diagnostic_H7'),hdsf(7)];matrix=[];marks=set()
for i,(k,title) in enumerate(fields):
 matrix.append([title,*[n(r[k]) for r in cases]])
 for j,r in enumerate(cases,1):
  if j in [1,4]:continue
  cur=current(r['horizon_days']);better=r[k]>cur[k]+1e-10 if 'hit' in k else r[k]<cur[k]-1e-10
  if better:marks.add((i,j))
rebuild(1,['统一指标','24h\n当前推荐','此前24h候选\nMoirai2小时Rank8','HDSF\n24小时','7天\n当前推荐','7天曲线最低\nMoirai2三小时CPU内存','HDSF\n7天'],matrix,[2.45,1.18,1.46,1.18,1.18,1.72,1.18],marks)
para(7,'浅绿色表示相对同跨度当前推荐的改善。HDSF两条曲线及日均、MAE和RMSE均改善；24小时峰值MAPE增加0.186307个百分点，a5及绝对偏差变差。7天峰值略改善，但P95、2%命中、a5及绝对偏差变差。各指标分别比较，不把单项最低称为全面最优。')
b1=a['bootstrap']['1'];b7=a['bootstrap']['7'];para(8,f'HDSF曲线相对当前推荐分别减少{n(b1["gain_pp"])}与{n(b7["gain_pp"])}个百分点。按起报日期重采样的探索95%区间分别为[{n(b1["paired_origin_exploratory_95_gain_pp"][0])}, {n(b1["paired_origin_exploratory_95_gain_pp"][1])}]与[{n(b7["paired_origin_exploratory_95_gain_pp"][0])}, {n(b7["paired_origin_exploratory_95_gain_pp"][1])}]，均跨零；同月TEST多轮查看且7天窗口重叠，缺少独立确认。')
set_cell(tables[2].rows[2].cells[1],'三条5分钟记录聚合为15分钟；评分无效目标排除。HDSF历史异常块取有效原始点均值，旧输入前向填补，最大输入差0.005167；有效评分真值与掩码完全相等。',fill='F5F7F9',alignment=WD_ALIGN_PARAGRAPH.CENTER);changes.append('table2 row2 input preprocessing disclosure')
para(13,'★及浅绿色表示114组不同预测中的该项最低或最高，含未采用候选和HDSF成员；完整142条记录保存在顶层CSV。各模型按相同起报、真值和质量掩码评分。HDSF开发WAPE使用拟合融合权重的样本，不代表独立验证；原项目公布的池化成绩不参与排名。')
fcurve=['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae','rmse','peak_mape','p95_ape','hit2_percent','hit5_percent'];keys1=['r32_selected_H1','r44_hdsf_H1','last_source','ridge','lgb','chronos','tirex','tf25','equal','r37_diagnostic_H1','r41_diagnostic_H1','r42_diagnostic_H1','r43_diagnostic_H1'];keys7=['r36_selected_H7','r44_hdsf_H7','r31_selected_H7','mean3_source','chronos','tirex','tf25_7','r38_diagnostic_H7','r40_diagnostic_H7','r41_diagnostic_H7','r42_diagnostic_H7','r43_diagnostic_H7']
def metric_table(index,H,keys,fields):
 rr=[];marks=set()
 for i,key in enumerate(keys):
  r=get(H,key);values=[label(r)]
  for j,k in enumerate(fields,1):
   v=r.get(k);win=v is not None and k!='development_wape_original_grid' and star(H,k,v);values.append('—' if v is None else n(v)+(' ★' if win else''))
   if win:marks.add((i,j))
  rr.append(values)
 ww=widths(tables[index])
 if index==4:ww[0]-=.05;ww[8]+=.05
 rebuild(index,[c.text for c in tables[index].rows[0].cells],rr,ww,marks)
metric_table(4,1,keys1,fcurve);metric_table(5,7,keys7,fcurve);metric_table(6,7,keys7,['daily_wape','relative_daily_error_reduction_percent','daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent'])
para(14,'24小时曲线最低更新为HDSF半日单纯形融合2.335113%，此前小时同场站Rank8候选为2.367483%。HDSF也改善MAE、RMSE、P95及2%/5%命中；峰值MAPE2.053383%高于当前1.867076%，不能称为每项最优。')
para(17,'7天曲线最低仍为三小时57域功率加真实a5 CPU内存的10%修正2.694096%。HDSF为2.701623%，低于当前推荐2.706501%，在58组7天预测中排第6；日均更优不等于7天曲线最优。')
para(19,'每日均值由同一672点曲线按每天96点计算，评分排除不完整有效日。最低值更新为HDSF1.868830%，此前季节残差10%候选为1.889163%，当前推荐1.904238%。HDSF也改善日均MAE/RMSE，但日均2%及5%命中仍未同时领跑。')
para(26,'上表为旧Moirai匹配CPU对照：同骨干、14天历史、3小时粒度，表内固定25%无日均协调；最终旧诊断为10%全域日均协调。开发略有改善，未形成独立确认增益。HDSF本次没有额外使用ClusterData CPU内存，不能将其改善归因于这些辅助通道。')
para(27,'官方ClusterData2019仅cell-a的12085个Parquet必要列已完整读取/审计，剩余下载0。真实CPU内存只覆盖a的5个PDU，共10条历史辅助通道；其余52域未伪造CPU。CPU单位NCU，功率为归一化利用率；作业去重与时间对齐限制保留，job-window equivalent不是CPU方差。')
groups=[]
for H,keys in [(1,['r32_selected_H1','r37_diagnostic_H1','r44_hdsf_H1']),(7,['r36_selected_H7','r40_diagnostic_H7','r44_hdsf_H7'])]:
 for key in keys:
  r=get(H,key);groups.append([f'{H}天 '+label(r),n(r['cella5_wape']),n(r['other52_wape']),f'{r["improved_pdus"]}/57',f'{r["improved_cells"]}/10',f'{r["improved_origins"]}/{9 if H==1 else 3}'])
grouphead=[c.text for c in tables[9].rows[0].cells];grouphead[0]='当前推荐与研究候选';rebuild(9,grouphead,groups)
para(29,'数据CC BY4.0；Moirai权重CC BY-NC4.0、代码Apache2.0；TimesFM3仅限非商业非生产研究；IBM PatchTST-FM-r2可选Apache2.0/OpenMDW1.0。来源、许可及哈希保留在round29/35/44。')
intervals=[[c.text for c in row.cells] for row in tables[10].rows[1:]]
for H in [1,7]:
 v=next(v for v in a['new_intervals'] if v['horizon_days']==H and v['method'].startswith('HDSF'));intervals.insert(1 if H==1 else 3,[f'{H}天 HDSF半日单纯形融合\n开发残差10/90经验80%区间',80,f'{v["coverage_percent"]:.3f}',n(v['mean_width_normalized']),n(v['winkler80'])])
rebuild(10,[c.text for c in tables[10].rows[0].cells],intervals)
para(31,'当前24小时沿用第33轮开发选定对称区间；当前7天、旧候选及HDSF用冻结开发残差10/90经验80%区间，不按测试校准。HDSF覆盖率75.568%及76.009%，未达到80%；附加经验区间不是HDSF原生概率输出，覆盖、宽度和Winkler不同时最优。')
para(32,'结论：HDSF在统一评价下使24小时曲线和7天日均达到最低，7天曲线仅排第6，部分峰值与分组指标退步，尚非单一方案全面或大幅领先。原三轮共同因子、注意力及变化量的负面结果保留。31天同月TEST多轮查看，缺少独立确认；正式推荐与旧选择记录不因测试排名改写。HDSF的75项科学及96项交付核查通过，PatchTST固定分批及原D1残差问题均已披露，详细复现见本目录HDSF专项汇报。')
footer=doc.sections[0].footer.paragraphs[0]
for run in footer.runs:
 if '2026年10月7日' in run.text:run.text=run.text.replace('2026年10月7日','2026年10月8日').replace('统一评价与新方法实测','统一评价与HDSF模型比较')
original=ZipFile(ref);replace={'word/document.xml':etree.tostring(doc._element,xml_declaration=True,encoding='UTF-8',standalone=True),'word/footer1.xml':etree.tostring(doc.sections[0].footer._element,xml_declaration=True,encoding='UTF-8',standalone=True)}
with ZipFile(out,'w',ZIP_DEFLATED) as target:
 for item in original.infolist():target.writestr(deepcopy(item),replace.get(item.filename,original.read(item.filename)))
assert sha(source)==evidence['source_sha256'];check=ZipFile(out);assert set(check.namelist())==set(original.namelist());assert all(check.read(n)==original.read(n) for n in original.namelist() if n not in replace)
(W/'authoring_log.json').write_text(json.dumps(dict(output=str(out),output_sha256=sha(out),source_unchanged=True,modified_parts=list(replace),preserve_only_parts_unchanged=True,edits=changes,data_sha256=sha(Q/'hdsf_comparison.json'),table_model_keys={'4':keys1,'5':keys7,'6':keys7},comparison_pool=142,unique_predictions=114),ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(dict(output=str(out),sha256=sha(out),preserved_parts=len(original.namelist())-2),ensure_ascii=False))
