from pathlib import Path
from zipfile import ZipFile
import json,hashlib,sys,datetime
from docx import Document
from lxml import etree
T=Path(sys.argv[1]);R=T/'experiment/round44';W=R/'docs/integrated_template';W.mkdir(parents=True,exist_ok=True);source=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx';ref=W/'reference.docx';sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest();expected='1a65995d312f2f53b441342283c33e917f1abb94a37270e44dc4641a97a409a4';assert sha(source)==expected;ref.write_bytes(source.read_bytes());d=Document(ref);z=ZipFile(ref)
parts=[dict(path=n,bytes=len(z.read(n)),sha256=hashlib.sha256(z.read(n)).hexdigest(),scope='editable' if n in ['word/document.xml','word/footer1.xml'] else'preserve_only') for n in z.namelist()]
record=dict(source=str(source),retained_reference=str(ref),source_sha256=expected,pages=9,section_count=1,source_QA=str(T/'experiment/round43/docs/report_qa/verification.json'),all_reference_pages_visually_inspected=True,parts=parts,paragraph_slots=[dict(index=i,style=p.style.name,text=p.text) for i,p in enumerate(d.paragraphs)],table_slots=[dict(index=i,rows=len(t.rows),columns=len(t.columns),headers=[c.text for c in t.rows[0].cells]) for i,t in enumerate(d.tables)],section_xml=etree.tostring(d.sections[0]._sectPr).decode(),field_scope='Footer PAGE only; Word readonly export refreshes displayed page numbers; field element preserved')
(W/'template_evidence.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
contract='''# HDSF加入原三项任务报告的模板执行约定

## 权威参考与固定证据
原文档和本目录reference.docx SHA256均为1a65995d312f2f53b441342283c33e917f1abb94a37270e44dc4641a97a409a4。原9页PNG、PDF及其哈希见experiment/round43/docs/report_qa/verification.json；本次已实际查看全部9页，复用与当前原DOCX哈希一致的缓存渲染。原文档与参考副本不改写。

## 页面与字体
单节，实际宽11.6903英寸、高8.2701英寸；L/R0.65、T0.55、B0.50英寸。原方向标记0保持，视觉为横向页面。保持原sectPr、页眉页脚距离、分页结构。Normal Calibri/微软雅黑10.5pt，1.1行距，段后7pt；Title21pt，Heading1 15pt、段后8pt、显式分页。全部黑色标题，原蓝灰表头与浅灰交替行，E9F4E5用于最优/改善标记。保持全部字体与字号，不为塞内容缩小。

## 表格与段落插槽
稳定定位为word/document.xml正文中的第0至33个非表格段落以及第0至10个表格，明细与完整原文快照在template_evidence.json。修改P0标题版名，P1/2总论与142/114记录范围，P3简明HDSF构成，P6/7/8多维比较说明，P13/14/17/19最优/任务结论，P26/27/29必要来源简述，P31/32最终区间与结论。其余段落保留，允许为添加HDSF所需的上下文措辞修正。正文原9章顺序不变，不另造布局。
T0总结增加HDSF列；T1多维矩阵增加两个跨度HDSF列，保留原当前/原研究列；T4/5/6三项任务对照各增加HDSF行，保留全部原模型；T9分组和T10区间各增加HDSF两行；T2补充HDSF历史异常输入差异并明确评分真值/掩码相等；T3原文历史锚点、T7三轮训练、T8 CPU匹配对照原样保留。新表按来源的行、单元格、段落及运行属性复制，不使用新风格。更新全部★/浅绿色到114组全局排名；多维矩阵绿色仍为对旧综合推荐改善。

## 包与关系保留
仅word/document.xml和word/footer1.xml允许改动；styles/stylesWithEffects、settings、theme、numbering、全部.rels、customXml、metadata、thumbnail及所有其他部件须逐字节相等。未发现正文SDT、图片、外链、脚注；仅页脚PAGE字段保留。用python-docx修改原本对象，随后只把允许的两个XML部件替换进原ZIP，其他原包部件直接复制字节。

## 容量与核查
保留原9页各章模式。新增总结列和多维列可重分配同一10.35英寸表宽，其他表列宽不变。内容超出页面时只简化新增/允许改写的说明或复用原分页模式，不静默改字号。查看每张终稿PNG，并对原/新页面做范围差异核对。核对原文档SHA、保留包部件、sectPr、原模型数值、HDSF所有统一指标、绿色/星标、表头和9章顺序；源原文档不变后才能交付。
'''
(W/'artifact.md').write_text(contract,encoding='utf-8');print('Source9pages inspected; templatecontract andpreserve-only package inventory saved')
