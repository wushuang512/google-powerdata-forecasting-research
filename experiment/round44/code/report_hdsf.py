from pathlib import Path
import sys,json,hashlib
T=Path(sys.argv[1]);B=T/'experiment';R=B/'round44';Q=R/'results'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
a=read(Q/'hdsf_comparison.json');s=a['hdsf_summary'];qa=read(Q/'science_verification.json');assert qa['all_passed'];assert hashlib.sha256((B/'round29/code/aligned_metrics.py').read_bytes()).hexdigest()==a['reference_audit']['canonical_code_sha256']
if '--validate-only' in sys.argv:print('Actual two horizons, immutable metrics, source andscientific verification complete');sys.exit(0)
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
doc=Document();sec=doc.sections[0];sec.page_width=Inches(11.69);sec.page_height=Inches(8.27);sec.top_margin=Inches(.55);sec.bottom_margin=Inches(.5);sec.left_margin=sec.right_margin=Inches(.65)
for k in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[k];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑');snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(snap)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.08;doc.styles['Normal'].paragraph_format.space_after=Pt(6);doc.styles['Title'].font.size=Pt(21);doc.styles['Heading 1'].font.size=Pt(15)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('2026年10月8日  HDSF复现与统一评价  ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
def p(t):return doc.add_paragraph(t)
def h(t):pp=doc.add_paragraph(t,'Heading 1');pp.paragraph_format.page_break_before=True
def table(headers,rows,widths,size=9.8):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w);t.rows[0].cells[j].text=str(headers[j])
 for rr in rows:
  cells=t.add_row().cells
  for j,v in enumerate(rr):cells[j].text=str(v)
 for i,row in enumerate(t.rows):
  tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
  if i==0:tr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
   cp.append(bd);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCEAF2' if i==0 else'F5F7F9' if i%2==0 else'FFFFFF');cp.append(shade)
   for pp in c.paragraphs:
    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=Pt(size*1.16)
    for run in pp.runs:run.font.size=Pt(size);run.bold=i==0
 return t
def n(v):return f'{v:.6f}'
short={'hdsf':'HDSF半日单纯形融合','chronos2_long':'Chronos-2 最多28天','tirex2':'TiRex-2 14天','timesfm25':'TimesFM2.5 200M 14天','toto2_joint':'Toto2.0 313M 同cell/设备类型联合14天','toto2_single':'Toto2.0 313M 单域14天','timesfm30':'TimesFM3.0 57域联合14天','timesfm30_long':'TimesFM3.0 57域联合最多21天','timesfm30_uni':'TimesFM3.0 单域最多21天','patchtstfm_long':'IBM PatchTST-FM-r2 385M 最多21天','local24':'四种子LightGBM残差50%＋加权滞后回归50%','local7':'水平/日内形状平滑（3/5/10天半衰期）'}
def prior(H):return next(r for r in a['unique_rows'] if r['horizon_days']==H and r['status']=='当前综合推荐')
def new(H):return s['scores'][str(H)]['hdsf']
def oldrows(H):return [r for r in a['rows'] if r['horizon_days']==H and not r['model'].startswith('r44_')]
def ranking(H,k):
 rr=[r for r in a['unique_rows'] if r['horizon_days']==H];return 1+sum(r[k]<new(H)[k]-1e-10 for r in rr),len(rr)
doc.add_paragraph('HDSF半日单纯形融合复现与统一评价','Title')
p('本次独立复现HDSF半日单纯形融合，三个预测任务与此前模型的成绩如下。训练保留原项目D14/D22配方；最终评分按9.28参考文档，每个供电域分别计算WAPE，再对57域等权平均。')
rr=[]
for H,k,title in [(1,'wape','未来24小时完整曲线'),(7,'wape','未来7天完整曲线'),(7,'daily_wape','未来7天每日平均负荷')]:
 cur=prior(H);best=min(oldrows(H),key=lambda r:r[k]);rank,total=ranking(H,k);rr.append([title,n(new(H)[k]),n(cur[k]),n(best[k]),n(cur[k]-new(H)[k]),f'{rank}/{total}'])
table(['预测任务','HDSF WAPE %↓','此前综合推荐 %↓','此前全部记录最低 %↓','相对推荐减少 pp↑','HDSF同口径排名'],rr,[2.30,1.60,1.65,1.75,1.65,1.40],10.5)
curves=[prior(H)['wape']-new(H)['wape'] for H in [1,7]];p(('两个跨度的曲线WAPE均低于此前综合推荐。' if all(x>0 for x in curves) else'曲线WAPE未在两个跨度同时超过此前综合推荐。')+'以下逐项列明改善与退步；原项目公布的池化分数不参与上述排名。单项最低包含旧轮次未采用的冻结诊断，不能把该名次当成独立验证。')
p('24小时开发A/B混合比例='+n(s['selection']['configs']['24h']['mix'])+'；7天='+n(s['selection']['configs']['7d']['mix'])+'。HDSF只是本次复现候选，旧部署入口保留；没有按测试成绩改成员池、重选比例或改写旧轮次选择。')
table(['此前综合推荐','明确方法'],[['24小时 R32',prior(1)['model_name']],['7天 R36',prior(7)['model_name']]],[1.70,8.65],10)
oldcount=len(a['rows'])-len(a['new_rows']);p(f'本次新增{len(a["new_rows"])}条成员及HDSF记录，与此前{oldcount}条记录按同一代码比较；精确相同的预测合并后共{len(a["unique_rows"])}组。原始记录、完整模型名称和所有补充指标保留在CSV。功率单位为归一化利用率，MAE/RMSE不是kW或kWh。')
h('未来24小时完整负荷曲线与主要指标')
fields=['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae','rmse','peak_mape','p95_ape','hit2_percent','hit5_percent'];rr=[]
for method in ['hdsf',*s['selection']['configs']['24h']['methods']]:
 r=next(x for x in a['new_rows'] if x['horizon_days']==1 and x['model']==f'r44_{method}_H1');rr.append([short[method],*[n(r[k]) for k in fields]])
for label,r in [('此前综合推荐 R32',prior(1)),('此前单项曲线最低（冻结诊断）',min(oldrows(1),key=lambda x:x['wape'])),('前一天同刻基线',next(x for x in oldrows(1) if x['model']=='last_source'))]:rr.append([label,*[n(r[k]) if r.get(k) is not None else'—' for k in fields]])
table(['清晰模型名称/历史长度','开发\nWAPE %↓','测试\nWAPE %↓','基线误差\n减少 %↑','MAE↓','RMSE↓','日最大值\nMAPE %↓','APE95\n分位 %↓','APE≤2%\n占比 %↑','APE≤5%\n占比 %↑'],rr,[3.30,.77,.77,.83,.71,.71,.86,.77,.80,.83],8.9)
p('开发起报D14/16/18/20，测试起报D22至D30，各次96点；测试49243个有效时点。日最大值仅比较幅值；峰值出现时刻作为独立补充。开发WAPE包含HDSF用于拟合权重的样本，不能代替留一起报验证。')
p('此前曲线最低诊断的完整模型名称：'+min(oldrows(1),key=lambda x:x['wape'])['model_name']+'。')
h('未来7天完整曲线与每日平均负荷')
methods=['hdsf',*s['selection']['configs']['7d']['methods']];rr=[];daily=[]
entries=[]
for method in methods:entries.append((short[method],next(x for x in a['new_rows'] if x['horizon_days']==7 and x['model']==f'r44_{method}_H7')))
entries += [('此前综合推荐 R36',prior(7)),('此前最低曲线冻结诊断',min(oldrows(7),key=lambda x:x['wape'])),('此前最低日均冻结诊断',min(oldrows(7),key=lambda x:x['daily_wape'])),('前三天同刻均值基线',next(x for x in oldrows(7) if x['model']=='mean3_source'))]
for label,r in entries:
 rr.append([label,*[n(r[k]) for k in ['wape','relative_curve_error_reduction_percent','mae','rmse','peak_mape','p95_ape','hit2_percent','hit5_percent']]])
 daily.append([label,*[n(r[k]) for k in ['daily_wape','relative_daily_error_reduction_percent','daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent']]])
table(['模型名称','曲线\nWAPE %↓','基线误差\n减少 %↑','MAE↓','RMSE↓','日最大值\nMAPE %↓','APE95\n分位 %↓','APE≤2%\n占比 %↑','APE≤5%\n占比 %↑'],rr,[3.35,.86,.92,.79,.79,.98,.88,.89,.89],9.1)
p('每次672点，测试D22/23/24三个相互重叠的7天窗口，共114902个有效时点。每日平均负荷另按完整96点评分，共1187个有效PDU日，不能与曲线WAPE横向比较。')
table(['同一预测曲线导出的每日均值','日均\nWAPE %↓','基线日均误差\n减少 %↑','日均MAE↓','日均RMSE↓','日均APE≤2%\n占比 %↑','日均APE≤5%\n占比 %↑'],daily,[3.35,1.17,1.45,1.05,1.05,1.17,1.11],9.1)
p('此前最低曲线诊断：'+min(oldrows(7),key=lambda x:x['wape'])['model_name']+'；此前最低日均诊断：'+min(oldrows(7),key=lambda x:x['daily_wape'])['model_name']+'。完整成绩见全部指标CSV。')
h('HDSF与此前综合推荐补充指标和区间')
rr=[]
for k,title in [('wape','曲线WAPE %↓'),('daily_wape','每日平均负荷WAPE %↓'),('mae','曲线MAE↓'),('rmse','曲线RMSE↓'),('peak_mape','每日最大值幅值MAPE %↓'),('p95_ape','时点APE95分位 %↓'),('hit2_percent','时点APE≤2%占比 %↑'),('hit5_percent','时点APE≤5%占比 %↑'),('daily_hit2_percent','日均APE≤2%占比 %↑'),('daily_hit5_percent','日均APE≤5%占比 %↑'),('absolute_bias_percent','整体绝对偏差 %↓'),('peak_time_mae_minutes','峰值时刻偏差 分钟↓'),('cella5_wape','cell-a五域WAPE %↓'),('other52_wape','其余52域WAPE %↓')]:rr.append([title,n(prior(1)[k]),n(new(1)[k]),n(prior(7)[k]),n(new(7)[k])])
table(['统一评价指标','此前24小时 R32','HDSF 24小时','此前7天 R36','HDSF 7天'],rr,[3.75,1.65,1.65,1.65,1.65],9.7)
iv=[x for x in a['new_intervals'] if x['method'].startswith('HDSF')];table(['附加经验80%区间','实际覆盖 %','平均宽度↓','Winkler80↓'],[[f'{x["horizon_days"]}天',n(x['coverage_percent']),n(x['mean_width_normalized']),n(x['winkler80'])] for x in iv],[3.75,2.20,2.20,2.20],9.7)
p('区间仅用各PDU开发预测残差的10/90分位数校准，再按同一覆盖、宽度和Winkler公式测测试集；属于附加经验区间，不是HDSF原生概率输出。没有用测试真值校准。')
h('评价口径审查与原项目公布成绩')
p('9.28参考文档“四 评估指标”的原文：'+a['reference_audit']['exact_statement'])
p('对域i先算WAPE_i = 100 × Σ有效点|预测−真值| / Σ有效点|真值|，最终WAPE = (1/57) × Σ_i WAPE_i。池化WAPE = 100 × Σ_i Σ有效点|预测−真值| / Σ_i Σ有效点|真值|；它等价于按各域实际值总量加权的WAPE，域权重不同。')
table(['项目','本次严格处理'],[['曲线WAPE','旧评价代码未修改；每域单独计算，再57域等权平均。'],['评分目标','57域31天15分钟数据；原HDSF输入自然排序，最终按旧序列标识重排；所有有效真值及掩码与旧数据逐点相等。'],['日均负荷','旧统一v2规则要求整日96点全部有效。24小时508个PDU日，7天1187个PDU日；项目池化日均保留异常日剩余有效点，有1197个样本。'],['模型训练及选权','忠实保留原HDSF的开发池化L1拟合与留一起报选择。训练损失可不同，但最终报告及模型排序均用统一评价口径。'],['其他9.28指标','有效时点MAE、APE≤2%/≤5%、APE95、每日最大值幅值误差、同任务基线改善率及80%区间覆盖均一致。'],['补充指标','RMSE、完整日日均MAE/RMSE/命中、峰值时刻偏差、分组和偏差统一计算；峰值时刻不属于原幅值指标。']],[2.05,8.30],10)
rr=[]
for task,H,k,label in [('24h',1,'wape','24小时曲线'),('7d',7,'wape','7天曲线'),('7d',7,'daily_wape','7天每日均值')]:
 pooled=s['source_scores'][task]['daily' if k=='daily_wape' else'curve']['hdsf']['wape_pct'];published={'24小时曲线':2.336179,'7天曲线':2.683159,'7天每日均值':1.837829}[label];rr.append([label,f'{published:.6f}',n(pooled),n(new(H)[k])])
table(['任务','原项目公布池化 %','本次原项目口径 %','本次统一口径 %'],rr,[3.40,2.30,2.30,2.35],10)
p('第二、三列用于复现核对；第四列才可与此前模型直接排序。软件版本及CPU数值环境可导致末位差异，不能把跨口径差异称为模型改进。')
h('原HDSF配方冻结记录及研究结论')
c1=s['selection']['configs']['24h'];c7=s['selection']['configs']['7d'];table(['项目','24小时','7天'],[['预测点/半日段','96 / 2','672 / 14'],['成员池','8个基础成员＋local24','4个基础成员＋local7'],['A/B结构','A不含local24；B包含local24','A不含local7；B包含local7'],['单纯形拟合','每段权重非负、和为1；L1，最多25轮','同左，每段最多20000个开发有效点'],['选权','41点A/B混合网格；留一起报容差0.005pp','同左；仅使用D14/15开发标签'],['开发拟合池化WAPE %',n(c1['development_wape']),n(c7['development_wape'])],['留一起报池化WAPE %',n(c1['loo_wape']),n(c7['loo_wape'])],['选中B比例',n(c1['mix']),n(c7['mix'])],['开发约束回退',str(c1['selection_policy']['constraint_fallback']),str(c7['selection_policy']['constraint_fallback'])]],[2.10,4.13,4.12],10)
p('local24采用原项目四种子101/202/303/404、LightGBM L1、15叶、350轮、学习率0.04、64个历史特征；与加权滞后回归各50%。开发只拟合D14，最终只拟合D22；测试起报更新已观测历史输入，不新增标签训练。原源码D1残差以整个训练前缀均值为基线，与后续日的前一天残差定义不一致。本次忠实保留并记录该问题；它不等于使用测试标签泄漏。')
p('测试前已固定成员池、两种跨度权重、局部模型、源码提交、数据、官方权重和环境指纹。真实科学核查'+str(qa['check_count'])+'项通过；额外运行原项目因果/融合/数据/指标测试，并对完整四种子模型做无缓存未来污染重训核对。未复用外部预测包。')
for H in [1,7]:
 b=a['bootstrap'][str(H)];g=b['gain_pp'];rel=100*(1-new(H)['wape']/prior(H)['wape']);p(f'{H}天曲线相对此前综合推荐减少{n(g)}个百分点，相对误差变化{n(rel)}%。按起报日期配对重采样的探索95%改善区间为[{n(b["paired_origin_exploratory_95_gain_pp"][0])}, {n(b["paired_origin_exploratory_95_gain_pp"][1])}]个百分点。')
p('May2019测试集已多轮查看，项目原成员设计也参考过历史测试成绩。仅31天、7天仅3个重叠窗口，缺少独立确认月份，预训练语料与PowerData2019的重合也未完全排除；开发拟合较好和当前测试排名不能证明独立泛化或全面显著提升。所有成员及较差结果作为复现证据保留。')
h('GPU实测与统一评价核查')
g=read(Q/'gpu_probe.json');table(['实际验证项目','结果'],[['显卡',g['device']+'，8GB显存'],['模型与输入','Toto2.0 313M同cell/设备类型联合，D22起报，14天历史，96点预测'],['CPU推理时间',f'{g["cpu_inference_seconds_from_log"]:.1f}秒（同一次原CPU复现日志，0.1秒精度）'],['GPU推理时间',n(g['inference_seconds'])+'秒（同步等待GPU完成）'],['约计推理加速',f'{g["approximate_speedup"]:.2f}倍，模型加载时间另计'],['GPU显存峰值',f'{g["peak_vram_allocated_bytes"]/2**30:.3f} GiB已分配；{g["peak_vram_reserved_bytes"]/2**30:.3f} GiB保留'],['GPU与CPU最大逐点差',f'{g["max_absolute_difference"]:.9g}（归一化功率）'],['CPU/GPU曲线WAPE',n(g['cpu_single_origin_metrics']['wape'])+'% / '+n(g['gpu_single_origin_metrics']['wape'])+'%，同一未改评价代码'],['GPU环境',g['torch_version']+'，CUDA'+g['cuda_version']+'，float32，TF32关闭']],[3.0,7.35],10.5)
p('复用本机已有CUDA Torch，在独立环境使用本轮模型依赖，没有另下载CUDA包或复制已有大权重。这是一次Toto2.0单起报的实际GPU加速和数值核对，未完成所有基础成员的GPU全测试，不能把约计倍数推广成整个HDSF流程加速倍数。')
p('GPU和CPU使用相同真值、有效点掩码、起报及统一评价代码。这个单起报中有一个PDU没有完整有效日，其57域等权日均WAPE在旧统一代码中未定义，记录为null；没有删除该域或改用池化来填出一个数字。完整HDSF的9个/3个起报结果没有该问题。')
p('正式三任务结果仍为完整CPU复现，GPU核对没有参与成员选择、超参数选择或改变冻结权重。后续加速可使用GPU，但必须继续记录实际环境、输入和数值差异；不将运行平台变化伪装成精度改进。内存保护仅约束本实验子进程，未关闭用户程序。')
h('复现文件环境与使用边界')
table(['文件或入口','用途'],[['HDSF统一指标结果.csv','本次全部16条成员/HDSF记录和清晰模型名称。'],['HDSF与此前模型全部指标比较.csv','包含此前'+str(len(a['rows'])-len(a['new_rows']))+'条及本次16条，全部统一评价指标。'],['各任务单项最优结果_含HDSF.csv','分别列24小时曲线、7天曲线、每日均值的单项最低/最高及方法身份。'],['HDSF_H1/H7_历史滚动预测.csv','完整历史测试曲线；真值及质量标签仅在冻结后附加用于核查。'],['HDSF_H1/H7_每日平均负荷.csv','各起报/PDU/目标日的96点均值及完整日有效标志。'],['32_HDSF半日单纯形融合复现.cmd','恢复或重读本轮研究；保留原冻结选择，不按测试名次重选。'],['experiment/round44/docs/protocol.json','固定源码提交、任务日期、算法配方及未改评价代码指纹。'],['experiment/round44/results','逐成员独立预测、冻结记录、真实评价、区间、审计与资源记录。']],[4.35,6.00],10)
p('实体保存位置：${PROJECT_ROOT}。所有新模型、环境、程序和产出均在D盘；桌面仍为原兼容入口。本次没有扩展下载ClusterData，仅补齐官方Toto2.0 313M及IBM PatchTST-FM-r2权重，四套已有权重复用。')
p('独立CPU环境使用Python3.12.14（项目示例3.13）、Torch2.8.0+cpu、Transformers4.57.6、LightGBM4.6.0。IBM源码固定3d92535ca4138648b9d1e2f3f1421b5dec4dc3aa，权重固定b125275f9204d37cbb81fe47b9cf5e08a521e829。上游wheel遗漏模型子包，改为同提交完整源码安装；构建标识含提交号，推理模型代码未修改。工具包声明的Torch2.10训练依赖未替换HDSF指定2.8，以实际导入和前向推理核查为准，不声称整套工具包依赖检查完全一致。')
p('PatchTST原57域一批的推理触发真实内存保护，失败证据保留；在开发融合及测试前登记固定每批8域、16线程、float32，不改变模型权重或HDSF配方。同一实际历史输入的8域一批与4域分批前向核对最大差小于2×10⁻⁶；未完成原57域一批的输出，不能宣称与它逐位完全一致。其他成员采用原运行配方。')
p('数据CC BY4.0署名要求继续保留。Google TimesFM3权重仅限非商业研究、非生产；IBM PatchTST-FM-r2允许选Apache2.0或OpenMDW1.0。GitHub项目未提供独立LICENSE，不能据此宣称源代码可任意再分发。各权重来源、实际SHA256、版本和许可材料均保留。')
p('模型输出为功率归一化利用率，不是物理kW。原HDSF采用非负凸权重，这只是组合结构约束，不等于能量守恒。未静默裁剪原基础模型或HDSF点预测；经验区间按旧统一流程限制到0至1。预测范围检查记录在本轮审计。')
p('项目来源：https://github.com/daifeizhai/PowerData_Prediction；固定源码提交648777507ebc57e7498ce61f4ae2236ab179e6a8。历史测试曲线不是今天的负荷预测，未把未来CPU、未来生产功率或目标真值输入模型。')
out=T/'HDSF半日单纯形融合复现与统一评价汇报.docx';doc.save(out);print(str(out))
