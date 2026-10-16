from pathlib import Path
import sys,json,hashlib,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
T=Path(sys.argv[1]);B=T/'experiment';R=B/'round40';Q=R/'results'
a=json.loads((Q/'continued_comparison.json').read_text(encoding='utf-8'));qa=json.loads((Q/'independent_verification.json').read_text(encoding='utf-8'))
assert a['all_checks_passed'] and len(a['rows'])==114 and qa['all_passed'] and qa['check_count']==140
assert not any(x['adopt'] for r in a['release_decisions'].values() for x in r.values())
out=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx'
if '--validate-only' in sys.argv:print('114 uniform records; 92 unique predictions; actual four rounds; no adopted upgrades');sys.exit(0)
backup=R/'docs/before_continued_report';backup.mkdir(parents=True,exist_ok=True)
if not (backup/out.name).exists():shutil.copyfile(out,backup/out.name)
layout=(B/'round35/code/report_curves_priority.py').read_text(encoding='utf-8');layout=layout[layout.index('doc=Document()'):layout.index("doc.add_paragraph('Google PowerData 2019 统一评价与新模型研究结果'")];exec(layout)
for r in a['rows']+a['unique_rows']:
 if r['model']=='r32_selected_H1':r['report_name']='24小时当前推荐\nChronos2 TiRex2 TimesFM3 Ridge TTM'
 if r['model']=='r36_selected_H7':r['report_name']='7天当前推荐\nMoirai2同场站10% 全域日均协调'
 if r['model']=='r31_selected_H7':r['report_name']='此前7天推荐\nTimesFM3 Ridge HGB TTM'
def percent(k):return 'hit' in k or k.startswith('relative')
doc.add_paragraph('Google PowerData 2019 三项预测任务与模型比较','Title')
p('本次继续探索Salesforce Moirai 2.0 R small，完成多分辨率输出头训练、季节变化量、相关域联合预测及真实CPU内存辅助输入四组实验。24小时曲线、7天曲线和7天日均的测试最低值都有更新，但分别来自不同配置；相对当前推荐改善约0.27%、0.46%、0.79%，没有达到大幅提升，也未找到同时稳健领跑两类曲线的统一方法。')
rr=[]
for task,H,k in [('未来24小时完整曲线',1,'wape'),('未来7天完整曲线',7,'wape'),('未来7天每日平均负荷',7,'daily_wape')]:
 cur=recommended(H);win=min(group(H),key=lambda r:r[k]);rr.append([task,n(cur[k]),n(win[k]),win['report_name'],f'{100*(1-win[k]/cur[k]):.3f}%'])
table(['任务','当前推荐\nWAPE %','全部实测最低\nWAPE %','最低值的明确模型','相对当前\n误差减少'],rr,[1.55,1.15,1.30,5.05,1.30],10.4)
p('上表最低值均为开发升级未通过的冻结研究候选。实验未按测试换候选、权重或部署；当前24小时与7天推荐仍为第32轮和第36轮。114条记录按同一规则复算，去掉相同预测的别名后为92组不同预测。最低值与正式推荐分别报告。')
p('模型名称说明：Moirai2指Salesforce Moirai 2.0 R small；Chronos2指Chronos-2；TTM指IBM Granite TinyTimeMixer R2；HGB指HistGradientBoosting。TimesFM 2.5与TimesFM3分别列示。本文组合名称明确给出骨干、输入方式、输出修正及权重，不以笼统“融合模型”代称。')
p('当前24小时方案：Chronos-2 25%＋TiRex-2 25%＋TimesFM3 50%，叠加Ridge残差校准、TimesFM3末层微调及TinyTimeMixer多任务输出头10%稀疏峰值修正。当前7天方案：Moirai2同场站14天历史中位数10%＋此前TimesFM3原生672点、Ridge日均及Fourier形状、HGB与TinyTimeMixer重建方案90%，再保持57域预测日均总量。当前方案是后续四轮的固定参照。')
p('历轮结论保留：第8c轮冻结权重微调WAPE2.374440%，比当时第8b轮2.374128%略差；完整CPU与稀疏MoE第9轮选择alpha0，未升级。第36轮当前7天方案的a5退步0.008072个百分点，超原0.005容差0.003072；依据用户允许次指标微小差距的偏好明确采用，原验收例外记录保留。后续四轮未再按测试变更采用。')
h('本次曲线最低候选与当前推荐的多维比较')
old1=recommended(1);new1=row(1,'r37_diagnostic_H1');old7=recommended(7);new7=row(7,'r40_diagnostic_H7')
fields=[('wape','曲线WAPE %↓'),('daily_wape','日均WAPE %↓'),('mae','曲线MAE↓'),('rmse','曲线RMSE↓'),('peak_mape','日最大值MAPE %↓'),('p95_ape','APE95分位 %↓'),('hit2_percent','APE≤2%占比 %↑'),('hit5_percent','APE≤5%占比 %↑'),('daily_mae','日均MAE↓'),('daily_rmse','日均RMSE↓'),('daily_hit2_percent','日均APE≤2%占比 %↑'),('daily_hit5_percent','日均APE≤5%占比 %↑'),('cella5_wape','a5曲线WAPE %↓'),('other52_wape','其余52域曲线WAPE %↓'),('absolute_bias_percent','绝对总偏差 %↓'),('peak_time_mae_minutes','峰值时刻误差 分钟↓')]
vals=[];marks=set()
for i,(k,title) in enumerate(fields):
 vals.append([title,n(old1[k]),n(new1[k]),n(old7[k]),n(new7[k])])
 for j,oldr,newr in [(2,old1,new1),(4,old7,new7)]:
  if newr[k]>oldr[k]+1e-10 if percent(k) else newr[k]<oldr[k]-1e-10:marks.add((i,j))
table(['统一指标','24h当前推荐','24h小时同场站Rank8\n研究最低候选','7天当前推荐','7天三小时CPU内存\n研究最低候选'],vals,[2.80,1.58,2.22,1.58,2.17],10,marks)
p('浅绿色表示相对当前推荐的改善。24小时候选每日均值2%命中略退，峰值幅值也不是单项最低。7天CPU内存候选峰值MAPE增加0.075195个百分点，时点2%命中下降0.046126个百分点，其余表列指标多数改善；这只是同月测试的描述性对比。')
p('24小时曲线改善0.006415个百分点，只有4／9个起报改善；按日期重采样的探索95%改善区间跨零。7天曲线改善0.012405个百分点，3／3个重叠起报改善。三个窗口和多轮查看过的同月测试不构成独立显著性证据；不能把小幅变化解释为“大幅突破”。')
h('9月28日评价方法一致性审查')
table(['项目','统一实现及核查结论'],[
 ['曲线WAPE','每PDU有效点绝对误差和／实际绝对值和，再对57域等权平均。与原文明确规定一致。'],
 ['粒度与异常点','三条5分钟样本平均成15分钟；输入仅历史前向填充，目标无效点排除。24小时49243个、7天114902个有效点。'],
 ['开发与测试','24小时开发14／16／18／20，测试22至30；7天开发14／15，测试22／23／24。在当日结束时起报。'],
 ['任务与基线','24小时一次96点，对照前一天同刻；7天一次672点，对照前三天同刻均值。改善率＝1−模型WAPE／该任务基线WAPE。'],
 ['9月28日指标','曲线WAPE、MAE、APE≤2%／≤5%占比、APE95分位、每日最大值幅值MAPE、7天每日均值WAPE及80%区间覆盖均列示。'],
 ['日均统一规则','排除含任一无效点的整日目标；真值与预测均取完整96点均值。24小时508个、7天1187个有效PDU日。'],
 ['补充评价','RMSE、日均MAE／RMSE／命中、峰值时刻、偏差、a5与其余52域、区间宽度和Winkler80采用同一实现。']],[2.05,8.30],10.2)
legacy={'mean3_source':2.262,'chronos':2.618,'tf25_7':2.204,'tirex':2.127,'equal':2.293};rr=[]
p('日均历史锚点与统一复算如下。')
for key,orig in legacy.items():
 r=row(7,key);rr.append([r['report_name'],f'{orig:.3f}',n(r.get('legacy_daily_wape',r['daily_wape_valid_slots_supplement'])),n(r['daily_wape'])])
table(['原文日均锚点模型','原文WAPE %','旧有效点日均 %','统一完整日WAPE %'],rr,[5.80,1.45,1.55,1.55],9.8)
p('原文未附完整日均实现。完整日规则复现固定基线、Chronos-2和TimesFM 2.5的三位小数锚点；TiRex及等权组合仍有小于0.001个百分点的差异，不能声称原代码全部复现。严格的新旧排序全部采用本次统一v2复算值，不混入历史引用。日均口径更新导致旧7天推荐由1.918249%变为1.922296%，不是模型变化。')
fields=['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae','rmse','peak_mape','p95_ape','hit2_percent','hit5_percent']
def curve_table(H,keys):
 data=[];marks=set()
 for i,key in enumerate(keys):
  r=row(H,key);rr=[r['report_name']]
  for j,k in enumerate(fields,1):
   v=r.get(k);yes=v is not None and k!='development_wape_original_grid' and star(H,k,v);rr.append('—' if v is None else n(v)+(' ★' if yes else''))
   if yes:marks.add((i,j))
  data.append(rr)
 table(['明确模型及配置','开发\nWAPE %↓','测试\nWAPE %↓','误差减少\n%↑','MAE↓','RMSE↓','日最大值\nMAPE %↓','P95\nAPE %↓','APE≤2%\n占比 %↑','APE≤5%\n占比 %↑'],data,[3.16,.79,.79,.80,.80,.80,.85,.77,.77,.82],9.1,marks)
h('未来24小时完整负荷曲线模型结果')
keys1=['r32_selected_H1','last_source','ridge','lgb','chronos','tirex','tf25','equal']+[f'r{n}_diagnostic_H1' for n in range(37,41)]
p('★及浅绿色表示92组不同预测中的该项最低或最高，包含未采用研究候选。完整114条原始结果保存在顶层CSV；研究候选均在测试前已按开发信息冻结，单项第一不等同于采用。')
curve_table(1,keys1)
p('本次曲线最低是小时同场站Rank8输出头10%及当前24h方案90%、全域日均协调，WAPE2.367483%；开发期曲线却比当前推荐差，因此不按测试名次替换。相关域联合的2.370007%也优于旧记录，但并未成为24小时最低。')
h('未来7天完整负荷曲线模型结果')
keys7=['r36_selected_H7','r31_selected_H7','mean3_source','chronos','tirex','tf25_7','tf3_uni21_7','equal']+[f'r{n}_diagnostic_H7' for n in range(37,41)]
p('每次直接预测672个15分钟点；三个测试起报窗口重叠，全部输入均来自起报前。三小时Moirai输出56点，再保留当前方案15分钟局部形状重建；最终评分粒度仍为15分钟，不在低频上评分后声称提高。')
curve_table(7,keys7)
p('7天曲线最低是三小时57域功率联合真实a5 CPU内存的10%修正，WAPE2.694096%；其完整开发期曲线不优于当前推荐。相关域联合候选2.695253%略高。当前7天推荐仍2.706501%，不能据此声称推荐已领跑所有实测配置。')
h('未来7天每日平均负荷模型结果')
rr=[];marks=set();f2=['daily_wape','relative_daily_error_reduction_percent','daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent']
for i,key in enumerate(keys7):
 r=row(7,key);vals=[r['report_name']]
 for j,k in enumerate(f2,1):
  yes=star(7,k,r[k]);vals.append(n(r[k])+(' ★' if yes else''))
  if yes:marks.add((i,j))
 rr.append(vals)
table(['明确模型及配置','日均\nWAPE %↓','日均误差\n减少 %↑','日均MAE↓','日均RMSE↓','日均APE≤2%\n占比 %↑','日均APE≤5%\n占比 %↑'],rr,[4.12,1.00,1.05,1.00,1.00,1.09,1.09],10,marks)
p('每日均值来自同一672点预测，每天平均96点；部署文件也按96点生成日均，不使用未来真实异常标记。日均最低为季节残差10%候选1.889163%，相对当前1.904238%减少约0.792%。它在开发期未满足全部升级要求，曲线2.700885%也不是7天最低，不能将两种配置成绩合并。')
p('曲线与日均单位都是归一化功率利用率，MAE和RMSE也为该单位；每日均值不等于日耗电kWh。峰值比较独立列示，用户允许峰值小幅差距不意味着可以根据已看过的测试重新选权重。')
h('四组新方向的开发证据和采用判断')
rr=[]
for nr in range(37,41):
 for H in [1,7]:
  r=row(H,f'r{nr}_diagnostic_H{H}');release=a['release_decisions'][str(nr)][str(H)];rr.append([f'{H}天 '+r['report_name'],n(r['development_wape_original_grid']),n(r['wape']),n(r['daily_wape']),n(r['peak_mape']),'DEV选中 TEST未通过' if release['dev_selected'] else'DEV未通过'])
table(['明确方法与跨度','开发曲线\nWAPE %↓','测试曲线\nWAPE %↓','日均\nWAPE %↓','峰值\nMAPE %↓','冻结采用结论'],rr,[4.40,1.00,1.00,1.00,1.00,1.96],10)
p('多分辨率组比较1小时与3小时、单变量与同场站、原始及Rank4／Rank8残差头。原始1140万参数骨干冻结，只训练3840或7680个输出适配参数，4轮固定优化，标签分别严格截止第14天与第22天。训练损失下降没有带来稳定的开发升级；小时Rank8的测试改善保留为负面开发与正面测试并存的结果。')
p('季节残差仅减去起报前最近3天的同刻均值，再加回同一历史轮廓；相关域组仅按训练期小时水平与一阶变化的相关性分成最多8域一组。所有候选先用开发早段选结构、后段选权重及日均协调；7天两起报重叠，不能称独立验证。第38轮24小时原生单变量在DEV选中，但测试2%命中退步约0.138个百分点、偏差退步约0.010492个百分点，超过登记门槛，未采用。')
p('140项独立复核通过：指标回算、实际重建重放、冻结SHA及时间顺序、缓存来源权重输入绑定、训练标签边界、零残差头等值、历史相关分组、控制模型回归及未来功率／CPU扰动检查。扰动检查验证输入包装边界，不证明不存在模型过拟合或预训练资料重叠。')
h('真实CPU内存对照与来源范围')
rr=[]
for v in a['cpu_dev_ablation']:
 rr.append([f'{v["horizon_days"]}天 '+('57域功率＋a5 CPU内存' if 'cpu_mem' in v['family'] else'同骨干仅57域功率'),n(v['development_wape']),n(v['development_daily_wape']),n(v['cella5_wape']),n(v['other52_wape'])])
table(['开发期匹配对照 25%修正 无日均协调','开发曲线WAPE %↓','开发日均WAPE %↓','a5 WAPE %↓','其余52 WAPE %↓'],rr,[4.00,1.58,1.58,1.58,1.58],10.5)
p('两组输入使用相同Moirai骨干、14天历史与3小时粒度。CPU内存对照在完整开发期略优于仅功率控制，但两者均未形成可靠升级；上表固定25%用于结构比较，最终冻结研究诊断为10%加全域日均协调，不能把两种设置混成一个分数。TEST只评分开发冻结候选，没有新增匹配控制的测试搜索，所以本轮尚无独立确认的CPU增益结论。')
p('实际读取为官方ClusterData2019 cell-a的12085个Parquet必要列，完整审计通过，剩余下载0。真实CPU及内存只覆盖a的5个PDU，共10条历史辅助通道；其余52域没有伪造CPU。CPU单位NCU，内存为归一化量，功率为利用率。共同输入仍受作业聚合去重键、时间对齐及可观测范围限制；第三任务字段是job-window equivalent，不是CPU方差。')
rr=[]
for H,nr in [(1,37),(7,40)]:
 for r in [recommended(H),row(H,f'r{nr}_diagnostic_H{H}')]:
  rr.append([f'{H}天 '+r['report_name'],n(r['cella5_wape']),n(r['other52_wape']),f'{r["improved_pdus"]}/57',f'{r["improved_cells"]}/10',f'{r["improved_origins"]}/{9 if H==1 else 3}'])
table(['当前推荐与曲线最低研究候选','a5 WAPE %↓','其余52 WAPE %↓','改善基线的域','改善基线的场站','改善基线的起报'],rr,[4.55,1.16,1.36,1.09,1.09,1.10],10)
p('分组改善数均相对于对应历史基线，不能解释为候选在每个域都超过当前推荐。全域日均协调只保持归一化预测总量，不是物理电量守恒；CPU辅助输出也不能当作kW。')
p('Google数据许可CC BY4.0；Moirai权重CC BY-NC4.0，代码Apache2.0；TimesFM3权重仅限非商业非生产研究。官方固定版本、来源、许可证和校验保存在round29及round35。')
h('预测区间和最终研究结论')
rr=[]
for H in [1,7]:
 v=next(v for v in a['intervals'] if v.get('method','').startswith('R33' if H==1 else'R36') and v['horizon_days']==H)
 rr.append([f'{H}天 当前推荐开发校准',80,f'{v["coverage_percent"]:.3f}',n(v['mean_width_normalized']),n(v['winkler80'])])
for nr in range(37,41):
 for H in [1,7]:
  v=next(v for v in a['intervals'] if v.get('method','').startswith(f'R{nr} ') and v['horizon_days']==H)
  rr.append([f'{H}天 '+row(H,f'r{nr}_diagnostic_H{H}')['report_name'],80,f'{v["coverage_percent"]:.3f}',n(v['mean_width_normalized']),n(v['winkler80'])])
table(['区间方法','标称覆盖 %','实际覆盖 %','平均宽度↓','Winkler80↓'],rr,[5.55,1.12,1.18,1.25,1.25],9.8)
p('当前24小时沿用第33轮开发选定对称区间；当前7天及新候选用冻结开发残差10／90分位经验80%区间，不按测试搜索校准。新候选区间围绕其各自预测，不用于替换当前推荐区间。覆盖、宽度和Winkler80并不同时最优，7天开发窗口重叠，经验区间无独立覆盖保证。')
p('结论：Moirai2方向具有小幅改善线索，小时适配头、相关域及真实CPU输入各有局部优势；四轮没有得到“同一方案稳健领跑所有曲线并大幅提升”的证据。继续在同一31天上追逐测试最低值会加重选择偏差。所有训练、负面开发结果和测试预测保留，正式推荐保持不变；确认更高精度仍需要独立月份或未参与探索的同类场站数据。')
p('所有实体文件位于D盘，桌面目录为兼容junction。新导出为2019年5月末历史复现，6月1至7无真值，不是今天实时电负荷。原始列字节、源聚合、权重和失败证据保留；压缩必要列84.78GiB、原对象407.55GiB、迁移逻辑105.481GiB为不同统计量，不能混称下载量。')
doc.save(out);print(json.dumps(dict(document=str(out),sha256=hashlib.sha256(out.read_bytes()).hexdigest()),ensure_ascii=False))
