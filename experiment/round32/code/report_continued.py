from pathlib import Path
import sys,json,hashlib,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
T=Path(sys.argv[1]);B=T/'experiment';R=B/'round32';Q=R/'results';a=json.loads((Q/'continued_comparison.json').read_text(encoding='utf-8'));assert a['all_checks_passed'] and len(a['rows'])==86
out=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx'
if '--validate-only' in sys.argv:print('86 uniform rows; R31/32 checks complete; audit ambiguity preserved');sys.exit(0)
backup=R/'docs/before_aligned_report';backup.mkdir(parents=True,exist_ok=True)
if not (backup/out.name).exists():shutil.copyfile(out,backup/out.name)
doc=Document();sec=doc.sections[0];sec.page_width=Inches(11.69);sec.page_height=Inches(8.27);sec.top_margin=Inches(.55);sec.bottom_margin=Inches(.5);sec.left_margin=sec.right_margin=Inches(.65)
for k in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[k];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑');pp=st.element.get_or_add_pPr();snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.1;doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(21);doc.styles['Title'].paragraph_format.space_after=Pt(10);doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 1'].paragraph_format.space_before=Pt(0);doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
for b in list(doc.styles.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('2026年10月7日  统一评价与新方法实测  ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
def p(t):return doc.add_paragraph(t)
def h(t):pp=doc.add_paragraph(t,'Heading 1');pp.paragraph_format.page_break_before=True
def table(headers,rows,widths,size=9.4,marks=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w);t.rows[0].cells[j].text=str(headers[j])
 for rr in rows:
  cells=t.add_row().cells
  for j,v in enumerate(rr):cells[j].text=str(v)
 for i,row in enumerate(t.rows):
  tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
  if i==0:tr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
   cp.append(bd);mar=OxmlElement('w:tcMar')
   for side,v in [('top',65),('bottom',65),('left',75),('right',75)]:e=OxmlElement('w:'+side);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');mar.append(e)
   cp.append(mar);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCEAF2' if i==0 else('E9F4E5' if marks and (i-1,j) in marks else'F5F7F9' if i%2==0 else'FFFFFF'));cp.append(shade)
   for pp in c.paragraphs:
    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=Pt(size*1.15)
    for run in pp.runs:run.font.size=Pt(size);run.bold=i==0 or bool(marks and (i-1,j) in marks)
 return t
def group(H):return [r for r in a['unique_rows'] if r['horizon_days']==H]
def row(H,key):return next(r for r in group(H) if key==r['model'] or key in r['aliases'])
def recommended(H):return next(r for r in group(H) if r['status']=='当前综合推荐')
def n(x):return f'{x:.6f}'
def star(H,k,v):return abs(v-(max if k.startswith('hit') or k.startswith('daily_hit') or k.startswith('relative') else min)(r[k] for r in group(H)))<1e-10
def label(r):return r['report_name']+('\n当前综合推荐' if r['status']=='当前综合推荐' else'\n未采用诊断' if '未采用' in r['status'] else'')
doc.add_paragraph('Google PowerData 2019 统一评价与新模型研究结果','Title')
p('统一9.28评价方法后，本次继续训练新架构并完成曲线重建。当前采用TinyTimeMixer多任务指导的分跨度重建：24小时在Chronos-2＋TiRex-2＋TimesFM3曲线上作10%稀疏峰值校正；7天在TimesFM3＋Ridge＋HGB曲线上加入10%零均值形状并保持日均。新方案多项改善，仍有微小退步，尚非所有指标第一。全部86条记录使用同一评价口径。')
rr=[]
for task,H,k in [('24小时完整曲线',1,'wape'),('7天完整曲线',7,'wape'),('7天每日平均负荷',7,'daily_wape')]:
 cur=recommended(H);win=min(group(H),key=lambda r:r[k]);rr.append([task,cur['report_name'],n(cur[k]),n(win[k]),win['report_name']+'\n'+win['status']])
table(['任务','当前综合推荐','推荐WAPE %','表内单项最低 %','单项最低的方法及身份'],rr,[1.35,3.30,1.13,1.18,3.40],10)
p('表内最低包含未采用诊断，只是统一口径下的实测排序。它不用于改写旧实验的开发冻结选择；单项最低与综合推荐分别列示。日均WAPE采用本次完整日规则，此前7天推荐日均由旧口径1.918249%变为1.922296%，这不是模型退步。')
tt=[]
for name in ['r30:TTM512 multitask head','r30:TTM512 multitask all']:
 r1=row(1,name);r7=row(7,name);tt.append([r1['report_name'],n(r1['wape']),n(r1['peak_mape']),n(r7['wape']),n(r7['daily_wape']),n(r7['peak_mape'])])
table(['直接训练的新架构','24小时曲线\nWAPE %','24小时峰值\nMAPE %','7天曲线\nWAPE %','7天日均\nWAPE %','7天峰值\nMAPE %'],tt,[4.15,1.26,1.26,1.26,1.26,1.26],10)
p('新架构有专项进步：解码器训练的24小时峰值误差1.371050%，低于此前曲线推荐的2.000202%；全参数训练的7天峰值误差1.293665%，低于此前推荐的1.976260%。但曲线和日均更差，开发期未通过替换条件，故未采用。它们也没有超过所有既有峰值对照，不能称为全面最优。')
h('当前重建方案与此前推荐的逐项比较')
old1=row(1,'best');new1=recommended(1);old7=row(7,'r26_diagnostic');new7=recommended(7)
fields_compare=[('wape','曲线WAPE %↓'),('mae','MAE↓'),('rmse','RMSE↓'),('peak_mape','每日最大值MAPE %↓'),('p95_ape','APE95分位 %↓'),('hit2_percent','APE≤2%占比 %↑'),('hit5_percent','APE≤5%占比 %↑'),('daily_wape','每日均值WAPE %↓'),('peak_time_mae_minutes','峰值时刻偏差 分钟↓'),('cella5_wape','a5曲线WAPE %↓'),('other52_wape','其余52域WAPE %↓')]
vals=[];marks=set()
for i,(k,title) in enumerate(fields_compare):
 vals.append([title,n(old1[k]),n(new1[k]),n(old7[k]),n(new7[k])])
 for j,oldr,newr in [(2,old1,new1),(4,old7,new7)]:
  gain=newr[k]>oldr[k]+1e-10 if k.startswith('hit') else newr[k]<oldr[k]-1e-10
  if gain:marks.add((i,j))
table(['统一评价指标','此前24小时推荐','新24小时推荐','此前7天推荐','新7天推荐'],vals,[3.45,1.72,1.72,1.72,1.72],10,marks)
p('浅绿色标出相对此前推荐的改善。24小时曲线、MAE、RMSE、日均、峰值、2%／5%命中和峰值时刻改善，P95不变；a5 WAPE退步约0.000812个百分点。7天曲线、MAE、RMSE、峰值、P95、5%命中及两组WAPE改善，完整日均保持相同；2%命中退步约0.029590个百分点、峰值时刻偏差增加约1.454分钟。均满足预先登记的微小退步范围，可以按用户的综合偏好替代旧推荐，不能说所有指标最优。')
p('第31轮按原DEV起报比较两种已训练TinyTimeMixer专家的四种固定重建，共每跨度8项；冻结7天解码器专家零均值形状10%，24小时仍保留原方案。第32轮每跨度只检查四个登记比例，冻结24小时稀疏峰值10%，7天未支持进一步扩大比例，保留第31轮10%。测试只评分冻结方案或开发最佳诊断，没有测试比例网格。')
p('24小时峰值目标＝0.9原预测最大值＋0.1 TinyTimeMixer预测最大值，在原最大值位置修正并限制其他点；7天曲线＝0.9原曲线＋0.1去日均后的TinyTimeMixer形状，再投影到原预测日均及0至1范围。这些是预测重建约束，不能当作物理容量或能量守恒。')
p('曲线改善幅度仍小：24小时约0.000229个百分点，7天约0.002822个百分点。起报重采样探索95%区间分别约[−0.000068,0.000514]与[−0.001102,0.005267]个百分点，均跨零；7天窗口还重叠，不能宣称独立显著或大幅改善。')
h('9月28日评价方法一致性审查')
table(['项目','统一实现及核查结果'],[
 ['曲线WAPE','每PDU有效点绝对误差和／实际绝对值和，再对57域等权平均。与原文明确规定一致。'],
 ['粒度与异常点','三条5分钟样本平均成15分钟；输入仅历史前向填充，目标无效点排除。24小时49243个有效点与原文一致。'],
 ['开发与测试','24小时开发14／16／18／20，测试22至30；7天开发14／15，测试22／23／24。均在当日结束时起报。'],
 ['基线与改善率','24小时为前一天同刻；7天为前三天同刻均值。改善率＝1−模型WAPE／同任务基线WAPE。'],
 ['时点及峰值','有效点MAE、APE≤2%／≤5%占比、APE95分位；每日最大值幅值MAPE。峰值时刻独立列为补充。'],
 ['日均异常日规则','旧实现保留异常日剩余有效点。统一v2排除含任一无效点的整日目标；真值与预测均取完整96点均值。'],
 ['区间及补充指标','实测80%区间覆盖，同时报告宽度和Winkler80。RMSE、偏差、分组、日均命中统一实现，属于补充指标。']],[2.05,8.30],10.2)
p('日均历史锚点与统一复算值如下。')
legacy={'mean3_source':2.262,'chronos':2.618,'tf25_7':2.204,'tirex':2.127,'equal':2.293};rr=[]
for key,orig in legacy.items():
 r=row(7,key);rr.append([r['report_name'],f'{orig:.3f}',n(r.get('legacy_daily_wape',r['daily_wape_valid_slots_supplement'])),n(r['daily_wape'])])
table(['原文日均锚点','原文WAPE %','旧有效点日均 %','统一完整日WAPE %'],rr,[5.80,1.45,1.55,1.55],9.7)
p('不能声称已完全复现原评价代码。完整日规则复现固定基线2.262%、Chronos-2的2.618%和TimesFM 2.5的2.204%；TiRex及等权组合仍有小于0.001个百分点的末位差异。原文未附日均数组、异常日实现和池化代码。因此原文历史表仅作锚点，新的严格比较全部使用同一v2重算值；规则依据异常目标排除原则和基线锚点确定，不依据新模型输赢。')
h('未来24小时完整曲线模型结果')
keys=[ 'r32_selected_H1','last_source','mean3_source','ridge','lgb','chronos','tirex','tf25','equal','r22_diagnostic','r30:TTM512 multitask head','r30:TTM512 multitask all']
fields=['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae','rmse','peak_mape','p95_ape','hit2_percent','hit5_percent']
def curve_table(H,keys):
 data=[];marks=set()
 for i,k in enumerate(keys):
  r=row(H,k);rr=[label(r)]
  for j,f in enumerate(fields,1):
   v=r.get(f);yes=v is not None and f!='development_wape_original_grid' and star(H,f,v);rr.append('—' if v is None else n(v)+(' ★' if yes else''))
   if yes:marks.add((i,j))
  data.append(rr)
 table(['明确模型及身份','开发\nWAPE %↓','测试\nWAPE %↓','误差减少\n%↑','MAE↓','RMSE↓','日最大值\nMAPE %↓','P95\nAPE %↓','APE≤2%\n占比 %↑','APE≤5%\n占比 %↑'],data,[3.16,.79,.79,.80,.80,.80,.85,.77,.77,.82],8.9,marks)
p('代表模型列在下表，全部86条原始记录见“模型各项指标完整结果.csv”。★及浅绿色指78组不同预测中的单项最优，未在本页列出的单项赢家仍在“各任务单项最优结果.csv”。测试为9个起报，开发列严格按9.28四个起报计算。')
curve_table(1,keys)
p('24小时新推荐保留Chronos-2 25%＋TiRex-2 25%＋TimesFM3 50%及Ridge与末层微调，加入TinyTimeMixer多任务解码器的10%稀疏峰值修正。测试WAPE 2.373899%，峰值MAPE 1.867076%。它仍不是峰值最佳；前一天同刻约1.068%。')
h('未来7天完整曲线模型结果')
keys7=['r31_selected_H7','mean3_source','chronos','tirex','tf25_7','tf3_uni21_7','equal','r19_selected','r22_diagnostic','r27_frozen','r30:TTM512 multitask head','r30:TTM512 multitask all']
p('每次预测672个15分钟点，测试起报22／23／24，均只使用各起报前历史。三个7天窗口重叠。开发为14／15两个起报，不能把曲线WAPE与每日均值WAPE混作同一个任务。')
curve_table(7,keys7)
p('7天新推荐在此前TimesFM3直接672点、Ridge1000时频总量校准和25% HGB时点稀疏峰值修正上，加入10% TinyTimeMixer零均值形状并保持原预测日均。曲线WAPE 2.720522%，峰值MAPE 1.949178%，没有大幅或独立显著提升证据。')
h('未来7天每日平均负荷模型结果')
rr=[];marks=set();fields2=['daily_wape','relative_daily_error_reduction_percent','daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent']
for i,k in enumerate(keys7):
 r=row(7,k);vals=[label(r)]
 for j,f in enumerate(fields2,1):
  yes=star(7,f,r[f]);vals.append(n(r[f])+(' ★' if yes else''))
  if yes:marks.add((i,j))
 rr.append(vals)
table(['明确模型及身份','日均\nWAPE %↓','日均误差\n减少 %↑','日均MAE↓','日均RMSE↓','日均APE≤2%\n占比 %↑','日均APE≤5%\n占比 %↑'],rr,[4.12,1.00,1.05,1.00,1.00,1.09,1.09],9.4,marks)
p('日均来自同一672点预测曲线，每天平均96点。评分排除不完整的目标日，共1187个有效PDU日；24小时任务相应为508个。输出部署日均使用完整96点预测，不使用未来真实异常标记。日均单位是归一化功率利用率，不是日耗电kWh。')
p('本次统一口径后的7天综合推荐日均WAPE为1.922296%，对应基线2.262403%。TinyTimeMixer解码器训练为2.165669%，全参数训练为2.208120%；它们的峰值专项改善没有转化为更好的日均。')
h('两种新建模路线的完整实验结论')
new=[]
for nround in [29,30]:
 summary=a['new_rounds'][str(nround)]
 for item in summary['selection']['ranking']:
  key=f'r{nround}:'+item['model'];r1=row(1,key);r7=row(7,key);new.append([r1['report_name'],f'{item["joint_cost"]:.4f}',n(r1['wape']),n(r1['peak_mape']),n(r7['wape']),n(r7['daily_wape']),n(r7['peak_mape'])])
table(['新方法及固定配置','开发综合代价↓','24小时曲线 %','24小时峰值 %','7天曲线 %','7天日均 %','7天峰值 %'],new,[4.08,1.01,1.05,1.05,1.05,1.05,1.06],9.3)
p('IBM TinyTimeMixer R2使用官方Apache 2.0权重，分别为512点和1024点历史，原生720点输出裁剪672点；前96点就是24小时预测。参数分别为1444880和4243536。第二路线从真实功率训练57域联合因子与傅里叶日内系数，2或4个因子、Ridge10或100，直接预测未来7个日轮廓，不在旧基础模型残差上继续细调。')
p('新架构训练固定4轮、批量16、AdamW、权重衰减0.01、梯度截断1、种子20261007。只训练解码器和头时学习率10⁻⁴；全参数为10⁻⁵。损失为相对时点SmoothL1加0.25日均绝对误差和0.25峰值绝对误差；重复目标时间按逆出现次数降低曲线训练权重。开发训练标签截至第14天，最终截至第22天。512点历史满足初始14天的监督长度，选择依据不是测试名次。')
p('开发综合代价是24小时曲线、24小时峰值、7天曲线、7天日均各自除以当前推荐，再等权平均。预先要求综合代价≤0.98且任何主项比值≤1.02；否则保留旧推荐。8项均未通过。排名及参数先冻结，之后才评分测试；没有根据本次测试较好分数改选配置。')
h('预测区间覆盖与补充评价')
iv=[v for v in a['intervals'] if (v.get('model','').startswith('TTM512 multitask') and v['horizon_days']==7) or '当前综合推荐' in v.get('method','') or v.get('method','').startswith('R33 ')];rr=[]
for v in iv:
 name=v.get('method',v.get('model'))
 if '当前综合推荐' in name:name=name.replace('当前综合推荐','此前推荐')
 if name.startswith('R33 '):name=('新24小时推荐 TinyTimeMixer稀疏峰值10% 对称80%经验区间' if v['horizon_days']==1 else'新7天推荐 TinyTimeMixer日均保持形状10% 非对称80%经验区间')
 if name.startswith('R31 '):name='新7天推荐 TinyTimeMixer日均保持形状10%经验区间'
 if v.get('model'):name='IBM TinyTimeMixer R2 多任务训练 '+('解码器和输出头' if v['model'].endswith('head') else'全参数')
 name+=f'  {v["horizon_days"]}天';rr.append([name,80,f'{v["coverage_percent"]:.3f}',n(v['mean_width_normalized']),n(v['winkler80'])])
table(['区间方法','标称覆盖 %','实际覆盖 %','平均宽度↓','Winkler80↓'],rr,[5.25,1.12,1.18,1.40,1.40],10)
p('点预测保持冻结。第33轮仅在开发期逐起报留一核查三种区间校准，先要求覆盖距80%不超过3个百分点，再选Winkler80较低者：24小时采用逐PDU绝对残差80分位的对称区间，7天采用10%和90%分位非对称区间。测试覆盖分别77.812%和78.905%，比旧推荐更接近80%且宽度更小；Winkler80仍略高，区间并非所有补充指标都更优。7天开发窗口重叠，没有独立覆盖保证。原生TiRex的56.6%仅历史引用，不能按测试再换校准。')
rr=[]
for H in [1,7]:
 for r in [recommended(H),row(H,'r30:TTM512 multitask head'),row(H,'r30:TTM512 multitask all')]:rr.append([f'{H}天 '+r['report_name'],f'{r["peak_time_mae_minutes"]:.2f}',n(r['absolute_bias_percent']),n(r['cella5_wape']),n(r['other52_wape']),f'{r["improved_pdus"]}/57  {r["improved_cells"]}/10  {r["improved_origins"]}/{9 if H==1 else 3}'])
table(['代表模型','峰值时刻偏差\n分钟↓','绝对总偏差 %↓','a5 WAPE %↓','其余52 WAPE %↓','改善基线的\n域 场站 起报'],rr,[4.18,1.25,1.20,1.18,1.18,1.36],9.5)
h('采用结论数据限制与复现入口')
p('结论一：曲线评价已确认与9.28明确公式一致；日均口径已统一复算，但原文缺少完整实现，历史末位仍不能全部核验。严格的新旧比较以本次v2保存数组复算表为准，历史引用不混入排序。Ridge和LightGBM对应目前已跑的具体参数，不声称逐值复现原文未交付的配置。Toto及直接日均Chronos的原预测数组未保存，不补编分数。旧模型和失败结果保留，原冻结实验不改写。')
p('结论二：两条新建模路线和新架构训练已真实完成，第31/32轮重建继续取得多项改善。按开发冻结和微小退步范围，24小时第32轮、7天第31轮替代旧综合推荐。仍存在少量退步和其他单项更低的模型，不声称所有指标第一。原第8b轮相比第5轮仅约0.000052个百分点改善、探索区间跨零，不能说显著或大幅。')
p('结论三：现有训练资料只有31天，开发7天仅两个起报；同月测试已经多轮查看，无法作为真正独立确认。新方法失败只说明本次配置和数据条件下没有综合优势，不能证明所有新架构都不适用、参数永远无法提高，或已经达到理论上限。公开预训练资料与本数据是否重叠未能完全排除。进一步确认全面提升需要独立月份或未参与探索的同类场站数据。')
p('此前24小时骨干参数：Chronos-2与TiRex-2各25%，TimesFM3 50%；Ridge残差正则100000、权重50%；TimesFM3末层学习率10⁻⁶、100步、教师约束0.1，候选与原方案按开发冻结50%组合。此前7天骨干参数：TimesFM3直接672点，Ridge1000修正日均和前三阶傅里叶形状，预测总量约束25%；HGB7叶60轮、学习率0.03、最小叶128、L2为10，配Ridge1000稀疏峰值头；日均增益1.5，新修正仅25%，保持a5并作cell预测总量约束。')
p('数据实体全部在D盘，桌面目录仍为兼容junction。Google功率为归一化利用率，不是kW。官方ClusterData cell-a 12085文件必要列已读完且审计通过，剩余下载0；不是全部八个集群或全部字段。真实CPU单位NCU，只覆盖a集群5域；本次新架构统一采用功率历史，未把未来CPU真值接入预测。条件单调或预测总量约束均不等于物理能量守恒。')
p('数据许可CC BY4.0；TimesFM3权重仅限非商业非生产研究。TinyTimeMixer官方模型卡和许可、固定版本代码及权重SHA保存在round29；CPU运行库独立安装，原CUDA环境保留。新研究预测为2019年5月末历史复现，6月1至7无真值，不能作为今天负荷。')
p('顶层“模型各项指标完整结果.csv”保留86条记录，“各任务单项最优结果.csv”给出逐项最优，“模型预测区间结果.csv”保留全部区间评价。19号入口仍对应当前推荐，22号为原始TinyTimeMixer训练研究输出，23号为新重建方案历史复现。评价审查、真实评分、冻结记录、官方来源和独立核查分别保存在round29至round33。')
p('官方新模型来源：https://huggingface.co/ibm-granite/granite-timeseries-ttm-r2；方法论文：https://arxiv.org/abs/2401.03955；实现：https://github.com/ibm-granite/granite-tsfm。')
doc.save(out);print(json.dumps(dict(document=str(out),sha256=hashlib.sha256(out.read_bytes()).hexdigest()),ensure_ascii=False))
