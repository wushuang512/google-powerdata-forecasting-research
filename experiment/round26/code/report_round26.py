from pathlib import Path
import sys,json,hashlib,shutil,ast
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
T=Path(sys.argv[1]);B=T/'experiment';R=B/'round26';Q=R/'results'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
a=read(B/'round20/results/original_metrics_complete.json');v=read(B/'round20/results/original_metrics_final_verification.json');r=read(Q/'summary.json');cmp=read(Q/'recommendation_comparison.json');dep=read(Q/'exports/verification.json');assert v['all_passed'] and dep['all_passed'] and len(cmp['improved'])==8 and r['selection']['horizons']['7']['chosen_alpha']==.25
for n in range(20,27):assert read(B/f'round{n}/results/independent_verification.json')['all_passed']
out=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx'
if '--validate-only' in sys.argv:print(json.dumps(dict(all_original_metrics=True,actual_rounds=list(range(20,27)),new_recommendation='round26',improved=8,minor_regressions=5)));sys.exit(0)
backup=R/'docs/before_report_update';backup.mkdir(parents=True,exist_ok=True)
if '--repair-layout' not in sys.argv:
 assert not (backup/out.name).exists(),'Never append twice'
 shutil.copyfile(out,backup/out.name);shutil.copyfile(B/'round19/docs/report_qa/verification.json',backup/'prior_visual_verification.json')
doc=Document(backup/out.name)
tree=ast.parse(Path(__file__).with_name('report_revision.py').read_text(encoding='utf-8'));node=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='table');exec(compile(ast.Module(body=[node],type_ignores=[]),'table_helper','exec'))
for sec in doc.sections:
 for p in sec.footer.paragraphs:
  for run in p.runs:run.text=run.text.replace('2026年10月5日','2026年10月6日')
doc.paragraphs[1].text='本汇报分别评价未来24小时曲线、未来7天曲线与每日平均负荷。10月6日完成第20至26轮优化，并补齐9.28原报告的全部指标。按用户认可的多项改善与极小取舍标准，第26轮作为当前7天推荐；未声称每项均第一。'
current7=r['scores']['7']['selected'];current1=r['scores']['1']['selected']
for i,k in [(2,'wape'),(3,'daily_wape')]:
 doc.tables[0].rows[i].cells[1].text='TimesFM3＋Ridge时频校准\nHGB时点残差＋稀疏峰值修正\n固定25%保守叠加'+(' 后逐日平均' if i==3 else '')
 doc.tables[0].rows[i].cells[2].text=f'{current7[k]:.6f}'
 for c in doc.tables[0].rows[i].cells:
  for p in c.paragraphs:
   p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(1)
   for rr in p.runs:rr.font.size=Pt(10.4)
doc.paragraphs[4].text='24小时保留Chronos-2＋TiRex-2＋TimesFM3三模型固定权重，加Ridge残差校准与TimesFM3末层微调。7天新推荐在第19轮基础上保守叠加双头修正；完整组成和系数见新增章节，编号只用于追溯。'
doc.paragraphs[6].text='七轮继续探索均保留真实记录。第26轮固定25%强度在评分前按开发期及新取舍标准冻结；相较第19轮，14项明细有8项改善、1项持平、5项极微小退步。前期表格保留为历史对照，最新完整口径见末七页。'
doc.paragraphs[7].text='同一推荐程序按已知跨度选分支：24小时沿用第8b轮，7天使用第26轮，7日日均取同一672点曲线平均。当前推荐最佳指多项表现与微小取舍的综合选择，不等于各项指标都第一。功率为归一化利用率，不是kW或kWh。'
def page(t):
 p=doc.add_paragraph(t,'Heading 1');p.paragraph_format.page_break_before=True
def p(t):doc.add_paragraph(t)
def row(H,key):return next(x for x in a['rows'] if x['horizon_days']==H and x['model']==key)
def fmt(k,x):
 return '—' if x is None else f'{x:.6f}' if k in ['mae','rmse','daily_mae','daily_rmse'] else f'{x:.4f}'
short={'current_selected':'当前推荐\nTimesFM3＋Ridge时频\nHGB时点与稀疏峰值25%' ,'r19_selected':'此前推荐\nTimesFM3＋Ridge时频\n全场总量约束25%','tf3_uni21_7':'TimesFM3 单变量\n最大21天历史直接672点','tf3':'TimesFM3 单变量\n最大21天历史直接96点','tirex':'TiRex-2 零样本\n14天历史 分段递推7天','chronos':'Chronos-2 零样本单变量','equal':'Chronos-2 50%\nTiRex-2 50%','ridge':'Ridge回归 alpha1000','lgb':'LightGBM 15叶200树','last_source':'前一天同刻','mean3_source':'前三天同刻均值'}
short1=dict(short,tirex='TiRex-2 零样本\n14天历史直接96点')
short7=dict(short,current_selected='TimesFM3＋Ridge时频校准\nHGB时点与稀疏峰值25% 当前推荐',r19_selected='TimesFM3＋Ridge时频校准\n全场总量约束25% 此前推荐')
page('九月二十八日原始指标与新采用标准')
p('本次对照参考材料原文件“9.28算力中心需求预测实验.docx”，逐项补齐评估指标、误差分布、区间和改善范围。原文件还包含海光寺日用电案例及历史长度、粒度、逐cell留出等消融；本轮聚焦Google，不把这些旧案例写成本轮重跑。')
table(['原文指标或检查','本汇报使用的具体口径'],[
 ['开发及测试WAPE','曲线分别按每个PDU计算绝对误差和／真值和，再对57个域等权平均。24小时开发对照用14／16／18／20；新优化还分别列前后四日。7天开发起报14／15。'],
 ['相对基线误差减少','24小时对前一天同刻，7天对前三天同刻均值。100×（基线WAPE−模型WAPE）／基线WAPE；负值表示误差增加。'],
 ['MAE','有效点评分，Google单位为归一化利用率。原文海光寺MAE为kWh，两个单位不直接比较。'],
 ['时点误差不超过2%及5%占比','有效点APE＝100×绝对误差／真值；分别统计APE≤2及≤5的比例，越高越好。与WAPE不同，不能用100减WAPE代替。'],
 ['时点APE的95分位','在同一任务有效点上计算P95，越低越好；不对各PDU的95分位再平均。'],
 ['每日最大值平均百分比误差','逐PDU逐预测日，在相同有效时刻取实际与预测最大值，求幅值MAPE。原文不评峰值时刻，本次另列分钟偏差。'],
 ['7天每日平均负荷WAPE','真值与预测日均均取同日有效时刻，逐PDU计算比值再等权平均。部署输出是完整96点日均，不使用未来真值掩码。'],
 ['80%区间实际覆盖率','统计有效真值在保存区间内的占比；同时列平均宽度与Winkler80，避免仅靠加宽区间提高覆盖率。'],
 ['改善的域 cell及起报数','与本任务原文基线逐域、逐cell、逐起报比较，7天窗口重叠，不当作3周独立测试。']],[2.12,4.55],10.3)
p('你最新认可的采用标准是多项改善并允许极小差距。第26轮在评分前登记：曲线／日均／峰值WAPE或MAPE退步不超过0.005个百分点，P95不超过0.01个百分点，MAE／RMSE不超过0.2%，2%及5%占比不下降超过0.05个百分点；分组与偏差也设上限。开发期至少两类指标改善后才采用固定25%。这是一项工程选择，不是统计显著性标准。')
page('按原口径补齐二十四小时评价')
p('测试为5月23至31日、57个PDU、9个起报，共49243个有效点。表中开发WAPE对应原文起报14／16／18／20；不会把测试分数放入开发列。主方案名称的完整结构见第2页。')
keys=['current_selected','tf3','tirex','chronos','ridge','lgb','last_source']
table(['明确模型或方法','开发WAPE %','测试WAPE %','相对基线减少 %','MAE'],[[('Chronos-2＋TiRex-2＋TimesFM3\nRidge校准＋TimesFM3末层微调' if k=='current_selected' else short1[k])]+[fmt(m,row(1,k).get(m)) for m in ['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae']] for k in keys],[3.00,.88,.88,1.08,.83],10.0)
table(['模型或方法','APE≤2%占比 %','APE≤5%占比 %','P95 APE %','日最大值MAPE %'],[[('当前三模型校准与末层微调方案' if k=='current_selected' else short1[k])]+[fmt(m,row(1,k)[m]) for m in ['hit2_percent','hit5_percent','p95_ape','peak_mape']] for k in ['current_selected','equal','tirex','last_source']],[2.65,1.03,1.03,.93,1.03],10.2)
c=row(1,'current_selected');p(f'当前方案相较前一天同刻曲线误差减少{c["relative_curve_error_reduction_percent"]:.2f}%，在57／57域、10／10cell和9／9起报上改善。2%误差占比{c["hit2_percent"]:.2f}%，5%占比{c["hit5_percent"]:.2f}%，不表示97.63%的时点“预测正确”。')
p('每日最大值MAPE仍为2.0002%，高于前一天同刻的1.0682%。本轮没有得到足以替换次日方案的开发证据，正式24小时入口保持原模型。完整58行跨任务表另外列RMSE、日均、偏差、分组及低估频率。')
page('按原口径补齐未来七天完整曲线评价')
p('3个起报分别在5月22／23／24日结束，目标为23至29、24至30、25至31日，共114902个有效点。所有对照使用同一目标掩码；原文7天基线为前三天同刻均值，不用前一天同刻替代。')
keys=['current_selected','r19_selected','tf3_uni21_7','tirex','equal','mean3_source']
table(['明确模型或方法','开发WAPE %','测试WAPE %','误差减少 %','MAE'],[[short7[k]]+[fmt(m,row(7,k)[m]) for m in ['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae']] for k in keys],[2.73,.92,.96,1.06,1.00],10.0)
table(['模型或方法','APE≤2%占比 %','APE≤5%占比 %','P95 APE %','日最大值MAPE %'],[[short7[k]]+[fmt(m,row(7,k)[m]) for m in ['hit2_percent','hit5_percent','p95_ape','peak_mape']] for k in ['current_selected','r19_selected','tirex','mean3_source']],[2.65,1.03,1.03,.93,1.03],10.0)
c=row(7,'current_selected');p(f'新推荐曲线相较前三天同刻基线误差减少{c["relative_curve_error_reduction_percent"]:.2f}%，在{c["improved_pdus"]}／57域、{c["improved_cells"]}／10cell及{c["improved_origins"]}／3起报改善。峰值MAPE仍高于前三天均值基线约1.3036%，并非每项对每个基线都胜出。')
p('开发窗口与测试分开报告。骨干最多用21天历史；早期不足21天时，只用当时已存在的数据。预测一次直接输出672点，TiRex-2七天对照采用分段递推。')
page('未来七天每日平均负荷比较')
p('每天的平均负荷从同一条7天完整曲线计算，单个部署起报输出57个PDU×7日＝399个日均。它不另用未来每日真值训练校准，也不把曲线WAPE作为日均WAPE。以下均为3个重叠测试起报的统一复算。')
table(['明确模型或方法','日均WAPE %','误差减少 %','日均MAE','日均RMSE'],[[short7[k]]+[fmt(m,row(7,k)[m]) for m in ['daily_wape','relative_daily_error_reduction_percent','daily_mae','daily_rmse']] for k in keys],[2.73,.92,1.06,.98,.98],10.2)
table(['方案','日均APE≤2%占比 %','日均APE≤5%占比 %'],[[short7[k]]+[fmt(m,row(7,k)[m]) for m in ['daily_hit2_percent','daily_hit5_percent']] for k in ['current_selected','r19_selected','mean3_source']],[3.10,1.78,1.79],10.2)
p(f'新推荐日均WAPE为{c["daily_wape"]:.6f}%，相较前三天同刻均值基线误差减少{c["relative_daily_error_reduction_percent"]:.2f}%。相较第19轮，日均WAPE与MAE改善，RMSE相对上升约0.0178%；按评分前登记的小幅取舍容限采纳，不能称三项都改进。')
p('评估日均时，真值与预测均取同日真值可用时刻，确保缺测不改变比较范围；实际部署没有未来真值掩码，日均取完整96点。单位是归一化功率利用率，不是每日用电量kWh；3个窗口覆盖的日期重叠，不是21个独立日期。')
page('当前推荐方案的组成与真实取舍')
p('完整名称为TimesFM3原生672点预测，加零截距Ridge1000傅里叶残差与预测总量约束，再以固定25%强度叠加HGB时点残差和Ridge稀疏峰值双头修正，并将其中日均修正增益设为1.5。a集群5域保持此前预测，其余按cell保持归一化预测总量。它是按你的多项改善与微小取舍偏好确定的当前推荐方案。')
before=r['scores']['7']['prior'];after=r['scores']['7']['selected']
labels={'wape':'曲线WAPE %','mae':'时点MAE','rmse':'时点RMSE','daily_wape':'日均WAPE %','peak_mape':'日最大值MAPE %','p95_ape':'P95时点APE %','hit2_percent':'APE≤2%占比 %','hit5_percent':'APE≤5%占比 %','daily_mae':'日均MAE','daily_rmse':'日均RMSE','peak_time_mae_minutes':'峰值时刻偏差 分钟','absolute_bias_percent':'绝对总偏差 %','cella5_wape':'a集群5域WAPE %','other52_wape':'其余52域WAPE %'}
table(['指标','第19轮原推荐','第26轮新推荐','新减原'],[[labels[k],fmt(k,before[k]),fmt(k,after[k]),f'{after[k]-before[k]:+.8f}' if k in ['mae','rmse','daily_mae','daily_rmse','absolute_bias_percent'] else f'{after[k]-before[k]:+.6f}'] for k in labels],[2.40,1.35,1.35,1.57],10.0)
p('14项中8项改善、1项持平、5项微小退步，相关指标并非14份独立证据。时点RMSE相对上升约0.0041%，日均RMSE约0.0178%；5%误差占比下降约0.00087个百分点，相当于少1个有效点达标；峰时偏差增加约15.8秒。都在评分前登记的容限内。')
p('时点头为HistGradientBoosting绝对损失、7叶、60轮、学习率0.03、最小叶128、L2为10。峰值头为Ridge1000，只在预测最高16点附近作保持日均的稀疏修正。时点修正10%后，峰值修正分支混入50%；保留a集群并作cell总量一致性。将该新增修正分成日均和日内形状后，日均增益1.5、形状1；最终75%第19轮输出＋25%该候选。所有“均值约束”来自预测，不是未来日均真值，也不是kW能量守恒。')
page('七轮继续优化的真实结果')
p('开发与最终训练的目标截止分别在第14天之前及第22天之前。第22轮六个小型时点头使用原生7天训练缓存，每6点抽取1点，相位随起报和步长固定轮换，并按重复目标日期反频率加权，不将重复标签当作新增独立样本。第20至25轮规则及选择记录全部保留，第26轮才登记微小取舍容限。')
methods={20:'TimesFM3学习峰值头\n全日均值保持投影',21:'TimesFM3学习峰值头\n最高负荷点稀疏投影',22:'HGB原生时点残差\n绝对或平方损失',23:'HGB时点残差\nRidge稀疏峰值双头',24:'双头加历史门控对照\n保留a集群及cell总量',25:'日均与形状增益分开\n修正日均增益1.5',26:'第19轮基础加固定25%\n保守双头与日均修正'}
rows=[]
for n in range(20,27):
 x=read(B/f'round{n}/results/summary.json');s=x['selection']['horizons']['7'];c=x['scores']['7']['candidate'];rows.append([str(n),methods[n],f'{s["chosen_alpha"]:g}']+[fmt(k,c[k]) for k in ['wape','daily_wape','peak_mape']])
table(['轮次','明确新方式','开发采用强度','固定候选测试曲线WAPE %','日均WAPE %','日最大值MAPE %'],rows,[.50,2.40,.70,1.15,.92,1.00],10.0)
p('表中是每轮评分前冻结的单个候选；第20至23轮强度0表示未采用，不能把其诊断低分说成正式升级。第24及25轮开发支持，但测试仍有取舍；第26轮按固定25%保守组合及新容限准入，未在测试上扫描权重。小幅低分不能证明大幅提升，也不能保证剩余指标必然被超过。')
p('此前等待3GiB是重模型加载余量。针对你的内存疑问，对缓存与小型回归头新增了独立保护：启动可用物理2.125GiB、提交6GiB，本实验Windows作业硬限额512MiB，物理低于1.5GiB或提交低于2GiB连续两次只停本实验。第21轮开发实际峰值提交约247MiB，重模型政策未降低。记录了评分前资源策略调整，未关闭你的其他程序。')
page('区间覆盖来源边界与交付入口')
p('80%是标称水平，实际覆盖率必须另测。下表为保存区间的同月有效点评分；经验区间来自开发残差，开发数据同时参与选择，7天窗口重叠，不能承诺跨月80%覆盖。Winkler80同时惩罚宽度和区间外误差，越低越好。')
iv=a['intervals'];names=['24小时当前推荐 经验区间','24小时Chronos-2 原生区间','7天第26轮推荐 经验区间','7天第19轮原推荐 经验区间','7天Chronos-2 原生区间','7天TiRex-2 原文历史引用']
table(['区间方法及来源','实际覆盖 %','平均宽度','Winkler80'],[[names[i],f'{x["coverage_percent"]:.3f}']+[('未保存' if x.get(k) is None else f'{x[k]:.6f}') for k in ['mean_width_normalized','winkler80']] for i,x in enumerate(iv)],[3.20,1.02,1.20,1.25],10.2)
p('TiRex-2的56.6%仅引用9.28原文；原生区间数组没有保存，不能冒充本轮重新测得或补出宽度。新推荐区间实际覆盖81.895%，不是新的独立覆盖保证。区间宽度单位也是归一化利用率。')
p('9.28原表的部分日均与模型版本分数与本次统一复算并不完全一致，例如7天前三天均值原表2.262%、本次2.257474%。原表对应日均数组及实现未随参考材料提供，部分模型配置已变化；不能把所有差额都归因于调优。原文值仅作历史引用，新表全部同目标、同掩码、同方法复算。')
p('实体目录仍为D盘，桌面为兼容junction。新19号推荐入口复现1天或7天：5472条次日点预测、38304条7天点预测、399个7天日均，均已真实导出。未来日期是2019年6月1日PT00时至6月7日PT23时45分，UTC对应6月1日07时至6月8日06时45分；缺6月真值，不能称为今天的实际电负荷。09号次日入口保留，17号旧推荐与18号研究替代用于追溯。')
p('Google只有5月31天；测试多轮查看、7天窗口重叠，且预训练与该月数据的重合未完整排除。当前结果仅为探索性研究推荐，无独立月份确认，不能声称全指标第一、大幅或显著提高。官方cell-a12085对象必要列已完整读取审计，剩余下载0；CPU为NCU，功率非kW。数据CC BY4.0，TimesFM3权重仅非商业非生产研究。')
p('完整表为9月28日指标完整模型对照.csv及9月28日指标预测区间对照.csv；逐轮协议、冻结、训练权重、预测与检查在experiment/round20至round26。最终文件清单记录当前树，PDF与页图仅用于内部排版核查。')
doc.save(out);print(json.dumps(dict(document=str(out),sha256=sha(out),appended_pages_expected=7),ensure_ascii=False))

