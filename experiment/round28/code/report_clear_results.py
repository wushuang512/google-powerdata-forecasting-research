"""Replace the accumulated report with a results-first, nonrepetitive comparison."""
from pathlib import Path
import sys,json,hashlib,shutil,math
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
T=Path(sys.argv[1]).resolve();B=T/'experiment';R=B/'round28';Q=R/'results'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
a=read(Q/'clear_comparison.json');assert len(a['rows'])==62 and all(read(B/f'round{n}/results/independent_verification.json')['all_passed'] for n in [27,28]);assert a['recommendation']['recommended7']=='round26'
out=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx';backup=R/'docs/before_clear_report'
if '--validate-only' in sys.argv:print(json.dumps(dict(rows=62,new_rounds_verified=True,recommended24='round8b',recommended7='round26')));sys.exit(0)
backup.mkdir(parents=True,exist_ok=True)
if not (backup/out.name).exists():shutil.copyfile(out,backup/out.name);shutil.copyfile(B/'round26/docs/report_qa/verification.json',backup/'prior_visual_verification.json')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(11.69);sec.page_height=Inches(8.27);sec.top_margin=Inches(.52);sec.bottom_margin=Inches(.50);sec.left_margin=sec.right_margin=Inches(.65)
for k in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[k];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.08;doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(21);doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(0);doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
doc.styles['Heading 2'].font.size=Pt(11.5);doc.styles['Heading 2'].paragraph_format.space_before=Pt(5);doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for k in ['Normal','Title','Heading 1','Heading 2']:
 pp=doc.styles[k].element.get_or_add_pPr();snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('2026年10月6日  模型指标与优化结果  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def p(t):return doc.add_paragraph(t)
def h(t,new=True):
 pp=doc.add_paragraph(t,'Heading 1');pp.paragraph_format.page_break_before=new;return pp
def table(headers,rows,widths,size=9.1,best_cells=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for j,v in enumerate(headers):t.rows[0].cells[j].text=str(v)
 for row in rows:
  cells=t.add_row().cells
  for j,v in enumerate(row):cells[j].text=str(v)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
   cp.append(bd);mar=OxmlElement('w:tcMar')
   for side,v in [('top',45),('bottom',45),('left',65),('right',65)]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');mar.append(e)
   cp.append(mar);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCEAF2' if i==0 else ('E9F4E5' if best_cells and (i-1,j) in best_cells else'F7F8FA' if i%2==0 else'FFFFFF'));cp.append(shade)
   for pp in c.paragraphs:
    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=Pt(size*1.14);snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp._p.get_or_add_pPr().append(snap)
    for run in pp.runs:run.font.size=Pt(size);run.bold=i==0 or bool(best_cells and (i-1,j) in best_cells)
 return t
def group(H):return [r for r in a['unique_rows'] if r['horizon_days']==H]
def recommend(H):return next(r for r in group(H) if r['status']=='当前综合推荐')
def value(k,x):return '—' if x is None else f'{x:.6f}' if k in ['wape','daily_wape','peak_mape','mae','rmse','daily_mae','daily_rmse'] else f'{x:.3f}'
def star(H,k,v):
 higher=k.startswith('hit') or k.startswith('daily_hit') or k.startswith('relative_');best=(max if higher else min)(r[k] for r in group(H));return abs(v-best)<1e-10
def label(r):
 status={'已评估对照':'对照','当前综合推荐':'综合推荐','此前综合推荐':'此前推荐','未采用诊断':'未采用诊断','日均峰值研究替代':'研究替代','开发冻结研究替代':'新研究替代','开发未支持 未采用诊断':'新未采用诊断'}[r['status']]
 return r['report_name']+f'\n{status}  胜基线 {r["improved_pdus"]}/57域 {r["improved_cells"]}/10cell {r["improved_origins"]}/{r["total_origins"]}起报'
doc.add_paragraph('Google PowerData 2019 模型指标比较与优化结论','Title')
p('比较未来24小时完整曲线、未来7天完整曲线及每日平均负荷。先给结果与最优标记，再说明本轮优化和采用依据。星号与浅绿色标出各任务已列对照中的单项最优，包含未采用诊断；它们与综合推荐分别列示。')
summary=[]
for task,H,k in [('24小时完整曲线',1,'wape'),('7天完整曲线',7,'wape'),('7天每日平均负荷',7,'daily_wape')]:
 rec=recommend(H);win=min(group(H),key=lambda r:r[k]);summary.append([task,rec['report_name'],value(k,rec[k]),value(k,win[k]),win['report_name']+'\n'+win['status']])
table(['预测任务','当前综合推荐','推荐WAPE %','单项最低WAPE %','实测单项最低的方法及身份'],summary,[1.35,3.32,1.10,1.20,3.42],10.5)
p('综合推荐仍为24小时Chronos-2＋TiRex-2＋TimesFM3加Ridge校准和末层微调，7天TimesFM3＋Ridge时频校准加HGB时点与稀疏峰值25%。单项最低的未采用诊断没有通过当轮开发采用规则，旧冻结选择保留。当前没有一个方法在全部指标中都第一。')
p('最新两轮已实际完成。LightGBM稳健时点方案降低7天日均和峰值，但曲线、RMSE及5%占比变差；限幅日均形状方案的日均、峰值时刻和偏差改善，但曲线、P95和峰值幅值变差。二者作为研究替代保留，未替换综合推荐。')
p('读表口径：WAPE、MAPE、P95及占比为%；MAE和RMSE为归一化功率利用率。↓越低越好，↑越高越好。开发列缺少对应保存数组时显示—。胜基线范围分别为PDU、cell和起报数；完全相同预测已合并，62个原始条目仍全部保存在CSV。')
p('24小时测试9个起报、57域、49243个有效点；7天测试3个重叠起报、114902个有效点。均为2019年5月回顾，缺独立确认月。表中的最优是同一评价口径的描述性结果，不能视作已证明跨月显著优势。')
curve_fields=['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae','rmse','peak_mape','p95_ape','hit2_percent','hit5_percent']
headers=['明确模型及身份\n胜对应基线的范围','开发\nWAPE %↓','测试\nWAPE %↓','误差减少\n%↑','MAE↓','RMSE↓','日最大值\nMAPE %↓','P95\nAPE %↓','APE≤2%\n占比 %↑','APE≤5%\n占比 %↑']
widths=[3.34,.78,.78,.77,.80,.80,.84,.72,.78,.78]
for H,chunk in [(1,12),(7,11)]:
 rows=group(H)
 starts=[0,12] if H==1 else list(range(0,len(rows),chunk))
 for part,start in enumerate(starts,1):
  h(('未来24小时完整曲线' if H==1 else'未来7天完整曲线')+f'模型结果 第{part}部分')
  if part==1:p('按测试曲线WAPE从低到高排列。24小时基线为前一天同刻，7天基线为前三天同刻均值。开发WAPE用原9.28对应起报；测试单项最低不用于修改冻结参数。')
  data=[];highlight=set()
  end=starts[part] if part<len(starts) else len(rows)
  for i,r in enumerate(rows[start:end]):
   row=[label(r)]
   for j,k in enumerate(curve_fields,1):
    vv=r.get(k);yes=vv is not None and k!='development_wape_original_grid' and star(H,k,vv);row.append(value(k,vv)+(' ★' if yes else''))
    if yes:highlight.add((i,j))
   data.append(row)
  table(headers,data,widths,8.7,highlight)
daily_fields=['daily_wape','relative_daily_error_reduction_percent','daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent']
rows=sorted(group(7),key=lambda r:r['daily_wape'])
for part,start in enumerate(range(0,len(rows),16),1):
 h(f'未来7天每日平均负荷模型结果 第{part}部分')
 if part==1:p('每个日均来自同一672点曲线。按日均WAPE从低到高排列；真值和预测使用同一有效时刻评分。部署日均取完整96点，单位不是kWh。')
 data=[];highlight=set()
 for i,r in enumerate(rows[start:start+16]):
  status={'已评估对照':'对照','当前综合推荐':'综合推荐','此前综合推荐':'此前推荐','未采用诊断':'未采用','日均峰值研究替代':'研究替代','开发冻结研究替代':'新研究替代','开发未支持 未采用诊断':'新未采用'}[r['status']]
  row=[r['report_name']+'（'+status+'）']
  for j,k in enumerate(daily_fields,1):
   yes=star(7,k,r[k]);row.append(value(k,r[k])+(' ★' if yes else''))
   if yes:highlight.add((i,j))
  data.append(row)
 table(['明确模型及身份','日均\nWAPE %↓','日均误差\n减少 %↑','日均MAE↓','日均RMSE↓','日均APE≤2%\n占比 %↑','日均APE≤5%\n占比 %↑'],data,[4.01,1.00,1.03,1.00,1.00,1.10,1.10],9.4,highlight)
h('峰值时刻偏差与分组表现')
p('以下为代表方法的额外指标，完整62条记录的所有字段见“模型各项指标完整结果.csv”；所有单项最优方法另列于“各任务单项最优结果.csv”。表中★仍按完整对照集合判断。')
extra=['peak_time_mae_minutes','absolute_bias_percent','cella5_wape','other52_wape','peak_underestimate_percent']
for H in [1,7]:
 doc.add_paragraph('24小时曲线' if H==1 else'7天曲线','Heading 2');g=group(H);chosen=[]
 for r in [recommend(H),min(g,key=lambda r:r['wape']),min(g,key=lambda r:r['peak_mape']),next(r for r in g if r['model']=='tirex'),next(r for r in g if (('last' in r['aliases'] or 'last_source' in r['aliases']) if H==1 else r['model']=='mean3_source'))]:
  if r['model'] not in [x['model'] for x in chosen]:chosen.append(r)
 rows=[];hi=set()
 for i,r in enumerate(chosen):
  row=[r['report_name']+'\n'+r['status']]
  for j,k in enumerate(extra,1):
   yes=k!='peak_underestimate_percent' and star(H,k,r[k]);row.append(value(k,r[k])+(' ★' if yes else''))
   if yes:hi.add((i,j))
  rows.append(row)
 table(['明确模型及身份','峰值时刻偏差\n分钟↓','绝对总偏差\n%↓','a集群5域\nWAPE %↓','其余52域\nWAPE %↓','峰值低估频率\n% 仅诊断'],rows,[4.25,1.25,1.22,1.22,1.22,1.22],9.5,hi)
h('预测区间实测与本轮优化取舍')
iv=a['intervals'];table(['区间方法及来源','标称覆盖 %','实际覆盖 %','平均宽度↓','Winkler80↓'],[[x['method'],80,f'{x["coverage_percent"]:.3f}',('未保存' if x.get('mean_width_normalized') is None else f'{x["mean_width_normalized"]:.6f}'),('未保存' if x.get('winkler80') is None else f'{x["winkler80"]:.6f}')] for x in iv],[5.70,1.10,1.10,1.25,1.20],10)
p('覆盖率应接近标称80%，不是越高越好；同时看宽度与Winkler80。TiRex-2的56.6%仅引用9.28原文，原区间数组未保存。其余按保存数组实测，经验区间没有独立覆盖保证。')
newrows=[]
for n in [26,27,28]:
 rr=read(B/f'round{n}/results/summary.json');c=rr['scores']['7']['selected'];newrows.append([('TimesFM3＋Ridge时频及HGB稀疏峰值25%' if n==26 else'TimesFM3＋Ridge＋HGB再加LightGBM L1时点10%' if n==27 else'TimesFM3＋Ridge＋HGB再加限幅日均形状10%'),rr['selection']['horizons']['7']['chosen_alpha'],f'{c["wape"]:.6f}',f'{c["daily_wape"]:.6f}',f'{c["peak_mape"]:.6f}',('当前综合推荐' if n==26 else'日均峰值改善 曲线分布退步' if n==27 else'日均时刻偏差改善 曲线峰值退步')])
table(['明确方案','本轮权重','曲线WAPE %','日均WAPE %','峰值MAPE %','实际取舍结论'],newrows,[4.42,.80,1.02,1.02,1.02,2.05],9.7)
p('第27轮4个LightGBM开发模型已真实拟合，最终仅重训冻结所需的L1 15叶头。第28轮复用该头，测试4个限幅与形状配置；两轮次日均未获得开发支持。参数先由开发期冻结，旧选择不按测试较低分改写。')
h('方法参数评价口径与研究结论')
p('24小时综合推荐：Chronos-2与TiRex-2各14天历史单变量，TimesFM3为14天同cell联合目标及production功率代理；基础权重25%／25%／50%。加Ridge残差校准，正则100000、权重50%；再加TimesFM3输出头及末层微调，学习率10⁻⁶、100步、教师约束0.1。该修正候选与先前模型按开发冻结50%混合，保留分阶段0至1截断。')
p('7天综合推荐：TimesFM3最多21天历史直接672点，零截距Ridge1000校准日均与前三阶傅里叶形状，再作全场归一化预测总量约束25%。叠加HGB绝对损失7叶时点头与Ridge1000稀疏峰值头，日均修正增益1.5，最终新增修正只占25%；a5保持原输出，其余按cell总量一致性。HGB为60轮、学习率0.03、最小叶128、L2为10；峰值头只改最高16个预测点附近。')
p('其他模型：Chronos-2三分支分别是production协变量微调、同cell目标协变量零样本、总量与分量微调；加TiRex-2日历，非零权重约36.748%／34.564%／27.891%／0.796%，再作3点平滑。峰值版加前日峰值形状校正。7天专用Chronos LoRA与分量约束组合还含自建Patch Transformer和历史曲线，具体冻结系数及源数组保存在原实验目录。')
p('最新校准：LightGBM L1 15叶、80树、学习率0.03、最小叶128、L2为20。第28轮将残差分为日均与25%日内形状，原始修正限于历史7日均值的0.5%，保持a5和cell预测总量，再以10%加入原推荐。该限制是预测正则，不是已知容量下的物理能量守恒。两轮均只用训练及开发信息选模型，保留真实负面结果。')
p('曲线WAPE先逐PDU计算绝对误差和／真值和，再对57域等权平均；MAE、RMSE、P95及2%／5%占比在同一有效点上计算。每日最大值MAPE用有效时刻的最大值幅值；日均WAPE先逐PDU比值再平均。24小时开发起报14／16／18／20、测试22至30；7天开发14／15、测试22／23／24，均指当日结束时起报。')
p('结论：综合推荐保留24小时第8b轮、7天第26轮。第27/28轮有专项进步，没有更全面的精度升级。单项最低可能属于未采用诊断或简单基线，不能把曲线最好等同峰值最好。5月测试被多轮查看，7天窗口重叠且只有31天历史，缺独立确认月；当前最优仅是表内实测比较，不能宣称大幅或独立显著提升，也未证明理论精度上限。')
p('数据与产出实体仍在D盘，桌面为兼容junction。官方cell-a12085个Parquet必要列读取及审计已完成，剩余下载0；真实CPU为NCU且仅覆盖a5域，功率为归一化利用率。Google数据CC BY4.0，TimesFM3仅非商业非生产研究。19号入口对应当前综合推荐；新研究替代输出在round27/28。June1至7的预测是2019年5月末历史复现，无六月真值，不是今天负荷。全部对照和单项最优CSV保存在顶层，过程与核查证据保留在实验目录。')
doc.save(out);print(json.dumps(dict(document=str(out),sha256=sha(out),results_first=True,history_not_appended=True),ensure_ascii=False))
