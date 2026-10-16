from pathlib import Path
import sys,json,hashlib,csv,datetime,subprocess,os,tomllib
sys.dont_write_bytecode=True
import numpy as np
import pandas as pd
from docx import Document
from pypdf import PdfReader
T=Path(sys.argv[1]).resolve();B=T/'experiment';R=B/'round28';Q=R/'results'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
a=read(Q/'clear_comparison.json');m=read(T/'交付文件清单.json');qa=read(R/'docs/report_qa/verification.json');checks={};doc=Document(T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx')
checks['two_new_real_experiments_finished']=all(read(B/f'round{n}/results/resource_status.json')['stage']=='completed' and read(B/f'round{n}/results/independent_verification.json')['all_passed'] for n in [27,28])
checks['frozen_choices_unchanged']=all(sha(B/f'round{n}/results/selection.json')==(B/f'round{n}/results/selection.sha256').read_text().strip() for n in [26,27,28])
checks['current_recommendation_and_negative_conclusions_clear']=a['recommendation']['recommended24']==m['recommended24']=='round8b' and a['recommendation']['recommended7']==m['recommended7']=='round26' and not a['recommendation']['new_universal_winner'] and not m['all_metric_goal_met'] and not m['independent_confirmation']
checks['all62_records_and58_unique_predictions']=len(a['rows'])==62 and len(a['unique_rows'])==58 and sum(r['horizon_days']==1 for r in a['unique_rows'])==26 and sum(r['horizon_days']==7 for r in a['unique_rows'])==32 and len(list(csv.DictReader((T/'模型各项指标完整结果.csv').open(encoding='utf-8-sig'))))==62
checks['csvs_and_original_reference_bound']=sha(T/'模型各项指标完整结果.csv')==a['csv_sha256'] and sha(T/'各任务单项最优结果.csv')==a['winners_csv_sha256'] and sha(T/'模型预测区间结果.csv')==a['interval_csv_sha256'] and sha(T/'参考材料/9.28算力中心需求预测实验.docx')==a['source_reference']['reference_sha256'] and len(a['source_reference']['metrics'])==11 and sha(B/'round20/results/original_metrics_complete.json')==a['original_metrics_json_sha256']
sys.path.insert(0,str(R/'code'));sys.path.insert(0,str(B/'code'));from source_metrics import evaluate;import common;d=common.load()
checks['all_sources_and_metrics_recomputed']=True
for row in a['rows']:
 H=row['horizon_days'];origins=list(range(22,31)) if H==1 else[22,23,24];src=row['source']
 if src is None:p=common.baseline(d,origins,H,kind='last' if row['model']=='last_source' else'mean3')
 else:
  checks['all_sources_and_metrics_recomputed'] &= sha(B/src['path'])==src['sha256']
  with np.load(B/src['path']) as z:p=z[src['array']].astype('float64')
 sc=evaluate(common,d,origins,p,H);checks['all_sources_and_metrics_recomputed'] &= hashlib.sha256(p.tobytes()).hexdigest()==row['prediction_digest'] and all(sc[k]==row[k] if isinstance(sc[k],str) else abs(sc[k]-row[k])<1e-10 for k in sc)
checks['exact_duplicate_merge_and_all_aliases_retained']=all(len({r['prediction_digest'] for r in a['unique_rows'] if r['horizon_days']==H})==len([r for r in a['unique_rows'] if r['horizon_days']==H]) for H in [1,7]) and sum(len(r['aliases']) for r in a['unique_rows'])==62
checks['all_metric_winners_mathematically_correct']=True
for w in a['winners']:
 H=1 if w['task'].startswith('24') else 7;g=[r for r in a['unique_rows'] if r['horizon_days']==H];k=w['metric'];best=(max if w['direction']=='越高越好' else min)(r[k] for r in g);keys=[r['model'] for r in g if abs(r[k]-best)<1e-10];checks['all_metric_winners_mathematically_correct'] &= abs(best-w['best_value'])<1e-12 and keys==w['keys']
def val(k,v):return '—' if v is None else f'{v:.6f}' if k in ['wape','daily_wape','peak_mape','mae','rmse','daily_mae','daily_rmse'] else f'{v:.3f}'
def mark(H,k,v):
 g=[r for r in a['unique_rows'] if r['horizon_days']==H];higher=k.startswith('hit') or k.startswith('daily_hit') or k.startswith('relative_');best=(max if higher else min)(r[k] for r in g);return abs(v-best)<1e-10
fields=['development_wape_original_grid','wape','relative_curve_error_reduction_percent','mae','rmse','peak_mape','p95_ape','hit2_percent','hit5_percent'];checks['all_word_curve_cells_and_stars_correct']=len(doc.tables)==12
for H,positions in [(1,[1,2]),(7,[3,4,5])]:
 g=[r for r in a['unique_rows'] if r['horizon_days']==H];cells=[row.cells for pos in positions for row in doc.tables[pos].rows[1:]];checks['all_word_curve_cells_and_stars_correct'] &= len(cells)==len(g)
 for row,cells_row in zip(g,cells):
  for j,k in enumerate(fields,1):
   v=row.get(k);star=v is not None and k!='development_wape_original_grid' and mark(H,k,v);checks['all_word_curve_cells_and_stars_correct'] &= cells_row[j].text==val(k,v)+(' ★' if star else'')
checks['all_word_daily_cells_and_stars_correct']=True;g=sorted([r for r in a['unique_rows'] if r['horizon_days']==7],key=lambda r:r['daily_wape']);cells=[row.cells for pos in [6,7] for row in doc.tables[pos].rows[1:]]
for row,cells_row in zip(g,cells):
 for j,k in enumerate(['daily_wape','relative_daily_error_reduction_percent','daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent'],1):checks['all_word_daily_cells_and_stars_correct'] &= cells_row[j].text==val(k,row[k])+(' ★' if mark(7,k,row[k]) else'')
checks['overview_recommendation_and_lowest_separate']=True
for i,(H,k) in enumerate([(1,'wape'),(7,'wape'),(7,'daily_wape')],1):
 g=[r for r in a['unique_rows'] if r['horizon_days']==H];cur=next(r for r in g if r['status']=='当前综合推荐');best=min(r[k] for r in g);cells=doc.tables[0].rows[i].cells;checks['overview_recommendation_and_lowest_separate'] &= cells[2].text==val(k,cur[k]) and cells[3].text==val(k,best) and '未采用诊断' in cells[4].text
checks['all11_final_pages_actually_reviewed']=qa['all_passed'] and qa['all_pages_visually_reviewed'] and qa['layout_passed'] and qa['pages']==11==len(PdfReader(qa['pdf']).pages)==m['word_pages'] and qa['page_numbers_reviewed']==list(range(1,12)) and sha(T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx')==qa['document_sha256']==m['report_sha256'] and all(sha(x['file'])==x['sha256'] for x in qa['images']) and sha(qa['pdf'])==qa['pdf_sha256']
checks['no_old_history_appended_and_no_title_border']=len(doc.paragraphs)<60 and doc.paragraphs[0].style.name=='Title' and not list(doc.styles.element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr'))
checks['new_exports_bounds_times_and_daily_means']=True
for n in [27,28]:
 ex=read(B/f'round{n}/results/exports/verification.json')
 for H in [1,7]:
  df=pd.read_csv(ex[f'file{H}']);daily=pd.read_csv(ex[f'daily_file{H}']);v=df[['forecast_normalized_power','lower80_empirical','upper80_empirical']];ts=pd.to_datetime(df['timestamp_UTC'],utc=True);df['date_PT']=df['timestamp_PT'].str[:10];mean=df.groupby(['pdu','cell','date_PT'],sort=False)['forecast_normalized_power'].mean().reset_index(name='daily_mean_normalized_power');expected=pd.date_range('2019-06-01T07:00:00Z',periods=H*96,freq='15min')
  ok=len(df)==57*H*96 and df['pdu'].nunique()==57 and not df.duplicated(['pdu','timestamp_PT']).any() and np.isfinite(v).all().all() and v.ge(0).all().all() and v.le(1).all().all() and (df['lower80_empirical']<=df['forecast_normalized_power']).all() and (df['forecast_normalized_power']<=df['upper80_empirical']).all() and pd.DatetimeIndex(sorted(ts.unique())).equals(expected) and len(daily)==57*H and np.allclose(mean['daily_mean_normalized_power'],daily['daily_mean_normalized_power'],rtol=0,atol=1e-15) and sha(ex[f'file{H}'])==ex[f'sha{H}']
  checks['new_exports_bounds_times_and_daily_means'] &= bool(ok)
env=dict(os.environ,PYTHONDONTWRITEBYTECODE='1',PYTHONIOENCODING='utf-8');checks['actual_current_and_new_alternative_entry_runs']=True
entries=[('19_7天当前推荐模型历史复现.cmd',26),('20_7天稳健时点校准研究对照历史复现.cmd',27),('21_7天限幅日均形状研究对照历史复现.cmd',28)]
for filename,n in entries:
 r=subprocess.run(['cmd.exe','/d','/c',str(T/filename)],capture_output=True,text=True,encoding='utf-8',env=env,check=True);info=json.loads(r.stdout);checks['actual_current_and_new_alternative_entry_runs'] &= info['curve_file']==str(B/f'round{n}/results/exports/selected_7day_curve.csv') and info['historical_replay']
checks['only_useless_own_layout_intermediates_removed']=read(Q/'layout_intermediate_cleanup.json')['Files']==48 and read(Q/'layout_intermediate_cleanup.json')['FinalVerifiedPagesPreserved'] and (R/'docs/before_clear_report/Google PowerData 2019 三项预测任务与模型比较汇报.docx').exists() and len(qa['images'])==11
checks['official_source_complete_no_new_download']=read(T/'google-clusterdata-2019/results/complete_projection_audit.json')['source_files']==12085 and m['actual_download_remaining']==0
checks['automation_paused']=tomllib.loads(Path('${USER_HOME}/.codex/automations/google/automation.toml').read_text(encoding='utf-8'))['status']=='PAUSED'
cmd="$p=Get-CimInstance Win32_Process | Where-Object { $_.Name -match '^python' -and $_.CommandLine -match 'GooglePowerData2019' -and $_.CommandLine -match '(run_.*guarded|extract_cpu|robust_point_experiment|bounded_level_experiment)' }; if($p){$p | Select-Object ProcessId,CommandLine | ConvertTo-Json -Compress}else{'[]'}";r=subprocess.run(['powershell.exe','-NoProfile','-Command',cmd],capture_output=True,text=True,check=True);checks['no_live_owned_research_workers']=json.loads(r.stdout)==[]
checks['inventory_unique_exists_and_bytes']=len(m['all_files'])==len({r['path'] for r in m['all_files']})==m['files'] and all((T/r['path']).is_file() and (T/r['path']).stat().st_size==r['bytes'] for r in m['all_files']) and sum(r['bytes'] for r in m['all_files'])==m['bytes']
actual={str(p.relative_to(T)) for p in T.rglob('*') if p.is_file() and p!=T/'交付文件清单.json'};checks['inventory_no_omissions_and_root_d']=actual=={r['path'] for r in m['all_files']} and str(T).startswith('D:')
out=dict(all_passed=all(bool(v) for v in checks.values()),checks={k:bool(v) for k,v in checks.items()},checked_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),files=m['files'],bytes=m['bytes'],word_pages=11,raw_records=62,unique_predictions=58,new_experiments=[27,28],recommended24='round8b',recommended7='round26',new_universal_winner=False,independent_confirmation=False)
assert out['all_passed'],json.dumps(out,ensure_ascii=False)
if '--record-in-manifest' in sys.argv:
 m.update(inventory_verified=True,final_delivery_all_passed=True,final_delivery_check_count=len(checks),independent_final_check_utc=out['checked_utc'],final_delivery_verification=out);(T/'交付文件清单.json').write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False))
