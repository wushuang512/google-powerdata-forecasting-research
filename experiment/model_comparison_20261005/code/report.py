from pathlib import Path
import sys,json,datetime
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
TOP=Path(sys.argv[1]).resolve();B=TOP/'experiment';O=B/'model_comparison_20261005'
S=json.loads((O/'results/summary.json').read_text());A=S['models'];C=json.loads((O/'cleanup_record.json').read_text(encoding='utf-8-sig'));V=json.loads((O/'results/verification.json').read_text())
assert S['all_passed'] and V['all_passed'] and C['state']=='completed'
OUT=TOP/'Google PowerData 2019 最佳模型多维度评估与文件整理汇报.docx'
if '--validate-only' in sys.argv:print('Completed comparison, cleanup and independent verification available');sys.exit(0)
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=sec.right_margin=Inches(.85)
for k in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[k];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
    pp=st.element.find(qn('w:pPr'))
    if pp is not None:
        for ch in list(pp):
            if ch.tag==qn('w:pBdr'):pp.remove(ch)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.12;doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(10);doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;foot.add_run('研究评估 2026年10月5日  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def p(t):return doc.add_paragraph(t)
def h(t,page=False):
    q=doc.add_paragraph(t,'Heading 1');q.paragraph_format.page_break_before=page;return q
def table(headers,rows,widths,size=10.5):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for j,x in enumerate(headers):t.rows[0].cells[j].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for j,x in enumerate(row):cells[j].text=str(x)
    for i,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if i==0:tr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
            pr.append(bd);mar=OxmlElement('w:tcMar')
            for side,n in [('top',65),('bottom',65),('left',85),('right',85)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCEAF2');pr.append(sh)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(1);q.paragraph_format.line_spacing=1.05
                q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in q.runs:r.font.size=Pt(size);r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def name(i):return A[i]['name'].replace('Chronos 与 TiRex 等权组合','Chronos 与 TiRex').replace('Chronos 2 历史协变量','Chronos 2 协变量')
def fmt(x,n=4):return f'{x:.{n}f}'
core=['best','r5','tf3','r4','r3','equal','tirex','chronoscov','tf25','chronos','r2curve','r2peak','ridge','lgb','last']
doc.add_paragraph('Google PowerData 2019 模型评估与文件整理汇报','Title')
p('本汇报面向后续论文与模型使用，统一比较当前正式最佳与此前较好方法的曲线、日均、峰值、误差尾部、时间与场站稳定性，并说明模型保留和删除依据。评估完成于2026年10月5日，全部实体文件仍位于D盘研究目录。')
h('主要结论')
gain=(A['last']['wape']-A['best']['wape'])/A['last']['wape']*100
p(f'继续使用 round8b 作为次日完整曲线的正式模型：全域曲线 WAPE 为 {A["best"]["wape"]:.6f}%，相较前一天同刻误差减少约 {gain:.2f}%。它相较第五轮只有约0.000052个百分点改善；与TimesFM 3的差异约0.003645个百分点，日期配对区间均跨零，不能声称显著领先。')
p(f'峰值目标有不同结论：前一天同刻峰值 MAPE 为 {A["last"]["peak_mape"]:.6f}%，明显低于当前最佳的 {A["best"]["peak_mape"]:.6f}%。当前最佳还不是峰值最优方案，保留简单基线和第二轮峰值侧重模型有实际比较价值。')
table(['主要方法','曲线 WAPE %','日均 WAPE %','峰值 MAPE %'],[[name(i),fmt(A[i]['wape']),fmt(A[i]['daily_wape']),fmt(A[i]['peak_mape'])] for i in ['best','r5','tf3','r4','equal','tirex','r2peak','lgb','last']],[2.35,1.35,1.35,1.52])
p('本次是历史预测的统一复算，没有重新训练、重新选超参数或依据测试名次改写部署。第十三轮候选虽有略低测试WAPE，开发后段不支持采用，因此仍作为未采用对照。完整CPU与MoE候选同样没有改变正式模型。')
p(f'按“太差的没必要保留”的要求，已删除 {C["deleted_files"]} 个文件，逻辑大小 {C["deleted_logical_bytes"]/2**30:.3f} GiB；保留共用权重、当前最佳、强对照、原始数据及必要结果记录。逻辑大小不等于实际磁盘释放量。')
h('评估范围与计算规则',True)
p('统一任务为15分钟粒度的次日96点预测，起报完成日为第22至30日，预测2019年5月23至31日，共57个PDU、9个日期和49,243个有效点。数据先按原规则处理坏测量标记，三条5分钟测量全部有效才构成有效15分钟目标；各方法在同一真值掩码上评分。')
table(['维度','定义与解释'],[
 ['曲线 WAPE','先计算每个PDU有效点绝对误差和除以真实值和，再对57域等权平均。是原实验主指标。'],
 ['日均 WAPE','同一有效时间点上计算真实与预测日均；先每域累计日均绝对误差比，再对域等权平均。'],
 ['峰值 MAPE','在真实有效时间点中比较每日最大值，对所有有效域日等权平均。'],
 ['P95 APE','全体有效点百分比绝对误差的95分位。反映尾部误差，不同于95%置信区间。'],
 ['RMSE 与偏差','RMSE使用归一化功率单位；有符号总偏差为预测减真实，负值表示整体偏低。'],
 ['稳定性','逐日、逐PDU、逐集群和预测步长分段；记录最差日与PDU误差90分位。'],
 ['区间与资源','仅评价已有开发期校准区间；统计主要权重的磁盘长度，不把它当成内存峰值或推理时延。']],[1.3,5.27])
p('归一化功率是容量利用率，不是kW。不同PDU容量未知，等权汇总与把全体点合并的WAPE均不能直接解释为物理电量误差。CPU单位为NCU，不能与归一化功率相加；本次没有把CPU预测当作真实功率。')
p('所有预测数组按原保存顺序与协议读取，并校验来源SHA256；前一天同刻数组与直接按历史构造的基线逐点相同。统一转成64位浮点进行复算，个别旧表日均指标可能有约百万分之一级的舍入差异，本表以统一计算为准。')
p('2019年5月开发与测试已多轮查看，缺少未查看过的确认月份。日期配对bootstrap使用5,000次共同日期重采样，保持同日57域一起抽取；区间是探索性、逐项区间，未做多重比较校正，不代表独立显著性证明。')
p('全部文件实体在 ${PROJECT_ROOT}，桌面同名目录仍为兼容链接。详细对照表、逐日逐域结果、来源哈希、清理计划和记录位于 experiment\\model_comparison_20261005。历史汇报已归档，开始阅读及全目录清单按当前实际保留内容更新。')
h('完整次日比较与误差尾部',True)
table(['方法','WAPE %','日均 %','峰值 %','P95 APE %'],[[name(i),fmt(A[i]['wape']),fmt(A[i]['daily_wape']),fmt(A[i]['peak_mape']),fmt(A[i]['p95_ape'])] for i in core],[2.2,1.04,1.04,1.04,1.25],10.3)
p(f'TimesFM 3的P95 APE为 {A["tf3"]["p95_ape"]:.4f}%，低于当前最佳的 {A["best"]["p95_ape"]:.4f}%；但其日均和峰值指标更差。因而各指标没有统一的绝对赢家，不能只依据曲线名次删除有其他优势的模型。')
p('第三轮主组合使用该轮明确标记的 new_curve 数组；第二轮曲线方案使用第二轮自己的 new_curve。名称相近的历史数组不能互相代替。全部15个主对照和3个近似持平或负面候选均在结果表中保留来源与数组名。')
p('第9、11、12、13、14轮的正式 selected 与当前最佳逐点一致。本报告将其记为保留旧最佳的实验结论，不把重复预测计作五个独立新模型或新的精度提升。第6、7、8轮结果并非与当前最佳逐点相同，其完整分数仍在历史评分目录。')
h('曲线与峰值的取舍',True)
doc.add_picture(str(O/'figures/curve_and_peak.png'),width=Inches(6.57))
p('图中两列指标均越小越好，排序固定为方法展示顺序，不把不同指标合成未经定义的总分。前一天同刻在曲线列较差，在峰值列较好；第二轮峰值方案以一定曲线代价换取较低的峰值误差。')
table(['方法','RMSE 利用率','总偏差 %','峰值低估频率 %','平均峰值低估 %'],[[name(i),fmt(A[i]['rmse'],5),fmt(A[i]['bias_percent']),fmt(A[i]['peak_underestimate_frequency_percent'],2),fmt(A[i]['peak_underestimate_mean_percent'])] for i in ['best','tf3','r2peak','lgb','last']],[1.85,1.16,1.1,1.24,1.22],10.3)
p('平均峰值低估以每个域日真实峰值为分母，正差计入低估，未低估的日计0。低估频率与低估幅度回答不同问题，不能只用频率判断供电风险。这里没有容量参数、成本函数或调度仿真，不能据此给出设备容量结论。')
p(f'当前最佳峰值发生时间平均误差约 {A["best"]["peak_time_mae_minutes"]:.1f} 分钟，前一天同刻约 {A["last"]["peak_time_mae_minutes"]:.1f} 分钟。该诊断按首个最大值计时，平坦曲线和并列峰会造成较大或不稳定的时间误差，不能把它等同于尖峰定位能力。')
h('逐日与场站稳定性',True)
doc.add_picture(str(O/'figures/day_stability.png'),width=Inches(6.57))
table(['对照方法','最佳赢日期','最佳赢PDU','最佳曲线改善95区间 pp'],[[name(i),f'{A[i]["best_wins_dates"]}/9',f'{A[i]["best_wins_domains"]}/57',f'[{A[i]["best_gain_over_candidate_day_bootstrap95_pp"][0]:.5f}, {A[i]["best_gain_over_candidate_day_bootstrap95_pp"][1]:.5f}]'] for i in ['r5','r4','tf3','equal','tirex','last']],[2.15,1.1,1.1,2.22],10.3)
p('改善区间定义为对照WAPE减去当前最佳WAPE，正值有利于当前最佳。赢日期和赢PDU仅为同月描述，不能当作互相独立的57次试验。当前最佳对前一天同刻在9个日期、57个PDU均改善曲线，但对TimesFM 3只赢4个日期，体现优势并不稳定。')
table(['方法','最差日 WAPE %','PDU WAPE P90 %','a集群5域 WAPE %','其余52域 WAPE %'],[[name(i),fmt(A[i]['worst_forecast_day_wape']),fmt(A[i]['domain_wape_p90']),fmt(A[i]['cella5_wape']),fmt(A[i]['other52_wape'])] for i in ['best','tf3','tf25','last']],[1.9,1.16,1.18,1.16,1.17],10.1)
p('TimesFM 2.5在a集群5域的曲线误差略低于当前最佳，但全域较差。不能因该子集排名事后切换部署。逐域、逐集群及0至6、6至12、12至18、18至24小时预测段明细均已导出，便于后续定位误差来源。')
h('预测区间与七天任务',True)
table(['已有经验区间','标称覆盖率 %','实际覆盖率 %','平均区间宽度 利用率'],[[a['model'],80,fmt(a['coverage_percent'],2),fmt(a['mean_width_util'],5)] for a in S['intervals']],[2.1,1.35,1.35,1.77])
p('当前最佳80%经验区间在这九天覆盖约82.78%的有效点，平均宽度约0.05741利用率。宽度沿用开发残差校准，未用测试真值重拟合；序列与场站相关、反复选模和分布变化使覆盖率不具备独立保证。没有区间输出的方法不填写虚构覆盖率。')
seven=S['seven_day']; wanted=[('experiment','last'),('experiment','ridge_1000'),('experiment','lgb_p0_l15_n200'),('round2','zero_uni'),('round2','tirex_uni'),('round2','old_equal_mix'),('round2','new_curve'),('round2','new_peak'),('round4','best_mix')]
labels={'last':'前一天同刻','ridge_1000':'Ridge','lgb_p0_l15_n200':'LightGBM','zero_uni':'Chronos 2','tirex_uni':'TiRex 2','old_equal_mix':'Chronos 与 TiRex','new_curve':'第二轮曲线方案','new_peak':'第二轮峰值方案','best_mix':'第四轮组合'}
rr=[next(a for a in seven if (a['round'],a['model'])==x) for x in wanted]
table(['已有七天方法','曲线 WAPE %','日均 WAPE %','峰值 MAPE %'],[[labels[a['model']],fmt(a['wape']),fmt(a['daily_wape']),fmt(a['peak_mape'])] for a in rr],[2.45,1.35,1.35,1.42])
p('七天任务为672点，只有起报第22、23、24日三个窗口，目标日期重叠，因此单列比较，不与次日分数混排。当前round8b没有同协议七天预测数组，本报告不推造它的七天分数，也不把逐日滚动次日预测拼成一次性七天预测。')
p('保留的TiRex七天曲线对照较好；第二轮峰值方案七天峰值误差约1.3621%，略低于前一天同刻的1.3768%。这只是既有探索性结果，不能把三个重叠窗口解释为可靠跨月优势。')
h('保留模型与部署成本',True)
base=B/'models'
assets=[('Chronos 2','chronos2-base/model.safetensors','通用基础权重'),('TiRex 2','tirex2-base/model.ckpt','单模型及七天对照'),('TimesFM 2.5','timesfm25-base/model.safetensors','单模型与a集群对照'),('TimesFM 3','timesfm3-base/model.safetensors','当前组合及强单模型对照')]
table(['主要权重','磁盘 MiB','保留用途'],[[n,fmt((base/rel).stat().st_size/2**20,2),why] for n,rel,why in assets]+[['当前最佳专用增量',fmt((B/'round8b/models/best_tail_lr1e6_teacher01_s100_final.pt').stat().st_size/2**20,2),'正式round8b组件'],['Ridge',fmt((base/'ridge_1000_final.pkl').stat().st_size/2**20,4),'传统轻量对照'],['LightGBM',fmt((base/'lgb_p0_l15_n200_final.pkl').stat().st_size/2**20,4),'传统对照与峰值取舍']],[1.85,1.15,3.57])
p('表中是主要权重文件长度，Python运行环境、数据、校准器和预测缓存另计；不能把MiB文件大小当作RAM或显存需求。当前组合保留Chronos、TiRex、TimesFM 3及专用增量，资源组成比单一TimesFM 3更复杂，但本次没有在统一条件下补测推理时延或运行峰值内存。旧日志不同运行条件的耗时不混成速度排名。')
p('当前可用内存不足3GiB的重模型准入门槛，本次仅做已保存预测的复算和部署缓存重建核对，未降低保护阈值、未加载基础模型，也未关闭用户程序。独立核对通过126项，包含保留权重哈希、统一指标和当前最佳预测的5472行、57域、数值及区间顺序。')
table(['使用入口','当前作用'],[['09 现有最佳微调','正式最佳次日预测入口 保留'],['02 TiRex 预测七天','既有七天对照 保留'],['04 第二轮峰值方案','峰值侧重对照 保留'],['16 直接预测共享状态','明确说明未采用权重已删除 转到09']],[2.6,3.97])
p('默认历史截至2019年5月31日，示例预测2019年6月1日00:00至23:45，共5472行。六月没有真实标签，这是历史复现推理，不是今天的电负荷。正式模型与原冻结选择均未更换，自动续跑继续停用。')
h('文件整理与最终研究判断',True)
table(['清理类别','删除依据','保留内容'],[
 ['第二轮四组开发权重','未采用且曲线 日均 峰值均被冻结协变量方案支配','入选模型与可用峰值变体'],
 ['第四轮十组开发权重','未采用且三个开发指标均被保留协变量模型支配','入选配置与较好协变量对照'],
 ['第十四轮专用权重','18候选前后开发曲线均劣化至少20% 正式采用0','协议 分数 预测数组及原审计'],
 ['源码字节码与旧页图','可重建缓存或已被最终版替代的草稿','运行依赖 正式报告与最终核对页']],[1.65,2.85,2.07],10.2)
p(f'本次删除共 {C["deleted_files"]} 个文件，逻辑长度 {C["deleted_logical_bytes"]:,} 字节，约 {C["deleted_logical_bytes"]/2**30:.3f} GiB。逐目标名单、删除前SHA256、分数依据和实际删除状态见本次清理记录。没有测量实际释放磁盘空间；硬链接、分配单位及其他写入使其可能不同。')
table(['近期未采用候选','测试曲线 WAPE %','测试日均 WAPE %','采用情况'],[[name(i),fmt(A[i]['wape'],6),fmt(A[i]['daily_wape'],6),'开发不支持替换'] for i in ['r13candidate','r9fusion','r14candidate']],[2.35,1.3,1.3,1.62],10.3)
p('第十四轮原123项独立验算是删除前的记录，仍作为历史证据保留。专用权重删除后可用保存预测复算分数，但不能再直接重放该轮神经网络或断点继续训练；本次不把旧验算宣称为删后模型完整性检查。旧报告原文记载当时保留状态，本次清理记录说明随后变化。')
p('当前输入和31天样本下，近期多类调整未支持可靠的大幅提升，符合工程平台判断，未证明理论误差极限。保留round8b、TimesFM 3、TiRex、传统模型和峰值基线，后续应在未查看的真实功率数据上预先冻结目标再验证。不能用新的最低测试小数点分数替代可靠提升证据。')
p('原始功率数据和官方cell a的12,085个Parquet必要列保留，真实CPU与内存完整数据只覆盖a集群5域，剩余下载为0。column_cache保存真实官方压缩列字节，不是可随意删除的安装缓存；JOB窗口、去重键不完整和时间对齐限制仍存在。Google数据为CC BY 4.0；保留的TimesFM 3权重仅限非商业、非生产用途。')
doc.save(OUT);print(OUT)
