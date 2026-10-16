"""Explicit 24h curve, 7day curve, and 7day daily-mean report revision."""
from pathlib import Path
import sys,json,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
TOP=Path(sys.argv[1]).resolve(); B=TOP/'experiment';O=B/'model_comparison_20261005'
R=json.loads((O/'results/task_revision/summary.json').read_text(encoding='utf-8'));A={a['model']:a for a in R['models']};Q={a['model']:a for a in R['seven_day']};DM={a['model']:a for a in R['daily_mean']}
C=json.loads((O/'cleanup_record.json').read_text(encoding='utf-8-sig'));V=json.loads((O/'results/verification.json').read_text(encoding='utf-8'))
assert R['all_passed'] and V['all_passed'] and C['state']=='completed'
assert R['actual_missing_tasks_completed'] and not R['missing_seven_day']
CHECK_COUNT=len(R['checks']); DETAIL_COUNT=len(R['seven_day'])*57*3*7
WIN_CURVE=R['best7curve'];WIN_MEAN=R['best7daily']
OUT=TOP/'Google PowerData 2019 三项预测任务与模型比较汇报.docx'
if '--validate-only' in sys.argv: print(f'Three tasks and {CHECK_COUNT} checks validated; no authoring');sys.exit(0)
backup=O/'docs/before_task_revision';backup.mkdir(parents=True,exist_ok=True)
previous=TOP/'Google PowerData 2019 最佳模型多维度评估与文件整理汇报.docx'
if previous.exists() and not (backup/previous.name).exists():shutil.copyfile(previous,backup/previous.name)
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.78);sec.left_margin=sec.right_margin=Inches(.80)
W=6.67
for k in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[k];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
 pp=st.element.find(qn('w:pPr'))
 if pp is not None:
  for ch in list(pp):
   if ch.tag==qn('w:pBdr'):pp.remove(ch)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.12;doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(10);doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;foot.add_run('2026年10月5日 修订汇报  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def p(t):return doc.add_paragraph(t)
def h(t,page=False):
 q=doc.add_paragraph(t,'Heading 1');q.paragraph_format.page_break_before=page;return q
def sub(t):return doc.add_paragraph(t,'Heading 2')
def table(headers,rows,widths,size=10.4):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for j,x in enumerate(headers):t.rows[0].cells[j].text=str(x)
 for row in rows:
  c=t.add_row().cells
  for j,x in enumerate(row):c[j].text=str(x)
 for i,row in enumerate(t.rows):
  tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
  if i==0:tr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for s in ['top','left','bottom','right']:
    e=OxmlElement('w:'+s);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
   pr.append(bd);mar=OxmlElement('w:tcMar')
   for s,n in [('top',60),('bottom',60),('left',85),('right',85)]:
    e=OxmlElement('w:'+s);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCEAF2');pr.append(sh)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(1);q.paragraph_format.line_spacing=1.05;q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    for rr in q.runs:rr.font.size=Pt(size);rr.bold=i==0
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
def n(i):return R['short_names'][i]
def f(x,d=4):return f'{x:.{d}f}'
def horizons(i,k):return [f(a[k],3) for a in sorted([x for x in R['horizons'] if x['model']==i],key=lambda x:x['forecast_day'])]
core=['best','r5','tf3','r4','r3','equal','tirex','chronoscov','tf25','chronos','r2curve','r2peak','ridge','lgb','last']
single7=['tf3_uni21_7','tirex','tf3cov','tf25_7','chronoscov','lgb','ridge','chronos','last','week']
combo7=['best_transfer7','r5_transfer7','r3_transfer7','r4','r2curve','r2peak','ensemble7','equal']
seven=single7+combo7
doc.add_paragraph('Google PowerData 2019 三项预测任务与模型比较汇报','Title')
p('本汇报分别评价未来24小时完整负荷曲线、未来7天完整负荷曲线、未来7天每日平均负荷。各表明确写出模型名称与处理方式，给出实际保存预测的统一复算结果，并说明模型保留、历史实验未采用和文件清理情况。')
h('三项任务的实际结论')
table(['预测任务','具体模型或方法','主指标 %'],[
 ['未来24小时完整曲线\n每PDU 96点',n('best'),f(A['best']['wape'],6)],
 ['未来7天完整曲线\n每PDU 672点',n(WIN_CURVE),f(Q[WIN_CURVE]['wape'],6)],
 ['未来7天每日平均负荷\n每PDU 7个日均值',n(WIN_MEAN)+' 的曲线逐日平均',f(DM[WIN_MEAN]['daily_wape'],6)]],[1.85,3.5,1.32],10.8)
p('24小时当前正式方案在文件中编号round8b。完整名称是“Chronos-2＋TiRex-2＋TimesFM3三模型固定权重预测，加Ridge残差校准与TimesFM3末层微调”。编号仅用于追溯目录，不代替模型名称；组成、权重和微调参数见后文。')
p(f'24小时曲线相较前一天同刻误差减少 {(A["last"]["wape"]-A["best"]["wape"])/A["last"]["wape"]*100:.2f}%。相较仅加Ridge残差校准的上一正式方案，改善约0.000052个百分点，探索性日期配对区间跨零，不能称为大幅或显著提高。')
p('缺失的五个主方案已实际补测，7天曲线和7天日均各有18个明确命名的对照。TimesFM3的21天历史单变量版在两项7天任务中误差最低。当前24小时最佳也已完成7天迁移版补测；Ridge等残差处理按固定起报历史逐日应用，名称明确标注“迁移”，不是另行训练的7天模型。')
p('目前没有一个方法在三项任务及峰值等所有指标中同时第一。7天日均结果由各模型的672点预测逐日取均值得到，没有另行训练专门的7日均值模型。功率为归一化利用率，所有日均值和绝对误差均使用该单位，不是kW或kWh。')
p('补测前冻结了权重、原参数和推理适配方式，随后实际运行672点推理，不重新训练，不按测试结果修改方案。2019年5月测试已多轮查看，且7天窗口重叠，缺少独立确认月；“最低误差”仅指这份对照中的描述性结果，不是独立显著性证明。')
p('曲线WAPE先逐PDU计算有效点绝对误差和除以真值绝对值和，再对57个PDU等权平均；不把不同容量的PDU视为已知kW负荷。7天曲线沿用相同评分口径，日均指标另按每日均值评分。')

h('当前正式模型的组成与微调参数',True)
p('先定义C为Chronos-2的14天历史单变量预测，T为TiRex-2的14天历史单变量预测，F为TimesFM3的14天历史同集群联合目标与历史production功率代理协变量预测。三者都只读取起报前的历史，不输入未来真实工作量。clip表示将输出限制到0至1。')
sub('三模型固定权重预测')
p('P4 = clip(0.25 × C + 0.25 × T + 0.50 × F, 0, 1)。这是第四轮冻结正式方案，也是当前正式模型的基础。TimesFM3此分支使用14天历史与功率代理协变量，与21天单变量TimesFM3对照是不同配置。')
table(['基础模型及输入','固定权重','对应保存记录'],[
 ['Chronos-2 零样本单变量\n历史14天1344点','25%','native / zero_uni'],
 ['TiRex-2 零样本单变量\n历史14天1344点','25%','native / tirex_uni'],
 ['TimesFM3 同集群联合功率目标\n历史14天与production功率代理','50%','tf3_cellprod']],[3.45,1,2.22])
sub('加入Ridge残差校准')
p('在P4基础上，用历史功率、历史production代理及已经观测到的过去预测误差训练Ridge残差回归。Ridge正则参数为100000。设其预测残差为r，校准分支K = clip(P4 + r, 0, 1)，正式输出P5 = clip(0.5 × P4 + 0.5 × K, 0, 1)。第五轮采用的是这个Ridge分支，稀疏MoE网络未被采用。')
sub('加入TimesFM3末层微调')
p('只微调TimesFM3输出头与最后一个Transformer层；AdamW学习率1×10⁻⁶，权重衰减0.01，100步，教师输出约束权重0.1，梯度范数上限1。训练使用截止第22天的数据，分位数及边界损失的具体实现保留在原实验代码中。其他基础模型与Ridge校准保持原方案。')
p('设微调后的TimesFM3预测为F*，候选J = clip(P5 + 0.5 × (F* − F), 0, 1)。开发期冻结采用权重0.5，当前正式输出P8b = clip(0.5 × P5 + 0.5 × J, 0, 1)。无截断时，相当于在P5上叠加0.25 × (F* − F)；实际预测严格保留上述分阶段截断。')
p('冻结配置为best_tail_lr1e6_teacher01_s100。测试数组的P5校准式和P8b最终加权式、三模型固定权重式均已复核。完整模型名称用于所有结果表，实验编号和数组键只作为追溯信息。')

h('其他加权方法的明确组成',True)
sub('Chronos2三分支与TiRex2日历分支')
p('第二轮“曲线”方法由以下四个非零分支构成，先按开发期冻结权重加权，再做3点局部平滑与0至1截断。其权重在7天评估中原样转用；7天权重不是根据测试结果重新学习。')
table(['分支模型与处理','权重 %'],[
 ['Chronos-2 历史production协变量全参数微调\n学习率1×10⁻⁶ 300步',f(36.7480924275998,6)],
 ['Chronos-2 同集群联合目标与历史production协变量\n零样本预测',f(34.56400544511919,6)],
 ['Chronos-2 总功率 production与余量分量联合微调\n学习率1×10⁻⁶ 300步 加分量一致性投影',f(27.89141255562963,6)],
 ['TiRex-2 日历分支',f(.7964895716513902,6)]],[5.27,1.4],10.2)
p('单变量LoRA与旧等权分支的权重为0；历史平滑分支约1.17×10⁻¹⁸，可视为数值零。分量约束是归一化功率的代数一致性与非负约束，不能解释为设备能量守恒。峰值版本在同一曲线之上，使用起报前一天的峰值作强度0.25的形状校正；截断前保持全天96点均值，不保证评分掩码下的日均完全不变。')
sub('Chronos2 TiRex2 TimesFM3与残差加权方法')
p('第三轮主方法的非零权重依次为：Chronos-2与TiRex-2等权预测17.170258%，Chronos-2对1减功率变换的预测1.228588%，在等权预测上加Ridge残差校准的分支39.969919%，在等权预测上加前一日已观测平均偏差0.25倍的分支10.135183%，TimesFM3的21天单变量分支31.496051%。其余分支为0或数值零。该方法不是第四轮25%／25%／50%的固定权重预测。')
sub('七天专用历史加权方法')
p('第一轮七天加权方法的权重为：Chronos-2 LoRA 100步分支48.238511%，Chronos-2分量约束分支34.005038%，Chronos-2历史production协变量分支15.332152%，历史曲线加趋势分支1.895472%，受约束Patch Transformer分支0.528828%。LightGBM和Chronos-2零样本分支为数值零。Patch Transformer是自建PatchTST思路网络，不是官方PatchTST复现。')
p('表格中使用“Chronos-2三分支＋TiRex-2日历”或列出三种基础模型及具体校准方式，均对应以上完整组成。没有将不同配置笼统合并为一个“融合模型”。')

h('未来24小时完整负荷曲线比较',True)
p('预测2019年5月23至31日，共9个起报、57个PDU、每次96点，49,243个有效评分点。曲线WAPE是主指标；日均、峰值与P95点误差用于展示其他表现，数值均为百分比，越低越好。')
table(['明确模型或方法名称','曲线\nWAPE','日均\nWAPE','峰值\nMAPE','P95点\nAPE'],[[n(i),f(A[i]['wape']),f(A[i]['daily_wape']),f(A[i]['peak_mape']),f(A[i]['p95_ape'])] for i in core],[2.90,.93,.93,.93,.98],10.1)
p(f'TimesFM3同集群历史代理配置单独24小时WAPE为{R["tf3cov24"]["wape"]:.6f}%，不是21天单变量配置。普通Ridge的alpha为1000，区别于残差校准的100000；LightGBM为15叶200树，TimesFM 2.5开启force_flip_invariance。')

h('24小时曲线的峰值与稳定性评估',True)
doc.add_picture(str(O/'figures/task_revision/24h_curve_peak.png'),width=Inches(W))
p(f'前一天同刻的峰值MAPE为{A["last"]["peak_mape"]:.6f}%，优于当前正式模型的{A["best"]["peak_mape"]:.6f}%。Chronos-2三分支加TiRex-2日历的峰值校正版为{A["r2peak"]["peak_mape"]:.6f}%，以曲线误差代价改善峰值。曲线最优不等于峰值最优。')
table(['对照方法','最佳赢\n日期数','最佳赢\nPDU数','最佳改善95%区间\n百分点'],[[n(i),f'{A[i]["best_wins_dates"]}/9',f'{A[i]["best_wins_domains"]}/57',f'[{A[i]["best_gain_over_candidate_day_bootstrap95_pp"][0]:.6f}, {A[i]["best_gain_over_candidate_day_bootstrap95_pp"][1]:.6f}]'] for i in ['r5','tf3','tirex','last']],[2.85,.70,.70,2.42],9.9)
p('区间定义为对照WAPE减当前正式WAPE。使用5000次共同日期重采样，保持同日57域一起抽取；是同月探索性区间，未做多重比较校正。对Ridge校准上一正式方案和TimesFM3单变量的区间跨零，优势不稳定。')
p('保留逐日、逐域、逐集群、预测小时段、RMSE、有符号偏差、峰值低估频率与幅度明细。当前正式经验80%区间测试覆盖约82.78%，仅为开发残差区间的同月回看，不承诺跨月覆盖。')

h('未来7天完整曲线的单模型与基线比较',True)
p('每次在起报时一次性预测未来672点。三个起报完成日为5月22、23、24日，目标窗口分别为5月23至29日、24至30日、25至31日。共57个PDU、114,902个有效评分点，包含重叠窗口中的重复目标时刻，不能看作三周独立测试。')
table(['明确模型或方法名称','曲线\nWAPE %','点误差\nMAE','峰值\nMAPE %'],[[n(i),f(Q[i]['wape']),f(Q[i]['mae'],5),f(Q[i]['peak_mape'])] for i in single7],[3.28,1.12,1.12,1.15],10.3)
p(f'TimesFM3的21天历史单变量版最低，为{Q[WIN_CURVE]["wape"]:.6f}%，相较TiRex-2的{Q["tirex"]["wape"]:.6f}%减少{(Q["tirex"]["wape"]-Q[WIN_CURVE]["wape"])/Q["tirex"]["wape"]*100:.2f}%曲线误差。相较前一天同刻重复7天曲线，误差减少{(Q["last"]["wape"]-Q[WIN_CURVE]["wape"])/Q["last"]["wape"]*100:.2f}%。这是既有冻结模型的新任务实测，不是新调参产生的提升。')
p('TimesFM3的21天单变量配置与14天同集群功率代理配置是两个明确分开的方法。TimesFM 2.5补测保留原1344点历史与flip不变性，只把允许的最大推理跨度设为768点以容纳672点输出，没有改动权重。')
p('新增TimesFM3、TimesFM 2.5与Chronos-2计算路径都先复现已保存的96点预测。最大绝对差均不超过1.20×10⁻⁷利用率，远低于预登记的1×10⁻⁵数值容差。原模型与原参数得到核对后才继续7天推理。')

h('未来7天完整曲线的加权与校准比较',True)
table(['明确模型或方法名称','曲线\nWAPE %','点误差\nMAE','峰值\nMAPE %'],[[n(i),f(Q[i]['wape']),f(Q[i]['mae'],5),f(Q[i]['peak_mape'])] for i in combo7],[3.40,1.09,1.09,1.09],10.3)
p(f'当前次日最佳的7天迁移版曲线WAPE为{Q["best_transfer7"]["wape"]:.6f}%，仅Ridge校准迁移版为{Q["r5_transfer7"]["wape"]:.6f}%。末层微调带来的7天变化约{Q["r5_transfer7"]["wape"]-Q["best_transfer7"]["wape"]:.6f}个百分点，仍然极小；两个方案都未超过21天单变量TimesFM3。')
p('原Ridge残差回归只为96点建立特征。本次预先固定迁移方式：每个预测日使用自己的基础预测曲线，但所有真实历史功率、production代理及已观测过去误差都停在同一个起报日；不逐日加入未来实测值。第三轮残差分支仅更新起报时已知的未来星期信息，权重和校准系数不变。')
p('当前末层微调迁移版仍按P5加0.5倍TimesFM3微调与原模型预测差构造候选，再用冻结权重0.5混合。TimesFM3两个分支都直接预测672点。它是明确登记的推理迁移版，不能与专门在7天目标上重新训练的模型混淆。')
p('7天峰值最小的主对照是Chronos-2三分支＋TiRex-2日历加前日峰值校正版，峰值MAPE为1.362089%；前一天同刻为1.376785%。此处只作描述，不基于3个重叠窗口宣称独立显著优势。')

h('七天完整曲线按预测日比较',True)
doc.add_picture(str(O/'figures/task_revision/seven_curve_horizon.png'),width=Inches(W))
table(['方法','D1','D2','D3','D4','D5','D6','D7'],[[n(i)]+horizons(i,'curve_wape') for i in ['tf3_uni21_7','best_transfer7','tirex','r4','last']],[2.32]+[.6214]*7,9.4)
p('D1至D7为起报后第1至7天，分列WAPE按同一3个起报评分，不要求逐日单调变差。两种跨度的起报集合不同，下一页仅展示配置覆盖，不能把差值视为同样本退化幅度。')

h('三项任务的模型覆盖与成绩对照',True)
def qkey(i):return R['seven_task_key'].get(i,i)
table(['具体模型与7天适配说明','24小时曲线\nWAPE %','7天曲线\nWAPE %','7天日均\nWAPE %'],[[n(qkey(i)),f(A[i]['wape']) if i!='tf3cov' else f(R['tf3cov24']['wape']),f(Q[qkey(i)]['wape']),f(DM[qkey(i)]['daily_wape'])] for i in ['best','r5','r3','tf3','tf25','tirex','r4','equal','tf3cov','chronoscov','last']],[3.19,1.16,1.16,1.16],10.1)
p('本表将未来24小时完整曲线、未来7天完整曲线、未来7天每日平均负荷并列展示，均有真实预测数组。24小时共9个起报，7天共3个重叠起报，各任务内统一评分，但跨任务差值不是相同样本的退化幅度。')
p('名称中带“7天迁移”的三项，24小时列为原已部署96点方案，7天两列为本次已冻结并验证的推理迁移方式。单变量TimesFM3与TimesFM 2.5则使用原参数直接输出672点。18个7天对照均已评估，缺失的五个主方案不再以“未评估”留空。')

h('未来7天每日平均负荷的单模型比较',True)
p('这项任务的输出是未来7天每天一个平均值，每PDU共7个值，评价一天的平均水平。它与672点曲线任务分开评分；日均误差可能因正负点误差抵消而更低，所以不能用日均准确率代替曲线准确率。')
p('为统一评分，对每个PDU、起报和预测日，真实日均与预测日均都在真实有效的15分钟时刻上取平均；共有1197个有效PDU起报日记录。日均WAPE先每PDU累计日均绝对误差除以真实日均绝对值和，再对57个PDU等权平均。预测部署另导出完整96点的均值，不利用未来真值掩码。')
table(['明确模型或方法名称','日均\nWAPE %','日均绝对\nMAE','日均\nRMSE'],[[n(i),f(DM[i]['daily_wape']),f(DM[i]['daily_mean_mae'],5),f(DM[i]['daily_mean_rmse'],5)] for i in sorted(single7,key=lambda i:DM[i]['daily_wape'])],[3.35,1.12,1.10,1.10],10.2)
p(f'TimesFM3的21天单变量曲线逐日平均WAPE最低，为{DM[WIN_MEAN]["daily_wape"]:.6f}%，相较TiRex-2的{DM["tirex"]["daily_wape"]:.6f}%减少{(DM["tirex"]["daily_wape"]-DM[WIN_MEAN]["daily_wape"])/DM["tirex"]["daily_wape"]*100:.2f}%日均误差。相较前一天同刻，日均误差减少{(DM["last"]["daily_wape"]-DM[WIN_MEAN]["daily_wape"])/DM["last"]["daily_wape"]*100:.2f}%。')
p('日均MAE与RMSE的单位为归一化利用率，WAPE为百分比。分数由同一18种曲线输出得到，不重新选择每日更好的模型，也不额外训练直接日均预测器。')

h('未来7天每日平均负荷的加权校准比较',True)
table(['明确模型或方法名称','日均\nWAPE %','日均绝对\nMAE','日均\nRMSE'],[[n(i),f(DM[i]['daily_wape']),f(DM[i]['daily_mean_mae'],5),f(DM[i]['daily_mean_rmse'],5)] for i in sorted(combo7,key=lambda i:DM[i]['daily_wape'])],[3.35,1.12,1.10,1.10],10.2)
p(f'当前次日最佳迁移版的日均WAPE为{DM["best_transfer7"]["daily_wape"]:.6f}%，仅Ridge校准迁移版为{DM["r5_transfer7"]["daily_wape"]:.6f}%，微调改变量同样很小。它们优于原三模型固定权重方案，但没有超过21天单变量TimesFM3。')
p(f'LightGBM日均{DM["lgb"]["daily_wape"]:.6f}%，优于三模型25%／25%／50%方案的{DM["r4"]["daily_wape"]:.6f}%；LightGBM的7天曲线WAPE却更差。这说明平均负荷任务与完整曲线任务应分开评价，不能只用其中一个指标决定全部任务。')
p('峰值形状校正版与原曲线版本的日均分数接近，原因是它在截断前保持全天均值。真实有效点掩码与截断仍会造成极小差异，并不意味着使用了未来日均真值。所有残差校准都只使用起报前已知的历史误差信息。')

h('七天每日平均负荷按预测日比较',True)
doc.add_picture(str(O/'figures/task_revision/seven_mean_horizon.png'),width=Inches(W))
table(['方法','D1','D2','D3','D4','D5','D6','D7'],[[n(i)]+horizons(i,'daily_mean_wape') for i in ['tf3_uni21_7','best_transfer7','tirex','r4','last']],[2.32]+[.6214]*7,9.4)
p('各列为预测第1至7天的日均WAPE百分比。由于每PDU分母不同，且预测起报窗口重叠，七列的简单算术平均不必等于上一页累计7天的主日均WAPE。所有数值均来自同一672点预测，不重新选择更好的每日模型。')
sub('每日平均值明细的核查内容')
p(f'seven_day_daily_means_detail.csv保留18个方法×57个PDU×3个起报×7天，共{DETAIL_COUNT}行记录。每行明确模型全名、PDU与集群、起报完成日、D1至D7、目标日期、有效点数、真实日均、预测日均和绝对误差。不能把这些PDU均值简单相加，解释为全中心物理总负荷。')
p(f'逐PDU逐起报逐日独立循环复算了全部18个方法的日均WAPE与MAE，与统一评分程序一致。三个任务、基线、原模型复现及组合公式共{CHECK_COUNT}项修订检查通过；原126项清理与评估核对证据继续保留。')

h('未采用方法与精度提升结论',True)
p('当前功率代理输入下，近期微调已经进入工程平台期。证据是多轮验证未支持可靠升级、正式模型相较Ridge校准上一方案的改变量极小，而不是证明理论精度上限。真实CPU特征和新网络实验已完成后仍保留旧最佳；应把负面结果写入研究结论。')
table(['已评估的未采用方法','24小时曲线\nWAPE %','采用结论'],[
 ['三日近期加权SmoothL1残差校准器\nh3 p2',f(A['r13candidate']['wape'],6),'开发后段不支持\n不按测试低分替换'],
 ['RevIN多尺度分解与算电Cross-Attention\n稀疏MoE真实CPU残差候选 10轮',f(A['r9fusion']['wape'],6),'开发选择保留旧最佳\n真实CPU仅a集群5域'],
 ['TiDE思路64维局部直接预测网络\n20轮 两固定种子平均',f(A['r14candidate']['wape'],6),'整体明显变差\n专用权重已删除']],[3.45,1.22,2.0],10.1)
p('第9、11、12、13、14轮正式selected预测与当前末层微调正式方案逐点一致，只表示保留原方案，不计为五种独立升级。第8c轮门控与权重探索冻结结果2.374440182003358%变差，未采用；不为得到更低测试数值回改冻结选择。')
p('补测发现21天单变量TimesFM3的7天成绩优于原TiRex-2，表明之前7天比较不完整；它并非本次调参得到的新模型，也不能据此宣称次日模型已大幅提升。缺少独立月份、容量对应的物理功率标签与更完整的算力功率映射，限制进一步结论。未来真值驱动的日偏移或逐点专家选择只能作诊断。')
sub('文件整理的实际处理')
p(f'此前已删除{C["deleted_files"]}个文件，逻辑长度{C["deleted_logical_bytes"]:,}字节，约{C["deleted_logical_bytes"]/2**30:.3f}GiB，物理释放量未测量。本次修订没有再次删除模型或数据。明显较差且未采用的专用权重、过期草稿与可重建字节码已清理；基础权重、强对照、当前正式模型、原始列字节、冻结协议、负面分数与预测数组仍保留。')
p('第十四轮专用权重已删，其删除前123项验算是历史证据，不能宣称当前还能直接重放该网络。16号入口已转到正式最佳并保留内存保护。此前十份正式Word已SHA256核对后移至历史汇报；本次修订前版本另保留作修改记录。')

h('数据范围资源与交付文件',True)
p('Google PowerData2019功率数据为2019年5月31天、15分钟粒度、57个PDU。官方ClusterData2019 cell-a共12085个Parquet对象的必要列已完整读取并审计；剩余下载为0。只覆盖cell-a必要列，不是全部集群或全部字段，真实CPU与内存对应a集群5个功率域。')
p('CPU为NCU单位，功率是归一化利用率。production字段仍是功率代理；资源第三任务字段job-window equivalent不是CPU variance。条件单调回归只约束其他变量相同条件下的CPU作用，不能宣称设备能量守恒。时间对齐、去重键与JOB窗口代理的具体限制保留在原完整来源汇报中。')
p('保留官方压缩列字节约84.78GiB，与原对象总规模约407.55GiB、整目录迁移时逻辑总量约105.481GiB是不同口径。迁移逐文件SHA256证明当时复制一致；后续报告、入口和清理允许改变目录，不能说当前目录仍逐字等于迁移快照。')
p('实体根目录为 ${PROJECT_ROOT}；桌面同名目录为兼容junction。09号入口复现当前正式次日模型，02号入口复现TiRex-2七天曲线。默认历史截至2019年5月31日，部署示例的未来日期是2019年6月1日及后续历史预测跨度，不是今天的实际中心负荷，六月缺少真值。')
p('研究接续与16号兼容入口保留资源保护。补测首次在TimesFM3加载时因可用物理内存低于1.5GiB被保护停止，证据保留；改为相同配置与权重直接在GPU初始化后恢复完成，未降精度或降低保护门槛。未关闭用户程序。自动续跑保持PAUSED。')
table(['修订交付内容','位置或文件名'],[
 ['24小时完整曲线多维度对照','24小时完整曲线模型比较.csv'],
 ['7天完整曲线独立对照','7天完整曲线模型比较.csv'],
 ['7天每日平均负荷独立对照','7天每日平均负荷模型比较.csv'],
 ['D1至D7曲线与日均分日误差','7天分预测日模型比较.csv'],
 [f'{DETAIL_COUNT}行每日平均值明细','experiment/model_comparison_20261005/\nresults/task_revision/seven_day_daily_means_detail.csv'],
 ['来源与冻结参数 修改验证','同一task_revision目录中的summary.json\n各轮selection.json与协议'],
 ['清理记录和历史报告','清理说明.md 清理记录.json 历史汇报']],[2.22,4.45],10.2)
p('Google数据许可为CC BY4.0，沿用已保存的官方许可核对记录；TimesFM3权重仅限非商业、非生产研究，不能把当前部署样例解释为生产使用授权。正式交付为此Word与可编辑对照表；PDF和页图用于内部排版核查。')

h('汇报完成后的两轮联合优化补充结果',True)
J={}
for rr in [15,16]:
 root=B/f'round{rr}'
 result=json.loads((root/'results/summary.json').read_text(encoding='utf-8'))
 verification=json.loads((root/'results/independent_verification.json').read_text(encoding='utf-8'))
 assert verification['all_passed'] and not result['all_metric_test_candidate_dominance']
 assert result['selection']['chosen_alpha']==0 and (root/'results/pipeline_complete.json').exists()
 J[rr]=result
p('按用户提出的“一个方式在各项评价中都最好”目标，先完成前三项任务汇报，再实际运行两种日均与峰值约束方式。每种方式各比较12个预登记开发配置，包括不调整对照；测试成绩没有用于回改选择。两轮均未达到全指标优于其他方法，不能写成已经全面最优。')
p('24小时基础使用Chronos-2＋TiRex-2＋TimesFM3，加Ridge残差校准与TimesFM3末层微调；7天基础使用21天历史单变量TimesFM3。Ridge日均与峰值预测头正则参数10／100只用训练期拟合；其配置没有胜出，不把它们写成已采用模型。')
table(['具体约束方案','24小时曲线\nWAPE %','24小时峰值\nMAPE %','7天曲线\nWAPE %','7天峰值\nMAPE %'],[
 ['上述既有基础方案',f(J[15]['scores24']['prior']['wape']),f(J[15]['scores24']['prior']['peak_mape']),f(J[15]['scores7']['prior']['wape']),f(J[15]['scores7']['prior']['peak_mape'])],
 ['基础预测＋历史峰值100%\n保持日均的整曲线仿射投影',f(J[15]['scores24']['candidate']['wape']),f(J[15]['scores24']['candidate']['peak_mape']),f(J[15]['scores7']['candidate']['wape']),f(J[15]['scores7']['candidate']['peak_mape'])],
 ['基础预测＋前日形状25%\n历史峰值50%与日均约束\n局部最小改变量投影',f(J[16]['scores24']['candidate']['wape']),f(J[16]['scores24']['candidate']['peak_mape']),f(J[16]['scores7']['candidate']['wape']),f(J[16]['scores7']['candidate']['peak_mape'])]],[3.15,.88,.88,.88,.88],10.0)
p('局部投影的7天峰值MAPE为1.181206%，低于原18个主对照中峰值校正版的1.362089%，但曲线WAPE从2.779707%升至2.845126%，次日曲线也变差。整曲线投影对曲线影响更大。改善峰值不能替代全指标改善，更不能称为大幅、可靠提升。')
p('选模先用4个开发期次日起报选择方法，再用后4个开发起报与2个重叠7天开发窗口检查共享调整强度。曲线WAPE、MAE、RMSE、日均WAPE、峰值MAPE、P95点APE、峰值时刻偏差、绝对总偏差、a集群5域及其余52域WAPE共10项指标，两个跨度分别比较最强对照。两轮冻结强度都为0，因此保留各自基础预测。')
p('日均约束保持完整96点预测均值；有效点评分仍可能产生微小变化，不是未来日均真值校准。开发7天窗口与次日开发日期重叠，同月测试此前已经查看，不能当作独立确认。新方式未修改09与02正式入口；实验selected7代表本轮TimesFM3基础方案，不代表02的TiRex-2入口已替换。')
p('两轮各16项独立检查通过，包括冻结文件、所有评分、未来标签扰动后训练系数不变、起报后真值扰动不影响特征，以及预测边界与投影可行性。完整配置、模型、负面结果及独立检查保存在experiment/round15与round16；本轮记录见根目录“联合优化补充实验记录.md”。目前没有证据支持“一个方法在所有评价中第一”。')
doc.save(OUT)
print(OUT)
