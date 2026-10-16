"""Append the actual bounded round17 outcome; preserve the reviewed prior report."""
from pathlib import Path
import sys,json,hashlib,shutil,ast
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
T=Path(sys.argv[1]);B=T/'experiment';R=B/'round17';Q=R/'results';O=B/'model_comparison_20261005'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
s=read(Q/'summary.json');v=read(Q/'independent_verification.json');dev=read(Q/'dev_summary.json');assert v['all_passed'] and (Q/'pipeline_complete.json').exists()
out=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx'
if '--validate-only' in sys.argv:print(json.dumps(dict(round17_complete=True,checks=len(v['checks']),alpha=s['selection']['chosen_alpha'],all_metric_winner=s['all_metric_test_candidate_dominance']),ensure_ascii=False));sys.exit(0)
backup=R/'docs/before_report_update';backup.mkdir(parents=True,exist_ok=True)
if '--repair-layout' not in sys.argv:
 assert not (backup/out.name).exists(),'Do not append twice'
 shutil.copyfile(out,backup/out.name);shutil.copyfile(O/'docs/report_qa_task_revision/verification.json',backup/'prior_visual_verification.json')
doc=Document(backup/out.name)
bestnew=min((a for a in dev['candidates'] if a['id']!='prior'),key=lambda a:a['worst_metric_ratio']);devprior=next(a for a in dev['candidates'] if a['id']=='prior')
# Reuse table formatting only, without executing the prior report builder.
tree=ast.parse(Path(__file__).with_name('report_revision.py').read_text(encoding='utf-8'));node=next(a for a in tree.body if isinstance(a,ast.FunctionDef) and a.name=='table');exec(compile(ast.Module(body=[node],type_ignores=[]),'existing_table_helper','exec'))
q=doc.add_paragraph('第17轮多跨度残差校准实际结果','Heading 1');q.paragraph_format.page_break_before=True
assert s['selection']['candidate']['id']=='prior' and s['selection']['chosen_alpha']==0,'Rewrite narrative for a different actual selection'
doc.add_paragraph('本轮未得到全指标领先方法。18个新增配置在次日开发筛选中不及原方案，冻结修正强度为0，未评分新增配置的测试集。下表区分开发比较与保留方案测试复核。')
doc.add_paragraph('新方式为Chronos-2 TiRex-2 TimesFM3加多跨度Ridge傅里叶残差校准。用历史7天统计、预测形状及历史production功率代理学习日均、前三阶傅里叶形状和峰值误差；正则100与1000意见相反时可关闭修正。骨干权重冻结，不使用未来真值。')
names={'wape':'曲线WAPE %','mae':'MAE 归一化功率','rmse':'RMSE 归一化功率','daily_wape':'日均WAPE %','peak_mape':'峰值MAPE %','p95_ape':'P95点APE %','peak_time_mae_minutes':'峰值时刻偏差 分钟','absolute_bias_percent':'绝对总偏差 %','cella5_wape':'a集群5域WAPE %','other52_wape':'其余52域WAPE %'}
rows=[]
for k,n in names.items():
 digits=5 if k in ['mae','rmse'] else 3
 rows.append([n]+[f'{x:.{digits}f}' for x in [devprior['scores'][k],bestnew['scores'][k],s['scores24']['prior'][k],s['scores7']['prior'][k]]])
table(['指标 越低越好','开发次日\n原方案','开发次日\n最优新增','测试次日\n保留方案','测试7天\n保留方案'],rows,[2.05,1.15,1.15,1.15,1.17],size=9.7)
c=s['selection']['candidate'];a=s['selection']['chosen_alpha'];candidate_id=c['id']
doc.add_paragraph(f'开发比较覆盖起报14至17。最优新增为正则1000、均值和形状修正、历史峰值50%、意见门控，仍有曲线与日均退步。全19项中原方案胜出，后四个次日起报与两个7天开发窗口均保留原方案，强度{a:g}。次日沿用第8b轮；7天沿用TimesFM3历史21天直接672点，日均来自同一曲线。')
doc.add_paragraph(f'训练缓存含15个起报、57域的直接672点预测。开发训练只用第14天之前标签，最终重拟合只用第22天之前标签，重叠目标按反频率赋权。次日训练残差采用固定零样本三模型代理，避免微调模型样本内误差。{len(v["checks"])}项实施核对通过；评分前的缓存字段预检失败已修复并留档。')
doc.add_paragraph('5月测试多轮查看，7天窗口重叠，缺少独立确认月，不能声称大幅、显著或全方位第一。完整候选与冻结记录见experiment/round17，开发分数和未采用原因见根目录第17轮多跨度残差优化记录.md。')
doc.save(out);print(json.dumps(dict(document=str(out),sha256=sha(out),appended_round=17),ensure_ascii=False))
