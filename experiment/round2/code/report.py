from pathlib import Path
import sys,json,csv
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

R=Path(sys.argv[1]);B=R.parent;TOP=B.parent
def read(name):return json.loads((R/'results'/name).read_text())
def label(k):
    return {'last':'前一天同刻','zero_uni':'Chronos-2 零样本','zero_cov':'Chronos-2 历史生产特征','zero_cellprod':'Chronos-2 集群与生产联合','tirex_uni':'TiRex-2 零样本','tirex_calendar':'TiRex-2 日历特征','old_equal_mix':'复现旧等权组合','round1_primary':'第一轮冻结组合','new_curve':'第二轮曲线组合','new_peak':'第二轮峰值修正','lora_uni_lr3e5_s300':'Chronos-2 单变量 LoRA','full_cov_lr1e6_s300':'Chronos-2 生产特征全量微调','full_parts_lr1e6_s300':'Chronos-2 分量一致性全量微调','history13':'历史曲线平滑'}.get(k,k)

def main():
    s=read('summary.json');sel=read('selection.json');v=read('verification.json');m=s['horizon1'];m7=s['horizon7'];c=s['comparison'];p=m['new_curve'];old=m['old_equal_mix'];goal=s['goal_achieved_exploratory']
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21.59);sec.page_height=Cm(27.94);sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for st in doc.styles:
        for e in list(st.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Calibri';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor.from_string('000000')
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.15;doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 2'].font.size=Pt(12)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;foot.add_run('Google PowerData 2019 · 第二轮研究  |  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='1';rr.append(tt);fld.append(rr);foot._p.append(fld)
    for r in foot.runs:r.font.size=Pt(8)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
    def txt(t,bold=False):
        q=doc.add_paragraph();q.add_run(t).bold=bold;return q
    def title(t):doc.add_heading(t.replace('.','').replace('：',' ').replace('，',' '),1)
    def page(t):doc.add_page_break();title(t)
    def table(headers,rows,widths):
        tb=doc.add_table(rows=1,cols=len(headers));tb.autofit=False
        for col,w in zip(tb.columns,widths):col.width=Cm(w)
        for cell,t,w in zip(tb.rows[0].cells,headers,widths):cell.text=str(t);cell.width=Cm(w)
        rep=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(rep)
        for row in rows:
            for cell,t,w in zip(tb.add_row().cells,row,widths):cell.text=str(t);cell.width=Cm(w)
        for i,row in enumerate(tb.rows):
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for j,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E1EDF1' if i==0 else ('F4F7F8' if i%2==0 else 'FFFFFF'));pr.append(shade)
                borders=OxmlElement('w:tcBorders')
                for side in ['top','bottom','left','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                pr.append(borders)
                mar=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),'75');x.set(qn('w:type'),'dxa');mar.append(x)
                pr.append(mar)
                for para in cell.paragraphs:
                    para.alignment=WD_ALIGN_PARAGRAPH.CENTER if j>0 and len(str(cell.text))<28 else WD_ALIGN_PARAGRAPH.LEFT
                    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.08
                    for run in para.runs:run.font.size=Pt(9);run.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    def fig(name,width=16.5):doc.add_picture(str(R/'figures'/name),width=Cm(width))
    fmt=lambda x:f'{x:.4f}%'
    doc.add_heading('Google PowerData 2019\n第二轮精度提升实验汇报',0)
    txt('2026 年 9 月 30 日  |  32 个成功候选配置 · 15 次开发期微调 · 冻结方案回测')
    title('1. 结论先行')
    conclusion='本轮达到预设的同月探索性提升标准。' if goal else '本轮未达到“显著优于旧最优方法”的预设标准。'
    txt(conclusion,bold=True)
    txt(f'第二轮曲线组合的次日测试 WAPE 为 {fmt(p["wape"])}。旧等权组合本次复现为 {fmt(old["wape"])}，相对误差变化为 {c["old_equal_mix"]["relative_reduction_percent"]:+.2f}%（正值表示降低）。旧报告记载的 2.416% 与此次复现差异为 {old["wape"]-2.416:+.4f} 个百分点。')
    txt(f'相对第一轮冻结组合，误差降低 {c["round1_primary"]["relative_reduction_percent"]:.2f}%；相对前一天同刻，误差降低 {c["last"]["relative_reduction_percent"]:.2f}%。这些比较不能替代与旧最优方法的比较。')
    txt(f'七天曲线 WAPE 为 {fmt(m7["new_curve"]["wape"])}，优于第一轮的 {fmt(m7["round1_primary"]["wape"])}，仍高于 TiRex-2 单模型的 {fmt(m7["tirex_uni"]["wape"])}。')
    keys=['last','old_equal_mix','round1_primary','zero_cov','new_curve','new_peak']
    table(['方法','曲线 WAPE','日均 WAPE','日峰值 MAPE'],[[label(k),fmt(m[k]['wape']),fmt(m[k]['daily_wape']),fmt(m[k]['peak_mape'])] for k in keys],[6.1,3.1,3.1,3.7])
    txt('判定规则：相对复现旧组合的 WAPE 至少降低 5%，且按连续 3 天与 cell 重采样的配对 95% 区间下限大于 0。该规则在查看第二轮回测分数前写入方案。')
    txt('结论边界：本月数据此前已经被查看，本轮属于同一公开数据上的适应性研究。即使区间支持改善，也不是独立新月份上的泛化证明。WAPE 不能直接解释为“100−WAPE 的准确率”。')

    page('2. 数据、切分与评价口径')
    txt('官方名称为 PowerData 2019 traces；观测期为 2019 年 5 月。10 个 cell、57 个供电域，每域 8,928 条 5 分钟记录。归一化功率利用率不是 kW/MW，无法据此直接计算机房电费或总容量。')
    table(['阶段','模型可训练数据','预测目标／用途'],[['开发训练','5 月 1—14 日','所有微调候选使用相同截止日'],['开发评价','冻结开发模型，仅上下文更新','5 月 15—22 日，连续 8 天'],['最终训练','5 月 1—22 日','重训已选择的模型，固定参数'],['次日回测','后续已观测历史可作上下文','5 月 23—31 日，共 9 次完整起报'],['七天回测','同一最终模型和次日组合权重','第 22、23、24 天结束起报，窗口重叠']],[3.2,5.5,7.3])
    txt('原始文件先按 time 排序，再将相邻 3 个点平均为 15 分钟。坏测量不进入评分；输入仅作因果前向填充。15 分钟标签要求 3 个原点均有效；次日回测共 49,243 个有效点。生产估计坏标记共 116,517 条（约 22.90%），用缺失掩码处理。')
    txt('官方 PDF 对质量标记存在文字歧义，本实验依据官方分析 Notebook 中 NOT bad_measurement_data 和 NOT bad_production_power_data 的筛选语义。原始压缩文件和 CSV 保留原字节，并再次核验了解压一致性。')
    txt('主指标：每个供电域先计算 Σ|预测−实际| / Σ|实际|，再对 57 个域取平均，乘以 100%。日均、日峰值使用实际与预测相同的有效时点；缺失数据不被当作零负荷。')
    txt('开发期先比较模型家族，再用前 4 天拟合权重、后 4 天检查统一／cell／域权重规则，最后用 8 天拟合冻结权重。家族选择已看全部开发期，因此后半段检查也不是完全独立验证。测试日期没有用于第二轮参数选择。')

    page('3. 尝试了什么，最终保留什么')
    txt('开发阶段成功评估 32 个基础配置：Chronos-2 的单变量、历史生产特征、日历特征、集群联合和全部域联合；不同学习率／步数的 LoRA 与全参数微调；同 cell 同伴序列；总量—生产—剩余分量联合微调；TiRex-2 的协变量、集群结构及 672／1,344／2,048 点上下文。')
    table(['组合成员','开发 WAPE','最终平均权重'],[[label(k),fmt(sel['development_scores'][k]['wape']),f'{100*sel["mean_weights"][k]:.2f}%'] for k in sel['pool']],[8,4,4])
    txt('最终采用统一的非负、和为 1 的权重。cell／单域权重未通过后半段开发比较；单变量 LoRA 和纯历史成员权重为零。TiRex-2 日历成员权重很小，因此复杂度增加带来的实用价值有限。')
    txt(f'最终曲线进行 {sel["smooth_width"]} 点局部平滑，再限制在 [0,1]。平滑是验证选择的连续性处理，不声称来自真实设备的爬坡限制。峰值版本另用前一天峰值作 {sel["peak_beta"]:.2f} 强度修正，修正前保持每日均值，随后截断。')
    txt('物理分量成员同时预测总功率 t、生产功率 p、剩余量 r，并做最小平方距离投影，使 t=p+r、p≥0、r≥0、t≤1。生产历史有效比例不足 20% 时保留总量预测，不强行使用分量投影。剩余量只是数学残差，不能解释为实测制冷功率。')
    txt('该方案是可复现的领域适配与约束组合，尚未论证为新的原创算法。全参数微调采用学习率 1×10⁻⁶、300 步、1,344 点上下文；开发／最终训练种子均为 20260930。')

    page('4. 次日结果：整体与逐日表现')
    fig('accuracy.png')
    txt(f'曲线组合的 95 分位点 APE 为 {fmt(p["p95_ape"])}；旧组合为 {fmt(old["p95_ape"])}。峰值修正版本的峰值 MAPE 为 {fmt(m["new_peak"]["peak_mape"])}，曲线 WAPE 为 {fmt(m["new_peak"]["wape"])}。曲线改善与峰值改善需要分别核对。')
    txt('图中统一保留冻结方案。某个单模型在测试中得分较低，也不据此重新设置组合权重；完整成员和对照分数均保存在 test24_metrics.csv。')

    page('5. 改善是否稳定，是否可以称为显著')
    table(['比较对象','相对误差降低','域／日／cell 胜数','配对 95% 区间（百分点）'],[[label(k),f'{z["relative_reduction_percent"]:+.2f}%',f'{z["domains_won"]}/57 · {z["days_won"]}/9 · {z["cells_won"]}/10',f'[{z["day_and_cell_95"][0]:+.4f}, {z["day_and_cell_95"][1]:+.4f}]'] for k,z in c.items()],[4.5,3.1,4.2,4.2])
    txt('区间针对“对照 WAPE−新组合 WAPE”，正值有利于新组合。5,000 次配对重采样，每次抽取连续 3 天块，并以 cell 为组抽样供电域；同时保存只按日块抽样的区间。这样保留部分时间和组内相关性，但 9 天、10 个 cell 仍然很少。')
    fig('cells.png')
    txt(conclusion,bold=True)
    txt('本轮与旧组合的配对区间跨过 0，而且曲线点估计更差，没有显著领先旧方法的证据。与第一轮比较，单按日块的区间为正，但加入 cell 抽样后仍跨 0；因此相对第一轮的 2.40% 改善也应称为小幅同月改善，不宜写成稳定泛化突破。')
    txt('前沿模型可能在公开数据上存在预训练重合，本次未排除。跨月份、跨季节和新机房推广需要新增从未用于调参的完整测试数据；第一轮 LightGBM 的留一 cell 实验不能代表本轮组合的跨 cell 泛化。')

    page('6. 七天预测与约束收益')
    fig('seven.png')
    table(['方法','七天曲线 WAPE','日均 WAPE','日峰值 MAPE'],[[label(k),fmt(m7[k]['wape']),fmt(m7[k]['daily_wape']),fmt(m7[k]['peak_mape'])] for k in ['tirex_uni','old_equal_mix','round1_primary','new_curve','new_peak']],[6.1,3.1,3.1,3.7])
    txt('本轮七天结果使用次日开发阶段冻结的相同权重，不再针对七天测试调参。TiRex-2 官方检查点最长输出 320 点，采用 320＋320＋32 点递推，只回填自身预测；不读取未来真实负荷。三个窗口重叠，不作独立周显著性推断。旧报告约 2.88% 的七天实现细节未完整取得，因此仅作历史参考；七天等权对照为本次递推实现。')
    parts='full_parts_lr1e6_s300';cov='full_cov_lr1e6_s300'
    txt(f'分量一致性微调单模型：次日 WAPE {fmt(m[parts]["wape"])}；生产特征微调单模型：{fmt(m[cov]["wape"])}。二者训练目标和输入结构不同，这不是只改变一个约束项的纯消融，不能把差异全部归因于物理约束。')
    txt('约束可保证输出可行性，但不保证误差更低。本次对 10,000 组随机分量验证了非负、加和与容量边界；对全部最终组合验证了非负和不超过归一化容量。约束收益与约束正确性分别记录。')

    page('7. 桌面交付、复现与明确建议')
    txt('全部工作继续保存在桌面 GooglePowerData2019_预测研究_20260929 文件夹。第一轮材料保留，第二轮集中在 experiment/round2。默认启动入口使用第二轮冻结曲线／峰值版本；第一轮七天入口另外保留，便于比较。')
    table(['文件／目录','内容'],[['本目录汇报 DOCX / PDF','本报告与可直接阅读的 PDF'],['google-powerdata-2019','官方原始 gzip、CSV、来源清单和 CC BY 4.0 说明'],['experiment/models','Chronos-2 与 TiRex-2 官方权重、第一轮模型'],['experiment/round2/models','15 组开发微调及最终选中模型的权重'],['experiment/round2/code','训练、组合、核验、图表、报告与本地推理代码'],['experiment/round2/results','开发／测试分数、逐点预测、权重、统计区间和日志'],['runtime/python 与 experiment/.venv','本机专用 Python 与已安装依赖'],['开始阅读.md／交付文件清单.json','操作说明、文件大小、SHA-256 和完整位置']],[6.6,9.4])
    txt('双击“第二轮预测未来一天.cmd”或“第二轮预测未来七天.cmd”。输出含曲线与峰值两个版本。默认以历史末日 2019 年 5 月 31 日为起点预测 6 月 1 日之后，不是当前真实日期的实时负荷，也没有新的真实标签可评分。')
    txt('已完成：未来值扰动不改变过去输入；训练和推理输入分别检查；冻结权重校验；保存预测重新组合与指标复算；桌面入口实际试运行。TiRex-2 使用官方 CPU 内核关闭即时编译，Chronos-2 使用本机 GPU。迁移机器或目录需重建环境或更新绝对路径。')
    txt(conclusion+' 次日曲线优先使用“最高次日曲线精度_旧组合.cmd”，七天曲线优先使用“TiRex单模型预测未来七天.cmd”。新组合保留用于研究曲线与峰值的取舍。后续应获取独立新月份，在方案冻结后一次性测试，避免重复查看相同日期造成选模偏差。',bold=True)
    txt('来源：Google 官方 github.com/google/cluster-data（PowerData2019.md、power_trace_analysis_colab.ipynb），原始桶 storage.googleapis.com/powerdata_2019；Amazon 官方 huggingface.co/amazon/chronos-2；NX-AI 官方 github.com/NX-AI/tirex-2 与 huggingface.co/NX-AI/TiRex-2。模型固定版本与哈希见交付清单，均保留 Apache 2.0 许可信息。')
    # Remove inherited paragraph rules after all content has been authored.
    for e in list(doc.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
    out=TOP/'Google PowerData 2019 第二轮精度提升汇报.docx';doc.save(out);print(out)

if __name__=='__main__':main()
