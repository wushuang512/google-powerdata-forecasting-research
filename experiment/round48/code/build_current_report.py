from pathlib import Path
from copy import deepcopy
import sys,json,csv,hashlib,zipfile,datetime,math
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
T=Path(sys.argv[1]); D=T/'experiment/round48/docs/final_report';a=json.loads((D/'report_data.json').read_text(encoding='utf-8'));ev=a['evidence'];ref=Path(ev['reference'])
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
assert sha(ref)==ev['reference_sha256'] and json.loads((D/'pre_authoring_verification.json').read_text(encoding='utf-8'))['all_passed']
doc=Document(ref);tp=deepcopy(doc.tables[4]._tbl.tblPr);body=doc._element.body
for x in list(body):
    if x.tag!=qn('w:sectPr'):body.remove(x)
width=doc.sections[0].page_width.inches-doc.sections[0].left_margin.inches-doc.sections[0].right_margin.inches
headings=[];table_data=[]
def p(text='',bold=False,size=None):
    x=doc.add_paragraph(style='Normal');x.paragraph_format.space_after=Pt(7);x.paragraph_format.line_spacing=1.1
    r=x.add_run(text);r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
    if size:r.font.size=Pt(size)
    return x
def head(text,first=False):
    x=doc.add_paragraph(text,style='Title' if first else 'Heading 1');x.paragraph_format.page_break_before=not first;x.paragraph_format.keep_with_next=True
    for r in x.runs:r.font.color.rgb=RGBColor(0,0,0)
    headings.append(text)
def sub(text):p(text,True,12)
def num(v):
    if v in [None,'']:return '—'
    try:
        f=float(v)
        if abs(f)<0.0000005:f=0.0
        return f'{f:.6f}' if math.isfinite(f) else '—'
    except (ValueError,TypeError):return str(v)
def table(headers,rows,widths=None,green=None):
    widths=widths or [width/len(headers)]*len(headers);tb=doc.add_table(rows=1,cols=len(headers));tb.autofit=False
    tb._tbl.remove(tb._tbl.tblPr);tb._tbl.insert(0,deepcopy(tp))
    for i,w in enumerate(widths):tb.columns[i].width=Inches(w)
    for ri,vals in enumerate([headers]+rows):
        rw=tb.rows[0] if ri==0 else tb.add_row();pr=rw._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for j,v in enumerate(vals):
            c=rw.cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for k,vv in [('top',45),('bottom',45),('left',70),('right',70)]:
                e=OxmlElement('w:'+k);e.set(qn('w:w'),str(vv));e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar);bd=OxmlElement('w:tcBorders')
            for k in ['top','left','bottom','right','insideH','insideV']:
                e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
            cp.append(bd);fill='DCEAF2' if ri==0 else 'F5F7F9' if ri%2==0 else 'FFFFFF';best=ri>0 and green and green(ri-1,j)
            if best:fill='E9F4E5'
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);cp.append(sh)
            x=c.paragraphs[0];x.paragraph_format.space_after=Pt(0);x.paragraph_format.space_before=Pt(0);x.paragraph_format.line_spacing=Pt(13);x.paragraph_format.keep_with_next=False
            text=str(v);numeric=all(ch in '0123456789.-+%/—' for ch in text)
            x.alignment=WD_ALIGN_PARAGRAPH.CENTER if j>0 and numeric else WD_ALIGN_PARAGRAPH.LEFT
            rr=x.add_run(text);rr.bold=ri==0 or bool(best);rr.font.name='Calibri';rr.font.size=Pt(10.5);rr.font.color.rgb=RGBColor(0,0,0);rr._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    x=p();x.paragraph_format.line_spacing=Pt(1);x.paragraph_format.space_after=Pt(5)
    table_data.append(dict(headers=headers,rows=rows));return tb
old={h:a['r47']['scores'][str(h)]['original_HDSF'] for h in [1,7]}
current={h:a['r48']['scores'][str(h)]['user_current_R47'] for h in [1,7]}
trial={h:a['r48']['scores'][str(h)]['DEV_frozen_R48_candidate'] for h in [1,7]}
hist=a['historical_rows'];hist.sort(key=lambda x:(int(x['horizon_days']),float(x['wape']),x['model']))
for i,x in enumerate(hist,1):x['record_id']=f'H{i:03d}'
names={'neural_only_fixed_ablation':'仅重训神经成员 保留原本地成员','local_only_fixed_ablation':'仅重训本地成员 保留原神经成员','all_refit':'神经与本地成员同时重训'}
def hrow(h,model):return next(x for x in hist if int(x['horizon_days'])==h and (x['model']==model or model in x.get('aliases',[])))
def metricrows(fields):
    out=[]
    for label,key in fields:
        out.append([label]+[num(z[h]['metrics'].get(key)) for h in [1,7] for z in [old,current,trial]])
    return out
head('Google PowerData 2019 当前推荐方案与全面精度优化研究汇报',True)
p('主要指标改善 当前推荐保留第47轮方案 第48轮尚不支持稳定升级',True,14)
p('2026年10月12日  原Google功率数据  57个PDU  沿用9月28日统一评价',size=10)
main=[]
def gain(v):return f'{0.0 if abs(v)<0.0000005 else v:+.6f}'
for label,h,k in [('24小时完整曲线 WAPE%',1,'wape'),('未来7天完整曲线 WAPE%',7,'wape'),('未来7天每日均值 WAPE%',7,'daily_wape'),('24小时每日均值 WAPE%',1,'daily_wape'),('24小时日最大值 MAPE%',1,'peak_mape'),('7天日最大值 MAPE%',7,'peak_mape')]:
    b,c,n=[z[h]['metrics'][k] for z in [old,current,trial]];main.append([label,num(b),num(c),num(n),gain(b-c),gain(c-n)])
table(['任务与主要指标↓','原HDSF','当前推荐 R47','新候选 R48','R47收益\n百分点','R48新增收益\n百分点'],main,[3,1.45,1.45,1.45,1.45,width-8.8])
p('当前推荐已按你的明确决定采用：24小时使用HDSF半日单纯形加PatchTST FM日均残差校正；7天使用HDSF半日单纯形加TimesFM3 Toto2 PatchTST FM零均值形状校正。每日均值由同一曲线计算。',True)
p('这一轮测试数字仍能微小下降，但未实现大幅、稳定或所有指标领先。R48后段开发收益不稳，24小时峰值及尾部指标微退步，故保留R47。原冻结验证门槛未通过的记录保留；用户采用是接受取舍的决定，不等同于通过独立验证。')
p('较低的研究数字仍有：R48的24小时曲线2.329603%；旧R39功率输入7天曲线2.695253%；第46轮仅替换神经成员的固定事后消融7天日均1.862836%。这些不等同于当前部署。原方案、推荐方案和研究候选的取舍如下文完整说明。')
p('阅读顺序：前部先看任务结果与模型名称；随后看完整指标、训练与改进原因、开发筛选、字段与数据处理。后部给出114组历史不同预测、28条底层模型前后对照、固定消融及全部已保存区间结果。—表示该项没有保存，绝不以0补齐。')
head('当前推荐模型的明确名称与组成')
table(['任务','完整方案名称','固定配置'],[
['未来24小时','HDSF半日单纯形加PatchTST FM日均残差校正','IBM Granite PatchTST-FM-r2 curve负荷损失重训\n日均残差权重1 日内形状权重0'],
['未来7天曲线及每日均值','HDSF半日单纯形加TimesFM3 Toto2 PatchTST FM零均值形状校正','第46轮神经成员前后差分\n日均残差权重0 日内形状权重0.5'],
['本轮24小时研究候选','上述24小时方案加同场站同类型相对校正收缩','同组历史相对校正与组中位数各50%'],
['本轮7天研究候选','上述7天方案加增强零均值形状校正','第46轮差分的形状权重从0.5提高至1']],[1.8,4.9,width-6.7])
p('HDSF是 Half Day Simplex Fusion：按12小时区间拟合非负且和为1的成员权重。24小时成员池包括Amazon Chronos-2、NX-AI TiRex-2、Google TimesFM2.5、Datadog Toto2.0 313M、Google TimesFM3及本地LightGBM与滞后回归。7天成员为TimesFM3、Toto2.0、IBM Granite PatchTST-FM-r2及历史水平与日内形状平滑。部分成员权重为0，成员池中出现不代表实际贡献。')
p('24小时新校正把PatchTST微调前后的预测差分分成每天均值和去均值的形状，只保留均值部分。因此整天上移或下移，主要修正水平偏差，不直接重塑峰值位置。7天只保留零均值部分，尽量维持原HDSF优秀的每日均值；这是日均几乎完全不变的直接原因。')
p('基础HDSF的开发权重与测试权重分别用D14和D22前可得信息拟合；新校正和筛选均只使用相应起报前的历史。修正后限制在[0,1]，仅与利用率目标范围一致，不表示物理电量守恒。')
head('所有点预测指标的直接对照')
p('六列依次是原HDSF、当前R47、本轮R48；左右分别对应24小时与7天。↓越小越好，↑越大越好。均值指标按完整有效日计算。')
fields=[('曲线 WAPE%↓','wape'),('曲线 MAE↓','mae'),('曲线 RMSE↓','rmse'),('APE第95百分位%↓','p95_ape'),('逐点APE≤2%占比%↑','hit2_percent'),('逐点APE≤5%占比%↑','hit5_percent'),('全局绝对偏差%↓','absolute_bias_percent'),('每日均值 WAPE%↓','daily_wape'),('每日均值 MAE↓','daily_mae'),('每日均值 RMSE↓','daily_rmse'),('日均APE≤2%占比%↑','daily_hit2_percent'),('日均APE≤5%占比%↑','daily_hit5_percent'),('日最大值 MAPE%↓','peak_mape'),('峰值时刻误差 分钟↓','peak_time_mae_minutes'),('cell-a五域 WAPE%↓','cella5_wape'),('其余52域 WAPE%↓','other52_wape'),('曲线较朴素基线误差减少%↑','relative_curve_error_reduction_percent'),('日均较朴素基线误差减少%↑','relative_daily_error_reduction_percent'),('改善朴素基线的PDU数↑','improved_pdus'),('改善朴素基线的cell数↑','improved_cells'),('改善朴素基线的起报数↑','improved_origins')]
table(['统一指标','24h 原HDSF','24h R47','24h R48','7天 原HDSF','7天 R47','7天 R48'],metricrows(fields),[3.08]+[(width-3.08)/6]*6)
p('有效曲线点24小时49243、7天114902；有效完整PDU日508、1187。改善域数和改善起报数均对朴素历史基线，不代表每个域都超过原HDSF。MAE和RMSE为归一化功率利用率，不是kW。',size=10)
head('预测区间和峰值低估的完整对照')
ir=[]
for h in [1,7]:
    for label,z in [('原HDSF',old),('当前R47',current),('本轮R48',trial)]:
        x=z[h]['interval'];ir.append([f'{h}天 {label}',num(x['nominal_percent']),num(x['coverage_percent']),num(x['mean_width_normalized']),num(x['winkler80']),num(x['below_percent']),num(x['above_percent'])])
table(['模型与跨度','标称覆盖%','实际覆盖%','平均宽度↓','Winkler80↓','低于下界%','高于上界%'],ir,[2.55]+[(width-2.55)/6]*6)
table(['额外峰值诊断','当前 R47','本轮 R48'],[[f'{h}天峰值低估比例%',num(current[h]['peak_underestimate_percent']),num(trial[h]['peak_underestimate_percent'])] for h in [1,7]],[5.1,2.5,width-7.6])
p('区间用开发期逐PDU残差的10%和90%经验分位数加到点预测，并裁剪到[0,1]。它是附加经验区间，不是基础模型原生概率头。当前两项实际覆盖约75.54%和75.99%，均低于标称80%；曲线更准不自动意味着区间校准更好。')
p('峰值低估比例统计有有效数据的PDU日中，预测最大值小于真实最大值的比例。低估比例与峰值MAPE含义不同；无幅度信息时不能单凭比例判断综合优劣。原历史表未保存的额外指标在附表中保留缺失。')
p('当前R47相对原HDSF曲线收益的日期重采样探索95%区间：24小时[-0.008432, 0.019481]个百分点，7天[-0.000543, 0.001348]个百分点，均跨零。测试月已反复查看，7天起报窗口重叠，不能将这些区间当成独立确认或显著性证明。',True)
head('未来24小时完整曲线与强模型比较')
models=['last_source','ridge','lgb','chronos','tirex','tf25','equal','r32_selected_H1','r37_diagnostic_H1','r44_timesfm30_long_H1','r44_hdsf_H1']
rs=[]
for m in models:
    x=hrow(1,m);rs.append([x['report_name'].replace('\n',' '),num(x['wape']),num(x['daily_wape']),num(x['peak_mape']),num(x['mae']),num(x['rmse']),num(x['p95_ape'])])
for label,z in [('当前R47 PatchTST日均残差校正',current),('R48 同组50%收缩 研究候选',trial)]:
    x=z[1]['metrics'];rs.append([label]+[num(x[k]) for k in ['wape','daily_wape','peak_mape','mae','rmse','p95_ape']])
table(['模型及固定配置','曲线WAPE%↓','日均WAPE%↓','峰值MAPE%↓','MAE↓','RMSE↓','P95 APE%↓'],rs,[4.1]+[(width-4.1)/6]*6)
p('当前R47比原HDSF曲线WAPE减少0.005491个百分点，约0.235%的相对误差改善；日均减少0.012152个百分点，峰值MAPE减少0.087887个百分点。P95 APE反而增加约0.017164个百分点，日均2%命中率也略降。领先主要指标不等于每项最优。')
head('未来7天曲线与每日平均负荷分别比较')
ms=['mean3_source','chronos','tirex','tf25_7','r31_selected_H7','r36_selected_H7','r37_diagnostic_H7','r38_diagnostic_H7','r39_diagnostic_H7','r40_diagnostic_H7','r44_hdsf_H7']
rs=[]
for m in ms:
    x=hrow(7,m);rs.append([x['report_name'].replace('\n',' '),num(x['wape']),num(x['daily_wape']),num(x['daily_mae']),num(x['daily_rmse']),num(x['peak_mape'])])
for label,z in [('当前R47 神经成员零均值形状0.5',current),('R48 神经成员零均值形状1 研究候选',trial)]:
    x=z[7]['metrics'];rs.append([label]+[num(x[k]) for k in ['wape','daily_wape','daily_mae','daily_rmse','peak_mape']])
table(['模型及固定配置','7天曲线\nWAPE%↓','每日均值\nWAPE%↓','日均MAE↓','日均RMSE↓','日最大值\nMAPE%↓'],rs,[4.45]+[(width-4.45)/5]*5)
p('当前R47的7天日均与原HDSF相同，曲线只微幅改善。历史R40含cell-a真实CPU/内存，曲线2.694096%仍更低，但每日均值更差；它属于不同输入范围的旧研究候选。本轮不使用这些额外数据。即便只看功率输入，历史R39曲线2.695253%也低于当前，不能宣称所有阶段曲线已经全面领跑。',True)
head('重训本地成员与重叠分段的结果')
p('第45轮已比较多种分段长度、重叠及全局权重收缩，并重训本地残差模型；开发期共保存70条试验记录。原训练实现中D1残差标签处理已审查，重训明确从D2起取标签，避免无可用历史的D1残差参与训练。')
table(['跨度','开发冻结的明确候选','曲线WAPE%','日均WAPE%','峰值MAPE%'],[[h,x['candidate_name'],num(x['candidate_score']['wape']),num(x['candidate_score']['daily_wape']),num(x['candidate_score']['peak_mape'])] for h,x in a['r45']['scores'].items()],[.6,5.5]+[(width-6.1)/3]*3)
p('24小时冻结12小时不重叠、7天冻结24小时50%重叠；两者都没有稳定优于原HDSF。上述结果说明重叠区间已实际尝试，不能只因它更平滑就断言更准。本地训练日志、参数、逐折结果、控制及失败证据均保留。')
table(['跨度','MAE','RMSE','P95 APE%','≤2%占比%','≤5%占比%','绝对偏差%'],[[h]+[num(x['candidate_score'][k]) for k in ['mae','rmse','p95_ape','hit2_percent','hit5_percent','absolute_bias_percent']] for h,x in a['r45']['scores'].items()],[.6]+[(width-.6)/6]*6)
table(['跨度','日均MAE','日均RMSE','日均≤2%','日均≤5%','a五域WAPE%','其余52域%','峰时分钟'],[[h]+[num(x['candidate_score'][k]) for k in ['daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent','cella5_wape','other52_wape','peak_time_mae_minutes']] for h,x in a['r45']['scores'].items()],[.6]+[(width-.6)/7]*7)
head('底层模型重训的范围与真实结果')
table(['架构','重训与选定方式','实际运行范围'],[
['Amazon Chronos-2','输出头和最后模块两类 每类40与80步\n开发选tail80','最长28天历史 H1'],['NX-AI TiRex-2','开发选head80','14天历史 H1 CPU'],['Google TimesFM2.5','开发选head80','14天历史 H1'],['Google TimesFM3','开发选tail40','14天多变量 21天多变量及单变量 H1/H7'],['Datadog Toto2.0 313M','开发选tail40','同场站同类型及单域 H1/H7'],['IBM Granite PatchTST-FM-r2','开发选tail40','21天历史 H7']],[2.3,4.6,width-6.9])
p('第46轮六个架构均实际反向传播更新；不是只训练本地LightGBM。12个开发训练与6个最终训练，共1320次参数更新，保留30个状态检查点；另有8次本地LightGBM种子拟合。原预训练权重不改写，派生权重独立保存。所选前后预测作同设备差分，以降低CPU/GPU数值差异的影响。')
p('配置名称中head表示只训练输出头，tail表示训练末端模块及输出头；s40和s80表示40步、80步检查点。各模型最终配置由开发期选择。',size=10)
rs=[]
for h in [1,7]:
    b=old[h]['metrics'];allref=next(x for x in a['ablation'] if int(x['horizon_days'])==h and x['model']=='all_neural_and_local_refit')
    rs.append([f'{h}天 原HDSF',num(b['wape']),num(b['daily_wape']),num(b['peak_mape'])]);rs.append([f'{h}天 全成员重训 旧权重',num(allref['wape']),num(allref['daily_wape']),num(allref['peak_mape'])])
table(['方案','曲线WAPE%↓','日均WAPE%↓','峰值MAPE%↓'],rs,[5.1]+[(width-5.1)/3]*3)
p('原开发升级门槛选择alpha0。全量替换没有稳定胜过旧组合：部分单模型改善，却被原HDSF的零权重或成员相互抵消；本地7天成员的变差还会侵蚀神经成员的收益。所有底层模型前后完整指标在附表保留。')
head('改进来自哪里及为何仍然有限')
p('HDSF最初超过旧方案的主要原因，是不同半日采用不同专家。24小时约86.52%的改善发生在后半天，原Chronos和Toto承担前半天，TimesFM3的不同模式承担后半天。同一成员池的固定消融支持分段优于等权和整段共享权重；属于事后机制解释，不能据此重新挑测试权重。')
p('7天优势主要在平均水平。此前固定分析中日均水平MSE下降6.35%，形状MSE反而上升1.14%；历史水平与日型平滑平均占29.01%权重，去掉会使日均变差。主要专家残差相关约0.93至0.96，互补空间有限。这也解释了曲线WAPE不一定同步领先。')
table(['第46轮固定消融 仅解释','24h曲线WAPE%','7天曲线WAPE%','7天日均WAPE%'],[
['原HDSF',num(old[1]['metrics']['wape']),num(old[7]['metrics']['wape']),num(old[7]['metrics']['daily_wape'])],
['仅重训神经成员 保留本地成员','2.337215','2.695437','1.862836'],
['仅重训本地成员 保留神经成员','2.335124','2.705615','1.875918'],
['神经与本地成员同时重训','2.337226','2.700025','1.870515']],[5.5]+[(width-5.5)/3]*3)
p('第47轮不再盲目全部替换，而是将水平与形状拆开。PatchTST针对负荷曲线损失重训，24小时只使用日均差分；7天保留原优秀水平，仅吸收神经成员零均值差分。R47另外训练Chronos与PatchTST两类损失配方，四个开发运行和一个最终PatchTST运行，共800次真实更新；不是所有架构完整从零重训。')
p('上表神经单独替换的7天事后数字更低，但本轮未拿它与R47的24小时测试赢家拼成新最佳。第47轮冻结选择及用户采用保持分离，所有消融不是独立预测上限。绝对误差是非线性的，单个成员改善也不能按权重简单相加。')
head('第48轮稳健校正的开发筛选与结论')
p('本轮登记16个结构候选，每个跨度8个；没有新增大模型梯度训练。尺度仅由过去7个完整有效日的日均差分MAD估计，设最小尺度并按预测时距扩展；同组校正按cell和MVPP类型分组，至少3个PDU才共享，绝不输入未来真实负荷。')
dev=[]
variant_names={'current':'当前R47','hard_cap_05':'日均硬限0.5倍MAD','hard_cap_1':'日均硬限1倍MAD','soft_cap_1':'日均软限1倍MAD','soft_cap_2':'日均软限2倍MAD','group_25':'同组相对校正收缩25%','group_50':'同组相对校正收缩50%','group_50_hard1':'同组收缩50%及硬限1倍','level_hard_quarter':'新增日均权重0.25 硬限','level_soft_half':'新增日均权重0.5 软限','level_group_half':'新增日均0.5 同组收缩','level_group_hard_half':'新增日均0.5 同组及硬限','shape075':'零均值形状权重0.75','shape1':'零均值形状权重1','shape_group075':'形状0.75 同组收缩'}
for x in a['dev48']:
    v=json.loads(x['early_metrics']);dev.append([x['horizon_days'],variant_names[x['variant']],num(v['wape']),num(v['daily_wape']),num(v['peak_mape']),'是' if x['early_eligible']=='True' else '否'])
table(['跨度 天','已登记结构','早段曲线WAPE%','早段日均WAPE%','早段峰值MAPE%','早段可选'],dev,[.7,3.2,1.65,1.65,1.65,width-8.85])
p('H1早段冻结同组50%收缩，H7冻结shape1。H1后段两个日期收益+0.000377和−0.001270个百分点，均值−0.000447；H7后段均值+0.000400、标准误0.000860，正收益日期占比不足。两项均未通过稳定升级门槛。',True)
p('TEST只评分冻结候选，没有依测试名次改参数。两条曲线相对当前仅减少0.000019和0.000157个百分点，日均未全面改善，其他指标有退步；本轮作为保留研究结果，不替换当前推荐。32项来源、未来扰动、复现、指标与区间核查通过。')
head('原始Google功率数据字段的含义')
table(['字段','类型与含义','本研究用途'],[
['time','UINT64 微秒 相对trace起点前600秒\n标记5分钟平均测量窗口起点 非Unix时间','排序和时间对齐\n600000000对应2019年5月1日00时PT'],
['cell','匿名集群标识 如a 至h以及新增场站','同组信息 不推断真实地址'],['pdu','匿名供电域 如pdu6或mvpp1\n与cell共同识别一条序列','目标序列与MVPP类别'],['measured_power_util','5分钟总功率除以本域容量 范围[0,1]\n包含机房冷却等非IT功率','唯一预测目标 及历史输入'],['production_power_util','生产负载对应的估算功率利用率\n由CPU利用率及机器闲忙功率等插值估计','原local24历史代理特征\n不是实测CPU NCU'],['bad_measurement_data','总功率质量标记\n官方Notebook以true为坏点','决定目标有效掩码和历史聚合'],['bad_production_power_data','生产估算功率的独立质量标记\n官方Notebook以true为坏点','仅处理生产代理 不自动否定总功率']],[2.3,5.2,width-7.5])
p('官方PDF对bad标记存在反向文字描述；本研究遵循官方Notebook的NOT bad过滤和实际代码，未根据测试成绩改标记。功率只有归一化利用率，容量未公布，不能恢复绝对kW或把日均值说成kWh。')
p('机器映射附表为machine_id、pdu、cell三列，覆盖a至h。有96616行、96536个全局机器ID、96614个cell与machine组合；2个组合对应多个PDU，没有生效日期，不能完整恢复机器迁移。新校正未把它扩展成全机器实时算力输入。')
head('从原始CSV到训练和评价的处理')
table(['处理步骤','具体操作与数据核查'],[
['来源保留','57个PDU压缩CSV以及机器映射文件原样保留；解压CSV只是格式展开，未改原始观测。508896条5分钟功率记录。'],['时间与粒度','按原始time排序，每3个5分钟点聚合为15分钟；得到57×31×96共169632个位置。相对time不擅自转换为Unix或本机时区。'],['质量与评分','每块仅对有效原点求均值；只要3点中有坏点，该15分钟目标不参与评分。24个部分有效块，169608个有效目标。评分掩码与旧方案一致。'],['历史输入','部分有效块可用其有效点均值作为上下文，输入缺口仅向前处理，禁用未来回填；异常上下文与旧处理最大差0.005167。'],['生产代理','生产坏点116517个；38556个不可用15分钟块。历史因果中位数可填20452块，剩18104块保持缺失处理，不将所有缺失改0。'],['本地特征','local24共64个特征：功率滞后、近3/5/7天均值中位数、趋势、运行统计、同cell/留一域历史、生产代理、日历与日内谐波、序列和cell编码。'],['训练与残差','仅训练边界以前标签；开发筛选后冻结，再测试。水平为每天96点均值；形状为减去均值后的96点残差，来源绑定数据和权重SHA。']],[1.6,width-1.6])
p('负荷特性：各域整体水平较稳定，日内波动小但有跨日偏移；原有效功率利用率范围约0.374333至0.842333。7天窗口短且相关，日均水平比细粒度形状更容易保留优势。只有一个月、57域并不等于拥有大量独立日期，这限制了复杂重训的可识别空间。')
p('本轮继续只用原Google PowerData2019。以前下载的ClusterData cell-a、其他数据以及其历史实验仍保存用于溯源，不加入本轮输入；HDSF的生产估算代理属于原功率CSV字段。')
head('所有评价指标的统一计算定义')
table(['指标','计算规则'],[
['曲线WAPE','每PDU有效点绝对误差和除以该PDU实际绝对值和，乘100%，再对57域等权平均。不是池化所有点后计算。'],['每日均值WAPE','预测和真值都取完整96点均值；排除任意无效点的整日目标，然后先PDU计算再等权平均。'],['MAE 和 RMSE','全部有效点池化，分别求绝对误差均值、均方误差平方根。每日均值版本只用完整有效PDU日。'],['P95 APE','所有有效点的绝对百分比误差第95百分位；APE为绝对误差除以实际绝对值乘100%。'],['2% 与 5%命中','有效点中APE≤2%或≤5%的比例；日均版本按完整有效PDU日同样计算。'],['每日最大值MAPE','有有效数据的PDU日内，对有效槽分别求预测和真实最大值，然后平均最大值APE。'],['峰时误差','有效槽中的首个最大值位置差的绝对值乘15分钟，跨PDU日平均；不做24小时循环距离。平台或平峰下敏感。'],['全局绝对偏差','100%乘全部有效点有符号误差和的绝对值，除以实际绝对值和。不是绝对误差均值。'],['分域WAPE','cell-a的5域与其余52域各按域WAPE等权平均，不把CPU换算成kW。'],['误差减少率','100%乘1减模型WAPE除以同任务朴素基线WAPE；日均与曲线用各自基线。'],['改善数量','分别数有多少PDU、cell及起报在WAPE上优于对应朴素基线。'],['覆盖率 宽度 Winkler80','真值在区间内的有效点比例；上界减下界的均值；80%区间宽度加出界距离的10倍罚项。越窄不是单独最优。']],[2.0,width-2.0])
p('9月28日原材料明确写明Google先对每条供电域分别计算WAPE，再对57条取平均。源项目公布的池化曲线分数不直接与本报告比较。历史仅有效槽均值WAPE另作诊断附表，主排名始终使用完整96点日均。',True)
head('时间划分与建模问题审查')
table(['范围','24小时','7天'],[['开发起报','D14 D16 D18 D20','D14 D15'],['测试起报','D22至D30 共9次','D22 D23 D24 共3次'],['预测长度','96点','672点'],['朴素曲线基线','前一天同刻 3.199450%','前三天同刻均值重复 3.211919%'],['朴素日均基线','1.783005%','2.262403%'],['有效评分点和完整日','49243点 508个PDU日','114902点 1187个PDU日']],[2.7,3.7,width-6.4])
p('未来输入检查：重训与新校正均检查起报后的数据扰动不影响该起报预测，核对控制缓存、模型实际参数更新、预测时间边界、所有值有限且有界、区间顺序及统一指标重算。第46轮475项科学核查、42项原生未来回放、30项参数更新核查；第47轮133项科学与36项运行核查；本轮32项核查及汇报前28项导出核查均通过。')
p('通过实现核查不能证明不存在统计过拟合。第47轮跨日期开发升级门槛未通过，用户是在看到真实结果后明确采用固定候选；原冻结选择和SHA未改。第48轮仍只用开发早段选结构、后段验稳定，再冻结评分测试，没有因测试微小改善改变部署。')
p('测试集反复被查看：即使每轮不按测试排名直接选参，长期方法设计仍受历史测试结果启发。缺少未看过的确认月份，独立显著性和大幅全面提升尚未成立。57个域相关、7天窗口重叠，不能把全部15分钟点视为相互独立样本。',True)
head('当前可运行入口与后续研究判断')
p('文件实体目录为${PROJECT_ROOT}，桌面旧路径保留兼容junction。当前推荐记录与源权重绑定，历轮数据、实际权重、协议、成功及负面结果均保留。')
table(['交付与用途','实际文件或入口'],[['24小时推荐历史复现','33_1天当前推荐HDSF校正历史复现.cmd'],['7天推荐历史复现','33_7天当前推荐HDSF校正历史复现.cmd'],['推荐方案与采用决定','当前推荐方案.json 及 experiment/round47/deployment/采用说明.md'],['曲线 CSV','experiment/round47/deployment/当前推荐_H1_历史滚动曲线.csv 共49248行\n同目录 当前推荐_H7_历史滚动曲线.csv 共114912行'],['每日均值 CSV','experiment/round47/deployment/当前推荐_H1_每日平均负荷.csv 共513行\n同目录 当前推荐_H7_每日平均负荷.csv 共1197行'],['本轮原始研究与核查','experiment/round48/results 及 experiment/round48/docs/protocol.json'],['完整机器可读指标','experiment/round48/docs/final_report 中全指标CSV和report_data.json']],[2.5,width-2.5])
p('上述入口是已审计历史缓存重构，用于复现2019年5月测试起报，并已核对预测与保存数组完全相同；底层原生模型推理在前两轮实际验证。它不是任意新上传数据的通用预测接口，也不是今天电负荷。导出保留原数字time、起报D日、目标D日和有效掩码，区间为开发经验区间。')
p('进一步全面提升仍有可能，但现有证据只支持局部、很小的改善，不能承诺做到全部最优。在同一月上继续密集试参，收益很容易落在日期波动内。下一步优先增加真正未见日期的确认；若仍限原数据，应先做嵌套按时间验证及复杂度更低的校正，对新结构的收益稳定性作证，再决定是否升级。',True)
p('来源与许可：Google PowerData2019及文档为CC BY4.0，保留署名、许可和处理说明。TimesFM3权重仅限非商业非生产研究；其他模型许可随原模型目录保留，不能把数据许可当成全部模型许可。官方字段文档在google-powerdata-2019/docs。HDSF参考项目为github.com/daifeizhai/PowerData_Prediction，冻结源commit 648777507ebc57e7498ce61f4ae2236ab179e6a8。')

# Full numeric appendices preserve all recorded metrics and readable model names.
for start in range(0,len(hist),18):
    chunk=hist[start:start+18];head(f'附表 历史模型名称索引 {start+1} 至 {start+len(chunk)}')
    p('H编号对应后面三组全指标表。同一预测数组的别名合并，完整142条原记录保留在顶层CSV。索引包含未采用及不同旧输入范围的候选，不以排名表示采用。旧名称里的当前均指当时的冻结参照。',size=10)
    vals=[]
    for i in range(0,len(chunk),2):
        pair=[]
        for x in chunk[i:i+2]:pair.extend([f"{x['record_id']}\n{x['horizon_days']}天",x['report_name'].replace('\n',' ')+'\n'+x['model']])
        if len(pair)==2:pair.extend(['',''])
        vals.append(pair)
    table(['编号 跨度','明确模型与原记录键','编号 跨度','明确模型与原记录键'],vals,[.65,4.5,.65,width-5.8])
families=[('曲线与尾部误差',['wape','mae','rmse','p95_ape','hit2_percent','hit5_percent','absolute_bias_percent','development_wape_original_grid'],['曲线WAPE%','MAE','RMSE','P95 APE%','≤2%占比%','≤5%占比%','绝对偏差%','原网格DEV\nWAPE%']),('每日均值与峰值',['daily_wape','daily_mae','daily_rmse','daily_hit2_percent','daily_hit5_percent','peak_mape','peak_time_mae_minutes','daily_wape_valid_slots_supplement'],['日均WAPE%','日均MAE','日均RMSE','日均≤2%\n占比%','日均≤5%\n占比%','峰值MAPE%','峰时误差\n分钟','旧有效槽\n日均WAPE%']),('分域和基线改善',['cella5_wape','other52_wape','relative_curve_error_reduction_percent','relative_daily_error_reduction_percent','improved_pdus','improved_cells','improved_origins'],['a五域WAPE%','其余52域\nWAPE%','曲线误差\n减少%','日均误差\n减少%','改善PDU','改善cell','改善起报'])]
for title,keys,labels in families:
    for start in range(0,len(hist),21):
        chunk=hist[start:start+21];head(f'附表 历史模型全部指标 {title} {start+1} 至 {start+len(chunk)}')
        p('模型完整名称见H编号索引。WAPE按57域等权；旧有效槽日均只作诊断，不参与主排名。开发原网格可能含拟合样本，不当成独立确认。',size=10)
        table(['编号']+labels,[[x['record_id']]+[num(x.get(k)) for k in keys] for x in chunk],[.6]+[(width-.6)/len(keys)]*len(keys))

native=a['native']+a['native47']
for i,x in enumerate(native,1):x['record_id']=f'N{i:02d}'
for start in range(0,len(native),14):
    head(f'附表 所有底层模型前后对照名称 {start+1} 至 {min(start+14,len(native))}')
    table(['编号','跨度 天','架构与模式','权重及训练配置'],[[x['record_id'],x['horizon_days'],x['model_name'],x['version']+' '+x.get('training_configuration',x.get('recipe',''))] for x in native[start:start+14]],[.6,.7,4.0,width-5.3])
for title,keys,labels in families:
    for h in [1,7]:
        chunk=[x for x in native if int(x['horizon_days'])==h];head(f'附表 底层模型全部指标 {h}天 {title}')
        p('N01至N24为第46轮同设备前后差分口径；N25至N28为第47轮新PatchTST曲线损失原生同GPU控制与重训预测。不同口径不得直接当成同一原生输出。',size=10)
        table(['编号']+labels,[[x['record_id']]+[num(x.get(k)) for k in keys] for x in chunk],[.6]+[(width-.6)/len(keys)]*len(keys))

abl=a['ablation']
def ablabel(s):
    if s.startswith('neural_family_only_'):
        return '仅重训'+{'chronos2':'Amazon Chronos-2','tirex2':'NX-AI TiRex-2','timesfm25':'Google TimesFM2.5','timesfm3':'Google TimesFM3','toto2':'Datadog Toto2.0','patchtst':'IBM PatchTST-FM r2'}[s.replace('neural_family_only_','')]
    return {'neural_only_all':'仅重训全部神经成员','local_refit_only':'仅重训本地成员','all_neural_and_local_refit':'神经与本地成员同时重训'}[s]
for start in range(0,len(abl),18):
    chunk=abl[start:start+18];head(f'附表 固定消融与改进来源 {start+1} 至 {start+len(chunk)}')
    p('全部为固定事后机制诊断，没有用于本轮按测试挑选方案；完整补充指标保存在固定消融全指标CSV。正收益表示相对原HDSF误差减少。',size=10)
    table(['跨度 天','固定消融结构','曲线WAPE%','日均WAPE%','峰值MAPE%','曲线收益\n百分点','日均收益\n百分点'],[[x['horizon_days'],ablabel(x['model']),num(x['wape']),num(x['daily_wape']),num(x['peak_mape']),num(x['curve_gain_pp']),num(x['daily_gain_pp'])] for x in chunk],[.65,3.6]+[(width-4.25)/5]*5)

ivs=a['intervals'][:]
for x in a['r45rows']:
    if x['role']=='DEV_selected_candidate':ivs.append(dict(horizon_days=x['horizon_days'],method='R45 '+x['model_name'],kind='开发残差经验区间',**{k:x.get(k) for k in ['nominal_percent','coverage_percent','mean_width_normalized','winkler80','below_percent','above_percent']}))
for h in [1,7]:
    for label,z in [('当前R47',current),('本轮R48',trial)]:ivs.append(dict(horizon_days=h,method=label+' 开发残差10/90经验区间',**z[h]['interval']))
def ivlabel(x):
    kind=x.get('kind','')
    tag='历史引用 原数组未保存' if '历史引用' in kind else '原生分位数' if kind=='保存原生分位数' else '经验或原保存校准区间'
    name=x.get('method') or x.get('model')
    assert name,'Interval model name missing'
    if not x.get('method'):
        name={'TTM512-720-r2':'IBM TinyTimeMixer R2 512点历史 原生720点裁剪','TTM1024-720-r2':'IBM TinyTimeMixer R2 1024点历史 原生720点裁剪','TTM512 multitask head':'IBM TinyTimeMixer R2 512点 多任务输出头训练','TTM512 multitask all':'IBM TinyTimeMixer R2 512点 多任务全参数训练'}.get(name,name.replace('JointFactorFourier','57PDU联合因子与Fourier直接预测'))
    return str(x['horizon_days'])+'天 '+name+'\n'+tag
for start in range(0,len(ivs),8):
    chunk=ivs[start:start+8];head(f'附表 历轮区间记录与本轮结果 {start+1} 至 {start+len(chunk)}')
    p('均为80%区间。标明历史引用的原数组未保存，不参加本轮实测排序；原生分位数与经验区间区分。覆盖率需与宽度和Winkler一起看，不给没有区间的点模型虚构结果。',size=10)
    table(['跨度与方法及记录类别','实际覆盖%','平均宽度↓','Winkler80↓','低于下界%','高于上界%'],[[ivlabel(x)]+[num(x.get(k)) for k in ['coverage_percent','mean_width_normalized','winkler80','below_percent','above_percent']] for x in chunk],[4.4]+[(width-4.4)/5]*5)

# Keep the source package and its opaque parts exactly, changing only authorized parts.
for footer in doc.sections[0].footer.paragraphs:
    for r in footer.runs:
        if '2026' in r.text:r.text='2026年10月12日 当前推荐与精度优化 '
st=doc.settings.element;e=st.find(qn('w:updateFields'))
if e is None:e=OxmlElement('w:updateFields');st.append(e)
e.set(qn('w:val'),'true')
work=D/'working.docx';doc.save(work);out=T/'Google PowerData 2019 当前推荐方案与全面精度优化研究汇报.docx'
allowed={'word/document.xml','word/footer1.xml','word/settings.xml'}
with zipfile.ZipFile(ref) as src,zipfile.ZipFile(work) as edited,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
    for info in src.infolist():z.writestr(info,edited.read(info.filename) if info.filename in allowed else src.read(info.filename))
with zipfile.ZipFile(out) as z:
    assert all(hashlib.sha256(z.read(x['path'])).hexdigest()==x['sha256'] for x in ev['parts'] if x['scope']=='preserve_only')
assert sha(ref)==ev['reference_sha256']
checkdoc=Document(out);assert len(checkdoc.sections)==1
assert str(checkdoc.sections[0]._sectPr.xml)==str(doc.sections[0]._sectPr.xml)
manifest=dict(document=str(out),document_sha256=sha(out),reference_unchanged=True,preserve_only_parts_equal=True,heading_count=len(headings),table_count=len(table_data),headings=headings,table_data=table_data,created_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),all_pages_visually_reviewed=False)
(D/'authoring_evidence.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
for file,records in [('历史不同预测全指标.csv',hist),('底层模型前后全指标.csv',native),('固定消融全指标.csv',abl),('全部区间全指标.csv',ivs)]:
    keys=list(dict.fromkeys(k for x in records for k in x))
    with (D/file).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=keys);w.writeheader();w.writerows({k:json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in x.items()} for x in records)
print(json.dumps(dict(document=str(out),headings=len(headings),tables=len(table_data),source_preserved=True),ensure_ascii=True))
