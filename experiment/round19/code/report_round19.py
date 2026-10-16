from pathlib import Path
import sys,json,hashlib,shutil,ast
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
T=Path(sys.argv[1]);B=T/'experiment';R=B/'round19';Q=R/'results'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
s=read(Q/'summary.json');v=read(Q/'independent_verification.json');d=read(Q/'dev_summary.json');r18=read(B/'round18/results/summary.json');d18=read(B/'round18/results/dev_summary.json');assert v['all_passed'] and read(B/'round18/results/independent_verification.json')['all_passed'];assert s['selection']['candidate']['id']=='global_peak0' and s['selection']['chosen_alpha']==.25
for h in [1,7]:assert read(Q/f'exports/export_{h}day_verification.json')['rows']==57*h*96
out=T/'Google PowerData 2019 三项预测任务与模型比较汇报.docx'
if '--validate-only' in sys.argv:print(json.dumps(dict(round19_verified=True,selected_alpha=.25,descriptive_bests=len(s['leading_metrics7']),document_not_modified=True)));sys.exit(0)
backup=R/'docs/before_report_update';backup.mkdir(parents=True,exist_ok=True)
if '--repair-layout' not in sys.argv:
 assert not (backup/out.name).exists(),'Do not append twice'
 shutil.copyfile(out,backup/out.name);shutil.copyfile(B/'model_comparison_20261005/docs/report_qa_task_revision/verification.json',backup/'prior_visual_verification.json')
doc=Document(backup/out.name)
tree=ast.parse(Path(__file__).with_name('report_revision.py').read_text(encoding='utf-8'));node=next(x for x in tree.body if isinstance(x,ast.FunctionDef) and x.name=='table');exec(compile(ast.Module(body=[node],type_ignores=[]),'table_helper','exec'))
# Refresh current opening; prior comparison tables remain clearly scoped historical controls.
doc.paragraphs[1].text='本汇报分别评价未来24小时曲线、未来7天曲线与每日平均负荷，并比较曲线、峰值和稳定性。第19轮按开发期冻结的7天新方法在回顾比较中有7项最低误差；详细模型、选参证据和局限见末三页。'
t=doc.tables[0]
for i,k in [(2,'wape'),(3,'daily_wape')]:
 t.rows[i].cells[1].text='TimesFM3原生672点＋Ridge时频残差＋全场预测总量约束'+(' 的曲线逐日平均' if i==3 else '')
 t.rows[i].cells[2].text=f'{s["scores7"]["selected"][k]:.6f}'
 for c in t.rows[i].cells:
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
   for rr in p.runs:rr.font.size=Pt(10.8)
doc.paragraphs[6].text='原18个7天对照已补测。第18轮原生7天残差候选出现曲线与日均改善，但开发期未获采用；第19轮加入全场预测总量约束后，开发期支持25%校准权重。旧表为前期对照，最新10维结果在第19轮结果表中列出。'
doc.paragraphs[7].text='按已知预测跨度选择分支可形成一套方法：24小时保留第8b轮，7天采用第19轮冻结研究配置。它在7天比较中有7项最低误差，尚未全指标第一。日均取同一曲线平均，功率单位为归一化利用率，不是kW或kWh。'
doc.paragraphs[8].text='参数按开发期冻结后才评分新候选，未按测试排名换参。5月测试已多轮查看，7天窗口重叠，缺少独立确认月；最低误差是该月的描述性结果，不能称为已证明大幅或显著提升。'
for p in doc.paragraphs:
 if p.text.startswith('TimesFM3的21天历史单变量版最低'):p.text='原18项7天对照中，'+p.text
 if p.text.startswith('TimesFM3的21天单变量曲线逐日平均WAPE最低'):p.text='原18项7天对照中，'+p.text
 if p.text.startswith('当前功率代理输入下，近期微调已经进入工程平台期'):p.text=p.text.replace('当前功率代理输入下，近期微调已经进入工程平台期','截至第17轮，24小时功率代理输入的近期微调已进入工程平台期',1)
def page(title):
 p=doc.add_paragraph(title,'Heading 1');p.paragraph_format.page_break_before=True
def para(t):doc.add_paragraph(t)
names={'wape':'曲线WAPE %','mae':'MAE 归一化功率','rmse':'RMSE 归一化功率','daily_wape':'日均WAPE %','peak_mape':'峰值MAPE %','p95_ape':'P95点APE %','peak_time_mae_minutes':'峰值时刻偏差 分钟','absolute_bias_percent':'绝对总偏差 %','cella5_wape':'a集群5域WAPE %','other52_wape':'其余52域WAPE %'}
def f(k,a):return f'{a[k]:.6f}' if k in ['mae','rmse'] else f'{a[k]:.4f}'
page('第18轮原生七天残差校准结果')
para('第18轮有曲线和日均进步，但未满足预登记的开发期采用规则。固定诊断候选在测试中有6项最低误差，仍未替换原基础；本轮真实负面选择记录继续保留。')
para('明确模型为TimesFM3历史21天直接672点预测，加零截距Ridge时频残差校准。Ridge正则1000，学习日均残差与前三阶傅里叶形状；无正则意见门控，系数修正上限为历史7天均值的5%。预训练骨干冻结。')
dw=next(x for x in d18['weights'] if x['alpha']==.25)
rows=[[names[k],f(k,d18['baseline_scores']),f(k,dw['scores']),f(k,r18['scores7']['diagnostic'])] for k in ['wape','daily_wape','peak_mape','absolute_bias_percent','cella5_wape']]
table(['指标 越低越好','开发7天\n原TimesFM3','开发7天\n25%候选','测试7天\n固定诊断候选'],rows,[2.12,1.43,1.48,1.64],10.2)
para('31个预登记配置包括原方案、18个既有残差方式和12个零截距方式。前一个7天开发起报选候选家族，后一个起报及两窗口汇总选权重。原规则要求不减少开发指标第一数量，候选从8项降至6项；a集群误差和总偏差也有退步，因此最终采用权重为0。诊断候选25%权重在测试前单独冻结。')
para('诊断曲线WAPE为2.7551%，日均WAPE为1.9382%；峰值MAPE上升到2.1830%，绝对总偏差升到0.4557%。这些结果提示残差中存在有用信息，但未经开发支持的测试低分不用于改写第18轮选择。')
para('开发校准器只使用第14天之前目标，最终校准器只使用第22天之前目标。训练原生预测缓存覆盖起报7至21，重叠目标按出现频率反比赋权。第17轮次日残差代理筛选没有胜出，本轮独立筛选原生7天分支。')
para('已核对20项实施检查，包含未来标签扰动、输入边界、原模型复现、冻结记录和全指标复算。最初工具提取预检失败发生在评分前，原失败证据已保留；没有按测试结果更换模型。')
page('第19轮七天多项指标领先的研究方案')
para('明确模型为TimesFM3原生672点预测，加零截距Ridge1000前三阶傅里叶残差校准及全场预测总量约束，采用权重25%。此权重由开发期选定，评分前已冻结；它在21个既有对照中取得7项描述性最低误差。')
table(['指标 越低越好','原TimesFM3\n7天基础','第18轮\n诊断候选','第19轮\n冻结选定','既有21对照\n各项最低'],[[names[k],f(k,s['scores7']['prior']),f(k,r18['scores7']['diagnostic']),f(k,s['scores7']['selected']),f(k,s['control_best7'])] for k in names],[2.05,1.12,1.12,1.18,1.20],10.0)
para('本轮第一的7项为曲线WAPE、MAE、RMSE、日均WAPE、P95点APE、a集群5域WAPE和其余52域WAPE。WAPE与MAE有关联，分组WAPE也由相同预测组成，不能把7项排名当作7次独立验证。对照包含18个原主方案、第15与16轮冻结诊断候选及第18轮诊断候选。')
prior=s['scores7']['prior'];new=s['scores7']['selected']
para(f'相较原TimesFM3，曲线误差相对减少{100*(1-new["wape"]/prior["wape"]):.2f}%，日均减少{100*(1-new["daily_wape"]/prior["daily_wape"]):.2f}%；峰值MAPE也由2.0280%降至1.9764%。绝对总偏差仍约0.1656%，增加约0.000023个百分点，源于有效点评分与完整网格约束的区别。')
para('峰值幅值、峰值时刻和绝对总偏差没有取得对照第一。24小时分支沿用第8b轮，WAPE仍为2.374128%，本轮没有宣称次日新提升，也没有用各日期或各指标的未来真值动态选择专家。')
dm=read(Q/'daily_mean_evaluation.json');assert len(dm['rows'])==22 and len(dm['new_leading_daily_metrics'])==3
para('单独复算每日平均负荷后，日均MAE为0.013362、日均RMSE为0.017124，连同日均WAPE均低于全部21个既有对照。单位为归一化利用率；这两项补充评价不参与选参，22种方法完整对照见原生七天调优每日平均负荷比较.csv。')
page('第19轮选择依据和可复现交付')
para('本轮事先登记8种新约束及不调整对照：全场、集群或单PDU保持每个7天窗口的预测总量，并比较是否保留原模型每日最大值。选择目标改为多项最优，同时限制曲线、日均、RMSE、P95、峰值、偏差和a集群退步；该新规则不改写第18轮。')
chosen=next(x for x in d['weights'] if x['alpha']==.25)
table(['开发起报或汇总','原曲线WAPE %','新曲线WAPE %','原日均WAPE %','新日均WAPE %'],[[label,f('wape',old),f('wape',new),f('daily_wape',old),f('daily_wape',new)] for label,old,new in [('起报14',d['configs'][0]['scores'],chosen['scores14']),('起报15',d['weights'][0]['scores15'],chosen['scores15']),('两窗口汇总',d['baseline_scores'],chosen['scores'])]],[1.47,1.30,1.30,1.30,1.30],10.3)
para('先按起报14、固定25%权重选全场总量约束，再用起报15与汇总选择0／10%／25%／50%／100%。25%是唯一通过全部门槛的非零权重，汇总有8项开发指标最低；50%及100%未通过，完整结果均保留。两个开发窗口重叠，不能视为独立重复。')
para('投影只对已有预测加公共偏移并截断到0至1，使57个PDU的7天归一化预测值总和等于原TimesFM3总和；最终取75%原预测加25%投影候选。它不是已知容量加权的总kW，不是能量守恒。保留每日峰值的候选没有被采用。')
ci=read(Q/'exploratory_date_bootstrap.json')['metrics']
para('按9个独特测试日期配对重采样5000次，新减原的探索性95%区间：曲线−0.0912至−0.0114个百分点，日均−0.1318至−0.0202个百分点，峰值−0.1258至0.0164个百分点。日期仍有序列依赖，测试多轮查看且基础模型曾由测试识别；区间不能用作独立显著性证明。')
para('统一研究入口按已知跨度路由：1天为Chronos-2 TiRex-2 TimesFM3固定权重加Ridge校准与TimesFM3末层微调；7天为本轮方案。两个分支属于一套选择程序，使用不同骨干配置。17号两个历史复现入口已实际导出5472条次日点预测、38304条7天点预测及399个7天日均值。')
para('输出是2019年5月末起报的历史推理复现，未来区间为6月1日00时至6月7日23时45分PT，UTC对应6月1日07时至6月8日06时45分；没有6月真值，不是今天的负荷预测。第18轮原CSV的UTC标签误写已留勘误，规范导出使用显式PT与UTC。7天经验区间来自两个开发窗口，不保证80%覆盖。')
para('18项第19轮实施核对通过。完整记录和多维对照见根目录第19轮多项最优研究记录.md、原生七天调优多维度对照.csv；已冻结代码、权重引用、参数、预测和负面结果在experiment/round18与round19。09与02原入口继续保留，未将旧TiRex-2入口静默换成新方案。')
doc.save(out);print(json.dumps(dict(document=str(out),sha256=sha(out),appended_rounds=[18,19]),ensure_ascii=False))
