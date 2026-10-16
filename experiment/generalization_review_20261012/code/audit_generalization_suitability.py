from pathlib import Path
import sys,json,csv,hashlib,datetime,collections
import openpyxl,numpy as np
T=Path(sys.argv[1]).resolve();R=T/'experiment/generalization_review_20261012';R.mkdir(parents=True,exist_ok=True)
def read(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):
    with p.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def save(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
wp=T/'海光寺-能耗能效报表_9.1-9.30电能统计V2.xlsx';h=sha(wp);w=openpyxl.load_workbook(wp,read_only=True,data_only=True);s=w['报表'];byday=collections.defaultdict(list)
for i,row in enumerate(s.iter_rows(values_only=True),1):
    if i>=21 and isinstance(row[0],str):byday[row[0]].append(list(row[:8]))
assert len(byday)==30 and all(len(v)==3 for v in byday.values())
eq={k:all(abs(v[0][j]-v[1][j]-v[2][j])<1e-6 for v in byday.values()) for k,j in [('IT',1),('cooling',2),('total',4)]}
assert all(eq.values());assert sha(wp)==h
sea=dict(source=str(wp),sha256=h,sheet='报表',range='A20:H110',date_min=min(byday),date_max=max(byday),unique_days=len(byday),numeric_data_rows=sum(map(len,byday.values())),records_per_day=3,aggregate_equals_two_subseries=eq,primary_daily_total_kwh_min=min(v[0][4] for v in byday.values()),primary_daily_total_kwh_max=max(v[0][4] for v in byday.values()),intraday_power_curve_available=False,possible_tasks=['next-day total energy kWh','seven future daily energy values'],not_supported=['96-point next-day power curve','672-point seven-day power curve','intraday peak power or time'],secondary_daily_average_power='Eday/24 only if eachreportedday isconfirmed acomplete24hourinterval; doesnotrestoreintradaycurve',independent_new_test=False,reason='SameSeaGuangSeptember2022workbook wasalready examined inpreviousprojectexperiments; 30datesdo notsupportlarge-scale fine-tuning orstrongindependentconfirmation.')
nm=read(T/'nlr-hpc-pue/manifest.json');na=read(T/'nlr-hpc-pue/results/quality_audit.json');bindings={}
for x in nm['files']:
    if x['path'].endswith('.parquet'):
        p=T/'nlr-hpc-pue'/x['path'];assert p.stat().st_size==x['bytes'] and sha(p)==x['sha256'];bindings[str(p)]=x['sha256']
power=na['files']['power'];monthly=list(csv.DictReader((T/'nlr-hpc-pue/results/monthly_coverage.csv').open(encoding='utf-8-sig')))
nlr=dict(source=str(T/'nlr-hpc-pue'),raw_file_hashes_verified=bindings,power_rows=power['rows'],timeline=power['timeline'],primary_target='facility_total_kw derived from same-timestamp fivepowercomponents; notPUE*IT asmodelinput',primary_fields=['it_power_kw','cooling_kw','hvac_kw','plug_and_light_kw','pump_kw'],primary_unit='kW',primary_common_field_start=max(power['fields'][k]['first_finite_ts'] for k in ['it_power_kw','cooling_kw','hvac_kw','plug_and_light_kw','pump_kw']),secondary_target='it_power_kw assessed separately; PUE assessed asaseparateefficiencytask',component_accounting=na['component_accounting_diagnostic'],no_extra_2_67kw_added=True,PUE_separate_efficiency_task=True,timezone_unspecified=True,usable_tasks=['15-minute grid H1=96,H7=672','samecurve full96valid dailymeans','same15minute-grid peak andtailmetrics'],requires=['causal timestampgrid andprefixedqualityrules','train/past-only scaling inverse tophysicalunits','nohardclipphysicalkW to0..1','nofutureobservedweather orfuturetruepowercomponents'],year_record_counts={str(y):sum(int(x['rows']) for x in monthly if x['dataset']=='power' and x['month'].startswith(str(y))) for y in [2022,2023,2024,2025]},candidate_split=dict(training='2016-06-12 to2022-12-31',development='2023',initial_external_test='2024',additional_untuned_confirmation='2025-01-01 tolastcompleteeligiblewindow',status='design proposal; noexternalmodelscorecomputed; boundariesadjustonlyforpredeclaredqualityeligibility,notmodelperformance'))
cp=T/'google-clusterdata-2019/aggregated/complete_usage_a_15min.npz';ca=read(T/'google-clusterdata-2019/results/complete_projection_audit.json')
with np.load(cp) as z:
    assert bool(z['complete']) and int(z['source_files'])==int(z['expected_source_files'])==12085
    ag=z['aggregate'];covered=[str(z['names'][i]) for i in range(57) if np.any(ag[i,:,:,0]>0)];placeholders=[str(z['names'][i]) for i in range(57) if np.all(ag[i]==0)]
assert len(covered)==5 and len(placeholders)==52 and set(covered)=={x['domain'] for x in ca['domains']}
cluster=dict(source=str(T/'google-clusterdata-2019'),complete_selected_column_cell_a=True,source_files=12085,period='May2019 sameoriginalGooglePowerDatamonth',actual_cpu_domains=covered,zero_placeholder_domains=placeholders,expected_slots_per_actual_domain=31*96,power_label_newly_independent=False,units=dict(cpu='NCU',memory='sum ofnormalizedmemoryusage'),appropriate_use=['samepowerlabel CPU/memorycovariate incrementalvalue matchedcontrols','workloadprediction asaseparate task','missingchannelrobustness withexplicitavailabilitymask'],inappropriate_use=['thirdindependentexternalpowerdataset','treat52placeholderdomainzerosasobservedCPU','callCPUforecast aphysicalkWpowerforecast'])
cfg=read(T/'experiment/round48/deployment/selected_configuration.json');assert cfg['scheme_id']=='R48_ROBUST_HDSF_USER_ADOPTED'
deployment={}
for h in [1,7]:
    q=T/'experiment/round48/deployment';ex=read(q/f'export_H{h}.json');cp=q/f'当前推荐_H{h}_历史滚动曲线.csv';dp=q/f'当前推荐_H{h}_每日平均负荷.csv';assert sha(cp)==ex['curve_sha256'] and sha(dp)==ex['daily_sha256']
    rr=list(csv.DictReader(cp.open(encoding='utf-8-sig')));dr=list(csv.DictReader(dp.open(encoding='utf-8-sig')))
    assert len(rr)==57*(9 if h==1 else 3)*96*h==ex['curve_rows'];assert len(dr)==len(rr)//96
    pred=np.array([float(x['prediction_normalized']) for x in rr]);actual=np.array([float(x['actual_normalized']) if x['evaluation_valid']=='1' else np.nan for x in rr]);mask=np.isfinite(actual);p=pred.reshape(57,-1);y=actual.reshape(57,-1);m=np.isfinite(y)
    cw=np.mean([100*abs(p[i][m[i]]-y[i][m[i]]).sum()/abs(y[i][m[i]]).sum() for i in range(57)]);assert abs(cw-ex['metrics']['wape'])<1e-12
    assert np.isfinite(pred).all() and ((pred>=0)&(pred<=1)).all() and len({(x['cell'],x['PDU']) for x in rr})==57 and int(mask.sum())==ex['metrics']['valid_points']
    deployment[str(h)]=dict(all_passed=True,curve_rows=len(rr),daily_rows=len(dr),independent_macro57_recompute=float(cw),configuration_sha256=sha(q/'selected_configuration.json'))
result=dict(created_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),assessment_completed=True,external_model_scoring_started=False,new_recommendation=cfg['scheme_id'],deployment_verification=deployment,sea=sea,nlr=nlr,cluster=cluster)
save(R/'suitability_audit.json',result)
design=dict(status='assessment and evaluation design; external model inference notyetstarted',current_reference_configuration_sha256=sha(T/'experiment/round48/deployment/selected_configuration.json'),
    outcomes_separate=['frozenweights transfer withdeterministicinputadapter','newdata train/development-only refit offrame; notzeroshot'],
    common_metrics=dict(WAPE='peractualseries sumabs(error)/sumabs(truth)*100 thenequalmacroN; GoogleN57unchanged; notmixdifferenttargetunits',daily='full96valid derivedmeans for15minpowerseries; originaldailyenergy staysdailyenergy',zero_APE='undefinedzeroAPE markedNA withcounts; notepsilon-scorehide',additional='MASE andRMSSE scales fromtrainingonly; denominatorzero=>undefined'),
    temporal='chronological only; nofuturelabels inhistory/features/normalization/calibration; noDEVlabeloverlapwithTEST; freezebeforeTEST',
    baseline=['previousday sametime','previousweek sametime','recent3day sametime mean','simpleRidge/LightGBM whenadequatetraininghistory'],
    uncertainty='pairedcalendar timeblock resampling at leastH7span; overlappingorigins andPDUcorrelation notIIDpoints; multiplemodelsearchandtestexposure disclosed',
    test_reports=['H1curve','H7curve','H7dailymean','leadD1toD7','peakMAPE/time/underestimation','MAE/RMSE/P95/hit2/hit5/bias','coverage/width/Winkler80','quarter/regime/qualitycoverage'],
    scope_decisions=dict(main_external_power_validation='NLR',small_sample_daily_method_transfer='SeaGuang',existingdataset_covariate_ablation='GoogleClusterData2019'),references=['https://otexts.com/fpp3/toolbox.html','https://robjhyndman.com/hyndsight/wape.html'])
save(R/'evaluation_design.json',design)
current=(T/'开始阅读.md').read_text(encoding='utf-8');(T/'开始阅读.md').write_text('''# 当前推荐与新验证安排

用户已明确采用第48轮固定候选为新推荐：24小时HDSF半日单纯形加PatchTST FM日均残差及同组50%收缩；7天HDSF半日单纯形加TimesFM3 Toto2 PatchTST FM零均值形状全量校正。曲线WAPE2.329603%及2.701184%，7天每日均值1.868830%。后段开发门槛未通过的事实保留，采用不改写为独立显著提升。

当前34、19和33入口已同步第48轮并实际复现两项历史CSV。原第47轮部署和58页Word保留为采用决定之前的历史快照。最新采用说明在experiment/round48/deployment/采用说明.md。

三个拟验证数据集的适用性审查在experiment/generalization_review_20261012/suitability_audit.json；科学评价设计在同目录evaluation_design.json。本次尚未计算外部模型成绩。NLR适合外部功率预测，海光寺仅30天日能耗适合日任务小样本诊断，ClusterData是同月辅助输入，不能算独立外部功率验证。

所有实体文件继续位于D盘；桌面旧目录为junction。统一公式按实际目标序列数等权，不把缺失填零或日能耗伪造为15分钟曲线。
''',encoding='utf-8')
(T/'当前进展.md').write_text('''# 当前进展

2026年10月12日，第48轮已由用户明确采用为当前推荐；原研究协议、冻结选择、R47配置和原报告没有改写。两项预测导出及独立WAPE重算通过。

海光寺、NLR、GoogleClusterData三个数据集的可用性和评价适用性审查完成，外部模型推理与评分尚未开始。拟按外部NLR功率、小样本海光寺日电量、同月Cluster协变量对照三类分别评估，不合并分数假称三数据集独立泛化。
''',encoding='utf-8')
print(json.dumps(dict(all_passed=True,current_scheme=cfg['scheme_id'],sea_days=30,nlr_power_rows=power['rows'],cluster_actual_domains=len(covered),cluster_missing_placeholders=len(placeholders),external_model_scores_computed=False)))
