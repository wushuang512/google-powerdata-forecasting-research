"""Standalone report for the user-supplied NLR protocol and added source controls."""
from pathlib import Path
import argparse,datetime,hashlib,json,importlib.util,csv
from docx.shared import Pt
from docx.oxml.ns import qn
ap=argparse.ArgumentParser();ap.add_argument('root');ap.add_argument('--validate-only',action='store_true');a=ap.parse_args();R=Path(a.root);F=R/'experiment/nlr_strict_protocol_20261013'
def read(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):
    with p.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
audit=read(F/'results/independent_verification.json');assert audit['all_passed']
raw_audit=read(F/'results/raw_stream_verification.json');assert raw_audit['all_passed']
assert raw_audit['processed_sha256']==sha(F/'data/processed.npz') and raw_audit['coverage_csv_sha256']==sha(F/'data/raw_interval_coverage.csv')
assert read(F/'results/diagnostic_tables_complete.json')['rows']==210
for rel,h in audit['bindings'].items():assert sha(R/rel)==h,rel
s=read(F/'results/summary.json');rows=s['rows'];assert len(rows)==42
q=read(F/'data/quality_report.json');sel=read(F/'results/selection.json');cfg=read(F/'docs/protocol.json');hist=read(F/'docs/historical_controls.json')
if a.validate_only:print(json.dumps(dict(all_passed=True,rows=len(rows),source_checks=audit['count'])));raise SystemExit()
spec=importlib.util.spec_from_file_location('report_style',R/'experiment/completed_main_summary_20261013/code/report_style.py');style=importlib.util.module_from_spec(spec);spec.loader.exec_module(style);rep=style.Report();doc=rep.doc;bound=[]
doc.core_properties.title='NLR严格协议重跑评估汇报';doc.core_properties.subject='新协议测试与固定源权重迁移分开解释'
original=rep.table
def table(*args,**kwargs):
    t=original(*args,**kwargs)
    for node in t._tbl.tblPr.findall('.//'+qn('w:top'))+t._tbl.tblPr.findall('.//'+qn('w:bottom')):
        if node.get(qn('w:w')) is not None:node.set(qn('w:w'),'45')
    p=doc.paragraphs[-1];p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.space_after=Pt(5)
    return t
rep.table=table
def new(title):rep.h(title).paragraph_format.page_break_before=True
rep.new=new
LABEL={'LightGBM':'LightGBM目标拟合','timesfm3':'TimesFM3原权重','toto2':'Toto2原权重','moirai2':'Moirai2原权重','HDSF':'HDSF四专家目标拟合','last':'前一天同刻','week':'上周同刻','mean3':'前三天同刻均值',
 'Google_Chronos_TiRex_equal':'Chronos2与TiRex2等权','Google_three_fixed':'三专家固定权重','Google_R5_Ridge':'三专家加Ridge校准','Google_R8b_finetuned':'第8轮TimesFM3微调组合','Google_HDSF_original':'原HDSF源权重迁移','Google_R48_robust':'第48轮HDSF源权重迁移'}
ORDER=['last','week','mean3','LightGBM','timesfm3','toto2','moirai2','HDSF']+hist['schemes']
def get(h,model,cohort='daily_origins'):return next(r for r in rows if r['horizon_days']==h and r['model']==model and r['cohort']==cohort)
def v(x,d=4):return '—' if x is None else str(x) if isinstance(x,int) else f'{x:.{d}f}'
def metrics_table(rr,fields,size=8.5):
    t=rep.table(['模型']+[title for k,title,digits in fields],[[LABEL[r['model']]]+[v(r.get(k),digits) for k,title,digits in fields] for r in rr],[1.9]+[5.2/len(fields)]*len(fields),size)
    for r in rr:
        for k,title,digits in fields:bound.append(dict(model=r['model'],horizon=r['horizon_days'],cohort=r['cohort'],key=k,value=r.get(k),display=v(r.get(k),digits)))
    return t
rep.p('NLR严格协议重跑评估汇报',True,18,'Title');rep.p('所有候选提前登记  独立目录重新计算  2026年10月13日',style='Subtitle')
rep.p('评价协议总体适合科学的探索性负荷预测比较。它按时间留出、限制历史填充、冻结后才评分，并同时报告7天重叠和非重叠窗口。正式研究仍需要核实源时区，并取得尚未查看的新月份或新场站。',True)
rep.p('本轮目标是五项功率各自15分钟均值之和，属于派生分项汇总，不是直接设施总表。与上轮单独NLR报告的默认标签、日期窗口和计分口径一致；此次重新生成预测并统一提前登记候选。更早的Google或其他NLR口径分数不混比。',size=10)
matrix=[]
for name in ORDER:
    u,w=get(1,name),get(7,name);matrix.append([LABEL[name]]+[v(u[k]) for k in ['curve_mape_percent','curve_wape_percent']]+[v(w[k]) for k in ['curve_mape_percent','curve_wape_percent','daily_mape_percent','daily_wape_percent']])
t=rep.table(['主要模型或方案','24h曲线\nMAPE %','24h曲线\nWAPE %','7d曲线\nMAPE %','7d曲线\nWAPE %','7d日均\nMAPE %','7d日均\nWAPE %'],matrix,[1.95]+[5.15/6]*6,8.6)
for j in range(1,7):
    best=min(float(row[j]) for row in matrix if row[j]!='—')
    for i,row in enumerate(matrix):
        if row[j]!='—' and abs(float(row[j])-best)<1e-9:
            for run in t.rows[i+1].cells[j].paragraphs[0].runs:run.font.bold=True
for name in ORDER:
    for h,keys in [(1,['curve_mape_percent','curve_wape_percent']),(7,['curve_mape_percent','curve_wape_percent','daily_mape_percent','daily_wape_percent'])]:
        r=get(h,name)
        for key in keys:bound.append(dict(model=name,horizon=h,cohort=r['cohort'],key=key,value=r[key],display=v(r[key])))
best1=min(ORDER,key=lambda n:get(1,n)['curve_wape_percent']);best7=min(ORDER,key=lambda n:get(7,n)['curve_wape_percent']);bestday=min(ORDER,key=lambda n:get(7,n)['daily_wape_percent'])
rep.p('粗体为本次已评价方法在测试中的描述性最低值。测试曲线WAPE最低：24小时为'+LABEL[best1]+'，7天为'+LABEL[best7]+'；7天日均WAPE最低为'+LABEL[bestday]+'。',True,10)
rep.p('没有一个方案在全部指标领跑。旧Google权重迁移的优势也不自动保留：本次原HDSF与第48轮源迁移的两项曲线WAPE均高于上述最低方案。第8轮的7天成绩来自已披露的扩展分支，不能改称为此前已经验证的七天推荐。',size=10)
rep.p('按共同的向前验证起报MAPE冻结选择：24小时为'+LABEL[sel['tasks']['1']['selected_method']]+'，7天为'+LABEL[sel['tasks']['7']['selected_method']]+'。全部十四个方案均在本轮测试前登记并参与这一选择。MAPE与WAPE权重不同，可能有不同名次；不按测试最低值改选，也不将100−WAPE称为准确率。',size=10)
rep.p('同一测试期：24小时328个完整窗口；7天251个完整窗口，按日滚动且可重叠。旧方案全部在此NLR数据重新推理，保留源参数并按14天输入适配；原1天组合的7天结果为扩展对照。',size=9.5)

new('模型定义与源推荐适配')
rep.table(['模型短名','本次实际实现'],[
 ['LightGBM目标拟合','历史功率、分项、温湿度统计和已知日历；两组叶数和正则配置初始只在训练期拟合，按验证起报MAPE选定，再固定配置重拟合训练加验证。没有未来真实天气或分项负荷。'],
 ['TimesFM3 Toto2 Moirai2','原预训练神经权重冻结，历史统一14天。初始缩放只用训练有效值95%分位数/0.8，最终缩放只用训练加验证；缩放不代表设施额定容量。本轮没有NLR神经梯度微调。'],
 ['HDSF四专家目标拟合','LightGBM、TimesFM3、Toto2、Moirai2为四个专家。原半日单纯形按绝对误差拟合，验证时进行向前时间交叉检验，并以MAPE选择全局与半日分支比例；拟合分数单列，不当作泛化。'],
 ['Chronos2与TiRex2等权','原强基线50%加50%，源模型原权重冻结。14天历史适配；七天TiRex原生320点分320加320加32，只反馈自身预测，不读取未来真值。'],
 ['三专家固定权重','Chronos2 25%、TiRex2 25%、TimesFM3 50%。沿用旧Google系数，重新生成NLR预测。原1天构型的7天应用标作延伸。'],
 ['三专家加Ridge校准','沿用第5轮Ridge权重及50%残差比例。Google生产代理不可用时用源缺失分支，NLR场站不伪造Google类别；既往预测误差仅用已结束的过去日期。7天按每个预测日使用当前已知历史作扩展。'],
 ['第8轮TimesFM3微调组合','上一组合加旧Google微调TimesFM3的差分，组件比例仍为源冻结50%。没有在本NLR测试上选步骤或调系数。7天分支为源1天模型延伸。'],
 ['原HDSF源权重迁移','保留第44轮专家参数、半日权重及源局部模型。原28天或21天历史在本次适配14天；同一设施的联合模式可退化为单变量，因此不同原组件可能一致。'],
 ['第48轮HDSF源权重迁移','在原HDSF之上，24小时用原PatchTST-FM日均差分，7天用原TimesFM3 Toto2 PatchTST-FM加权且逐日零均值的形状差分。只有一个设施，场站组收缩不启用。']], [1.8,5.3],9.7)
rep.p('本轮所有旧推荐方案均提前登记，在新目录从历史输入重新推理，与主要模型一同参与共同验证交叉检验选模。没有复制上轮预测缓存。旧权重的14天输入、单设施和缺失字段适配已预先指定，不称为原Google多域信息条件的原样复现。',size=10)

for H in [1,7]:
    new('未来24小时完整曲线与日均指标' if H==1 else'未来7天完整曲线与每日均值指标')
    rr=[get(H,n) for n in ORDER];rep.p(f'统一目标为派生分项汇总功率，误差单位kW。起报{rr[0]["origins"]}次，曲线点{rr[0]["points"]}个；MAPE有效点{rr[0]["mape_n"]}个，排除{rr[0]["mape_excluded"]}个。全部模型使用同一严格完整窗口。',size=10)
    metrics_table(rr,[('curve_mape_percent','曲线\nMAPE %',4),('curve_rmse','曲线RMSE\nkW',2),('curve_wape_percent','曲线\nWAPE %',4),('relative_curve_mape_percent_reduction_percent','MAPE相对\n基线下降 %',2),('relative_curve_wape_percent_reduction_percent','WAPE相对\n基线下降 %',2)])
    metrics_table(rr,[('daily_mape_percent','日均\nMAPE %',4),('daily_rmse','日均RMSE\nkW',2),('daily_wape_percent','日均\nWAPE %',4),('relative_daily_mape_percent_reduction_percent','MAPE相对\n基线下降 %',2),('relative_daily_wape_percent_reduction_percent','WAPE相对\n基线下降 %',2)])
    rep.p('相对误差降低率使用固定基线：24小时前一天同刻，7天前三天同刻均值。负值表示退步。MAPE排除绝对值不超过1e−8的标签，按每个起报先计算再等权平均；全零分母无定义时输出null。日均先按每96点聚合，再计算对应误差。',size=9.5)

new('峰值与其他原有诊断指标')
for H in [1,7]:
    rep.sub('24小时任务' if H==1 else'7天任务')
    metrics_table([get(H,n) for n in ORDER],[('p95_ape_percent','P95\nAPE %',3),('hit2_percent','命中2%\n%',2),('hit5_percent','命中5%\n%',2),('absolute_bias_percent','绝对\n偏差 %',3),('peak_mape_percent','峰值\nMAPE %',3),('peak_time_mae_minutes','峰时误差\n分钟',1),('peak_underestimate_percent','峰值低估\n%',2)],8.2)
rep.p('峰时取首个最大值所在15分钟格。峰值、日均与曲线的优胜方案不一定相同；完整MAE、日均命中率和计数保留在全指标CSV。此轮没有新增区间校准，不能用其他口径的旧区间覆盖率填表。',size=10)

new('其他误差与日均命中指标')
for H in [1,7]:
    rep.sub('24小时任务' if H==1 else'7天任务')
    metrics_table([get(H,n) for n in ORDER],[('curve_mae','曲线MAE\nkW',2),('daily_mae','日均MAE\nkW',2),('daily_hit2_percent','日均命中\n2% %',2),('daily_hit5_percent','日均命中\n5% %',2),('relative_curve_rmse_reduction_percent','曲线RMSE\n下降 %',2),('relative_daily_rmse_reduction_percent','日均RMSE\n下降 %',2)],8.5)
rep.p('曲线MAPE主口径是起报等权。点池化MAPE及有效点、有效起报、日均排除计数保存在全指标CSV，作为核对补充；完整窗口且零值排除数相同时两种MAPE相等。这里的命中率表示误差落入给定比例范围的点或日均比例，不能称为预测准确率。',size=10)

new('7天非重叠日历起报检查')
rr=[get(7,n,'nonoverlapping_calendar7day') for n in ORDER]
rep.p(f'从测试起始日按固定日历每7天取一次起报，再应用同样的历史及标签完整规则，共{rr[0]["origins"]}个窗口。相邻被评价窗口至少间隔7天，目标日期不重叠；仍属于同一设施的相关时间序列，不当作独立随机实验。',size=10)
metrics_table(rr,[('curve_mape_percent','曲线\nMAPE %',4),('curve_rmse','曲线RMSE\nkW',2),('curve_wape_percent','曲线\nWAPE %',4),('daily_mape_percent','日均\nMAPE %',4),('daily_rmse','日均RMSE\nkW',2),('daily_wape_percent','日均\nWAPE %',4)])
rep.p('逐日起报7天窗口会让同一日期在不同预测步长下重复计入评分，适合描述每天部署一次的多步预测。非重叠检查提供另一种时间覆盖；起报数量和日期不同，不能将两个分数的差简单解释成算法变好或变差。',size=10)

new('未来七天逐预测日曲线WAPE')
lead_rows=list(csv.DictReader((F/'results/新协议逐预测日测试指标.csv').open(encoding='utf-8-sig')))
lead_matrix=[]
for name in ORDER:
    rr=sorted([r for r in lead_rows if r['horizon_days']=='7' and r['cohort']=='daily_origins' and r['model']==name],key=lambda r:int(r['forecast_day']))
    assert len(rr)==7
    lead_matrix.append([LABEL[name]]+[v(float(r['curve_wape_percent'])) for r in rr])
    for r in rr:bound.append(dict(model=name,horizon=7,cohort='daily_origins',forecast_day=int(r['forecast_day']),key='curve_wape_percent',value=float(r['curve_wape_percent']),display=v(float(r['curve_wape_percent']))))
rep.table(['具体模型']+['第'+str(k)+'天 %' for k in range(1,8)],lead_matrix,[1.9]+[5.2/7]*7,8.4)
rep.p('这些是同一条672点预测曲线按未来第1至7天拆分后的成绩，每格评价对应96点曲线，不是每日平均负荷的误差。各预测日共享同一批251个起报。预测过程中没有读取当日起报之后的真实功率或天气；滚动更新只发生在下一次实际起报之前。',size=10)

new('验证选择与科学性审查')
validation=[]
for H in [1,7]:
    task=sel['tasks'][str(H)]
    for n in ORDER:
        score=task['cv_scores'][n];validation.append([str(H)+'天',LABEL[n],v(score['curve_mape_percent']),v(score['curve_wape_percent']),'冻结推荐' if n==task['selected_method'] else'比较'])
rep.table(['任务','候选','向前检验MAPE %','向前检验WAPE %','结论'],validation,[.5,2.1,1.6,1.6,1.3],9)
new('选模冻结与科学性边界')
rep.p('验证交叉检验分三次向前留出，融合拟合只用更早的验证日期；7天标签越过留出起点的样本剔除。每折LightGBM配置也只在此前验证日期选择。随后按选定比例拟合全部验证融合权重，在测试前同时冻结两个预测时距。',size=10)
rep.table(['任务','半日分支比例','融合样本内MAPE %','融合样本内WAPE %'],[[str(H)+'天',v(sel['tasks'][str(H)]['alpha']),v(sel['tasks'][str(H)]['fusion_in_sample_score']['curve_mape_percent']),v(sel['tasks'][str(H)]['fusion_in_sample_score']['curve_wape_percent'])] for H in [1,7]],[.8,1.8,2.25,2.25],9)
rep.p('半日分支比例由向前验证选择：24小时为'+v(sel['tasks']['1']['alpha'])+'，7天为'+v(sel['tasks']['7']['alpha'])+'。比例为0表示实际退回全局单纯形，不能归因于12小时分段。融合样本内分数使用拟合过的标签，只作诊断。',size=10)
rep.p('这套协议科学之处是时间顺序、因果输入、共享评价窗口、分层处理融合选模及冻结配置。其限制是：源时区未验证，严格缺测筛选只评价可用日期；已查看的2023至2025区间和未知基础模型预训练语料重叠，使此次不能称为新的严格盲测。全部十四个方案本轮都在首次测试推理之前登记并冻结；这不会消除之前看过同一测试时期的事实。',size=10)
rep.p('如果最终研究重点是曲线WAPE领先，应在下一份尚未评分的新协议中预先指定WAPE主选模、MAPE辅助约束。当前已冻结的MAPE规则保持不变，避免看过测试成绩后换目标。时间交叉验证原则见文末资料。',size=10)

new('原始字段与处理边界')
rep.table(['字段','含义与处理'],[
 ['ts','原始时间戳无时区。临时按UTC解释，质量报告保留timezone_verified false；不能据此证明真实UTC或当地自然日。'],
 ['it_power_kw','数据中心IT设备功率，kW。可另建IT功率预测协议，本轮未运行；实际目标是五项派生汇总。'],
 ['cooling_kw','室外冷却设备风机和伴热等功率，包含专用过滤泵的冷却负荷记录。'],
 ['hvac_kw','暖通设备功率，包括电气房间相关风机及补风机组。'],
 ['plug_and_light_kw','照明与插座，以及相关辅助负荷，kW。'],
 ['pump_kw','数据中心及冷却水循环泵等功率，kW。按记录值处理，不另外添加2.67kW。'],
 ['outdoor_air_temp','本地实际天气字段；官方说明为华氏度，先换算C=(F−32)×5/9再聚合。'],
 ['outdoor_air_humidity','相对湿度百分比；不把百分数误当0至1比例。'],
 ['facility_power_kw','五项各自区间均值的派生合计；任何一项区间均值缺失则缺失，不是直接测量的设施总功率。']], [1.55,5.55],9.5)
rep.p('先删除配置日期范围外记录，包括天气1970年时间。固定宽泛过滤范围为单项功率0至20000kW、温度−70至60摄氏度、湿度0至100%；这些是本轮提前设置的质量界限，不是已核实的设施额定容量。重复时间戳取均值，再按左闭右开15分钟区间取有限记录均值，时间标注区间起点。',size=10)
rep.p('有效区间不要求15条分钟记录都齐全，原始有限记录数单独保存。只对历史协变量最多前向填充4格；目标及历史目标不填充，无后向填充或双向插值。第一和最后完整96格目标日确定边界，内部缺测日期保留在时间轴。',size=10)
rep.p('已独立从原始分钟Parquet重新累计有限值和计数，核对五项合计、温度转换、异常值屏蔽及天气1970年记录剔除，结果与处理数组一致。质量报告的跳过数按24小时和7天两个任务合计；各任务逐日起报和跳过原因另见窗口资格CSV。',size=9.5)
rep.table(['时期','日期','合格24小时起报','合格7天起报'],[[phase,q['split_dates'][phase][0]+' 至 '+q['split_dates'][phase][1],str(q['eligible_origins'][phase+'1']),str(q['eligible_origins'][phase+'7'])] for phase in ['train','validation','test']],[.9,3.2,1.5,1.5],9)
rep.p('训练、验证、测试按连续完整日历网格约60%、20%、20%划分，内部缺测不删日期。窗口全部标签必须留在所属划分内，缺失历史或标签的窗口按固定规则跳过，数量及原因保留；预测起点t最多读取t−15分钟区间。实际未来IT、制冷和天气观测都不进入模型。',size=10)
rep.p('未来标签是否存在只用于定义可评分窗口，没有输入未来标签的数值。严格完整窗口会偏向无长缺测的日期；正式部署应先依据历史产生预测，再单独标注未来真值的评分资格。',size=9.5)
rep.p('原始Parquet与许可保留，处理后15分钟CSV、区间记录覆盖CSV、窗口跳过清单和模型预测都单独保存。TimesFM3原权重及其Google微调权重的非商业研究用途限制继续适用。所有当前Google部署文件保持原样。',size=10)
rep.sub('资料与复核入口')
for title,url in [('官方NLR数据目录与字段说明','https://data.nlr.gov/submissions/300'),('Hyndman与Athanasopoulos 时间序列交叉验证','https://otexts.com/fpp3/toolbox.html#time-series-cross-validation'),('预测误差和MAPE适用性','https://otexts.com/fpp3/toolbox.html#evaluating-point-forecast-accuracy')]:
    p=rep.p(title,size=9.5);rep.link(p,'  原文',url)
rep.p('完整数值见根目录“NLR严格重跑_新协议主要模型测试全指标.csv”及配套Python表。曲线和日均MAPE、RMSE、WAPE与相对基线下降比例均已导出。本轮为重新生成的预测结果，未用旧Google分数填表。复核证据见本轮results与data目录。',size=9.5)
out=R/'NLR严格协议重跑评估汇报.docx';doc.save(out);re=style.Document(out)
assert [[c.text for row in t.rows for c in row.cells] for t in re.tables]==[[c.text for row in t.rows for c in row.cells] for t in doc.tables]
save(F/'docs/document_content_verification.json',dict(all_passed=True,utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),document=str(out),document_sha256=sha(out),bound_numeric_values=bound,table_count=len(doc.tables),source_audit_sha256=sha(F/'results/independent_verification.json'),rows=42,results_first=True,all_candidates_pre_registered_before_this_test=True))
print(json.dumps(dict(all_passed=True,document=str(out),tables=len(doc.tables),numeric_fields=len(bound)),ensure_ascii=False))
